"""
Оформление месячного отчёта WB для человека — в том же стиле, что отчёт Ozon (report_design.py).

Видимые листы «Сводка» и «Заказы» строятся из финансового отчёта WB (wb reports/<ГГГГ-ММ>.json) при каждом
формировании Excel. Прежняя раскладка «Заказы» с блоком показателей в P:Q уходит на скрытый лист «Данные»
(report_book.DATA_SHEETS) — её читает AI-ассистент.

Деньги по заказу (srid) — фактические из детализации WB:
- цена продажи — сколько заплатил покупатель (retailAmount, уже со скидкой WB);
- скидка WB — ваша цена со скидкой продавца (retailPriceWithDisc × кол-во) − цена продажи. Её WB оплачивает сам
  из своей комиссии, поэтому в расчёте она идёт плюсом;
- комиссия WB — ваша цена − к перечислению − эквайринг: сколько WB удержал по ставке (commissionPercent, 38–41,5%).
  Цена продажи + скидка WB − комиссия − эквайринг = к перечислению ровно;
- начислено WB — всё, что WB перечислит по заказу за вычетом логистики и прочих удержаний (как в «Данных»);
- прибыль — начислено WB − себестоимость.
"""
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from scripts import wb_finance as finance
from scripts.report_book import CHARGES, DATA_SHEETS, ORDERS
from scripts.report_design import (
    ACCENT, ACCENT_SOFT, BG, DANGER, DANGER_SOFT, DATE, INT, MONEY, MONEY0, MUTED, PCT, PRIMARY, PRIMARY_SOFT,
    RUB, RUB_KOP, RUB_SIGNED, TEXT, fill, header_row, section, table, thin, widths,
)

STATUS = {"продажа": "Продажа", "возврат": "Возврат"}
VISUAL_SHEETS = ["Сводка", ORDERS, CHARGES]
# Понятные названия операций WB, не привязанных к заказам месяца.
CHARGE_NAMES = {
    "Логистика": "Логистика по заказам других месяцев",
    "Доставка": "Логистика по заказам других месяцев",
    "Коррекция стоимости доставки": "Коррекция логистики",
    "Хранение": "Хранение на складах WB",
    "Коррекция хранения": "Коррекция хранения",
    "Штраф": "Штрафы",
    "Разовое изменение срока перечисления денежных средств": "Досрочный вывод денег",
}
EXTRA_SHEETS = {  # лист, который пишет wb_monthly_report -> {колонка: формат}
    "Приёмка": {"Стоимость приёмки, ₽": MONEY, "Количество, шт.": INT},
    "Хранение": {"Стоимость хранения, ₽": MONEY},
    "Воронка продаж": {"Просмотры карточки": INT, "В корзину": INT, "Заказы": INT, "Сумма заказов, ₽": RUB,
                       "Выкупы": INT, "Сумма выкупов, ₽": RUB, "Отмены": INT, "В корзину, %": '0.0"%"',
                       "Корзина → заказ, %": '0.0"%"', "Выкуп, %": '0.0"%"'},
}


def _n(value):
    return float(finance.number(value or 0))


def effect(row):
    """Сколько строка детализации добавила к «Итого к оплате» (в сумме по всем строкам — ровно bankPaymentSum).
    У возврата forPay положительный, но WB его забирает: сверено с bankPaymentSum на возвратах 11.2025."""
    pay = -_n(row.get("forPay")) if str(row.get("sellerOperName") or "").strip().lower() == "возврат" else _n(row.get("forPay"))
    return (pay - _n(row.get("deliveryService")) - _n(row.get("paidStorage"))
            - _n(row.get("paidAcceptance")) - _n(row.get("deduction")) - _n(row.get("penalty"))
            + _n(row.get("additionalPayment")) - _n(row.get("paymentSchedule")))


def scheme(office):
    """FBS — товар уехал со своего склада, FBW — со склада WB. По названию склада в детализации."""
    office = str(office or "").strip()
    if not office:
        return ""
    return "FBS" if "поставщика" in office.lower() or office.endswith(" МП") else "FBW"


def _date(text):
    try:
        return datetime.strptime(str(text)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


# ── Данные ─────────────────────────────────────────────────────────────────

def orders(data, costs, returns_by_srid=None):
    """Строка на заказ (srid) с продажей или возвратом в этом месяце."""
    returns_by_srid = returns_by_srid or {}
    by_srid = {}
    for row in data.get("details", []):
        srid = str(row.get("srid") or "")
        if not srid:
            continue
        o = by_srid.setdefault(srid, {
            "srid": srid, "date": None, "name": "", "article": "", "scheme": "", "ops": [], "qty": defaultdict(float),
            "price": 0.0, "seller_price": 0.0, "wb_discount": 0.0, "commission": 0.0, "acquiring": 0.0,
            "logistics": 0.0, "accrued": 0.0})
        o["logistics"] -= _n(row.get("deliveryService"))
        o["accrued"] += effect(row)
        operation = str(row.get("sellerOperName") or "").strip().lower()
        if operation not in STATUS:
            continue
        sign = -1 if operation == "возврат" else 1
        if STATUS[operation] not in o["ops"]:
            o["ops"].append(STATUS[operation])
        o["qty"][str(row.get("nmId") or "")] += sign * _n(row.get("quantity"))
        price, pay, acquiring = _n(row.get("retailAmount")), _n(row.get("forPay")), _n(row.get("acquiringFee"))
        seller = _n(row.get("retailPriceWithDisc")) * _n(row.get("quantity")) or price
        o["price"] += sign * price
        o["seller_price"] += sign * seller
        o["wb_discount"] += sign * (seller - price)
        o["commission"] -= sign * (seller - pay - acquiring)
        o["acquiring"] -= sign * acquiring
        o["name"] = o["name"] or str(row.get("title") or "")
        o["article"] = o["article"] or str(row.get("vendorCode") or "")
        o["scheme"] = o["scheme"] or scheme(row.get("officeName"))
        day = _date(row.get("saleDt") or row.get("rrDate"))
        if day and (o["date"] is None or day < o["date"]):
            o["date"] = day
    result = []
    for o in by_srid.values():
        if not o["ops"]:
            continue  # логистика и хранение по заказам другого месяца — «Расходы вне заказов»
        cost, missing = 0.0, []
        for nm_id, qty in o["qty"].items():
            # qty < 0 — возврат продажи прошлого месяца: товар вернулся на склад, себестоимость возвращается.
            if qty:
                unit = costs.get(nm_id)
                if unit is None:
                    missing.append(nm_id)
                else:
                    cost += float(unit) * qty
        record = returns_by_srid.get(o["srid"]) or {}
        notes = [f"Причина возврата: {record['reason']}"] if record.get("reason") else []
        if missing:
            notes.append("нет себестоимости в wb_costs.xlsx")
        result.append({
            "date": o["date"], "srid": o["srid"], "name": o["name"], "article": o["article"], "scheme": o["scheme"],
            "status": " и ".join(o["ops"]).capitalize(), "qty": sum(o["qty"].values()),
            "price": round(o["price"], 2), "seller_price": round(o["seller_price"], 2),
            "wb_discount": round(o["wb_discount"], 2), "commission": round(o["commission"], 2), "acquiring": round(o["acquiring"], 2),
            "logistics": round(o["logistics"], 2), "accrued": round(o["accrued"], 2), "cost": None if missing else -round(cost, 2),
            "profit": None if missing else round(o["accrued"] - cost, 2), "note": "; ".join(notes) or None,
        })
    return sorted(result, key=lambda o: (o["date"] or date.min, o["srid"]))


def weekly_payouts(data):
    """Начисления к выплате по неделям месяца (отчёты WB ежедневные, выплата — раз в неделю)."""
    weeks = defaultdict(lambda: [None, None, 0.0])
    for report in data.get("reports", []):
        start, end = _date(report.get("dateFrom")), _date(report.get("dateTo"))
        if not start:
            continue
        monday = start - timedelta(days=start.weekday())
        week = weeks[monday]
        week[0] = min(filter(None, [week[0], start]))
        week[1] = max(filter(None, [week[1], end or start]))
        week[2] += _n(report.get("bankPaymentSum"))
    return [(first, last, round(amount, 2)) for _monday, (first, last, amount) in sorted(weeks.items())]


# ── Лист «Сводка» ──────────────────────────────────────────────────────────

def _rub(kopecks):
    return kopecks / 100


def _rubles(kopecks):
    """«192 759 ₽» — для текста; число в ячейке форматирует Excel."""
    return f"{kopecks / 100:,.0f} ₽".replace(",", " ")


def build_settlement(ws, row, settlement, key):
    """Блок «Расчёты с WB за этот месяц». Возвращает следующую свободную строку."""
    section(ws, row, "Расчёты с WB за этот месяц", span=8)
    month = (settlement or {}).get("months", {}).get(key)
    ws.merge_cells(f"B{row + 1}:I{row + 1}")
    status = ws[f"B{row + 1}"]
    if not month:
        status.value = ("Сверка выплат ещё не проводилась — она обновится при следующей загрузке месяца в «Бизнес-сводке»"
                        if not settlement else "Итоги WB за этот месяц ещё не попали в сверку выплат")
        color, bg = MUTED, BG
    elif month["done"]:
        status.value, color, bg = "✔  WB полностью рассчитался за месяц", PRIMARY, PRIMARY_SOFT
    else:
        status.value = f"⏳  WB ещё не рассчитался: осталось получить {_rubles(month['left'])}"
        color, bg = ACCENT, ACCENT_SOFT
    status.font = Font(bold=True, size=12, color=color)
    status.alignment = Alignment(indent=1, vertical="center")
    for c in ws[f"B{row + 1}:I{row + 1}"][0]:
        c.fill = fill(bg)
    ws.row_dimensions[row + 1].height = 28
    r = row + 3
    if month:
        for c, text in zip("BCDF", ["", "Начислено", "Закрыто", "Осталось получить"]):
            ws[f"{c}{r}"] = text
            ws[f"{c}{r}"].font = Font(size=9, bold=True, color=MUTED)
        ws.merge_cells(f"D{r}:E{r}")
        for i, (kind, due) in enumerate(month["due"].items(), start=r + 1):
            done = month["closed"][kind]
            ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"], ws[f"F{i}"] = kind, _rub(due), _rub(done), _rub(due - done)
            ws.merge_cells(f"D{i}:E{i}")
            for c in "CDF":
                ws[f"{c}{i}"].number_format = RUB
                ws[f"{c}{i}"].alignment = Alignment(horizontal="left")
            ws[f"F{i}"].font = Font(bold=True, color=PRIMARY if due == done else ACCENT)
            for c in "BCDEF":
                ws[f"{c}{i}"].border = Border(bottom=thin)
        r += 1 + len(month["due"])
        notes = []
        if month["profit"] is not None and month["profit"] < 0:
            notes.append(f"Месяц убыточный: себестоимость больше начисленного на {_rubles(-month['profit'])}, "
                         "закрывать по прибыли нечего.")
        if month["reason"]:
            notes.append(f"Себестоимость и прибыль не разделены: {month['reason']}.")
        if not month["due"]:
            notes.append("За месяц WB больше удержал, чем начислил: этот минус уменьшил выплаты по прошлым месяцам.")
        for note in notes:
            ws[f"B{r}"] = note
            ws[f"B{r}"].font = Font(size=9, italic=True, color=ACCENT)
            r += 1
    if settlement:
        checked = datetime.fromisoformat(settlement["checked_at"])
        ws[f"B{r}"] = (f"Всего выведено с WB с начала магазина: {_rubles(settlement['withdrawn'])} · на балансе WB "
                       f"{_rubles(settlement['balance'])}, из них к выводу {_rubles(settlement['for_withdraw'])} · "
                       f"сверка от {checked:%d.%m.%Y}")
        ws[f"B{r}"].font = Font(size=9, color=MUTED)
        ws[f"B{r + 1}"] = ("Выведено = всё начисленное WB по отчётам с начала магазина − текущий баланс: дат отдельных "
                           "выводов WB не отдаёт. Деньги закрывают месяцы по порядку, сначала себестоимость, потом прибыль.")
        ws[f"B{r + 1}"].font = Font(size=9, color=MUTED)
        r += 2
    return r


def commission_rate(rows):
    """Средняя ставка комиссии WB от ваших цен со скидкой продавца (доля) или None."""
    seller = sum(o["seller_price"] for o in rows)
    return -sum(o["commission"] for o in rows) / seller if seller else None


def commission_label(rows):
    rate = commission_rate(rows)
    return f"Комиссия WB по ставке {rate * 100:.1f}%".replace(".", ",") if rate is not None else "Комиссия WB"


def profit_lines(summary, rows):
    """Статьи «Из чего сложилась прибыль»: выручка, расходы WB, себестоимость. Сумма — прибыль (или None).

    Расходы — отрицательные числа. «Прочие операции WB» — остаток до «К выплате» (обычно ноль)."""
    revenue = float(summary.get("revenue") or 0)
    payout = float(summary.get("payout") or 0)
    known = [
        ("Скидка WB покупателям (оплатил WB)", round(sum(o["wb_discount"] for o in rows), 2)),
        (commission_label(rows), round(sum(o["commission"] for o in rows), 2)),
        ("Эквайринг", round(sum(o["acquiring"] for o in rows), 2)),
        ("Логистика", -float(summary.get("logistics") or 0)),
        ("Хранение на складах WB", -float(summary.get("storage") or 0)),
        ("Приёмка поставок", -float(summary.get("acceptance") or 0)),
        ("Удержания", -float(summary.get("deductions") or 0)),
        ("Штрафы", -float(summary.get("penalties") or 0)),
        ("Доплаты WB", float(summary.get("additional") or 0)),
        ("Досрочный вывод денег", -float(summary.get("fast_withdraw") or 0)),
    ]
    other = round(payout - revenue - sum(value for _label, value in known), 2)
    costs = summary.get("costs")
    return ([("Выручка (заплатили покупатели)", revenue)] + known + [("Прочие операции WB", other)]
            + [("Себестоимость товаров", -float(costs) if costs is not None else None)])


def order_stats(data, rows):
    """Блок «Заказы»: с продажей, по схемам, возвраты, средняя скидка WB покупателю (доля или None)."""
    sales = [o for o in rows if "Продажа" in o["status"]]
    by_scheme = defaultdict(int)
    for o in sales:
        by_scheme[o["scheme"] or "не указано"] += 1
    paid = [(_n(r.get("retailPriceWithDisc")), _n(r.get("retailAmount"))) for r in data.get("details", [])
            if str(r.get("sellerOperName") or "").strip().lower() == "продажа" and _n(r.get("retailPriceWithDisc"))]
    spp = 1 - sum(a for _p, a in paid) / sum(p for p, _a in paid) if paid else None
    return {"sales": len(sales), "returns": sum(1 for o in rows if "Возврат" in o["status"]),
            "by_scheme": dict(sorted(by_scheme.items())), "spp": spp}


def build_summary(ws, title_month, created, data, summary, rows, settlement=None):
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 2, "B": 40, "C": 16, "D": 14, "E": 3, "F": 30, "G": 16, "H": 16, "I": 16, "J": 2})
    revenue = float(summary.get("revenue") or 0)
    payout = float(summary.get("payout") or 0)
    profit = summary.get("profit")
    profit = float(profit) if profit is not None else None
    sales = [o for o in rows if "Продажа" in o["status"]]

    ws["B2"] = f"Wildberries · {title_month}"
    ws["B2"].font = Font(size=22, bold=True, color=TEXT)
    ws["B3"] = f"{len(rows)} заказов · продано {len(sales)} · отчёт сформирован {created:%d.%m.%Y}"
    ws["B3"].font = Font(size=11, color=MUTED)

    good = profit is None or profit >= 0
    cards = [
        ("Выручка", revenue, RUB, f"{len(sales)} заказов с продажей", PRIMARY_SOFT, TEXT),
        ("Чистая прибыль", profit if profit is not None else "—", RUB_SIGNED,
         "после всех расходов и себестоимости" if profit is not None else "не посчитана — см. ниже",
         PRIMARY_SOFT, PRIMARY if good else DANGER),
        ("Маржа", profit / revenue if profit is not None and revenue else "—", PCT, "чистая прибыль / выручка",
         PRIMARY_SOFT, TEXT),
        ("К выплате от WB", payout, RUB, "итог финансовых отчётов WB за месяц", ACCENT_SOFT, ACCENT),
    ]
    month = (settlement or {}).get("months", {}).get(f"{data['year']}-{data['month']:02d}")
    if month:
        left = _rub(month["left"])
        cards[3] = ("WB ещё должен", left, RUB, "месяц закрыт выплатами" if month["done"] else "по сверке выплат ниже",
                    PRIMARY_SOFT if month["done"] else ACCENT_SOFT, PRIMARY if month["done"] else ACCENT)
    for (label, value, fmt, note, bg, color), col in zip(cards, ["B", "C", "F", "G"]):
        span = {"B": "B", "C": "D", "F": "F", "G": "I"}[col]
        for r, (text, font) in enumerate([(label, Font(size=10, color=MUTED)), (value, Font(size=20, bold=True, color=color)),
                                          (note, Font(size=9, color=MUTED))], start=5):
            ws.merge_cells(f"{col}{r}:{span}{r}")
            c = ws[f"{col}{r}"]
            c.value, c.font = text, font
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            if r == 6:
                c.number_format = fmt
            for cc in ws[f"{col}{r}:{span}{r}"][0]:
                cc.fill = fill(bg)
    ws.row_dimensions[6].height = 34

    # Из чего сложилась прибыль: строки в сумме дают «К выплате» минус себестоимость.
    section(ws, 9, "Из чего сложилась прибыль", span=3)
    lines = profit_lines(summary, rows)
    for c, text in zip("BCD", ["Статья", "Сумма", "Доля от выручки"]):
        ws[f"{c}10"] = text
        ws[f"{c}10"].font = Font(size=9, bold=True, color=MUTED)
        ws[f"{c}10"].alignment = Alignment(horizontal="left" if c == "B" else "right")
    row = 11
    for label, value in lines:
        if value is not None and abs(value) < 0.005 and row > 11 and label != "Себестоимость товаров":
            continue  # пустые статьи не показываем
        ws[f"B{row}"], ws[f"C{row}"] = label, value if value is not None else "—"
        ws[f"C{row}"].number_format = RUB if row == 11 else MONEY0
        ws[f"D{row}"] = f"=IF(OR($C$11=0,ISTEXT(C{row})),0,ABS(C{row})/$C$11)" if row > 11 else 1
        ws[f"D{row}"].number_format = PCT
        ws[f"D{row}"].font = Font(color=MUTED)
        for c in "BCD":
            ws[f"{c}{row}"].border = Border(bottom=thin)
        row += 1
    ws[f"B{row}"] = "Чистая прибыль"
    ws[f"C{row}"] = f"=SUM(C11:C{row - 1})" if profit is not None else "—"
    ws[f"D{row}"] = f"=IF($C$11=0,0,C{row}/$C$11)" if profit is not None else None
    ws[f"C{row}"].number_format, ws[f"D{row}"].number_format = RUB_SIGNED, PCT
    for c in "BCD":
        ws[f"{c}{row}"].font = Font(bold=True, size=12, color=PRIMARY if good else DANGER)
        ws[f"{c}{row}"].fill = fill(PRIMARY_SOFT)
    ws.conditional_formatting.add(f"D12:D{row - 1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                   end_value=1, color="E8B9B2", showValue=True))
    note_row = row + 1
    for message in summary.get("notes", []):
        ws.merge_cells(f"B{note_row}:D{note_row}")
        ws[f"B{note_row}"] = message
        ws[f"B{note_row}"].font = Font(size=9, italic=True, color=ACCENT)
        ws[f"B{note_row}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[note_row].height = 30
        note_row += 1
    check_row = note_row

    # Заказы
    section(ws, 9, "Заказы", col=6, span=4)
    stats = order_stats(data, rows)
    spp = stats["spp"]
    order_lines = [("Заказов с продажей", stats["sales"], INT)]
    order_lines += [(f"  {name}", count, INT) for name, count in stats["by_scheme"].items()]
    order_lines += [("Возвраты", stats["returns"], INT), ("Средний чек", float(summary.get("average") or 0), RUB)]
    if spp is not None:
        order_lines.append(("Средняя скидка WB покупателю", spp, PCT))
    r = 10
    for label, value, fmt in order_lines:
        ws[f"F{r}"], ws[f"G{r}"] = label, value
        ws[f"G{r}"].number_format = fmt
        ws[f"F{r}"].border = ws[f"G{r}"].border = Border(bottom=thin)
        if label.startswith("  "):
            ws[f"F{r}"].font = ws[f"G{r}"].font = Font(color=MUTED)
        r += 1
    if spp is not None:
        ws[f"F{r}"] = "Скидку WB оплачивает WB, на вашу выплату она не влияет."
        ws[f"F{r}"].font = Font(size=9, italic=True, color=MUTED)

    # Расчёты с WB
    row = build_settlement(ws, max(check_row, r + 2) + 2, settlement, f"{data['year']}-{data['month']:02d}")

    # Начисления по неделям
    row += 2
    section(ws, row, "Начисления WB к выплате по неделям", span=8)
    for c, text in zip("BCD", ["Неделя", "К выплате", ""]):
        ws[f"{c}{row + 1}"] = text
        ws[f"{c}{row + 1}"].font = Font(size=9, bold=True, color=MUTED)
    ws[f"C{row + 1}"].alignment = Alignment(horizontal="right")
    weeks = weekly_payouts(data)
    for j, (first, last, amount) in enumerate(weeks, start=row + 2):
        ws[f"B{j}"] = f"{first:%d.%m} — {last:%d.%m.%Y}"
        ws[f"C{j}"] = amount
        ws[f"C{j}"].number_format = RUB_KOP
        for c in "BCD":
            ws[f"{c}{j}"].border = Border(bottom=thin)
    last = row + 2 + len(weeks)
    ws[f"B{last}"] = "Сколько WB начислил к выплате за дни недели по финансовым отчётам этого месяца."
    ws[f"B{last}"].font = Font(size=9, color=MUTED)

    # Топ товаров
    row = last + 3
    section(ws, row, "Товары: что принесло прибыль", span=8)
    by_article = defaultdict(lambda: {"name": "", "qty": 0.0, "revenue": 0.0, "profit": 0.0, "known": True})
    for o in sales:
        a = by_article[o["article"]]
        a["name"] = a["name"] or o["name"]
        a["qty"] += o["qty"]
        a["revenue"] += o["price"]
        if o["profit"] is None:
            a["known"] = False
        else:
            a["profit"] += o["profit"]
    top = sorted(by_article.items(), key=lambda kv: (kv[1]["known"], kv[1]["profit"]), reverse=True)
    for c, text in zip("BCDEFGH", ["Товар", "Артикул", "Продано, шт", "", "Выручка", "Прибыль", "Маржа"]):
        ws[f"{c}{row + 1}"] = text
        ws[f"{c}{row + 1}"].font = Font(size=9, bold=True, color=MUTED)
        ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="left" if c in "BC" else "right")
    for j, (article, a) in enumerate(top[:10], start=row + 2):
        name = a["name"] or ""
        ws[f"B{j}"], ws[f"C{j}"], ws[f"D{j}"] = (name[:44] + "…") if len(name) > 45 else name, article, a["qty"]
        ws[f"F{j}"] = a["revenue"]
        ws[f"G{j}"] = a["profit"] if a["known"] else "—"
        ws[f"H{j}"] = f"=IF(OR(F{j}=0,ISTEXT(G{j})),0,G{j}/F{j})" if a["known"] else None
        ws[f"C{j}"].alignment = Alignment(horizontal="left")
        ws[f"D{j}"].number_format, ws[f"F{j}"].number_format = INT, RUB
        ws[f"G{j}"].number_format, ws[f"H{j}"].number_format = RUB_SIGNED, PCT
        for c in "BCDEFGH":
            ws[f"{c}{j}"].border = Border(bottom=thin)
    end = row + 1 + min(len(top), 10)
    if top:
        ws.conditional_formatting.add(f"G{row + 2}:G{end}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                                         color="7FB8AD", showValue=True))
    ws[f"B{end + 1}"] = (f"Топ-10 из {len(top)} товаров по прибыли с проданных заказов. Все заказы — на листе «Заказы»."
                         if top else "Продаж в этом месяце нет.")
    ws[f"B{end + 1}"].font = Font(size=9, color=MUTED)

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ── Лист «Заказы» ──────────────────────────────────────────────────────────

def build_orders(ws, rows):
    heads = ["Дата продажи", "Номер заказа (srid)", "Товар", "Артикул", "Схема", "Статус", "Кол-во", "Цена продажи",
             "Скидка WB (оплатил WB)", "Комиссия WB", "Эквайринг", "Логистика", "Начислено WB", "Себестоимость", "Прибыль",
             "Комментарий"]
    header_row(ws, 1, heads)
    for i, o in enumerate(rows, start=2):
        values = [o["date"], o["srid"], o["name"], o["article"], o["scheme"], o["status"], o["qty"], o["price"],
                  o["wb_discount"], o["commission"], o["acquiring"], o["logistics"], o["accrued"], o["cost"], o["profit"],
                  o["note"]]
        for col, v in enumerate(values, start=1):
            ws.cell(row=i, column=col, value=v)
        ws.cell(row=i, column=1).number_format = DATE
        ws.cell(row=i, column=1).alignment = Alignment(horizontal="left")
        ws.cell(row=i, column=7).number_format = INT
        for col in range(8, 16):
            ws.cell(row=i, column=col).number_format = RUB_KOP if col == 15 else MONEY
    last = len(rows) + 1
    widths(ws, {"A": 13, "B": 24, "C": 46, "D": 11, "E": 8, "F": 18, "G": 9, "H": 14, "I": 15, "J": 14, "K": 12, "L": 13,
                "M": 15, "N": 15, "O": 14, "P": 40})
    if not rows:
        # Таблица Excel без строк данных — повреждённый файл; пустой месяц — только шапка и пояснение.
        ws.cell(row=2, column=1, value="Заказов в этом месяце нет").font = Font(italic=True, color=MUTED)
        return
    table(ws, f"A1:P{last}", "ЗаказыWB", "TableStyleLight8")
    total = last + 1
    ws.cell(row=total, column=1, value="Итого (по видимым строкам)").font = Font(bold=True)
    for col in range(7, 16):
        letter = get_column_letter(col)
        c = ws.cell(row=total, column=col, value=f"=SUBTOTAL(109,{letter}2:{letter}{last})")
        c.number_format = INT if col == 7 else (RUB_KOP if col == 15 else MONEY)
        c.font = Font(bold=True)
    for col in range(1, 17):
        ws.cell(row=total, column=col).fill = fill(PRIMARY_SOFT)
        ws.cell(row=total, column=col).border = Border(top=Side(style="medium", color=PRIMARY))
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "D2"
    ws.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="equal", formula=['"Возврат"'],
                                                            font=Font(color=DANGER), fill=fill(DANGER_SOFT)))
    ws.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="equal", formula=['"Продажа и возврат"'],
                                                            font=Font(color=MUTED), fill=fill(BG)))
    ws.conditional_formatting.add(f"I2:I{last}", CellIsRule(operator="greaterThan", formula=["0"],
                                                            font=Font(color=PRIMARY)))
    ws.conditional_formatting.add(f"O2:O{last}", CellIsRule(operator="lessThan", formula=["0"],
                                                            font=Font(color=DANGER, bold=True), fill=fill(DANGER_SOFT)))
    ws.conditional_formatting.add(f"N2:O{last}", FormulaRule(formula=['ISBLANK($N2)'], fill=fill(ACCENT_SOFT)))


# ── Сборка ─────────────────────────────────────────────────────────────────

# ── Лист «Расходы вне заказов» ─────────────────────────────────────────────

def charges(data):
    """Операции WB, не относящиеся к заказам с продажей или возвратом в этом месяце, с влиянием на выплату."""
    order_srids = {str(r.get("srid") or "") for r in data.get("details", [])
                   if str(r.get("sellerOperName") or "").strip().lower() in STATUS}
    result = []
    for row in data.get("details", []):
        srid = str(row.get("srid") or "")
        if srid and srid in order_srids:
            continue
        operation = str(row.get("sellerOperName") or "").strip() or "Без названия"
        comment = str(row.get("bonusTypeName") or "") or None
        name = CHARGE_NAMES.get(operation, operation)
        if operation in ("Логистика", "Доставка") and "отмен" in (comment or "").lower():
            name = "Логистика отменённых заказов (туда и обратно)"
        article = str(row.get("vendorCode") or "")
        result.append({"name": name, "operation": operation, "amount": round(effect(row), 2),
                       "date": _date(row.get("rrDate") or row.get("dateFrom")), "article": article,
                       "srid": srid if article else "",  # у хранения и вывода денег srid — служебный, не номер заказа
                       "comment": comment})
    return result


def build_charges(ws, items):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Расходы WB вне заказов"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = ("Хранение, приёмка, штрафы, досрочный вывод денег, логистика отменённых заказов и заказов других месяцев. "
                "Всё это WB уже вычел из «К выплате», в прибыли на «Сводке» учтено.")
    ws["A2"].font = Font(size=10, color=MUTED)
    groups = defaultdict(lambda: [0.0, 0])
    for item in items:
        groups[item["name"]][0] += item["amount"]
        groups[item["name"]][1] += 1
    paid = {name: g for name, g in groups.items() if abs(g[0]) >= 0.005}
    header_row(ws, 4, ["Вид расхода", "Сумма", "Операций"])
    i = 5
    for name, (total, count) in sorted(paid.items(), key=lambda kv: kv[1][0]):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=round(total, 2)).number_format = MONEY
        ws.cell(row=i, column=3, value=count).number_format = INT
        for col in range(1, 4):
            ws.cell(row=i, column=col).border = Border(bottom=thin)
        i += 1
    if not paid:
        ws.cell(row=i, column=1, value="Расходов вне заказов в этом месяце нет").font = Font(italic=True, color=MUTED)
        i += 1
    ws.cell(row=i, column=1, value="Итого").font = Font(bold=True)
    c = ws.cell(row=i, column=2, value=f"=SUM(B5:B{i - 1})" if paid else 0)
    c.number_format, c.font = MONEY, Font(bold=True)
    for col in range(1, 4):
        ws.cell(row=i, column=col).fill = fill(PRIMARY_SOFT)
    free = {name: g[1] for name, g in groups.items() if name not in paid}
    if free:
        i += 1
        ws.cell(row=i, column=1, value="Без суммы, справочно: " + "; ".join(
            f"{name} — {count}" for name, count in sorted(free.items()))).font = Font(size=9, color=MUTED)
    start = i + 3
    ws.cell(row=start - 1, column=1, value="Операции с суммой").font = Font(size=13, bold=True)
    header_row(ws, start, ["Вид расхода", "Сумма", "Дата", "Артикул", "Номер заказа (srid)", "Операция WB"])
    rows = sorted((item for item in items if abs(item["amount"]) >= 0.005),
                  key=lambda item: (item["date"] or date.min, item["name"]))
    for j, item in enumerate(rows, start=start + 1):
        values = [item["name"], item["amount"], item["date"], item["article"] or None, item["srid"] or None,
                  item["comment"] or item["operation"]]
        for col, value in enumerate(values, start=1):
            ws.cell(row=j, column=col, value=value)
        ws.cell(row=j, column=2).number_format = MONEY
        ws.cell(row=j, column=3).number_format = DATE
        ws.cell(row=j, column=3).alignment = Alignment(horizontal="left")
        ws.cell(row=j, column=6).font = Font(color=MUTED)
        for col in range(1, 7):
            ws.cell(row=j, column=col).border = Border(bottom=thin)
    widths(ws, {"A": 44, "B": 14, "C": 12, "D": 11, "E": 44, "F": 50})


# ── Приёмка, хранение, воронка ─────────────────────────────────────────────

def style_extra(ws, formats):
    """Оформить лист, записанный pandas (шапка в первой строке): шапка, форматы, ширина, таблица, итоги."""
    heads = [c.value for c in ws[1]]
    last = ws.max_row
    header_row(ws, 1, heads)
    ws.freeze_panes = "C2"
    for col, head in enumerate(heads, start=1):
        letter = get_column_letter(col)
        fmt = formats.get(head)
        width = 46 if head == "Название" else max(12, min(22, len(str(head)) + 4))
        ws.column_dimensions[letter].width = width
        if fmt:
            for row in range(2, last + 1):
                ws.cell(row=row, column=col).number_format = fmt
    if last < 2 or ws.tables:
        return  # пустой лист или уже оформлен (перерисовка готового отчёта)
    table(ws, f"A1:{get_column_letter(len(heads))}{last}", "WB" + "".join(ch for ch in ws.title if ch.isalnum()),
          "TableStyleLight8")
    total = last + 1
    ws.cell(row=total, column=1, value="Итого").font = Font(bold=True)
    for col, head in enumerate(heads, start=1):
        fmt = formats.get(head)
        if fmt in (MONEY, RUB, INT):
            letter = get_column_letter(col)
            c = ws.cell(row=total, column=col, value=f"=SUBTOTAL(109,{letter}2:{letter}{last})")
            c.number_format, c.font = fmt, Font(bold=True)
        ws.cell(row=total, column=col).fill = fill(PRIMARY_SOFT)
        ws.cell(row=total, column=col).border = Border(top=Side(style="medium", color=PRIMARY))


def saved_returns(ws):
    """Причины и статусы возвратов, уже записанные в «Данные» (колонки «Номер заказа», «… возврата (WB)»)."""
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows, ()))
    if "Номер заказа" not in header or "Причина возврата (WB)" not in header:
        return {}
    number, reason = header.index("Номер заказа"), header.index("Причина возврата (WB)")
    status = header.index("Статус возврата (WB)") if "Статус возврата (WB)" in header else None
    return {str(r[number]): {"reason": r[reason], "status": r[status] if status is not None and len(r) > status else None}
            for r in rows if len(r) > reason and r[number] and (r[reason] or (status is not None and len(r) > status and r[status]))}


def apply(wb, title, data, costs, returns_by_srid=None, settlement=None):
    """Спрятать прежний лист «Заказы» под именем «Данные» и построить «Сводку» и «Заказы» по данным WB.

    returns_by_srid=None — взять причины возвратов из «Данных» (перерисовка готового отчёта без запросов к WB).
    """
    hidden = DATA_SHEETS[ORDERS]
    if hidden not in wb.sheetnames:
        wb[ORDERS].title = hidden
    wb[hidden].sheet_state = "hidden"
    if returns_by_srid is None:
        returns_by_srid = saved_returns(wb[hidden])
    for name in VISUAL_SHEETS:
        if name in wb.sheetnames and name != hidden:
            wb.remove(wb[name])
    summary = finance.summarize(data, costs)
    rows = orders(data, costs, returns_by_srid)
    created = datetime.now()
    build_summary(wb.create_sheet("Сводка", 0), title, created, data, summary, rows, settlement)
    build_orders(wb.create_sheet(ORDERS, 1), rows)
    build_charges(wb.create_sheet(CHARGES, 2), charges(data))
    for name, formats in EXTRA_SHEETS.items():
        if name in wb.sheetnames:
            style_extra(wb[name], formats)
    wb.move_sheet(hidden, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(hidden))
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Сводка"
    wb["Сводка"].sheet_properties.tabColor = PRIMARY
    wb.active = 0

