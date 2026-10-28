"""
Оформление месячного отчёта Ozon для человека: листы «Сводка», «Заказы», «Реклама», «Расходы вне заказов».

Данные остаются в прежней раскладке на скрытых листах «Данные», «Данные кампаний», «Данные расходов»
(см. report_book.py) — их читают программа, сверка выплат, расчёт цен, ABC/XYZ и AI-ассистент.
Видимые листы строятся из скрытых заново при каждом изменении отчёта и ничего своего не хранят.

apply_design(книга, «Сентябрь 2026») — превратить отчёт старого вида в новый (или перерисовать новый).
Из консоли: python scripts/report_design.py — перевести в новый вид все отчёты старого вида в reports/Ozon/.
"""
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    from scripts.report_book import CAMPAIGNS, CHARGES, DATA_SHEETS, ORDERS, data_sheet, is_designed
except ModuleNotFoundError:
    from report_book import CAMPAIGNS, CHARGES, DATA_SHEETS, ORDERS, data_sheet, is_designed

# Палитра приложения
PRIMARY, PRIMARY_SOFT = "0E6B5C", "E7F2EF"
ACCENT, ACCENT_SOFT = "A56B36", "F6EDE4"
DANGER, DANGER_SOFT = "A24438", "F7E6E3"
TEXT, MUTED, BORDER, BG = "1D1D1F", "6E6E73", "D2D2D7", "F5F5F7"

RUB = '#,##0" ₽";[Color10]-#,##0" ₽";"—"'
RUB_SIGNED = '#,##0" ₽";[Red]-#,##0" ₽";"0 ₽"'
RUB_KOP = '#,##0.00" ₽";[Red]-#,##0.00" ₽";"—"'
MONEY = '#,##0.00" ₽";-#,##0.00" ₽";"—"'  # расходы: минус — норма, не красим
MONEY0 = '#,##0" ₽";-#,##0" ₽";"0 ₽"'
PCT = '0.0%;[Red]-0.0%;"—"'
INT = '#,##0;-#,##0;"—"'
DATE = "DD.MM.YYYY"

STATUS = {
    "delivered": "Доставлен", "cancelled": "Отменён", "returned": "Возврат", "delivering": "В пути",
    "awaiting_packaging": "Ожидает сборки", "ожидает расчёта": "Ожидает расчёта Ozon",
}
CAMPAIGN_STATE = {
    "archived": "В архиве", "inactive": "Остановлена", "CAMPAIGN_STATE_INACTIVE": "Остановлена",
    "CAMPAIGN_STATE_RUNNING": "Работает", "running": "Работает",
}
CAMPAIGN_TYPE = {"SKU": "Оплата за клик", "SEARCH_PROMO": "Продвижение в поиске", "ALL_SKU_PROMO": "Все товары",
                 "REF_VK": "Внешний трафик (VK)"}

thin = Side(style="thin", color=BORDER)


def fill(color):
    return PatternFill("solid", fgColor=color)


def num(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


# ── Чтение данных отчёта ───────────────────────────────────────────────────

def read_source(wb):
    ws = data_sheet(wb, ORDERS)
    orders = []
    for r in ws.iter_rows(min_row=2, max_col=13, values_only=True):
        if not r[1]:
            continue
        status = str(r[0] or "").strip()
        note = next((v for v in (r[6], r[7], r[10]) if isinstance(v, str) and v.strip() not in ("", "-")), None)
        orders.append({
            "status": STATUS.get(status, status), "raw_status": status, "number": r[1], "name": r[2],
            "article": r[3], "qty": num(r[4]), "price": num(r[5]), "commission": num(r[6]), "logistics": num(r[7]),
            "accrued": num(r[8]), "cost": num(r[9]), "profit": num(r[10]), "date": to_date(r[11]), "scheme": r[12],
            "note": note,
        })
    summary = {}
    for label, value in ws.iter_rows(min_col=16, max_col=17, values_only=True):
        if isinstance(label, str):
            summary.setdefault(label.strip(), value)
    payouts = []
    for r in ws.iter_rows(min_row=2, min_col=19, max_col=21, values_only=True):
        if isinstance(r[1], (int, float)):
            payouts.append({"date": to_date(r[0]), "amount": r[1], "what": r[2]})
    settle = {"cost": (ws["W2"].value, ws["X2"].value, ws["Y2"].value),
              "profit": (ws["W5"].value, ws["X5"].value, ws["Y5"].value),
              "closed": ws["X7"].value, "paid_total": ws["X9"].value}
    charges_ws, campaigns_ws = data_sheet(wb, CHARGES), data_sheet(wb, CAMPAIGNS)
    charges = [list(r) for r in charges_ws.iter_rows(min_row=2, values_only=True) if r[0]] if charges_ws else []
    campaigns = [list(r) for r in campaigns_ws.iter_rows(min_row=2, values_only=True) if r[0]] if campaigns_ws else []
    return orders, summary, payouts, settle, charges, campaigns


# ── Общие элементы ─────────────────────────────────────────────────────────

def widths(ws, mapping):
    for col, width in mapping.items():
        ws.column_dimensions[col].width = width


def header_row(ws, row, headers, start_col=1):
    for i, text in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=text)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill(PRIMARY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def section(ws, row, title, col=2, span=8):
    c = ws.cell(row=row, column=col, value=title)
    c.font = Font(size=13, bold=True, color=TEXT)
    for i in range(span):
        ws.cell(row=row, column=col + i).border = Border(bottom=Side(style="medium", color=PRIMARY))
    ws.row_dimensions[row].height = 24


def table(ws, ref, name, style="TableStyleLight1"):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)


# ── Лист «Сводка» ──────────────────────────────────────────────────────────

def build_summary(ws, title_month, created, orders, summary, payouts, settle, charges):
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 2, "B": 40, "C": 16, "D": 14, "E": 3, "F": 30, "G": 16, "H": 16, "I": 16, "J": 2})
    revenue = num(summary.get("Общая выручка")) or 0
    profit = num(summary.get("Чистая прибыль")) or 0
    delivered = [o for o in orders if o["raw_status"] == "delivered"]

    ws["B2"] = f"Ozon · {title_month}"
    ws["B2"].font = Font(size=22, bold=True, color=TEXT)
    ws["B3"] = (f"{len(orders)} заказов · доставлено {len(delivered)} · "
                f"отчёт сформирован {created:%d.%m.%Y}")
    ws["B3"].font = Font(size=11, color=MUTED)

    # Карточки
    # Сверки ещё нет, если отчёт сформирован, а сверить выплаты не удалось (нет связи с Ozon).
    settled = num(settle["cost"][0]) is not None
    owed = (num(settle["cost"][2]) or 0) + (num(settle["profit"][2]) or 0) if settled else None
    cards = [
        ("Выручка", revenue, RUB, f"{len(delivered)} доставленных заказов", PRIMARY_SOFT, TEXT),
        ("Чистая прибыль", profit, RUB_SIGNED, "после всех расходов и себестоимости", PRIMARY_SOFT, PRIMARY if profit >= 0 else DANGER),
        ("Маржа", profit / revenue if revenue else 0, PCT, "чистая прибыль / выручка", PRIMARY_SOFT, TEXT),
        ("Ozon ещё должен", owed if settled else "—", RUB,
         "сверка выплат ещё не проводилась" if not settled else
         "месяц закрыт выплатами" if settle["closed"] == "Да" else "по сверке выплат ниже",
         PRIMARY_SOFT if owed == 0 else ACCENT_SOFT, PRIMARY if owed == 0 else ACCENT),
    ]
    for (label, value, fmt, note, bg, color), col in zip(cards, ["B", "C", "F", "G"]):
        span = {"B": "B", "C": "D", "F": "F", "G": "I"}[col]
        for r, (text, font) in enumerate([(label, Font(size=10, color=MUTED)), (value, Font(size=20, bold=True, color=color)),
                                          (note, Font(size=9, color=MUTED))], start=5):
            ws.merge_cells(f"{col}{r}:{span}{r}")
            c = ws[f"{col}{r}"]
            c.value, c.font = text, font
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            if r == 6:
                c.number_format = fmt
            for cc in ws[f"{col}{r}:{span}{r}"][0]:
                cc.fill = fill(bg)
    ws.row_dimensions[6].height = 34

    # Из чего сложилась прибыль
    section(ws, 9, "Из чего сложилась прибыль", span=3)
    commission = sum(o["commission"] or 0 for o in orders)
    logistics = sum(o["logistics"] or 0 for o in orders)
    cost = sum(o["cost"] or 0 for o in orders)
    # Как в отчёте: заказы, ожидающие расчёта Ozon (прибыль не число), в прибыль не входят.
    accrued = sum(o["accrued"] or 0 for o in orders if o["profit"] is not None)
    other_orders = accrued - revenue - commission - logistics
    promo = -(num(summary.get("Продвижение Ozon")) or 0)
    external = -(num(summary.get("Внешний маркетинг")) or 0)
    storage = -(num(summary.get("Расход хранения FBO")) or 0)
    outside = -(num(summary.get("Расходы Ozon вне заказов")) or 0)
    lines = [
        ("Выручка (доставленные заказы)", revenue),
        ("Комиссия Ozon", commission),
        ("Логистика", logistics),
        ("Отмены, возвраты и прочее по заказам", other_orders),
        ("Себестоимость товаров", cost),
        ("Реклама в Ozon", promo),
        ("Внешний маркетинг", external),
        ("Хранение на складах Ozon", storage),
        ("Расходы Ozon вне заказов", outside),
    ]
    for c, text in zip("BCD", ["Статья", "Сумма", "Доля от выручки"]):
        ws[f"{c}10"] = text
        ws[f"{c}10"].font = Font(size=9, bold=True, color=MUTED)
        ws[f"{c}10"].alignment = Alignment(horizontal="left" if c == "B" else "right")
    row = 11
    for label, value in lines:
        ws[f"B{row}"], ws[f"C{row}"] = label, value
        ws[f"C{row}"].number_format = RUB if row == 11 else MONEY0
        ws[f"D{row}"] = f"=IF($C$11=0,0,ABS(C{row})/$C$11)" if row > 11 else 1
        ws[f"D{row}"].number_format = PCT
        ws[f"D{row}"].font = Font(color=MUTED)
        for c in "BCD":
            ws[f"{c}{row}"].border = Border(bottom=thin)
        row += 1
    ws[f"B{row}"] = "Чистая прибыль"
    ws[f"C{row}"] = f"=SUM(C11:C{row - 1})"
    ws[f"D{row}"] = f"=IF($C$11=0,0,C{row}/$C$11)"
    ws[f"C{row}"].number_format, ws[f"D{row}"].number_format = RUB_SIGNED, PCT
    for c in "BCD":
        ws[f"{c}{row}"].font = Font(bold=True, size=12, color=PRIMARY if profit >= 0 else DANGER)
        ws[f"{c}{row}"].fill = fill(PRIMARY_SOFT)
    check_row = row
    ws.conditional_formatting.add(f"D12:D{row - 1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                   end_value=1, color="E8B9B2", showValue=True))

    # Заказы
    section(ws, 9, "Заказы", col=6, span=4)
    counts = defaultdict(int)
    for o in orders:
        counts[o["status"]] += 1
    order_lines = [("Всего", len(orders))] + [(s, counts[s]) for s in
                   ("Доставлен", "Отменён", "Возврат", "В пути", "Ожидает расчёта Ozon") if counts[s]]
    r = 10
    for label, value in order_lines:
        ws[f"F{r}"], ws[f"G{r}"] = label, value
        ws[f"G{r}"].number_format = INT
        ws[f"F{r}"].border = ws[f"G{r}"].border = Border(bottom=thin)
        r += 1
    avg = num(summary.get("Средний чек")) or 0
    ws[f"F{r}"], ws[f"G{r}"] = "Средний чек", avg
    ws[f"G{r}"].number_format = RUB
    ws[f"F{r}"].border = ws[f"G{r}"].border = Border(bottom=thin)
    if counts["Ожидает расчёта Ozon"]:
        ws[f"F{r + 1}"] = "Заказы, ожидающие расчёта Ozon, в прибыль пока не входят."
        ws[f"F{r + 1}"].font = Font(size=9, italic=True, color=MUTED)

    # Выплаты
    row = max(check_row, r + 2) + 3
    section(ws, row, "Расчёты с Ozon за этот месяц", span=8)
    closed = settle["closed"] == "Да"
    ws.merge_cells(f"B{row + 1}:I{row + 1}")
    status = ws[f"B{row + 1}"]
    status.value = ("Сверка выплат ещё не проводилась — она обновится при следующем формировании любого отчёта"
                    if not settled else "✔  Ozon полностью рассчитался за месяц" if closed
                    else f"⏳  Ozon ещё не рассчитался: осталось получить {owed:,.0f} ₽".replace(",", " "))
    status.font = Font(bold=True, size=12, color=MUTED if not settled else PRIMARY if closed else ACCENT)
    for c in ws[f"B{row + 1}:I{row + 1}"][0]:
        c.fill = fill(PRIMARY_SOFT if closed else ACCENT_SOFT)
    status.alignment = Alignment(indent=1, vertical="center")
    ws.row_dimensions[row + 1].height = 28
    for c, text in zip("BCD", ["", "Начислено", "Закрыто"]):
        ws[f"{c}{row + 3}"] = text
    ws[f"E{row + 3}"] = None
    ws.merge_cells(f"D{row + 3}:E{row + 3}")
    ws[f"F{row + 3}"] = "Осталось получить"
    for c in "BCDF":
        ws[f"{c}{row + 3}"].font = Font(size=9, bold=True, color=MUTED)
    for i, (label, key) in enumerate([("Себестоимость", "cost"), ("Прибыль", "profit")], start=row + 4):
        due, done, left = settle[key]
        ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"], ws[f"F{i}"] = label, due, done, left
        ws.merge_cells(f"D{i}:E{i}")
        for c in "CDF":
            ws[f"{c}{i}"].number_format = RUB
            ws[f"{c}{i}"].alignment = Alignment(horizontal="left")
        ws[f"F{i}"].font = Font(bold=True, color=PRIMARY if not left else ACCENT)
        for c in "BCDEF":
            ws[f"{c}{i}"].border = Border(bottom=thin)
    r = row + 7
    ws[f"B{r}"] = "Выплаты, которые закрыли этот месяц"
    ws[f"B{r}"].font = Font(bold=True, color=TEXT)
    for c, text in zip("BCD", ["Дата выплаты", "Сумма", "На что ушла"]):
        ws[f"{c}{r + 1}"].alignment = Alignment(horizontal="right" if c == "C" else "left")
        ws[f"{c}{r + 1}"] = text
        ws[f"{c}{r + 1}"].font = Font(size=9, bold=True, color=MUTED)
    for j, p in enumerate(payouts, start=r + 2):
        what = re.sub(r"(\d)\.(\d{2}) ₽", r"\1,\2 ₽", str(p["what"] or ""))
        ws[f"B{j}"], ws[f"C{j}"], ws[f"D{j}"] = p["date"], p["amount"], what
        ws[f"B{j}"].number_format, ws[f"C{j}"].number_format = DATE, RUB_KOP
        ws[f"B{j}"].alignment = Alignment(horizontal="left")
        for c in "BCDEFGHI":
            ws[f"{c}{j}"].border = Border(bottom=thin)
    if not payouts:
        ws[f"B{r + 2}"] = "Выплат за этот месяц пока не было"
        ws[f"B{r + 2}"].font = Font(italic=True, color=MUTED)
    last = r + 2 + max(len(payouts), 1)
    ws[f"B{last}"] = f"Всего Ozon выплатил с начала магазина: {num(settle['paid_total']) or 0:,.0f} ₽".replace(",", " ")
    ws[f"B{last}"].font = Font(size=9, color=MUTED)

    # Топ товаров
    row = last + 3
    section(ws, row, "Товары: что принесло прибыль", span=8)
    by_article = defaultdict(lambda: {"name": "", "qty": 0, "revenue": 0.0, "profit": 0.0})
    for o in delivered:
        a = by_article[str(o["article"])]
        a["name"] = a["name"] or o["name"]
        a["qty"] += o["qty"] or 0
        a["revenue"] += o["price"] or 0
        a["profit"] += o["profit"] or 0
    top = sorted(by_article.items(), key=lambda kv: kv[1]["profit"], reverse=True)
    heads = ["Товар", "Артикул", "Продано, шт", "", "Выручка", "Прибыль", "Маржа"]
    for c, text in zip("BCDEFGH", heads):
        ws[f"{c}{row + 1}"] = text
        ws[f"{c}{row + 1}"].font = Font(size=9, bold=True, color=MUTED)
        ws[f"{c}{row + 1}"].alignment = Alignment(horizontal="left" if c in "BC" else "right")
    for j, (article, a) in enumerate(top[:10], start=row + 2):
        name = a["name"] or ""
        ws[f"B{j}"], ws[f"C{j}"], ws[f"D{j}"] = (name[:44] + "…") if len(name) > 45 else name, article, a["qty"]
        ws[f"F{j}"], ws[f"G{j}"] = a["revenue"], a["profit"]
        ws[f"H{j}"] = f"=IF(F{j}=0,0,G{j}/F{j})"
        ws[f"C{j}"].alignment = Alignment(horizontal="left")
        ws[f"D{j}"].number_format, ws[f"F{j}"].number_format = INT, RUB
        ws[f"G{j}"].number_format, ws[f"H{j}"].number_format = RUB_SIGNED, PCT
        for c in "BCDEFGH":
            ws[f"{c}{j}"].border = Border(bottom=thin)
    end = row + 1 + min(len(top), 10)
    if top:
        ws.conditional_formatting.add(f"G{row + 2}:G{end}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                                         color="7FB8AD", showValue=True))
    ws[f"B{end + 1}"] = (f"Топ-10 из {len(top)} товаров по прибыли с доставленных заказов. Все заказы — на листе «Заказы»."
                         if top else "Доставленных заказов в этом месяце нет.")
    ws[f"B{end + 1}"].font = Font(size=9, color=MUTED)

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ── Лист «Заказы» ──────────────────────────────────────────────────────────

def build_orders(ws, orders):
    heads = ["Дата отгрузки", "Номер заказа", "Товар", "Артикул", "Схема", "Статус", "Кол-во", "Цена продажи",
             "Комиссия Ozon", "Логистика", "Начислено Ozon", "Себестоимость", "Прибыль", "Комментарий"]
    header_row(ws, 1, heads)
    for i, o in enumerate(sorted(orders, key=lambda o: (o["date"] or date.min, o["number"])), start=2):
        values = [o["date"], o["number"], o["name"], o["article"], o["scheme"], o["status"], o["qty"], o["price"],
                  o["commission"], o["logistics"], o["accrued"], o["cost"], o["profit"], o["note"]]
        for col, v in enumerate(values, start=1):
            ws.cell(row=i, column=col, value=v)
        ws.cell(row=i, column=1).number_format = DATE
        ws.cell(row=i, column=4).number_format = "0"
        ws.cell(row=i, column=7).number_format = INT
        for col in range(8, 14):
            ws.cell(row=i, column=col).number_format = RUB_KOP if col == 13 else MONEY
    last = len(orders) + 1
    if not orders:
        # Таблица Excel без строк данных — повреждённый файл; пустой месяц — только шапка и пояснение.
        ws.cell(row=2, column=1, value="Заказов в этом месяце нет").font = Font(italic=True, color=MUTED)
        widths(ws, {"A": 14, "B": 20, "C": 46})
        return
    table(ws, f"A1:N{last}", "Заказы", "TableStyleLight8")
    total = last + 1
    ws.cell(row=total, column=1, value="Итого (по видимым строкам)").font = Font(bold=True)
    for col in range(7, 14):
        letter = get_column_letter(col)
        c = ws.cell(row=total, column=col, value=f"=SUBTOTAL(109,{letter}2:{letter}{last})")
        c.number_format = INT if col == 7 else (RUB_KOP if col == 13 else MONEY)
        c.font = Font(bold=True)
    for col in range(1, 15):
        ws.cell(row=total, column=col).fill = fill(PRIMARY_SOFT)
        ws.cell(row=total, column=col).border = Border(top=Side(style="medium", color=PRIMARY))
    widths(ws, {"A": 14, "B": 20, "C": 46, "D": 11, "E": 9, "F": 21, "G": 10, "H": 15, "I": 16, "J": 14, "K": 16,
                "L": 16, "M": 14, "N": 42})
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "D2"
    rng = f"F2:F{last}"
    for text, color, bg in [("Отменён", MUTED, BG), ("Возврат", DANGER, DANGER_SOFT),
                            ("Ожидает расчёта Ozon", ACCENT, ACCENT_SOFT), ("В пути", ACCENT, ACCENT_SOFT)]:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                      font=Font(color=color), fill=fill(bg)))
    ws.conditional_formatting.add(f"A2:N{last}", FormulaRule(formula=['$F2="Отменён"'], font=Font(color="9A9AA0")))
    ws.conditional_formatting.add(f"M2:M{last}", CellIsRule(operator="lessThan", formula=["0"],
                                                            font=Font(color=DANGER, bold=True), fill=fill(DANGER_SOFT)))


# ── Лист «Расходы вне заказов» ─────────────────────────────────────────────

def build_charges(ws, charges):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Расходы Ozon вне заказов"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = "Списания, которые не относятся к конкретному заказу: кросс-докинг, приёмка поставок, ранний вывод денег и т. п."
    ws["A2"].font = Font(size=10, color=MUTED)
    groups = defaultdict(lambda: [0.0, 0, ""])
    for r in charges:
        g = groups[r[1]]
        g[0] += num(r[4]) or 0
        g[1] += 1
        g[2] = r[5] or ""
    # Обе таблицы в одних колонках: A — вид расхода, B — сумма, дальше подробности.
    header_row(ws, 4, ["Вид расхода", "Сумма", "Операций", "", "", "Как учтено в прибыли"])
    i = 5
    for name, (total, count, how) in sorted(groups.items(), key=lambda kv: kv[1][0]):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=total).number_format = MONEY
        ws.cell(row=i, column=3, value=count).number_format = INT
        ws.cell(row=i, column=6, value=how).font = Font(color=MUTED)
        for col in range(1, 7):
            ws.cell(row=i, column=col).border = Border(bottom=thin)
        i += 1
    ws.cell(row=i, column=1, value="Итого").font = Font(bold=True)
    c = ws.cell(row=i, column=2, value=f"=SUM(B5:B{i - 1})" if groups else 0)
    c.number_format, c.font = MONEY, Font(bold=True)
    for col in range(1, 7):
        ws.cell(row=i, column=col).fill = fill(PRIMARY_SOFT)
    start = i + 3
    ws.cell(row=start - 1, column=1, value="Все операции").font = Font(size=13, bold=True)
    header_row(ws, start, ["Вид расхода", "Сумма", "Дата", "Номер поставки / кампании", "Код Ozon", "Как учтено в прибыли"])
    for j, r in enumerate(sorted(charges, key=lambda r: str(r[0])), start=start + 1):
        values = [r[1], num(r[4]), to_date(r[0]), None if r[3] == "—" else r[3], r[2], r[5]]
        for col, v in enumerate(values, start=1):
            ws.cell(row=j, column=col, value=v)
        ws.cell(row=j, column=2).number_format = MONEY
        ws.cell(row=j, column=3).number_format = DATE
        ws.cell(row=j, column=3).alignment = Alignment(horizontal="left")
        ws.cell(row=j, column=6).font = Font(color=MUTED)
        for col in range(1, 7):
            ws.cell(row=j, column=col).border = Border(bottom=thin)
    widths(ws, {"A": 54, "B": 14, "C": 12, "D": 28, "E": 10, "F": 50})


# ── Лист «Реклама» ─────────────────────────────────────────────────────────

def build_campaigns(ws, campaigns):
    heads = ["Кампания", "Состояние", "Тип", "Расход", "Показы", "Клики", "CTR", "Цена клика", "Заказы, шт",
             "Заказы, ₽", "ДРР"]
    ws["A1"] = "Реклама в Ozon за месяц"
    ws["A1"].font = Font(size=16, bold=True)
    active = [c for c in campaigns if (num(c[8]) or 0) > 0]
    idle = [c for c in campaigns if (num(c[8]) or 0) <= 0]
    spend = sum(num(c[8]) or 0 for c in active)
    sales = sum(num(c[14]) or 0 for c in active)
    ws["A2"] = (f"С расходом — {len(active)} из {len(campaigns)} кампаний · расход {spend:,.0f} ₽ · "
                f"заказы с рекламы {sales:,.0f} ₽").replace(",", " ")
    ws["A2"].font = Font(size=10, color=MUTED)
    header_row(ws, 4, heads)
    row = 5
    for c in sorted(active, key=lambda c: -(num(c[8]) or 0)) + sorted(idle, key=lambda c: str(c[1])):
        values = [c[1], CAMPAIGN_STATE.get(c[2], c[2]), CAMPAIGN_TYPE.get(c[4], c[4]), num(c[8]), num(c[9]),
                  num(c[10]), (num(c[11]) or 0) / 100, num(c[12]), num(c[13]), num(c[14]), (num(c[15]) or 0) / 100]
        for col, v in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=v)
        for col, fmt in zip(range(4, 12), [RUB_KOP, INT, INT, PCT, RUB_KOP, INT, RUB, PCT]):
            ws.cell(row=row, column=col).number_format = fmt
        if c in idle:
            ws.row_dimensions[row].outlineLevel = 1
            ws.row_dimensions[row].hidden = True
            for col in range(1, 12):
                ws.cell(row=row, column=col).font = Font(color="9A9AA0")
        for col in range(1, 12):
            ws.cell(row=row, column=col).border = Border(bottom=thin)
        row += 1
    if idle:
        note = ws.cell(row=row, column=1, value=f"+ {len(idle)} кампаний без расхода — скрыты (нажмите «+» слева, чтобы показать)")
        note.font = Font(size=9, italic=True, color=MUTED)
    ws.sheet_properties.outlinePr.summaryBelow = True
    widths(ws, {"A": 46, "B": 13, "C": 20, "D": 13, "E": 10, "F": 9, "G": 8, "H": 12, "I": 11, "J": 13, "K": 8})
    ws.freeze_panes = "B5"


VISUAL_SHEETS = ["Сводка", ORDERS, "Реклама", CHARGES]


def render(wb, title):
    """Перерисовать видимые листы по скрытым листам данных (книга уже нового вида)."""
    orders, summary, payouts, settle, charges, campaigns = read_source(wb)
    for name in VISUAL_SHEETS:
        if name in wb.sheetnames and name not in DATA_SHEETS.values():
            wb.remove(wb[name])
    created = wb.properties.created or datetime.now()
    builders = [
        lambda ws: build_summary(ws, title, created, orders, summary, payouts, settle, charges),
        lambda ws: build_orders(ws, orders),
        lambda ws: build_campaigns(ws, campaigns),
        lambda ws: build_charges(ws, charges),
    ]
    for index, (name, build) in enumerate(zip(VISUAL_SHEETS, builders)):
        build(wb.create_sheet(name, index))
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Сводка"
    wb["Сводка"].sheet_properties.tabColor = PRIMARY
    wb.active = 0


def apply_design(wb, title):
    """Отчёт старого вида: спрятать листы данных под именами «Данные…» и построить оформление.
    Отчёт нового вида: только перерисовать оформление."""
    if not is_designed(wb.sheetnames):
        for sheet, hidden_name in DATA_SHEETS.items():
            if sheet in wb.sheetnames:
                wb[sheet].title = hidden_name
        for hidden_name in DATA_SHEETS.values():
            if hidden_name in wb.sheetnames:
                wb[hidden_name].sheet_state = "hidden"
    render(wb, title)


def apply_design_file(path, title=None):
    try:
        from scripts.file_io import save_workbook_atomic
    except ModuleNotFoundError:
        from file_io import save_workbook_atomic
    path = Path(path)
    wb = load_workbook(path)
    try:
        apply_design(wb, title or path.stem)
        save_workbook_atomic(wb, path)
    finally:
        wb.close()


def main() -> None:
    """Перевести в новый вид все месячные отчёты Ozon старого вида (без запросов к Ozon)."""
    try:
        from scripts.app_paths import DATA_ROOT
        from scripts.payout_settlement import find_reports
    except ModuleNotFoundError:
        from app_paths import DATA_ROOT
        from payout_settlement import find_reports
    converted = 0
    for _month, path in find_reports(DATA_ROOT):
        with_names = load_workbook(path, read_only=True)
        try:
            designed = is_designed(with_names.sheetnames)
        finally:
            with_names.close()
        if designed:
            continue
        apply_design_file(path)
        converted += 1
        print(f"✅ {path.name}")
    print(f"Переведено в новый вид: {converted}")


if __name__ == "__main__":
    main()
