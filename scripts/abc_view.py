"""
Чтение готового ABC/XYZ-отчёта (Excel из ABC_XYZ_analytics_report.py) для экрана в приложении.

Собирает по каждому артикулу класс ABC, класс XYZ, общий код, прибыль, долю прибыли,
продажи по месяцам и коэффициент вариации, плюс расшифровку кодов из листа «Итог».
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

NOT_ENOUGH = "Недостаточно данных"
XYZ_STAT_COLUMNS = {"Общий итог", "Медиана", "Среднее значение", "Стандартное отклонение", "Коэффициент вариации", "Оценка"}
MONTHS_RU = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]


def _article(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _is_article(value: Any) -> bool:
    """Артикул или комплект артикулов («1114, 1102»), но не строка «Общая прибыль» и не заголовок."""
    return isinstance(value, (int, float)) or bool(re.fullmatch(r"\d+(\s*,\s*\d+)*", str(value).strip()))


ABC_TEXT = {"A": "Высокий оборот", "B": "Средний оборот", "C": "Низкий оборот"}
XYZ_TEXT = {"X": "стабильный спрос", "Y1": "слабая сезонность", "Y2": "сезонный спрос", "Y3": "сильная сезонность",
            "Y": "умеренные колебания спроса", "Z": "нерегулярный спрос"}


def _number(value: Any) -> Optional[float]:
    try:
        return None if value is None or value == "" else float(value)
    except (TypeError, ValueError):
        return None


def _month_key(name: str) -> tuple:
    match = re.match(r"(\S+)\s+(\d{4})", name.strip())
    if not match or match.group(1) not in MONTHS_RU:
        return (9999, 99)
    return (int(match.group(2)), MONTHS_RU.index(match.group(1)) + 1)


def report_period_title(path: Path) -> str:
    """«Май 2026-Август 2026.xlsx» → «Май 2026 — Август 2026»."""
    return re.sub(r"(\d{4})-(?=\D)", r"\1 — ", path.stem)


def _read_summary(ws) -> tuple:
    items: Dict[str, Dict[str, Any]] = {}
    legend: Dict[str, Dict[str, str]] = {}
    mode = None
    for row in ws.iter_rows(values_only=True):
        first = row[0] if row else None
        if first == "Артикул":
            mode = "items"
            continue
        if first == "Код":
            mode = "legend"
            continue
        if mode == "items":
            article = _article(first)
            if article is None:
                mode = None
                continue
            xyz = str(row[2] or "").strip()
            code = str(row[3] or "").strip()
            items[article] = {
                "article": article,
                "abc": str(row[1] or "").strip(),
                "xyz": xyz,
                "code": code if code and code != NOT_ENOUGH else None,
                "enough_data": xyz != NOT_ENOUGH,
            }
        elif mode == "legend" and first:
            title = str(row[1] or "").strip()
            recs = [line.strip(" •\t") for line in str(row[2] or "").splitlines() if line.strip(" •\t")]
            legend[str(first).strip()] = {"title": title.split("\n")[0], "detail": " ".join(title.split("\n")[1:]), "recommendations": recs}
    return items, legend


def _read_abc(ws, items: Dict[str, Dict[str, Any]]) -> None:
    header = None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Артикул":
            header = [str(v or "").strip() for v in row]
            continue
        if header is None:
            continue
        article = _article(row[0])
        if article is None:
            header = None
            continue
        if not _is_article(row[0]):
            continue  # строка «Общая прибыль» и легенда
        values = dict(zip(header, row))
        item = items.setdefault(article, {"article": article, "abc": str(values.get("Оценка") or ""), "xyz": "", "code": None, "enough_data": False})
        item["profit"] = _number(values.get("Сумма по полю Прибыль"))
        item["share"] = _number(values.get("Доля в общей прибыли"))
        item["cum_share"] = _number(values.get("Накопительная доля"))


def _read_xyz(ws, items: Dict[str, Dict[str, Any]], months: set) -> None:
    header = None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Артикул":
            header = [str(v or "").strip() for v in row]
            continue
        if header is None:
            continue
        article = _article(row[0])
        if article is None or not _is_article(row[0]):
            header = None
            continue
        values = dict(zip(header, row))
        item = items.get(article)
        if item is None:
            continue
        sales = {}
        for name in header[1:]:
            if not name or name in XYZ_STAT_COLUMNS or name.startswith("Винз."):
                continue
            months.add(name)
            # Если продажи «винзоризованы» (сглажен выброс), показываем исходное число — так понятнее.
            sales[name] = _number(values.get(name)) or 0.0
        item["sales"] = sales
        item["total_qty"] = _number(values.get("Общий итог"))
        item["cv"] = _number(values.get("Коэффициент вариации"))


def load_abc_report(path: Path) -> Dict[str, Any]:
    """Читает ABC/XYZ-отчёт. Возвращает period, built_at, months (по возрастанию), items, legend, total_profit."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if "Итог" not in wb.sheetnames:
            raise ValueError("в файле нет листа «Итог» — это не ABC/XYZ-отчёт")
        items, legend = _read_summary(wb["Итог"])
        months: set = set()
        if "ABC" in wb.sheetnames:
            _read_abc(wb["ABC"], items)
        if "XYZ" in wb.sheetnames:
            _read_xyz(wb["XYZ"], items, months)
    finally:
        wb.close()
    ordered_months = sorted(months, key=_month_key)
    rows = []
    for item in items.values():
        item.setdefault("profit", None)
        item.setdefault("share", None)
        item.setdefault("cum_share", None)
        item.setdefault("total_qty", None)
        item.setdefault("cv", None)
        item["sales"] = [item.get("sales", {}).get(month, 0.0) for month in ordered_months]
        info = legend.get(item["code"] or "")
        if info:
            item["title"] = info["title"]
        elif not item["enough_data"]:
            item["title"] = "Мало продаж для оценки спроса"
        else:
            item["title"] = f"{ABC_TEXT.get(item['abc'], '')}, {XYZ_TEXT.get(item['xyz'], '')}".strip(", ")
        item["recommendations"] = info["recommendations"] if info else []
        rows.append(item)
    rows.sort(key=lambda item: -(item["profit"] or 0.0))
    total_profit = sum(item["profit"] or 0.0 for item in rows)
    return {
        "path": str(path),
        "period": report_period_title(path),
        "built_at": datetime.fromtimestamp(path.stat().st_mtime),
        "months": ordered_months,
        "items": rows,
        "legend": legend,
        "total_profit": total_profit,
    }


def xyz_group(xyz: str) -> str:
    """Столбец матрицы: X, Y1, Y2, Y3, Y (без деления), Z или «—» для мало данных."""
    value = (xyz or "").strip()
    if value in ("X", "Y1", "Y2", "Y3", "Y", "Z"):
        return value
    return "—"


def matrix_counts(items: List[Dict[str, Any]]) -> Dict[tuple, List[str]]:
    """{(ABC, XYZ-столбец): [артикулы]} для матрицы."""
    cells: Dict[tuple, List[str]] = {}
    for item in items:
        cells.setdefault((item["abc"], xyz_group(item["xyz"])), []).append(item["article"])
    return cells
