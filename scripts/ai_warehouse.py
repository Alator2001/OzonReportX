# -*- coding: utf-8 -*-
"""
Аналитическая база AI-ассистента: все данные программы в одном SQLite-файле.

Источники (месячные отчёты Ozon и WB, costs.xlsx, ABC/XYZ, планы поставки, анализ органики,
снимки товаров, журнал автоакций, база поиска) раскладываются по таблицам с понятными колонками.
Каждый источник пересобирается, только когда изменились его файлы. База — производная
(`.cache/ai/warehouse.db`): её можно удалить, она соберётся заново.

`CATALOG` — словарь данных: что лежит в таблице, откуда, что значит каждая колонка.
По нему модель понимает, что спрашивать (`describe`), а инструменты строят запросы.

    python scripts/ai_warehouse.py            # собрать/обновить и показать каталог
"""

import json
import re
import sqlite3
import sys
from contextlib import closing
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

try:
    from scripts.file_lock import exclusive_file
except ModuleNotFoundError:
    from file_lock import exclusive_file
try:
    from scripts.report_book import CAMPAIGNS, CHARGES, ORDERS, data_sheet_name
except ModuleNotFoundError:
    from report_book import CAMPAIGNS, CHARGES, ORDERS, data_sheet_name
try:
    from scripts.report_paths import ozon_report_files, wb_report_files
except ModuleNotFoundError:
    from report_paths import ozon_report_files, wb_report_files
try:
    from scripts.app_paths import DATA_ROOT
except ModuleNotFoundError:
    from app_paths import DATA_ROOT

ROOT = DATA_ROOT
WAREHOUSE_PATH = Path(".cache") / "ai" / "warehouse.db"
# Меняйте при изменении CATALOG: старая база удаляется и собирается заново.
SCHEMA_VERSION = 3

MONTHS_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]

PERIOD = "Месяц отчёта, YYYY-MM"
ARTIKUL = "Артикул продавца (offer_id), текст"
SIGN_NOTE = "расходы со знаком минус"

# Таблица → описание, колонки (имя → (тип SQLite, описание)), колонка даты для диапазона в каталоге.
CATALOG: Dict[str, Dict[str, Any]] = {
    "ozon_orders": {
        "description": "Заказы Ozon построчно из месячных отчётов (лист «Заказы»). Месяц — по дате отгрузки. "
                       "Прибыль строки = начисление Ozon + себестоимость; " + SIGN_NOTE + ".",
        "source": "reports/Ozon/<Год>/<Месяц Год>.xlsx",
        "date_column": "period",
        "columns": {
            "period": ("TEXT", PERIOD),
            "status": ("TEXT", "Статус заказа: delivered, cancelled, delivering, returned и т. п."),
            "order_number": ("TEXT", "Номер отправления"),
            "product_name": ("TEXT", "Название товара"),
            "artikul": ("TEXT", ARTIKUL + "; в отправлении из нескольких товаров — артикулы через запятую "
                                "(«1101, 1105»), суммы строки — по всему отправлению"),
            "quantity": ("REAL", "Количество, шт."),
            "sale_price": ("REAL", "Цена продажи, ₽"),
            "commission": ("REAL", "Комиссия Ozon за продажу, ₽ (минус)"),
            "logistics": ("REAL", "Логистика, включая операционные ошибки продавца, ₽ (минус)"),
            "accrual": ("REAL", "Сумма начисления Ozon по заказу, ₽"),
            "cost": ("REAL", "Себестоимость, ₽ (минус)"),
            "profit": ("REAL", "Прибыль по заказу, ₽"),
            "ship_date": ("TEXT", "Дата отгрузки, YYYY-MM-DD"),
            "scheme": ("TEXT", "Схема: FBS или FBO"),
        },
    },
    "ozon_month_metrics": {
        "description": "Итоговые показатели месяца Ozon (блок P/Q листа «Заказы»): выручка, чистая прибыль, "
                       "рентабельность, расходы на продвижение, комиссии, логистика, хранение, баланс, выплаты. "
                       "Одна строка — один показатель; это официальные итоги месяца, их не нужно пересчитывать.",
        "source": "reports/Ozon/<Год>/<Месяц Год>.xlsx",
        "date_column": "period",
        "columns": {
            "period": ("TEXT", PERIOD),
            "metric": ("TEXT", "Название показателя, как в отчёте (например «Общая выручка», «Чистая прибыль»)"),
            "value": ("REAL", "Значение; NULL, если показатель текстовый"),
            "text": ("TEXT", "Текстовое значение (справки), если показатель не число"),
        },
    },
    "ozon_other_expenses": {
        "description": "Расходы Ozon вне заказов (лист «Расходы вне заказов»): хранение, кросс-докинг, приёмка, "
                       "ранний вывод средств и т. п.; " + SIGN_NOTE + ". Есть не во всех месяцах.",
        "source": "reports/Ozon/<Год>/<Месяц Год>.xlsx",
        "date_column": "date",
        "columns": {
            "period": ("TEXT", PERIOD),
            "date": ("TEXT", "Дата операции, YYYY-MM-DD"),
            "kind": ("TEXT", "Вид расхода"),
            "ozon_code": ("TEXT", "Код операции Ozon"),
            "reference": ("TEXT", "Номер поставки или кампании"),
            "amount": ("REAL", "Сумма, ₽ (минус)"),
            "net_profit_note": ("TEXT", "Как расход учтён в чистой прибыли"),
        },
    },
    "ozon_ad_campaigns": {
        "description": "Рекламные кампании Ozon за месяц (лист «Кампании»): расход, показы, клики, заказы, ДРР.",
        "source": "reports/Ozon/<Год>/<Месяц Год>.xlsx",
        "date_column": "period",
        "columns": {
            "period": ("TEXT", PERIOD),
            "campaign_id": ("TEXT", "ID кампании"),
            "name": ("TEXT", "Название кампании"),
            "state": ("TEXT", "Состояние кампании"),
            "payment_type": ("TEXT", "Тип оплаты"),
            "object_type": ("TEXT", "Тип объекта (SKU, SEARCH_PROMO и т. п.)"),
            "budget": ("REAL", "Бюджет, ₽"),
            "daily_budget": ("REAL", "Дневной бюджет, ₽"),
            "weekly_budget": ("REAL", "Недельный бюджет, ₽"),
            "spend": ("REAL", "Расход за месяц, ₽"),
            "impressions": ("REAL", "Показы"),
            "clicks": ("REAL", "Клики"),
            "ctr_pct": ("REAL", "CTR, %"),
            "avg_cpc": ("REAL", "Средняя цена клика, ₽"),
            "orders": ("REAL", "Заказы с рекламы, шт."),
            "orders_revenue": ("REAL", "Заказы с рекламы, ₽"),
            "drr_pct": ("REAL", "ДРР (доля рекламных расходов от заказов), %"),
        },
    },
    "products": {
        "description": "Справочник товаров Ozon из costs.xlsx (лист «Основной»): себестоимость, расчётные "
                       "и текущие цены, ожидаемая рентабельность. Текущее состояние на момент синхронизации цен.",
        "source": "costs.xlsx",
        "columns": {
            "artikul": ("TEXT", ARTIKUL),
            "cost": ("REAL", "Себестоимость, ₽"),
            "min_price": ("REAL", "Минимальная цена продажи (по минимальной марже), ₽"),
            "desired_price": ("REAL", "Желательная цена продажи (по желаемой марже), ₽"),
            "current_price": ("REAL", "Текущая цена на Ozon, ₽"),
            "promo_price": ("REAL", "Цена с учётом акций и скидок, ₽"),
            "expected_margin_pct": ("REAL", "Текущая ожидаемая рентабельность, %"),
            "expected_net_profit": ("REAL", "Ожидаемая чистая прибыль с единицы, ₽"),
            "logistics_per_unit": ("REAL", "Логистика за штуку в прошлом месяце, ₽"),
        },
    },
    "product_promotions": {
        "description": "Участие товаров в акциях Ozon (costs.xlsx, лист «Акции»): цена в акции и бустинг.",
        "source": "costs.xlsx",
        "columns": {
            "artikul": ("TEXT", ARTIKUL),
            "promotion": ("TEXT", "Название акции"),
            "promo_price": ("REAL", "Цена товара в акции, ₽"),
            "boost_pct": ("REAL", "Бустинг в акции, %"),
            "min_price": ("REAL", "Минимальная цена продажи товара, ₽"),
        },
    },
    "abc_xyz": {
        "description": "ABC/XYZ-классы товаров из готовых отчётов. Отчётов несколько, за разные периоды: "
                       "выбирайте нужный report_period (обычно самый свежий и длинный). "
                       "ABC — вклад в прибыль (A — главные), XYZ — стабильность спроса (X — стабильный).",
        "source": "ABC&XYZ reports/<период>.xlsx",
        "columns": {
            "report_period": ("TEXT", "Период отчёта, как в имени файла («Сентябрь 2025-Август 2026»)"),
            "period_start": ("TEXT", "Первый месяц периода, YYYY-MM"),
            "period_end": ("TEXT", "Последний месяц периода, YYYY-MM"),
            "artikul": ("TEXT", ARTIKUL),
            "abc": ("TEXT", "Класс ABC"),
            "xyz": ("TEXT", "Класс XYZ (с подклассом, например Y2)"),
            "abcxyz": ("TEXT", "Общая оценка, например AX"),
        },
    },
    "fbo_supply_plan": {
        "description": "Планы поставки на склады FBO: сколько штук каждого артикула везти в каждую зону. "
                       "Один план на дату расчёта (plan_date).",
        "source": "stocks reports/Поставка FBO <дата>.json",
        "date_column": "plan_date",
        "columns": {
            "plan_date": ("TEXT", "Дата расчёта плана, YYYY-MM-DD"),
            "arrival": ("TEXT", "Ожидаемая дата прихода поставки"),
            "artikul": ("TEXT", ARTIKUL),
            "zone": ("TEXT", "Зона складов FBO"),
            "quantity": ("REAL", "К поставке в зону, шт."),
            "total": ("REAL", "К поставке по артикулу во все зоны, шт."),
            "per_day": ("REAL", "Прогноз продаж в день, шт."),
            "own_stock": ("REAL", "Остаток на своём складе (FBS), шт."),
            "unit_cost": ("REAL", "Себестоимость за шт., ₽"),
            "cost": ("REAL", "Себестоимость поставки по артикулу, ₽"),
        },
    },
    "fbs_needs": {
        "description": "Потребность в остатке на своём складе (FBS) из плана поставки: спрос, нехватка.",
        "source": "stocks reports/Поставка FBO <дата>.json",
        "date_column": "plan_date",
        "columns": {
            "plan_date": ("TEXT", "Дата расчёта плана, YYYY-MM-DD"),
            "artikul": ("TEXT", ARTIKUL),
            "demand": ("REAL", "Спрос на FBS, шт."),
            "need": ("REAL", "Нужно на FBS, шт."),
            "stock": ("REAL", "Остаток FBS, шт."),
            "shortage": ("REAL", "Не хватает на FBS, шт."),
            "cost": ("REAL", "Себестоимость нехватки, ₽"),
        },
    },
    "organic_products": {
        "description": "Последний анализ органических продаж (раздел «Анализ») по товарам: заказы органика/реклама "
                       "за текущее и прошлое окно, остаток, цены, индекс цен, акции, контент-рейтинг, "
                       "главная проблема и советы. Окно — period_start…period_end.",
        "source": "analytics history/organic_last.json",
        "date_column": "created",
        "columns": {
            "created": ("TEXT", "Когда выполнен анализ"),
            "period_start": ("TEXT", "Начало текущего окна, YYYY-MM-DD"),
            "period_end": ("TEXT", "Конец текущего окна, YYYY-MM-DD"),
            "artikul": ("TEXT", ARTIKUL),
            "sku": ("INTEGER", "SKU Ozon"),
            "orders": ("REAL", "Заказы за окно, шт."),
            "orders_prev": ("REAL", "Заказы за прошлое окно, шт."),
            "ad_orders": ("REAL", "Из них с рекламы, шт."),
            "organic": ("REAL", "Органические заказы за окно, шт."),
            "organic_prev": ("REAL", "Органические заказы за прошлое окно, шт."),
            "organic_change": ("REAL", "Изменение органики, доля (0.5 = +50%)"),
            "lost_orders": ("REAL", "Оценка потерянных заказов, шт."),
            "stock": ("REAL", "Свободный остаток, шт."),
            "days_since_sale": ("REAL", "Дней с последней продажи"),
            "avg_price": ("REAL", "Средняя цена продажи за окно, ₽"),
            "avg_price_prev": ("REAL", "Средняя цена продажи за прошлое окно, ₽"),
            "price_now": ("REAL", "Текущая цена, ₽"),
            "min_price": ("REAL", "Минимальная цена, ₽"),
            "price_index_color": ("TEXT", "Индекс цен Ozon: GREEN/YELLOW/RED/WITHOUT_INDEX"),
            "price_index_value": ("REAL", "Значение индекса цен"),
            "promotions_count": ("INTEGER", "В скольких акциях участвует"),
            "promotions": ("TEXT", "Названия акций через «; »"),
            "content_rating": ("REAL", "Контент-рейтинг карточки, 0–100"),
            "content_missing": ("TEXT", "Что добавить в карточку, через «; »"),
            "local_share": ("REAL", "Доля заказов, доставленных из кластера покупателя, %"),
            "recommended_supply": ("REAL", "Сколько Ozon советует поставить на FBO, шт."),
            "searched_users": ("REAL", "Сколько покупателей искали по запросам с товаром"),
            "main_problem": ("TEXT", "Главная проблема товара (название шага диагностики)"),
            "main_status": ("TEXT", "Статус главной проблемы: bad/warn/ok/info"),
            "diagnosis": ("TEXT", "Диагностика по шагам: «шаг [статус]: текст», через перевод строки"),
            "advice": ("TEXT", "Советы через перевод строки"),
            "priority": ("REAL", "Приоритет внимания: чем больше, тем срочнее"),
        },
    },
    "organic_store": {
        "description": "Показатели магазина из последнего анализа органики: заказы, органика, реклама, "
                       "сезонное изменение год к году, рекламная воронка.",
        "source": "analytics history/organic_last.json",
        "columns": {
            "metric": ("TEXT", "Показатель: orders, orders_prev, ad_orders, organic, organic_prev, season_change, ad_funnel_* …"),
            "value": ("REAL", "Значение"),
        },
    },
    "product_snapshots": {
        "description": "Ежедневные снимки товаров (по одному в день, когда запускался анализ): цена, остаток, "
                       "индекс цен, число акций, контент-рейтинг. Для динамики во времени.",
        "source": "analytics history/snapshots/<дата>.json",
        "date_column": "date",
        "columns": {
            "date": ("TEXT", "Дата снимка, YYYY-MM-DD"),
            "artikul": ("TEXT", ARTIKUL),
            "price": ("REAL", "Цена, ₽"),
            "stock": ("REAL", "Остаток, шт."),
            "price_index": ("TEXT", "Индекс цен"),
            "promotions_count": ("INTEGER", "Число акций"),
            "content_rating": ("REAL", "Контент-рейтинг"),
        },
    },
    "auto_actions_log": {
        "description": "Журнал изменений автодобавления в акции и минимальной цены (раздел «Акции»).",
        "source": "actions history/auto_actions_log.jsonl",
        "date_column": "at",
        "columns": {
            "at": ("TEXT", "Когда, ISO-время"),
            "action": ("TEXT", "Действие"),
            "artikul": ("TEXT", ARTIKUL),
            "product_id": ("INTEGER", "product_id Ozon"),
            "confirmed": ("INTEGER", "1 — Ozon подтвердил изменение"),
            "details": ("TEXT", "Остальные поля записи, JSON"),
            "error": ("TEXT", "Ошибка, если была"),
        },
    },
    "wb_orders": {
        "description": "Заказы Wildberries построчно из месячных отчётов WB (лист «Заказы»).",
        "source": "reports/WB/<Год>/WB <Месяц Год>.xlsx",
        "date_column": "period",
        "columns": {
            "period": ("TEXT", PERIOD),
            "status": ("TEXT", "Статус"),
            "order_number": ("TEXT", "Номер заказа"),
            "product_name": ("TEXT", "Название товара"),
            "artikul": ("TEXT", "Артикул продавца WB"),
            "quantity": ("REAL", "Количество, шт."),
            "sale_price": ("REAL", "Цена продажи, ₽"),
            "accrual": ("REAL", "Сумма начисления WB, ₽"),
            "cost": ("REAL", "Себестоимость, ₽ (минус)"),
            "profit": ("REAL", "Прибыль, ₽"),
            "date": ("TEXT", "Дата, YYYY-MM-DD"),
            "return_reason": ("TEXT", "Причина возврата"),
            "return_status": ("TEXT", "Статус возврата"),
        },
    },
    "wb_month_metrics": {
        "description": "Итоговые показатели месяца Wildberries (блок показателей листа «Заказы» отчёта WB).",
        "source": "reports/WB/<Год>/WB <Месяц Год>.xlsx",
        "date_column": "period",
        "columns": {
            "period": ("TEXT", PERIOD),
            "metric": ("TEXT", "Название показателя"),
            "value": ("REAL", "Значение"),
            "text": ("TEXT", "Текстовое значение, если не число"),
        },
    },
    "wb_products": {
        "description": "Справочник товаров Wildberries (wb_costs.xlsx): себестоимость.",
        "source": "wb_costs.xlsx",
        "columns": {
            "nm_id": ("TEXT", "Артикул WB (nmId)"),
            "artikul": ("TEXT", "Артикул продавца"),
            "name": ("TEXT", "Название"),
            "cost": ("REAL", "Себестоимость, ₽"),
        },
    },
    "search_products": {
        "description": "Товары в базе поиска: связь SKU Ozon с артикулом продавца. Используйте для соединения "
                       "search_* таблиц (по sku) с остальными (по artikul).",
        "source": "search data/search.db",
        "columns": {
            "sku": ("INTEGER", "SKU Ozon"),
            "artikul": ("TEXT", ARTIKUL),
            "product_id": ("INTEGER", "product_id Ozon"),
            "name": ("TEXT", "Название"),
            "category": ("TEXT", "Категория"),
        },
    },
    "search_keywords": {
        "description": "Поисковые запросы покупателей, связанные с товарами магазина, и их кластеры.",
        "source": "search data/search.db",
        "columns": {
            "keyword_id": ("INTEGER", "ID запроса"),
            "query": ("TEXT", "Текст запроса"),
            "cluster_id": ("INTEGER", "ID кластера (группа запросов с одними леммами)"),
            "cluster": ("TEXT", "Название кластера"),
            "frequency": ("REAL", "Частотность запроса (покупателей за период сбора)"),
            "priority": ("INTEGER", "Приоритет запроса"),
        },
    },
    "search_sku_keywords": {
        "description": "Какие запросы относятся к какому товару; is_target — целевой запрос, score — оценка приоритета.",
        "source": "search data/search.db",
        "columns": {
            "sku": ("INTEGER", "SKU Ozon"),
            "keyword_id": ("INTEGER", "ID запроса"),
            "is_target": ("INTEGER", "1 — целевой запрос, отслеживается"),
            "score": ("REAL", "Оценка приоритета запроса для товара"),
        },
    },
    "search_query_stats": {
        "description": "Ежедневная статистика по запросу для товара: сколько покупателей искали, видели товар, "
                       "позиция, заказы. Без подписки Premium заполнено в основном searched_users.",
        "source": "search data/search.db",
        "date_column": "date",
        "columns": {
            "date": ("TEXT", "Дата, YYYY-MM-DD"),
            "sku": ("INTEGER", "SKU Ozon"),
            "keyword_id": ("INTEGER", "ID запроса"),
            "searched_users": ("INTEGER", "Покупателей искали"),
            "viewed_users": ("INTEGER", "Покупателей видели товар"),
            "view_conversion": ("REAL", "Конверсия в просмотр"),
            "position": ("REAL", "Средняя позиция"),
            "orders": ("INTEGER", "Заказы"),
            "gmv": ("REAL", "Сумма заказов, ₽"),
        },
    },
    "search_sku_daily": {
        "description": "Ежедневные итоги поиска по товару: спрос, просмотры, позиция, заказы и выручка.",
        "source": "search data/search.db",
        "date_column": "date",
        "columns": {
            "date": ("TEXT", "Дата, YYYY-MM-DD"),
            "sku": ("INTEGER", "SKU Ozon"),
            "searched_users": ("INTEGER", "Покупателей искали по запросам с товаром"),
            "viewed_users": ("INTEGER", "Покупателей видели товар"),
            "view_conversion": ("REAL", "Конверсия в просмотр"),
            "position": ("REAL", "Средняя позиция"),
            "gmv": ("REAL", "Сумма заказов, ₽"),
            "ordered_units": ("REAL", "Заказано, шт."),
            "revenue": ("REAL", "Выручка, ₽"),
        },
    },
    "search_positions": {
        "description": "Замеры позиции товара в выдаче Ozon по запросу и региону (парсер выдачи).",
        "source": "search data/search.db",
        "date_column": "date",
        "columns": {
            "date": ("TEXT", "Дата замера, YYYY-MM-DD"),
            "sku": ("INTEGER", "SKU Ozon"),
            "keyword_id": ("INTEGER", "ID запроса"),
            "region": ("TEXT", "Регион"),
            "position": ("INTEGER", "Позиция в выдаче"),
            "organic_position": ("INTEGER", "Позиция без учёта рекламы"),
            "is_ad": ("INTEGER", "1 — рекламное место"),
            "found": ("INTEGER", "1 — товар найден в выдаче"),
        },
    },
    "search_card_snapshots": {
        "description": "Снимки карточки товара: название и описание на дату — для отслеживания изменений.",
        "source": "search data/search.db",
        "date_column": "date",
        "columns": {
            "date": ("TEXT", "Дата снимка, YYYY-MM-DD"),
            "sku": ("INTEGER", "SKU Ozon"),
            "name": ("TEXT", "Название карточки"),
            "description": ("TEXT", "Описание карточки"),
        },
    },
}


# Представления «самое свежее»: модели не нужно каждый раз выяснять, какой отчёт или снимок последний.
# Колонки — как у базовой таблицы.
VIEWS: Dict[str, Dict[str, str]] = {
    "abc_xyz_latest": {
        "base": "abc_xyz",
        "description": "ABC/XYZ-классы из основного свежего отчёта: самый длинный период среди отчётов, закончившихся "
                       "не раньше чем за 2 месяца до самого свежего (короткий отчёт за 1–2 месяца для ABC ненадёжен).",
        "sql": """SELECT * FROM abc_xyz WHERE report_period = (
                    SELECT report_period FROM abc_xyz
                    WHERE period_end >= strftime('%Y-%m', (SELECT MAX(period_end) FROM abc_xyz) || '-01', '-2 months')
                    ORDER BY (CAST(substr(period_end, 1, 4) AS INTEGER) * 12 + CAST(substr(period_end, 6, 2) AS INTEGER))
                           - (CAST(substr(period_start, 1, 4) AS INTEGER) * 12 + CAST(substr(period_start, 6, 2) AS INTEGER)) DESC,
                             period_end DESC
                    LIMIT 1)""",
    },
    "product_snapshot_latest": {
        "base": "product_snapshots",
        "description": "Последний снимок товаров: цена, остаток, индекс цен, число акций.",
        "sql": "SELECT * FROM product_snapshots WHERE date = (SELECT MAX(date) FROM product_snapshots)",
    },
    "fbo_supply_latest": {
        "base": "fbo_supply_plan",
        "description": "Последний план поставки на FBO (артикул × зона).",
        "sql": "SELECT * FROM fbo_supply_plan WHERE plan_date = (SELECT MAX(plan_date) FROM fbo_supply_plan)",
    },
    "fbs_needs_latest": {
        "base": "fbs_needs",
        "description": "Нехватка на своём складе (FBS) по последнему плану поставки.",
        "sql": "SELECT * FROM fbs_needs WHERE plan_date = (SELECT MAX(plan_date) FROM fbs_needs)",
    },
}


# ---------- чтение значений ----------

def _num(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not text or text in ("—", "-"):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _date(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = _text(value)
    return text[:10] if text else None


def period_from_name(name: str) -> Optional[str]:
    """«Сентябрь 2026» → «2026-09»."""
    match = re.search(r"([А-Яа-яЁё]+)\s+(\d{4})", name)
    if not match:
        return None
    month = match.group(1).capitalize()
    if month not in MONTHS_RU:
        return None
    return f"{match.group(2)}-{MONTHS_RU.index(month) + 1:02d}"


def _sheet_rows(ws) -> List[Tuple[Any, ...]]:
    return [tuple(row) for row in ws.iter_rows(values_only=True)]


def _header_index(header: Sequence[Any]) -> Dict[str, int]:
    return {str(cell).strip(): i for i, cell in enumerate(header) if cell is not None and str(cell).strip()}


def _rows_by_header(rows: List[Tuple[Any, ...]], mapping: Dict[str, Sequence[str]]) -> Iterable[Dict[str, Any]]:
    """Строки листа как словари: колонка базы → значение, по первому совпавшему заголовку."""
    if not rows:
        return
    index = _header_index(rows[0])
    columns = {}
    for target, names in mapping.items():
        for name in names:
            if name in index:
                columns[target] = index[name]
                break
    for row in rows[1:]:
        values = {target: (row[i] if i < len(row) else None) for target, i in columns.items()}
        if any(value is not None for value in values.values()):
            yield values


def _metrics_block(rows: List[Tuple[Any, ...]], first_metric: str = "Общая выручка") -> List[Tuple[str, Any]]:
    """Блок «показатель — значение» справа от таблицы заказов: ищем колонку с первым показателем."""
    name_col = None
    for row in rows[:5]:
        for i, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip() == first_metric:
                name_col = i
                break
        if name_col is not None:
            break
    if name_col is None:
        return []
    result = []
    for row in rows:
        name = row[name_col] if name_col < len(row) else None
        value = row[name_col + 1] if name_col + 1 < len(row) else None
        if isinstance(name, str) and name.strip():
            result.append((name.strip(), value))
    return result


def _metric_rows(period: str, block: List[Tuple[str, Any]]) -> List[Tuple[Any, ...]]:
    rows = []
    for name, value in block:
        number = _num(value)
        rows.append((period, name, number, None if number is not None else _text(value)))
    return rows


# ---------- загрузчики источников ----------

class Source:
    """Источник: файлы, по которым видно изменение, и функция загрузки в свои таблицы."""

    def __init__(self, name: str, tables: Sequence[str], files: Callable[[Path], List[Path]], load: Callable[[Path, sqlite3.Connection], None]):
        self.name = name
        self.tables = list(tables)
        self.files = files
        self.load = load


def _insert(conn: sqlite3.Connection, table: str, rows: Iterable[Sequence[Any]]) -> None:
    columns = list(CATALOG[table]["columns"])
    placeholders = ",".join("?" for _ in columns)
    conn.executemany(f"INSERT INTO {table} ({','.join(columns)}) VALUES ({placeholders})", rows)


def _monthly_reports(root: Path, wb: bool) -> List[Path]:
    result = []
    for path in (wb_report_files(root) if wb else ozon_report_files(root)):
        if path.name.startswith("WB ") == wb and period_from_name(path.stem):
            result.append(path)
    return result


ORDER_COLUMNS = {
    "status": ["Статус"],
    "order_number": ["Номер заказа"],
    "product_name": ["Название товара"],
    "artikul": ["Артикул"],
    "quantity": ["Количество шт."],
    "sale_price": ["Цена продажи"],
    "commission": ["Комиссия за продажу Ozon"],
    "logistics": ["Логистика (Включает операционные ошибки продавца)", "Логистика"],
    "accrual": ["Сумма начисления"],
    "cost": ["Себестоимость"],
    "profit": ["Прибыль"],
    "ship_date": ["Дата отгрузки"],
    "scheme": ["Схема"],
}

EXPENSE_COLUMNS = {
    "date": ["Дата"],
    "kind": ["Вид расхода"],
    "ozon_code": ["Код Ozon"],
    "reference": ["Номер (поставки/кампании)"],
    "amount": ["Сумма"],
    "net_profit_note": ["В чистой прибыли"],
}

CAMPAIGN_COLUMNS = {
    "campaign_id": ["ID кампании"],
    "name": ["Название кампании"],
    "state": ["Состояние"],
    "payment_type": ["Тип оплаты"],
    "object_type": ["Тип объекта"],
    "budget": ["Бюджет (руб.)"],
    "daily_budget": ["Дневной бюджет (руб.)"],
    "weekly_budget": ["Недельный бюджет (руб.)"],
    "spend": ["Расход за период (руб.)"],
    "impressions": ["Показы"],
    "clicks": ["Клики"],
    "ctr_pct": ["CTR (%)"],
    "avg_cpc": ["Средняя цена клика (руб.)"],
    "orders": ["Заказы (шт.)"],
    "orders_revenue": ["Заказы (руб.)"],
    "drr_pct": ["ДРР (%)"],
}

TEXT_FIELDS = {"status", "order_number", "product_name", "artikul", "scheme", "kind", "ozon_code", "reference",
               "net_profit_note", "campaign_id", "name", "state", "payment_type", "object_type",
               "return_reason", "return_status", "nm_id", "promotion"}
DATE_FIELDS = {"ship_date", "date"}


def _convert(table: str, values: Dict[str, Any], extra: Dict[str, Any]) -> Tuple[Any, ...]:
    row = []
    for column in CATALOG[table]["columns"]:
        if column in extra:
            row.append(extra[column])
            continue
        value = values.get(column)
        if column in DATE_FIELDS:
            row.append(_date(value))
        elif column in TEXT_FIELDS:
            row.append(_text(value))
        else:
            row.append(_num(value))
    return tuple(row)


def _open_book(path: Path):
    return closing(load_workbook(path, read_only=True, data_only=True))


def load_ozon_reports(root: Path, conn: sqlite3.Connection) -> None:
    for path in _monthly_reports(root, wb=False):
        period = period_from_name(path.stem)
        with _open_book(path) as book:
            orders_sheet = data_sheet_name(book.sheetnames, ORDERS)
            if orders_sheet:
                rows = _sheet_rows(book[orders_sheet])
                _insert(conn, "ozon_orders", (
                    _convert("ozon_orders", values, {"period": period})
                    for values in _rows_by_header(rows, ORDER_COLUMNS)
                    if _text(values.get("order_number")) or _text(values.get("artikul"))
                ))
                _insert(conn, "ozon_month_metrics", _metric_rows(period, _metrics_block(rows)))
            charges_sheet = data_sheet_name(book.sheetnames, CHARGES)
            if charges_sheet:
                _insert(conn, "ozon_other_expenses", (
                    _convert("ozon_other_expenses", values, {"period": period})
                    for values in _rows_by_header(_sheet_rows(book[charges_sheet]), EXPENSE_COLUMNS)
                    if _num(values.get("amount")) is not None
                ))
            campaigns_sheet = data_sheet_name(book.sheetnames, CAMPAIGNS)
            if campaigns_sheet:
                _insert(conn, "ozon_ad_campaigns", (
                    _convert("ozon_ad_campaigns", values, {"period": period})
                    for values in _rows_by_header(_sheet_rows(book[campaigns_sheet]), CAMPAIGN_COLUMNS)
                    if _text(values.get("campaign_id"))
                ))


WB_ORDER_COLUMNS = {
    "status": ["Статус"],
    "order_number": ["Номер заказа"],
    "product_name": ["Название товара"],
    "artikul": ["Артикул"],
    "quantity": ["Количество шт."],
    "sale_price": ["Цена продажи"],
    "accrual": ["Сумма начисления"],
    "cost": ["Себестоимость"],
    "profit": ["Прибыль"],
    "date": ["Дата"],
    "return_reason": ["Причина возврата (WB)"],
    "return_status": ["Статус возврата (WB)"],
}


def load_wb_reports(root: Path, conn: sqlite3.Connection) -> None:
    for path in _monthly_reports(root, wb=True):
        period = period_from_name(path.stem)
        with _open_book(path) as book:
            # Отчёт нового вида: прежний лист «Заказы» лежит скрытым под именем «Данные» (report_book).
            sheet = data_sheet_name(book.sheetnames, ORDERS)
            if not sheet:
                continue
            rows = _sheet_rows(book[sheet])
            _insert(conn, "wb_orders", (
                _convert("wb_orders", values, {"period": period})
                for values in _rows_by_header(rows, WB_ORDER_COLUMNS)
                if _text(values.get("order_number")) or _text(values.get("artikul"))
            ))
            _insert(conn, "wb_month_metrics", _metric_rows(period, _metrics_block(rows)))


PRODUCT_COLUMNS = {
    "artikul": ["Артикул"],
    "cost": ["Себестоимость"],
    "min_price": ["Минимальная цена продажи"],
    "desired_price": ["Желательная цена продажи"],
    "current_price": ["Текущая цена на Ozon"],
    "promo_price": ["Цена с учётом акций и скидок"],
    "expected_margin_pct": ["Текущая ожидаемая рентабельность"],
    "expected_net_profit": ["Ожидаемая чистая прибыль"],
    "logistics_per_unit": ["Логистика за шт., ₽ (прошлый месяц)"],
}
BOOST_SUFFIX = " — бустинг, %"


def load_costs(root: Path, conn: sqlite3.Connection) -> None:
    path = root / "costs.xlsx"
    if not path.exists():
        return
    with _open_book(path) as book:
        main = book["Основной"] if "Основной" in book.sheetnames else book.worksheets[0]
        _insert(conn, "products", (
            _convert("products", values, {})
            for values in _rows_by_header(_sheet_rows(main), PRODUCT_COLUMNS)
            if _text(values.get("artikul"))
        ))
        if "Акции" not in book.sheetnames:
            return
        rows = _sheet_rows(book["Акции"])
        if not rows:
            return
        index = _header_index(rows[0])
        promotions = [name for name in index if name not in ("Артикул", "Минимальная цена продажи") and not name.endswith(BOOST_SUFFIX)]
        result = []
        for row in rows[1:]:
            artikul = _text(row[index["Артикул"]]) if "Артикул" in index else None
            if not artikul:
                continue
            min_price = _num(row[index["Минимальная цена продажи"]]) if "Минимальная цена продажи" in index else None
            for promotion in promotions:
                price = _num(row[index[promotion]])
                boost_col = index.get(promotion + BOOST_SUFFIX)
                boost = _num(row[boost_col]) if boost_col is not None else None
                if price is None and boost is None:
                    continue
                result.append((artikul, promotion, price, boost, min_price))
        _insert(conn, "product_promotions", result)


def _abc_files(root: Path) -> List[Path]:
    folder = root / "ABC&XYZ reports"
    return sorted(p for p in folder.glob("*.xlsx") if not p.name.startswith(("~", "."))) if folder.is_dir() else []


def load_abc_xyz(root: Path, conn: sqlite3.Connection) -> None:
    for path in _abc_files(root):
        parts = path.stem.split("-")
        start = period_from_name(parts[0]) if parts else None
        end = period_from_name(parts[-1]) if parts else None
        with _open_book(path) as book:
            if "Итог" not in book.sheetnames:
                continue
            mapping = {"artikul": ["Артикул"], "abc": ["Оценка по ABC"], "xyz": ["Оценка по XYZ"], "abcxyz": ["Общая оценка ABCXYZ"]}
            _insert(conn, "abc_xyz", (
                (path.stem, start, end, _text(v.get("artikul")), _text(v.get("abc")), _text(v.get("xyz")), _text(v.get("abcxyz")))
                for v in _rows_by_header(_sheet_rows(book["Итог"]), mapping)
                if _text(v.get("artikul"))
            ))


def _supply_files(root: Path) -> List[Path]:
    folder = root / "stocks reports"
    return sorted(folder.glob("Поставка FBO *.json")) if folder.is_dir() else []


def _plan_date(path: Path) -> Optional[str]:
    match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", path.stem)
    return f"{match.group(3)}-{match.group(2)}-{match.group(1)}" if match else None


def load_supply(root: Path, conn: sqlite3.Connection) -> None:
    for path in _supply_files(root):
        data = json.loads(path.read_text(encoding="utf-8"))
        plan_date = _plan_date(path) or _date(data.get("created"))
        arrival = _text(data.get("arrival"))
        rows = []
        for item in data.get("supply") or []:
            for zone, quantity in (item.get("by_zone") or {}).items():
                rows.append((plan_date, arrival, _text(item.get("article")), zone, _num(quantity), _num(item.get("total")),
                             _num(item.get("per_day")), _num(item.get("own_stock")), _num(item.get("unit_cost")), _num(item.get("cost"))))
        _insert(conn, "fbo_supply_plan", rows)
        _insert(conn, "fbs_needs", (
            (plan_date, _text(item.get("article")), _num(item.get("demand")), _num(item.get("need")),
             _num(item.get("stock")), _num(item.get("shortage")), _num(item.get("cost")))
            for item in ((data.get("fbs") or {}).get("rows") or [])
        ))


def _organic_file(root: Path) -> List[Path]:
    path = root / "analytics history" / "organic_last.json"
    return [path] if path.exists() else []


def load_organic(root: Path, conn: sqlite3.Connection) -> None:
    for path in _organic_file(root):
        data = json.loads(path.read_text(encoding="utf-8"))
        created = _text(data.get("created"))
        current = (data.get("periods") or {}).get("current") or [None, None]
        rows = []
        for item in data.get("products") or []:
            locality = item.get("locality") or {}
            local_share = None
            if _num(locality.get("total")):
                local_share = round(100.0 * (_num(locality.get("local")) or 0) / _num(locality.get("total")), 1)
            price_index = item.get("price_index") or {}
            actions = item.get("actions") or []
            steps = item.get("steps") or []
            rows.append((
                created, current[0], current[1], _text(item.get("offer_id")), item.get("sku"),
                _num(item.get("orders")), _num(item.get("orders_prev")), _num(item.get("ad_orders")),
                _num(item.get("organic")), _num(item.get("organic_prev")), _num(item.get("organic_change")),
                _num(item.get("lost_orders")), _num(item.get("stock")), _num(item.get("days_since_sale")),
                _num(item.get("avg_price")), _num(item.get("avg_price_prev")), _num(item.get("price_now")),
                _num(item.get("min_price")), _text(price_index.get("color")), _num(price_index.get("value")),
                len(actions), "; ".join(str(a) for a in actions) or None, _num(item.get("content_rating")),
                "; ".join(str(m) for m in item.get("content_missing") or []) or None, local_share,
                _num(locality.get("recommended_supply")), _num((item.get("search") or {}).get("searched")),
                _text(item.get("main_problem")), _text(item.get("main_status")),
                "\n".join(f"{s.get('name')} [{s.get('status')}]: {s.get('text')}" for s in steps if isinstance(s, dict)) or None,
                "\n".join(str(a) for a in item.get("advice") or []) or None, _num(item.get("priority")),
            ))
        _insert(conn, "organic_products", rows)
        metrics = []
        for key, value in (data.get("store") or {}).items():
            if isinstance(value, dict):
                metrics.extend((f"{key}_{k}", _num(v)) for k, v in value.items())
            elif _num(value) is not None:
                metrics.append((key, _num(value)))
        _insert(conn, "organic_store", metrics)


def _snapshot_files(root: Path) -> List[Path]:
    folder = root / "analytics history" / "snapshots"
    return sorted(folder.glob("*.json")) if folder.is_dir() else []


def load_snapshots(root: Path, conn: sqlite3.Connection) -> None:
    for path in _snapshot_files(root):
        data = json.loads(path.read_text(encoding="utf-8"))
        _insert(conn, "product_snapshots", (
            (path.stem, str(artikul), _num(v.get("price")), _num(v.get("stock")), _text(v.get("index")),
             v.get("actions") if isinstance(v.get("actions"), int) else None, _num(v.get("rating")))
            for artikul, v in data.items() if isinstance(v, dict)
        ))


def _actions_log_file(root: Path) -> List[Path]:
    path = root / "actions history" / "auto_actions_log.jsonl"
    return [path] if path.exists() else []


def load_actions_log(root: Path, conn: sqlite3.Connection) -> None:
    for path in _actions_log_file(root):
        rows = []
        for line in path.read_text(encoding="utf-8").splitlines():
            try:
                item = json.loads(line)
            except json.JSONDecodeError:
                continue
            if not isinstance(item, dict):
                continue
            details = {k: v for k, v in item.items() if k not in ("at", "action", "offer_id", "product_id", "confirmed", "error")}
            rows.append((_text(item.get("at")), _text(item.get("action")), _text(item.get("offer_id")), item.get("product_id"),
                         1 if item.get("confirmed") else 0, json.dumps(details, ensure_ascii=False) if details else None,
                         _text(item.get("error"))))
        _insert(conn, "auto_actions_log", rows)


def _wb_costs_file(root: Path) -> List[Path]:
    path = root / "wb_costs.xlsx"
    return [path] if path.exists() else []


def load_wb_costs(root: Path, conn: sqlite3.Connection) -> None:
    for path in _wb_costs_file(root):
        with _open_book(path) as book:
            mapping = {"nm_id": ["nmId"], "artikul": ["Артикул продавца"], "name": ["Название"], "cost": ["Себестоимость, руб."]}
            _insert(conn, "wb_products", (
                _convert("wb_products", v, {}) for v in _rows_by_header(_sheet_rows(book.worksheets[0]), mapping)
                if _text(v.get("nm_id")) or _text(v.get("artikul"))
            ))


def _search_file(root: Path) -> List[Path]:
    path = root / "search data" / "search.db"
    return [path] if path.exists() else []


SEARCH_QUERIES = {
    "search_products": "SELECT sku, offer_id, product_id, name, category FROM products",
    "search_keywords": "SELECT k.id, k.query, k.cluster_id, c.name, k.frequency, k.priority FROM keywords k LEFT JOIN clusters c ON c.id = k.cluster_id",
    "search_sku_keywords": "SELECT sku, keyword_id, is_target, score FROM sku_keywords",
    "search_query_stats": "SELECT date, sku, keyword_id, searched_users, viewed_users, view_conversion, position, orders, gmv FROM query_stats",
    "search_sku_daily": "SELECT date, sku, searched_users, viewed_users, view_conversion, position, gmv, ordered_units, revenue FROM sku_daily",
    "search_positions": "SELECT p.date, p.sku, p.keyword_id, r.name, p.position, p.organic_position, p.is_ad, p.found FROM positions p LEFT JOIN regions r ON r.id = p.region_id",
    "search_card_snapshots": "SELECT date, sku, name, description FROM card_snapshots",
}


def load_search(root: Path, conn: sqlite3.Connection) -> None:
    for path in _search_file(root):
        # Только чтение: база поиска принадлежит сборщику, её не трогаем и не блокируем на запись.
        with closing(sqlite3.connect(f"file:{path}?mode=ro", uri=True, timeout=30)) as source:
            for table, query in SEARCH_QUERIES.items():
                try:
                    rows = source.execute(query).fetchall()
                except sqlite3.OperationalError:
                    continue  # старая схема без нужной таблицы или колонки
                _insert(conn, table, rows)


SOURCES = [
    Source("ozon_reports", ["ozon_orders", "ozon_month_metrics", "ozon_other_expenses", "ozon_ad_campaigns"],
           lambda root: _monthly_reports(root, wb=False), load_ozon_reports),
    Source("wb_reports", ["wb_orders", "wb_month_metrics"], lambda root: _monthly_reports(root, wb=True), load_wb_reports),
    Source("costs", ["products", "product_promotions"], lambda root: [p for p in [root / "costs.xlsx"] if p.exists()], load_costs),
    Source("abc_xyz", ["abc_xyz"], _abc_files, load_abc_xyz),
    Source("supply", ["fbo_supply_plan", "fbs_needs"], _supply_files, load_supply),
    Source("organic", ["organic_products", "organic_store"], _organic_file, load_organic),
    Source("snapshots", ["product_snapshots"], _snapshot_files, load_snapshots),
    Source("actions_log", ["auto_actions_log"], _actions_log_file, load_actions_log),
    Source("wb_costs", ["wb_products"], _wb_costs_file, load_wb_costs),
    Source("search", list(SEARCH_QUERIES), _search_file, load_search),
]


# ---------- сборка ----------

def _signature(root: Path, files: List[Path]) -> str:
    parts = []
    for path in files:
        try:
            stat = path.stat()
        except OSError:
            continue
        parts.append([str(path.relative_to(root)), stat.st_mtime_ns, stat.st_size])
    return json.dumps(parts, ensure_ascii=False)


def _create_schema(conn: sqlite3.Connection) -> None:
    conn.execute("CREATE TABLE IF NOT EXISTS _meta (key TEXT PRIMARY KEY, value TEXT)")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS _sources (name TEXT PRIMARY KEY, signature TEXT, loaded_at TEXT, "
        "files INTEGER, error TEXT)"
    )
    for table, spec in CATALOG.items():
        columns = ", ".join(f"{name} {kind}" for name, (kind, _) in spec["columns"].items())
        conn.execute(f"CREATE TABLE IF NOT EXISTS {table} ({columns})")
        for column in ("period", "artikul", "sku", "date"):
            if column in spec["columns"]:
                conn.execute(f"CREATE INDEX IF NOT EXISTS {table}_{column} ON {table}({column})")
    for view, spec in VIEWS.items():
        conn.execute(f"DROP VIEW IF EXISTS {view}")
        conn.execute(f"CREATE VIEW {view} AS {spec['sql']}")
    conn.execute("INSERT OR REPLACE INTO _meta (key, value) VALUES ('schema_version', ?)", (str(SCHEMA_VERSION),))


def _schema_version(path: Path) -> Optional[int]:
    try:
        with closing(sqlite3.connect(path)) as conn:
            row = conn.execute("SELECT value FROM _meta WHERE key = 'schema_version'").fetchone()
            return int(row[0]) if row else None
    except (sqlite3.Error, ValueError):
        return None


def warehouse_path(root: Optional[Path] = None) -> Path:
    return Path(root or ROOT) / WAREHOUSE_PATH


def refresh(root: Optional[Path] = None) -> Dict[str, str]:
    """
    Пересобирает источники, файлы которых изменились. Возвращает {источник: "updated" | "fresh" | "error: …"}.
    Ошибка одного источника не мешает остальным: его таблицы остаются пустыми, причина — в _sources.error.
    """
    root = Path(root or ROOT)
    path = warehouse_path(root)
    path.parent.mkdir(parents=True, exist_ok=True)
    result: Dict[str, str] = {}
    with exclusive_file(path):
        if path.exists() and _schema_version(path) != SCHEMA_VERSION:
            path.unlink()
        with closing(sqlite3.connect(path, timeout=30)) as conn:
            with conn:
                _create_schema(conn)
            known = {name: signature for name, signature in conn.execute("SELECT name, signature FROM _sources")}
            for source in SOURCES:
                files = source.files(root)
                signature = _signature(root, files)
                if known.get(source.name) == signature:
                    result[source.name] = "fresh"
                    continue
                error = None
                try:
                    with conn:
                        for table in source.tables:
                            conn.execute(f"DELETE FROM {table}")
                        source.load(root, conn)
                except Exception as exc:  # битый файл не должен ломать остальные источники
                    error = f"{type(exc).__name__}: {exc}"
                with conn:
                    if error:
                        for table in source.tables:
                            conn.execute(f"DELETE FROM {table}")
                    conn.execute(
                        "INSERT OR REPLACE INTO _sources (name, signature, loaded_at, files, error) VALUES (?, ?, ?, ?, ?)",
                        # С ошибкой подпись не сохраняем: следующая сборка попробует снова.
                        (source.name, None if error else signature, datetime.now().isoformat(timespec="seconds"), len(files), error),
                    )
                result[source.name] = f"error: {error}" if error else "updated"
    return result


def connect(root: Optional[Path] = None, refresh_first: bool = True) -> sqlite3.Connection:
    """Соединение только для чтения: модель и инструменты не могут изменить базу."""
    if refresh_first:
        refresh(root)
    path = warehouse_path(root)
    conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA query_only = ON")
    return conn


def describe(conn: sqlite3.Connection) -> List[Dict[str, Any]]:
    """Каталог для модели: таблица, описание, строк, диапазон дат, колонки, ошибка источника."""
    errors = {}
    loaded = {}
    for name, error, loaded_at in conn.execute("SELECT name, error, loaded_at FROM _sources"):
        for source in SOURCES:
            if source.name == name:
                for table in source.tables:
                    errors[table] = error
                    loaded[table] = loaded_at
    result = []
    for table, spec in CATALOG.items():
        rows = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        item: Dict[str, Any] = {
            "table": table,
            "description": spec["description"],
            "source": spec["source"],
            "rows": rows,
            "loaded_at": loaded.get(table),
            "columns": {name: desc for name, (_, desc) in spec["columns"].items()},
        }
        date_column = spec.get("date_column")
        if date_column and rows:
            first, last = conn.execute(f"SELECT MIN({date_column}), MAX({date_column}) FROM {table}").fetchone()
            item["range"] = [first, last]
        if errors.get(table):
            item["error"] = errors[table]
        result.append(item)
    for view, spec in VIEWS.items():
        base = next(item for item in result if item["table"] == spec["base"])
        rows = conn.execute(f"SELECT COUNT(*) FROM {view}").fetchone()[0]
        result.append({**{k: v for k, v in base.items() if k not in ("range", "rows")},
                       "table": view, "description": spec["description"], "rows": rows, "view_of": spec["base"]})
    return result


def main() -> None:
    root = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT
    for name, status in refresh(root).items():
        print(f"{name}: {status}", flush=True)
    with closing(connect(root, refresh_first=False)) as conn:
        for item in describe(conn):
            span = " … ".join(str(x) for x in item.get("range", [])) if item.get("range") else ""
            error = f"  ОШИБКА: {item['error']}" if item.get("error") else ""
            print(f"{item['table']:<24} {item['rows']:>7}  {span}{error}", flush=True)


if __name__ == "__main__":
    main()
