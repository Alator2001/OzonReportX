"""
Сверка выплат Ozon с месячными отчётами: сколько денег реально пришло и что они покрыли.

Начисления каждого месяца (себестоимость, затем чистая прибыль) — это «долг» Ozon перед продавцом.
Выплаты (реальные переводы из отчёта о балансе) гасят этот долг по порядку, начиная с самого
старого месяца (FIFO). Результат пишется в каждый месячный отчёт на лист «Заказы»:
  S:U — какие выплаты (дата, сумма) закрыли этот месяц и что именно (себестоимость/прибыль);
  W:Y — себестоимость и прибыль месяца, сколько закрыто выплатами и сколько осталось закрыть.
Когда обе дельты равны нулю, Ozon полностью рассчитался за месяц.

Отчёт о балансе отдаёт выплаты суммой за период, без даты, поэтому дату каждой выплаты находим
делением периода пополам до одного дня. Найденные выплаты закрытых месяцев кэшируются в
«balance reports/payouts.json», и при следующей пересборке API спрашивается только за текущий месяц.
"""

import json
import os
import re
import sys
from calendar import monthrange
from contextlib import closing
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

try:
    from scripts.balance_report import get_balance_report, BALANCE_REPORTS_DIR
    from scripts.file_io import save_workbook_atomic
except ModuleNotFoundError:
    from balance_report import get_balance_report, BALANCE_REPORTS_DIR
    from file_io import save_workbook_atomic

MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
          "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
REPORT_NAME = re.compile(r"^(%s) (\d{4})\.xlsx$" % "|".join(MONTHS))
COST_LABEL = "Итоговая себестоимость"
PROFIT_LABEL = "Чистая прибыль"
COST, PROFIT = "Себестоимость", "Прибыль"
FIRST_COLUMN, LAST_COLUMN = 19, 25  # S..Y
MONEY_FORMAT = "#,##0.00"


def _kopecks(value) -> int:
    return int(round(float(value or 0) * 100))


def _rub(kopecks: int) -> float:
    return kopecks / 100


# ── Выплаты ──────────────────────────────────────────────────────────────────

def _paid_between(day_from: date, day_to: date, fetch) -> int:
    """Сумма выплат за период в копейках (положительное число)."""
    data = fetch(day_from.isoformat(), day_to.isoformat())
    payments = (data.get("total") or {}).get("payments") or []
    return -sum(_kopecks(p.get("value")) for p in payments)


def _find_payouts(day_from: date, day_to: date, fetch, paid=None) -> list[tuple[date, int]]:
    """Делит период пополам, пока каждая выплата не окажется в пределах одного дня."""
    if paid is None:
        paid = _paid_between(day_from, day_to, fetch)
    if paid == 0:
        return []
    if day_from == day_to:
        return [(day_from, paid)]
    middle = day_from + (day_to - day_from) // 2
    left = _paid_between(day_from, middle, fetch)
    return (_find_payouts(day_from, middle, fetch, left)
            + _find_payouts(middle + timedelta(days=1), day_to, fetch, paid - left))


def _month_payouts(year: int, month: int, last_day: date, fetch) -> list[tuple[date, int]]:
    """Выплаты месяца по дням. Отчёт о балансе принимает период не длиннее 30 дней."""
    first = date(year, month, 1)
    end = min(date(year, month, monthrange(year, month)[1]), last_day)
    segments = [(first, min(end, first + timedelta(days=29)))]
    if end > segments[0][1]:
        segments.append((segments[0][1] + timedelta(days=1), end))
    result = []
    for start, stop in segments:
        result.extend(_find_payouts(start, stop, fetch))
    return result


def load_payouts(start: date, today: date, cache_path: Path, fetch=get_balance_report) -> list[tuple[date, int]]:
    """Все выплаты Ozon с месяца start по today включительно, с кэшем закрытых месяцев."""
    try:
        cache = json.loads(cache_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        cache = {}
    months = cache.get("months") or {}
    payouts = []
    year, month = start.year, start.month
    while (year, month) <= (today.year, today.month):
        key = f"{year}-{month:02d}"
        closed = date(year, month, monthrange(year, month)[1]) < today
        if closed and key in months:
            found = [(date.fromisoformat(d), int(k)) for d, k in months[key]]
        else:
            found = _month_payouts(year, month, today, fetch)
            if closed:
                months[key] = [[d.isoformat(), k] for d, k in found]
        payouts.extend(found)
        month += 1
        if month == 13:
            year, month = year + 1, 1
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(json.dumps({"months": months}, ensure_ascii=False, indent=1), encoding="utf-8")
    return payouts


# ── Начисления из отчётов ────────────────────────────────────────────────────

def find_reports(reports_dir: Path) -> list[tuple[date, Path]]:
    """Месячные отчёты Ozon («Сентябрь 2026.xlsx») по порядку; отчёты WB и временные файлы пропускаются."""
    found = []
    for path in Path(reports_dir).glob("*.xlsx"):
        match = REPORT_NAME.match(path.name)
        if match:
            found.append((date(int(match.group(2)), MONTHS.index(match.group(1)) + 1, 1), path))
    return sorted(found)


def _orders_sheet(workbook):
    return workbook["Заказы"] if "Заказы" in workbook.sheetnames else workbook.active


def read_obligations(path: Path) -> dict[str, int]:
    """Себестоимость и чистая прибыль месяца из блока показателей P:Q, в копейках."""
    with closing(load_workbook(path, data_only=True, read_only=True)) as workbook:
        values = {}
        for label, value in _orders_sheet(workbook).iter_rows(min_col=16, max_col=17, values_only=True):
            if isinstance(label, str) and isinstance(value, (int, float)):
                values.setdefault(label.strip(), value)
    if COST_LABEL not in values or PROFIT_LABEL not in values:
        raise ValueError(f"в отчёте {path.name} нет показателей «{COST_LABEL}» и «{PROFIT_LABEL}»")
    return {COST: abs(_kopecks(values[COST_LABEL])), PROFIT: _kopecks(values[PROFIT_LABEL])}


# ── Распределение ────────────────────────────────────────────────────────────

def allocate(obligations: list[tuple[date, dict[str, int]]], payouts: list[tuple[date, int]]):
    """Гасит начисления выплатами по порядку: месяц за месяцем, внутри месяца сначала себестоимость.

    Убыточный месяц (отрицательная прибыль) не создаёт долга по прибыли.
    Возвращает ({месяц: [(дата, сумма, на что, сумма всей выплаты)]}, {месяц: {вид: закрыто}}, остаток).
    """
    buckets = [[month, kind, max(amounts[kind], 0)] for month, amounts in obligations for kind in (COST, PROFIT)]
    parts = {month: [] for month, _ in obligations}
    closed = {month: {COST: 0, PROFIT: 0} for month, _ in obligations}
    index = 0
    unallocated = 0
    for paid_on, amount in sorted(payouts):
        left = amount
        while left > 0:
            while index < len(buckets) and buckets[index][2] == 0:
                index += 1
            if index == len(buckets):
                unallocated += left
                break
            month, kind, due = buckets[index]
            part = min(left, due)
            buckets[index][2] -= part
            left -= part
            parts[month].append((paid_on, part, kind, amount))
            closed[month][kind] += part
    return parts, closed, unallocated


# ── Запись в отчёты ──────────────────────────────────────────────────────────

def _money(cell, kopecks: int):
    cell.value = _rub(kopecks)
    cell.number_format = MONEY_FORMAT


def write_settlement(path: Path, amounts: dict[str, int], parts, closed: dict[str, int],
                     total_paid: int, unallocated: int) -> None:
    workbook = load_workbook(path)
    try:
        ws = _orders_sheet(workbook)
        for row in ws.iter_rows(min_col=FIRST_COLUMN, max_col=LAST_COLUMN):
            for cell in row:
                cell.value = None

        ws["S1"], ws["T1"], ws["U1"] = "Дата выплаты Ozon", "Сумма из выплаты", "На что ушло"
        for row, (paid_on, part, kind, whole) in enumerate(parts, start=2):
            ws.cell(row=row, column=19, value=paid_on).number_format = "DD.MM.YYYY"
            _money(ws.cell(row=row, column=20), part)
            note = kind if part == whole else f"{kind} (часть выплаты {_rub(whole):,.2f} ₽)".replace(",", " ")
            ws.cell(row=row, column=21, value=note)
        if not parts:
            ws["S2"] = "Выплат за этот месяц пока не было"

        done = True
        for top, kind in ((1, COST), (4, PROFIT)):
            due = max(amounts[kind], 0)
            left = due - closed[kind]
            done = done and left == 0
            ws.cell(row=top, column=23, value=kind)
            ws.cell(row=top, column=24, value=f"{kind}: закрыто выплатами Ozon")
            ws.cell(row=top, column=25, value=f"{kind}: осталось закрыть")
            _money(ws.cell(row=top + 1, column=23), amounts[kind] if kind == PROFIT else due)
            _money(ws.cell(row=top + 1, column=24), closed[kind])
            _money(ws.cell(row=top + 1, column=25), left)
        if amounts[PROFIT] < 0:
            ws["W6"] = "Месяц убыточный: закрывать по прибыли нечего"

        ws["W7"] = "Ozon рассчитался за месяц"
        ws["X7"] = "Да" if done else "Нет"
        ws["W9"] = "Всего выплачено Ozon с начала магазина"
        _money(ws["X9"], total_paid)
        if unallocated:
            ws["W10"] = "Выплачено сверх начисленного по отчётам"
            _money(ws["X10"], unallocated)
        save_workbook_atomic(workbook, path)
    finally:
        workbook.close()


def update_reports(reports_dir, today: date | None = None, fetch=get_balance_report) -> dict:
    """Пересчитывает сверку выплат во всех месячных отчётах Ozon. Возвращает итоги за всё время."""
    reports_dir = Path(reports_dir)
    today = today or date.today()
    reports = find_reports(reports_dir)
    if not reports:
        return {}
    obligations = [(month, read_obligations(path)) for month, path in reports]
    cache_path = reports_dir.parent / BALANCE_REPORTS_DIR / "payouts.json"
    payouts = load_payouts(reports[0][0], today, cache_path, fetch)
    parts, closed, unallocated = allocate(obligations, payouts)
    total_paid = sum(amount for _, amount in payouts)
    for (month, path), (_, amounts) in zip(reports, obligations):
        write_settlement(path, amounts, parts[month], closed[month], total_paid, unallocated)
    open_months = [m for m, a in obligations
                   if closed[m][COST] < a[COST] or closed[m][PROFIT] < max(a[PROFIT], 0)]
    return {
        "paid": _rub(total_paid),
        "cost": _rub(sum(a[COST] for _, a in obligations)),
        "profit": _rub(sum(a[PROFIT] for _, a in obligations)),
        "unallocated": _rub(unallocated),
        "first_open_month": open_months[0] if open_months else None,
    }


def print_summary(summary: dict) -> None:
    if not summary:
        print("ℹ️ Месячных отчётов Ozon нет — сверять выплаты не с чем.")
        return
    print(f"💸 Выплачено Ozon с начала магазина: {summary['paid']:,.2f} ₽".replace(",", " "))
    print(f"🧾 Себестоимость за всё время: {summary['cost']:,.2f} ₽".replace(",", " "))
    print(f"📈 Чистая прибыль за всё время: {summary['profit']:,.2f} ₽".replace(",", " "))
    first_open = summary["first_open_month"]
    if first_open:
        print(f"⏳ Ozon ещё не рассчитался полностью, начиная с месяца: {MONTHS[first_open.month - 1]} {first_open.year}")
    else:
        print("✅ Ozon рассчитался по всем месяцам")


def main() -> None:
    reports_dir = Path(os.path.dirname(os.path.abspath(__file__))).parent / "reports"
    print("💸 Сверяем выплаты Ozon с отчётами...")
    print_summary(update_reports(reports_dir))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"❌ {exc}")
        sys.exit(1)
