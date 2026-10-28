"""Monthly Excel export for the WB business summary, mirroring Monthly_sales_report.py's layout.

The per-order sheet with the indicators block (the layout below) is the data sheet: it ends up hidden as «Данные»,
and wb_report_design.apply draws the visible «Сводка» and «Заказы» on top, in the style of the Ozon report.
"""
from collections import Counter
from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from scripts.file_io import atomic_output_path, excel_writer
from scripts import wb_finance as finance
from scripts.report_paths import MONTHS, wb_report_path
from scripts import wb_report_design, wb_settlement
from scripts.file_io import save_workbook_atomic
from scripts.report_paths import wb_report_files

COLUMNS = ["Статус", "Номер заказа", "Название товара", "Артикул",
           "Количество шт.", "Цена продажи", "Сумма начисления",
           "Себестоимость", "Прибыль", "Дата",
           "Причина возврата (WB)", "Статус возврата (WB)"]


def report_path(root, year, month):
    return wb_report_path(root, year, month, create=True)


def _number(value):
    return float(value) if isinstance(value, Decimal) else value


def build(root, data, costs, returns_by_srid=None, acceptance=None, storage=None, funnel=None):
    """Write the per-order sheet plus a business-indicators block, matching summarize()'s totals.

    returns_by_srid (optional): srid -> WB Returns and Item Movement Report record. That report
    is a rolling last-31-days logistics snapshot (dev.wildberries.ru order-returns), not tied to
    a calendar month, so it only enriches orders it happens to overlap with — it carries no
    monetary amount and never changes the financial totals computed by finance.summarize().

    acceptance/storage/funnel (optional): wb_acceptance.summarize()/wb_storage.summarize()/
    wb_funnel.summarize() results for the same month — per-item breakdowns written as extra
    sheets. Funnel is organic views/cart/orders/buyouts (not advertising); none of the three
    ever change the financial totals computed by finance.summarize().

    Not passed but already in the existing file of this month — kept: return reasons/statuses and the
    «Приёмка»/«Хранение»/«Воронка продаж» sheets, so rebuilding a month never loses what was loaded earlier.
    """
    output_file = report_path(root, data["year"], data["month"])
    kept_returns, kept_sheets = _saved_extras(output_file)
    returns_by_srid = returns_by_srid or kept_returns
    summary = finance.summarize(data, costs)
    rows, residual = finance.order_rows(data, costs)
    for row in rows:
        match = returns_by_srid.get(row["Номер заказа"])
        row["Причина возврата (WB)"] = (match or {}).get("reason") or ""
        row["Статус возврата (WB)"] = (match or {}).get("status") or ""
    frame = pd.DataFrame(rows, columns=COLUMNS)
    with atomic_output_path(output_file) as temporary:
        with excel_writer(temporary) as writer:
            frame.to_excel(writer, sheet_name="Заказы", index=False)
            if acceptance and acceptance.get("by_nm"):
                _by_nm_frame(acceptance["by_nm"], "Стоимость приёмки, ₽", extra_key="count",
                            extra_label="Количество, шт.").to_excel(writer, sheet_name="Приёмка", index=False)
            if storage and storage.get("by_nm"):
                _by_nm_frame(storage["by_nm"], "Стоимость хранения, ₽",
                            extra_key="vendor_code", extra_label="Артикул").to_excel(writer, sheet_name="Хранение", index=False)
            if funnel and funnel.get("by_nm"):
                _funnel_frame(funnel["by_nm"]).to_excel(writer, sheet_name="Воронка продаж", index=False)
            fresh = {"Приёмка": acceptance, "Хранение": storage, "Воронка продаж": funnel}
            for name, kept in kept_sheets.items():
                if not (fresh[name] and fresh[name].get("by_nm")):
                    kept.to_excel(writer, sheet_name=name, index=False)
        _write_indicators(temporary, summary, residual, returns_by_srid, acceptance, storage, funnel)
        workbook = load_workbook(temporary)
        try:
            wb_report_design.apply(workbook, f"{MONTHS[data['month'] - 1]} {data['year']}", data, costs,
                                   returns_by_srid, wb_settlement.settle(root, costs))
            workbook.save(temporary)
        finally:
            workbook.close()
    return output_file


def _saved_extras(path):
    """Из прежнего файла месяца: причины и статусы возвратов по srid и листы приёмки/хранения/воронки (без «Итого»)."""
    if not path.exists():
        return {}, {}
    try:
        workbook = load_workbook(path, read_only=True)
        try:
            names = workbook.sheetnames
            data_sheet = workbook["Данные" if "Данные" in names else "Заказы"]
            returns = wb_report_design.saved_returns(data_sheet)
        finally:
            workbook.close()
        sheets = {}
        for name in ("Приёмка", "Хранение", "Воронка продаж"):
            if name in names:
                frame = pd.read_excel(path, sheet_name=name)
                sheets[name] = frame[frame.iloc[:, 0].astype(str) != "Итого"]
        return returns, sheets
    except Exception:
        return {}, {}  # повреждённый или чужой файл — собираем с нуля


def _by_nm_frame(by_nm, amount_label, extra_key=None, extra_label=None):
    columns = ["nmId", "Название", amount_label] + ([extra_label] if extra_label else [])
    rows = []
    for nm_id, entry in sorted(by_nm.items(), key=lambda item: item[1]["total"], reverse=True):
        row = [nm_id, entry.get("subject", ""), _number(entry["total"])]
        if extra_key:
            row.append(entry.get(extra_key, ""))
        rows.append(row)
    return pd.DataFrame(rows, columns=columns)


def _funnel_frame(by_nm):
    columns = ["nmId", "Название", "Артикул", "Просмотры карточки", "В корзину", "Заказы",
               "Сумма заказов, ₽", "Выкупы", "Сумма выкупов, ₽", "Отмены",
               "В корзину, %", "Корзина → заказ, %", "Выкуп, %"]
    rows = []
    for entry in sorted(by_nm.values(), key=lambda item: item["views"], reverse=True):
        rows.append([entry["nmId"], entry["title"], entry["vendor_code"], entry["views"], entry["cart"],
                     entry["orders"], _number(entry["order_sum"]), entry["buyouts"], _number(entry["buyout_sum"]),
                     entry["cancels"], entry["add_to_cart_pct"], entry["cart_to_order_pct"], entry["buyout_pct"]])
    return pd.DataFrame(rows, columns=columns)


def _write_indicators(filename, summary, residual, returns_by_srid=None, acceptance=None, storage=None, funnel=None):
    returns_by_srid = returns_by_srid or {}
    workbook = load_workbook(filename)
    try:
        sheet = workbook["Заказы"]
        pairs = [
            ("Общая выручка", summary.get("revenue")),
            ("Прибыль по отчёту", summary.get("profit")),
            ("Себестоимость", summary.get("costs")),
            ("Net Margin, %", summary.get("margin")),
            ("Валовая прибыль", summary.get("gross")),
            ("Операционные расходы", summary.get("operating")),
            ("Итог WB к выплате", summary.get("payout")),
            ("Средний чек продажи", summary.get("average")),
            ("Продажи (уник. srid)", summary.get("sales")),
            ("Возвраты (уник. srid)", summary.get("returns")),
            ("Логистика", summary.get("logistics")),
            ("Хранение", summary.get("storage")),
            ("Приёмка", summary.get("acceptance")),
            ("Удержания", summary.get("deductions")),
            ("Штрафы", summary.get("penalties")),
            ("Доплаты", summary.get("additional")),
            ("Начисления вне заказов периода (см. примечание)", _number(residual)),
        ]
        for index, (label, value) in enumerate(pairs, start=1):
            sheet.cell(row=index, column=16, value=label)
            sheet.cell(row=index, column=17, value=_number(value) if value is not None else "—")
        for offset, note in enumerate(summary.get("notes", []), start=len(pairs) + 2):
            sheet.cell(row=offset, column=16, value="Внимание")
            sheet.cell(row=offset, column=17, value=note)
        note_row = len(pairs) + 2 + len(summary.get("notes", []))
        sheet.cell(row=note_row, column=16, value="Справка")
        sheet.cell(row=note_row, column=17,
                   value="«Начисления вне заказов периода» — логистика/хранение по заказам, "
                         "проданным в другом отчётном месяце; учтены в «Итог WB к выплате», "
                         "но не привязаны к строке заказа в этом отчёте.")
        row_cursor = note_row + 2
        if returns_by_srid:
            reasons = Counter(str(record.get("reason") or "не указана") for record in returns_by_srid.values())
            sheet.cell(row=row_cursor, column=16, value="Возвраты — топ причин (последние 31 день, WB Returns Report)")
            row_cursor += 1
            for reason, count in sorted(reasons.items(), key=lambda item: item[1], reverse=True)[:10]:
                sheet.cell(row=row_cursor, column=16, value=reason)
                sheet.cell(row=row_cursor, column=17, value=count)
                row_cursor += 1
            sheet.cell(row=row_cursor, column=16, value="Справка")
            sheet.cell(row=row_cursor, column=17,
                       value="Причина и статус возврата — из отдельного отчёта WB о движении товара (последние "
                             "31 день), не из финансового отчёта. Заполнены только у заказов этого месяца, "
                             "которые попали в оба отчёта; это логистические данные без суммы.")
            row_cursor += 2
        if acceptance is not None:
            sheet.cell(row=row_cursor, column=16, value="Приёмка по товарам (лист «Приёмка»), сумма")
            sheet.cell(row=row_cursor, column=17, value=_number(acceptance.get("total", 0)))
            row_cursor += 1
        if storage is not None:
            sheet.cell(row=row_cursor, column=16, value="Хранение по товарам (лист «Хранение»), сумма")
            sheet.cell(row=row_cursor, column=17, value=_number(storage.get("total", 0)))
            row_cursor += 1
        if acceptance is not None or storage is not None:
            sheet.cell(row=row_cursor, column=16, value="Справка")
            sheet.cell(row=row_cursor, column=17,
                       value="Детализация по товарам — из отдельных отчётов WB (Acceptance Expenses / Paid "
                             "Storage) за тот же период. Суммы могут немного отличаться от строк «Приёмка»/"
                             "«Хранение» выше (те берутся из финансового отчёта) — методика и момент расчёта "
                             "у WB отличаются между отчётами.")
            row_cursor += 2
        if funnel is not None:
            sheet.cell(row=row_cursor, column=16, value="Воронка продаж (лист «Воронка продаж») — просмотры")
            sheet.cell(row=row_cursor, column=17, value=funnel.get("views", 0))
            row_cursor += 1
            sheet.cell(row=row_cursor, column=16, value="Воронка продаж — в корзину")
            sheet.cell(row=row_cursor, column=17, value=funnel.get("cart", 0))
            row_cursor += 1
            sheet.cell(row=row_cursor, column=16, value="Воронка продаж — заказы")
            sheet.cell(row=row_cursor, column=17, value=funnel.get("orders", 0))
            row_cursor += 1
            sheet.cell(row=row_cursor, column=16, value="Воронка продаж — выкупы")
            sheet.cell(row=row_cursor, column=17, value=funnel.get("buyouts", 0))
            row_cursor += 1
            sheet.cell(row=row_cursor, column=16, value="Конверсия: просмотр → корзина → заказ → выкуп, %")
            sheet.cell(row=row_cursor, column=17,
                       value=f"{funnel.get('add_to_cart_pct')} / {funnel.get('cart_to_order_pct')} / {funnel.get('buyout_pct')}")
            row_cursor += 1
            sheet.cell(row=row_cursor, column=16, value="Справка")
            sheet.cell(row=row_cursor, column=17,
                       value="Воронка продаж — органический трафик (просмотры карточки, добавления в корзину, "
                             "заказы, выкупы) из отдельного отчёта WB Sales Funnel, не из рекламных инструментов "
                             "и не из финансового отчёта. Не влияет на суммы и прибыль.")
        for column, width in zip("ABCDEFGHIJKL", (18, 34, 45, 22, 12, 12, 14, 12, 12, 12, 26, 22)):
            sheet.column_dimensions[column].width = width
        sheet.column_dimensions["P"].width = 46
        sheet.column_dimensions["Q"].width = 40
        workbook.save(filename)
    finally:
        workbook.close()




def refresh_month(root, data, costs, settlement=None):
    """Перерисовать «Сводку», «Заказы» и «Расходы вне заказов» отчёта месяца без запросов к WB.

    Отчёт нового вида перерисовывается на месте: «Данные», «Приёмка», «Хранение», «Воронка продаж» и причины
    возвратов сохраняются. Отчёта нет или он старого вида (без «Данных») — собирается заново."""
    path = wb_report_path(root, data["year"], data["month"])
    if settlement is None:
        settlement = wb_settlement.settle(root, costs)
    if not path.exists():
        return build(root, data, costs)
    workbook = load_workbook(path)
    try:
        if "Данные" not in workbook.sheetnames:
            workbook.close()
            return build(root, data, costs)
        wb_report_design.apply(workbook, f"{MONTHS[data['month'] - 1]} {data['year']}", data, costs, None, settlement)
        save_workbook_atomic(workbook, path)
        return path
    finally:
        workbook.close()


def refresh_saved(root):
    """Перерисовать все сохранённые Excel-отчёты WB по wb reports/<ГГГГ-ММ>.json и свежей сверке выплат.
    Отчёт, для которого нет JSON этого токена, остаётся как есть."""
    costs = finance.read_costs(root)
    settlement = wb_settlement.settle(root, costs)
    done = []
    for path in wb_report_files(root):
        stem = path.stem.removeprefix("WB ").split()
        if len(stem) != 2 or stem[0] not in MONTHS or not stem[1].isdigit():
            continue
        data = finance.load_report(root, int(stem[1]), MONTHS.index(stem[0]) + 1)
        if data and data.get("reports"):
            done.append(refresh_month(root, data, costs, settlement))
    return done


def stale_months(root, saved):
    """Месяцы из сверки выплат, которых нет в wb reports/ или которые WB с тех пор дополнил (итог расходится > 1 ₽)."""
    result = []
    for key, total in sorted(saved["months"].items()):
        data = finance.load_report(root, int(key[:4]), int(key[5:7]))
        kept = sum(wb_settlement._kopecks(r.get("bankPaymentSum")) for r in (data or {}).get("reports", []))
        if not data or abs(kept - total) > wb_settlement.STALE_TOLERANCE:
            result.append(key)
    return result


def download_history(root, token, progress=print, cancel=None):
    """Загрузить финансовые отчёты WB за все месяцы с начала магазина (недостающие и устаревшие),
    обновить сверку выплат и собрать Excel-отчёты за каждый месяц. Финансовый API — запрос в минуту,
    поэтому месяц грузится 2–3 минуты. Возвращает пути собранных отчётов."""
    client = finance.FinanceClient(token, cancel, progress)
    try:
        saved = wb_settlement.refresh(root, client)
        months = stale_months(root, saved)
        progress(f"Нужно загрузить месяцев: {len(months)} — {', '.join(months) or 'все уже загружены'}")
        for index, key in enumerate(months, 1):
            progress(f"[{index}/{len(months)}] {key}: загрузка финансовых отчётов WB…")
            data = client.fetch(int(key[:4]), int(key[5:7]))
            finance.summarize(data, finance.read_costs(root))
            finance.ensure_costs(root, data["details"])
            finance.save_report(root, data)
    finally:
        client.session.close()
    costs = finance.read_costs(root)
    built = []
    for key in sorted(saved["months"]):
        data = finance.load_report(root, int(key[:4]), int(key[5:7]))
        if data and data.get("reports"):
            built.append(build(root, data, costs))
    return built


if __name__ == "__main__":
    import sys
    from scripts.app_paths import DATA_ROOT
    from scripts.marketplace_settings import read_settings
    if "--history" in sys.argv:
        # python -m scripts.wb_monthly_report --history — все месяцы с начала магазина
        for done in download_history(DATA_ROOT, read_settings(DATA_ROOT).get("WB_API_TOKEN") or "",
                                     progress=lambda message: print(message, flush=True)):
            print(f"✅ {done.name}", flush=True)
    else:
        for done in refresh_saved(DATA_ROOT):
            print(f"✅ {done.name}")
