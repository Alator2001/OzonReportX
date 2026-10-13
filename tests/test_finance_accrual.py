"""Regressions for the retirement of finance/transaction/list."""
from decimal import Decimal
import unittest
from unittest.mock import Mock, patch

import test_technical_fixes as technical
from test_technical_fixes import monthly


def money(amount):
    return {"amount": str(amount), "currency": "RUB"}


def accrual(number, amount, sale=None, commission=None):
    row = {"unit_number": number, "total_amount": money(amount)}
    if sale is not None:
        row["posting"] = {"products": [{"commission": {
            "seller_price": money(sale), "sale_commission": money(commission),
            "sale_amount": money(99999), "bonus": money(888),
        }, "delivery": {"total_accrued": money(-100)}}]}
    return row


class AccrualTests(unittest.TestCase):
    setUp = technical.OfflineTests.setUp

    def session(self, *pages):
        session = Mock()
        session.post.side_effect = [Mock(json=Mock(return_value=p)) for p in pages]
        return session

    def test_days_and_cursors_preserve_sales_returns_and_signed_totals(self):
        session = self.session(
            {"accruals": [accrual("order", "674.82", 1640, "-852.8")], "last_id": "page2"},
            {"accruals": [accrual("order", "-24.6"), accrual("other", 7)], "last_id": ""},
            {"accruals": [accrual("order", "-787.2", -1640, "852.8"),
                          accrual("", 123)], "last_id": ""},
        )
        result = monthly.get_transactions_by_posting("2026-09-01T00:00:00Z", "2026-09-02T23:59:59Z", session)
        self.assertEqual(sum(x["amount"] for x in result["order"]), Decimal("-136.98"))
        self.assertEqual(sum(x["accruals_for_sale"] for x in result["order"]), 0)
        self.assertEqual(sum(x["sale_commission"] for x in result["order"]), 0)
        self.assertEqual(set(result), {"order", "other"})
        calls = session.post.call_args_list
        self.assertEqual([c.kwargs["json"] for c in calls], [
            {"date": "2026-09-01", "last_id": ""},
            {"date": "2026-09-01", "last_id": "page2"},
            {"date": "2026-09-02", "last_id": ""},
        ])
        self.assertTrue(all(c.args[0].endswith("/v1/finance/accrual/by-day") for c in calls))

    def test_repeated_cursor_or_malformed_response_fails(self):
        for page in ({}, {"accruals": [], "last_id": "again"},
                     {"accruals": [accrual("order", "NaN")], "last_id": ""}):
            with self.subTest(page=page):
                session = self.session({"accruals": [], "last_id": "again"}, page)
                with self.assertRaises(RuntimeError):
                    monthly.get_transactions_by_posting("2026-09-01", "2026-09-01", session)

    def test_empty_day_is_a_valid_empty_result(self):
        session = self.session({"accruals": [], "last_id": ""})
        self.assertEqual(monthly.get_transactions("absent", "2026-09-01", "2026-09-01", session), [])

    def test_report_fetches_finances_once_and_keeps_formulas(self):
        postings = [{"posting_number": number, "status": "delivered", "__schema": "FBO",
                     "products": [{"offer_id": "sku", "name": "product", "quantity": 1}]}
                    for number in ("a", "b")]
        finances = {number: [{"amount": Decimal(650), "sale_commission": Decimal(-200),
                              "accruals_for_sale": Decimal(1000), "has_commission_data": True}]
                    for number in ("a", "b")}
        with patch.object(monthly, "get_transactions_by_posting", return_value=finances) as fetch, \
                patch.object(monthly, "load_cost_map", return_value={"sku": 100}), \
                patch.object(monthly, "_safe_save_excel", return_value="unused") as save:
            monthly.to_excel(postings, "2026-09-01", "2026-09-02", 9, 2026,
                             output_file="unused", session=Mock())
        fetch.assert_called_once()
        row = save.call_args.args[0].iloc[0]
        self.assertEqual(row["Цена продажи"], 1000)
        self.assertEqual(row["Комиссия за продажу Ozon"], -200)
        # Логистика — расход, со знаком минус, как комиссия: 1000 − 200 − 150 = 650 начислено.
        self.assertEqual(row["Логистика (Включает операционные ошибки продавца)"], -150)
        self.assertEqual(row["Прибыль"], 550)

    def test_charges_outside_postings_are_collected_not_attached(self):
        def non_item(number, amount, type_id):
            row = {"unit_number": number, "total_amount": money(amount), "accrued_category": "NON_ITEM",
                   "date": "2026-09-07", "non_item_fee": {"type_id": type_id}}
            if not number:
                row.pop("unit_number")
            return row
        session = self.session({"accruals": [
            accrual("0115-1", "650", 1000, "-200"),
            non_item("2000064870634", "-550.8", 12),          # кросс-докинг по номеру поставки
            non_item(None, "-71.71", 18),                       # ранний вывод — без номера
            non_item("0115-1", "-29.7", 93),                    # штраф по отправлению — остаётся у заказа
        ], "last_id": ""})
        charges = []
        result = monthly.get_transactions_by_posting("2026-09-07", "2026-09-07", session, account_charges=charges)
        self.assertEqual(sum(x["amount"] for x in result["0115-1"]), Decimal("620.3"))
        self.assertNotIn("2000064870634", result)
        self.assertEqual([(c["type_id"], c["amount"]) for c in charges], [(12, Decimal("-550.8")), (18, Decimal("-71.71"))])

    def test_profit_subtracts_charges_outside_orders_and_revenue_counts_only_sold(self):
        from openpyxl import Workbook, load_workbook
        path = self.directory / "report.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Заказы"
        ws.append(["Статус", "Номер", "Название", "Артикул", "Кол-во", "Цена продажи", "Комиссия", "Логистика", "Начислено", "Себестоимость", "Прибыль"])
        ws.append(["delivered", "a", "", 1, 1, 1000, -200, -150, 650, -100, 550])
        ws.append(["cancelled", "b", "", 1, 1, 900, "-", "-", -20, 0, -20])
        ws.append(["delivering", "c", "", 1, 1, 800, "-", "-", 0, 0, "-"])
        wb.save(path)
        monthly.write_account_charges_sheet(str(path), [
            {"date": "2026-09-07", "type_id": 12, "unit_number": "2000064870634", "amount": Decimal("-550.8")},
            {"date": "2026-09-15", "type_id": 18, "unit_number": "", "amount": Decimal("-71.71")},
            {"date": "2026-09-17", "type_id": 41, "unit_number": "42244585", "amount": Decimal("-62.52")},  # реклама — учтена отдельно
        ])
        monthly.calc_business_indicators(str(path), ozon_promotion_cost_override=0, external_marketing_cost_override=0)
        result = load_workbook(path, data_only=True)
        sheet = result["Заказы"]
        metrics = {sheet.cell(row=r, column=16).value: sheet.cell(row=r, column=17).value for r in range(1, 40)}
        self.assertEqual(metrics["Общая выручка"], 1000)                     # без отмены и заказа в пути
        self.assertEqual(metrics["Средний чек"], 1000)
        self.assertAlmostEqual(metrics["Чистая прибыль"], 550 - 20 - 550.8 - 71.71)
        self.assertAlmostEqual(metrics["Расходы Ozon вне заказов"], 622.51)
        self.assertEqual(metrics["Валовая прибыль"], 900)
        self.assertEqual(result["Расходы вне заказов"].max_row, 4)

    def test_accrual_window_looks_ahead_21_days_past_month_end(self):
        """Real data: the sale-with-commission accrual lands 2-19 days after shipment (median ~6),
        not just for orders shipped at month-end — postings stay scoped to the calendar month, but
        accruals must be fetched with a lookahead so this lag doesn't show up as "ожидает расчёта"."""
        postings = [{"posting_number": "a", "status": "delivered", "__schema": "FBO",
                     "products": [{"offer_id": "sku", "quantity": 1}]}]
        with patch.object(monthly, "get_transactions_by_posting", return_value={}) as fetch, \
                patch.object(monthly, "load_cost_map", return_value={}), \
                patch.object(monthly, "_safe_save_excel", return_value="unused"):
            monthly.to_excel(postings, "2026-08-01T00:00:00Z", "2026-08-31T23:59:59Z", 8, 2026,
                             output_file="unused", session=Mock())
        fetch.assert_called_once()
        called_date_from, called_date_to = fetch.call_args.args[0], fetch.call_args.args[1]
        self.assertEqual(called_date_from, "2026-08-01T00:00:00Z")
        self.assertEqual(called_date_to, "2026-09-21T23:59:59+00:00")
