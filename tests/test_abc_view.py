"""Offline checks for reading ABC/XYZ reports into the app screen."""
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import abc_view as av  # noqa: E402


def build_report(path: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Итог"
    summary.append(["Артикул", "Оценка по ABC", "Оценка по XYZ", "Общая оценка ABCXYZ"])
    summary.append([1102, "A", "X", "AX"])
    summary.append([1307, "C", "Y", "CY"])
    summary.append([1201, "B", "Недостаточно данных", "Недостаточно данных"])
    summary.append([None, None, None, None])
    summary.append(["Код", "Интерпретация", "Управленческие рекомендации"])
    summary.append(["AX", "Опора бизнеса\nВысокий оборот", "• Высокий уровень сервиса\n• Минимальные страховые запасы"])
    abc = wb.create_sheet("ABC")
    abc.append(["Артикул", "Сумма по полю Прибыль", "Доля в общей прибыли", "Накопительная доля", "Оценка"])
    abc.append([1102, 800.0, 0.8, 0.8, "A"])
    abc.append([1201, 150.0, 0.15, 0.95, "B"])
    abc.append([1307, 50.0, 0.05, 1.0, "C"])
    abc.append(["1114, 1102", 0.0, 0.0, 1.0, "C"])
    abc.append(["Общая прибыль", 1000.0, None, None, None])
    xyz = wb.create_sheet("XYZ")
    xyz.append(["Нерегулярный спрос"])
    xyz.append(["Артикул", "Июль 2026", "Июнь 2026", "Общий итог", "Среднее значение", "Стандартное отклонение", "Коэффициент вариации", "Оценка"])
    xyz.append([1307, 2, 1, 3, 1.5, 0.5, 0.33, "Y"])
    xyz.append([])
    xyz.append(["Сглаженный спрос"])
    xyz.append(["Артикул", "Июль 2026", "Винз. Июль 2026", "Июнь 2026", "Винз. Июнь 2026", "Общий итог", "Медиана",
                "Среднее значение", "Стандартное отклонение", "Коэффициент вариации", "Оценка"])
    xyz.append([1102, 6, None, 5, None, 11, 5.5, 5.5, 0.5, 0.09, "X"])
    wb.save(path)


class AbcViewTests(unittest.TestCase):
    def test_reads_items_months_and_legend(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "Июнь 2026-Июль 2026.xlsx"
            build_report(path)
            report = av.load_abc_report(path)
        self.assertEqual(report["period"], "Июнь 2026 — Июль 2026")
        self.assertEqual(report["months"], ["Июнь 2026", "Июль 2026"])
        self.assertEqual([item["article"] for item in report["items"]], ["1102", "1201", "1307", "1114, 1102"])   # по прибыли, без строки итога
        self.assertAlmostEqual(report["total_profit"], 1000.0)
        top = report["items"][0]
        self.assertEqual((top["code"], top["sales"], top["total_qty"]), ("AX", [5.0, 6.0], 11.0))
        self.assertEqual(top["title"], "Опора бизнеса")
        self.assertEqual(top["recommendations"], ["Высокий уровень сервиса", "Минимальные страховые запасы"])
        by_article = {item["article"]: item for item in report["items"]}
        self.assertEqual(by_article["1307"]["title"], "Низкий оборот, умеренные колебания спроса")   # кода нет в легенде
        self.assertFalse(by_article["1201"]["enough_data"])
        self.assertIsNone(by_article["1201"]["code"])

    def test_matrix_groups(self):
        items = [{"article": "1", "abc": "A", "xyz": "X"}, {"article": "2", "abc": "A", "xyz": "X"},
                 {"article": "3", "abc": "C", "xyz": "Недостаточно данных"}]
        self.assertEqual(av.matrix_counts(items), {("A", "X"): ["1", "2"], ("C", "—"): ["3"]})


if __name__ == "__main__":
    unittest.main()
