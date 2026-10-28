"""WB payout settlement: withdrawn = everything WB accrued − current balance, closing months oldest first."""
import json
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from scripts import wb_finance as finance  # noqa: E402
from scripts import wb_monthly_report, wb_settlement as ws  # noqa: E402

TOKEN = "token"


def listed(day, amount):
    return {"reportId": hash(day) & 0xFFFFFFF, "currency": "RUB", "dateFrom": day, "dateTo": day,
            "bankPaymentSum": str(amount)}


def month_data(year, month, payout, cost_nm="1", qty=1):
    report = {"reportId": 1, "currency": "RUB", "dateFrom": f"{year}-{month:02d}-05", "dateTo": f"{year}-{month:02d}-05",
              "retailAmountSum": str(payout), "bankPaymentSum": str(payout), "deliveryServiceSum": "0",
              "paidStorageSum": "0", "paidAcceptanceSum": "0", "deductionSum": "0", "penaltySum": "0",
              "additionalPaymentSum": "0"}
    sale = {"reportId": 1, "rrdId": 1, "currency": "RUB", "sellerOperName": "Продажа", "srid": f"s{month}",
            "nmId": int(cost_nm), "vendorCode": "A", "quantity": qty, "retailAmount": str(payout),
            "forPay": str(payout), "saleDt": f"{year}-{month:02d}-05"}
    return {"version": 1, "year": year, "month": month, "account": finance.identity(TOKEN),
            "fetched_at": "x", "reports": [report], "details": [sale]}


class SettlementTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        (self.root / ".env").write_text(f"WB_API_TOKEN={TOKEN}\n", encoding="utf-8")

    def client(self, reports, balance, for_withdraw=0):
        client = Mock(token=TOKEN)
        client.list_reports.return_value = reports
        client.balance.return_value = {"current": Decimal(str(balance)), "for_withdraw": Decimal(str(for_withdraw))}
        return client

    def save(self, *months):
        for data in months:
            finance.save_report(self.root, data)

    def test_refresh_sums_reports_by_month_and_keeps_balance(self):
        client = self.client([listed("2026-08-03", "100.10"), listed("2026-08-20", 50), listed("2026-09-01", -5)], 30, 20)
        saved = ws.refresh(self.root, client, date(2026, 10, 27))
        client.list_reports.assert_called_once_with(ws.HISTORY_START, "2026-10-27")
        self.assertEqual(saved["months"], {"2026-08": 15010, "2026-09": -500})
        self.assertEqual((saved["balance"], saved["for_withdraw"]), (3000, 2000))
        self.assertEqual(ws.load(self.root)["months"], saved["months"])

    def test_other_token_settlement_is_ignored(self):
        ws.refresh(self.root, self.client([listed("2026-08-03", 100)], 0))
        (self.root / ".env").write_text("WB_API_TOKEN=other\n", encoding="utf-8")
        self.assertIsNone(ws.load(self.root))
        self.assertIsNone(ws.settle(self.root, {}))

    def test_withdrawn_closes_oldest_month_cost_then_profit(self):
        self.save(month_data(2026, 8, 1000), month_data(2026, 9, 600))
        # Начислено 1600, на балансе 900 → выведено 700: август себестоимость 400 + прибыль 300 из 600.
        ws.refresh(self.root, self.client([listed("2026-08-05", 1000), listed("2026-09-05", 600)], 900))
        result = ws.settle(self.root, {"1": Decimal(400)})
        aug, sep = result["months"]["2026-08"], result["months"]["2026-09"]
        self.assertEqual(result["withdrawn"], 70000)
        self.assertEqual(aug["due"], {ws.COST: 40000, ws.PROFIT: 60000})
        self.assertEqual(aug["closed"], {ws.COST: 40000, ws.PROFIT: 30000})
        self.assertEqual((aug["left"], aug["done"]), (30000, False))
        self.assertEqual((sep["closed"], sep["left"]), ({ws.COST: 0, ws.PROFIT: 0}, 60000))

    def test_month_without_details_is_one_accrued_sum(self):
        ws.refresh(self.root, self.client([listed("2025-09-05", 500)], 0))
        month = ws.settle(self.root, {})["months"]["2025-09"]
        self.assertEqual((month["due"], month["done"]), ({ws.ACCRUED: 50000}, True))
        self.assertIn("не загружен", month["reason"])

    def test_stale_month_is_not_split(self):
        self.save(month_data(2026, 10, 1000))
        ws.refresh(self.root, self.client([listed("2026-10-05", 1000), listed("2026-10-26", 300)], 1300))
        month = ws.settle(self.root, {"1": Decimal(400)})["months"]["2026-10"]
        self.assertEqual(month["due"], {ws.ACCRUED: 130000})
        self.assertIn("загрузите месяц заново", month["reason"])

    def test_loss_month_owes_only_what_was_accrued(self):
        self.save(month_data(2026, 8, 300))
        ws.refresh(self.root, self.client([listed("2026-08-05", 300)], 0))
        month = ws.settle(self.root, {"1": Decimal(400)})["months"]["2026-08"]
        self.assertEqual(month["due"], {ws.COST: 30000, ws.PROFIT: 0})
        self.assertEqual((month["profit"], month["done"]), (-10000, True))

    def test_negative_month_counts_as_paid_against_earlier_months(self):
        # Август +1000, сентябрь −100 (только хранение), баланс 0: выведено 900, но август закрыт целиком.
        ws.refresh(self.root, self.client([listed("2026-08-05", 1000), listed("2026-09-05", -100)], 0))
        result = ws.settle(self.root, {})
        self.assertEqual(result["withdrawn"], 90000)
        self.assertTrue(result["months"]["2026-08"]["done"])
        self.assertEqual(result["months"]["2026-09"]["due"], {})


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        (self.root / ".env").write_text(f"WB_API_TOKEN={TOKEN}\n", encoding="utf-8")

    def texts(self, path):
        return [c.value for r in load_workbook(path)["Сводка"].iter_rows() for c in r if isinstance(c.value, str)]

    def test_summary_shows_settlement_and_refresh_keeps_extra_sheets(self):
        data = month_data(2026, 8, 1000)
        finance.save_report(self.root, data)
        storage = {"total": Decimal(5), "by_nm": {"1": {"total": Decimal(5), "subject": "Доски", "vendor_code": "A"}}}
        path = wb_monthly_report.build(self.root, data, {"1": Decimal(400)}, {"s8": {"reason": "Брак"}}, storage=storage)
        self.assertTrue(any(t.startswith("Сверка выплат ещё не проводилась") for t in self.texts(path)))

        client = Mock(token=TOKEN)
        client.list_reports.return_value = [listed("2026-08-05", 1000)]
        client.balance.return_value = {"current": Decimal(0), "for_withdraw": Decimal(0)}
        ws.refresh(self.root, client)
        (self.root / "wb_costs.xlsx").unlink(missing_ok=True)
        finance.ensure_costs(self.root, data["details"])
        book = load_workbook(self.root / "wb_costs.xlsx")
        book["Себестоимость WB"]["D2"] = 400
        book.save(self.root / "wb_costs.xlsx")

        self.assertEqual(wb_monthly_report.refresh_saved(self.root), [path])
        book = load_workbook(path)
        self.assertIn("Хранение", book.sheetnames)
        self.assertEqual(book["Заказы"]["P2"].value, "Причина возврата: Брак")
        self.assertTrue(any(t.startswith("✔  WB полностью рассчитался") for t in self.texts(path)))
        self.assertEqual(book["Сводка"]["G5"].value, "WB ещё должен")


class FinanceClientTests(unittest.TestCase):
    def test_list_reports_pages_and_rejects_repeats(self):
        client = finance.FinanceClient(TOKEN)
        page = [listed("2026-08-05", 1), listed("2026-08-06", 2)]
        with unittest.mock.patch.object(client, "request", side_effect=[page]) as request:
            self.assertEqual(client.list_reports("2024-01-29", "2026-10-27"), page)
        self.assertEqual(request.call_args[0][1]["dateTo"], "2026-10-27T23:59:59")
        with unittest.mock.patch.object(client, "request", side_effect=[page + page[:1]]):
            with self.assertRaises(finance.FinanceError):
                client.list_reports("2024-01-29", "2026-10-27")

    def test_balance(self):
        client = finance.FinanceClient(TOKEN, session=Mock())
        client.session.get.return_value = Mock(status_code=200, json=lambda: {"current": 34399.61, "for_withdraw": 25458.9})
        self.assertEqual(client.balance(), {"current": Decimal("34399.61"), "for_withdraw": Decimal("25458.9")})
        client.session.get.return_value = Mock(status_code=403)
        with self.assertRaises(finance.FinanceError):
            client.balance()


if __name__ == "__main__":
    unittest.main()


class HistoryTests(unittest.TestCase):
    def test_downloads_only_missing_and_stale_months_and_builds_reports(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / ".env").write_text(f"WB_API_TOKEN={TOKEN}\n", encoding="utf-8")
            finance.save_report(root, month_data(2026, 8, 1000))      # август загружен и совпадает
            finance.save_report(root, month_data(2026, 9, 500))       # сентябрь WB с тех пор дополнил
            client = Mock(token=TOKEN, session=Mock())
            client.list_reports.return_value = [listed("2026-07-05", 300), listed("2026-08-05", 1000),
                                                listed("2026-09-05", 500), listed("2026-09-20", 100)]
            client.balance.return_value = {"current": Decimal(0), "for_withdraw": Decimal(0)}
            client.fetch.side_effect = lambda year, month: month_data(year, month, {7: 300, 9: 600}[month])
            with unittest.mock.patch.object(finance, "FinanceClient", return_value=client):
                built = wb_monthly_report.download_history(root, TOKEN, progress=lambda _m: None)
            self.assertEqual([c.args for c in client.fetch.call_args_list], [(2026, 7), (2026, 9)])
            self.assertEqual([p.name for p in built], ["WB Июль 2026.xlsx", "WB Август 2026.xlsx", "WB Сентябрь 2026.xlsx"])
            self.assertTrue(all(p.parent == root / "reports" / "WB" / "2026" for p in built))
            self.assertEqual(wb_monthly_report.stale_months(root, ws.load(root)), [])
