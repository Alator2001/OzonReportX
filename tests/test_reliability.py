"""Offline tests for pacing, inter-process locking and preserving user data."""
from collections import deque
from datetime import datetime, timedelta, timezone, date
from email.utils import format_datetime
import importlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import unittest
from unittest.mock import Mock, patch

import pandas as pd
import requests
from openpyxl import load_workbook
import test_technical_fixes as technical
from test_technical_fixes import ROOT, ui_function, monthly, balance, pricing, recommended, ai_chat, file_io
from file_lock import exclusive_file, FileBusyError, file_signature
from report_http import ReportSession, retry_after_seconds
from ui_log import SessionLog
import fbo_supply_report as fbo
import ABC_XYZ_analytics_report as abc


class ReliabilityTests(unittest.TestCase):
    setUp = technical.OfflineTests.setUp

    def response(self, status, data=None, retry_after=None):
        return Mock(status_code=status, headers={} if retry_after is None else {"Retry-After": retry_after},
                    json=Mock(return_value=data))

    def test_429_waits_then_retries_same_payload(self):
        clock = {"now": 0.0}
        session = ReportSession()
        responses = [self.response(429, retry_after="12"), self.response(200, {"result": []})]
        payload = {"offset": 400}
        with patch.object(requests.Session, "request", side_effect=responses) as request, \
                patch("report_http.time.monotonic", side_effect=lambda: clock["now"]), \
                patch("report_http.time.sleep", side_effect=lambda seconds: clock.__setitem__("now", clock["now"] + seconds)):
            result = session.post("https://api-seller.ozon.ru/v3/posting/fbo/list", json=payload)
        self.assertEqual(result.status_code, 200)
        self.assertGreaterEqual(clock["now"], 12)
        self.assertEqual(request.call_count, 2)
        self.assertEqual(request.call_args_list[0], request.call_args_list[1])
        responses[0].close.assert_called_once()
        session.close()

    def test_repeated_429_stops_without_a_partial_success(self):
        session = ReportSession(max_retries=2)
        with patch.object(requests.Session, "request", return_value=self.response(429)) as request, \
                patch("report_http.time.sleep"):
            with self.assertRaisesRegex(RuntimeError, "429"):
                session.post("https://api-seller.ozon.ru/v3/posting/fbo/list", json={"offset": 400})
        self.assertEqual(request.call_count, 3)
        session.close()

    def test_long_retry_after_does_not_retry_early(self):
        session = ReportSession(max_wait=30)
        with patch.object(requests.Session, "request", return_value=self.response(429, retry_after="3600")) as request, \
                patch("report_http.time.sleep") as sleep:
            with self.assertRaises(RuntimeError):
                session.get("https://api-seller.ozon.ru/read")
        request.assert_called_once()
        sleep.assert_not_called()
        session.close()

    def test_retry_after_accepts_http_date_and_invalid_header(self):
        header = format_datetime(datetime.now(timezone.utc) + timedelta(seconds=30), usegmt=True)
        self.assertGreater(retry_after_seconds(header), 28)
        self.assertIsNone(retry_after_seconds("invalid"))

    def test_successful_requests_are_spaced(self):
        clock = {"now": 0.0}
        session = ReportSession(min_interval=0.5)
        with patch.object(requests.Session, "request", return_value=self.response(200)), \
                patch("report_http.time.monotonic", side_effect=lambda: clock["now"]), \
                patch("report_http.time.sleep", side_effect=lambda seconds: clock.__setitem__("now", clock["now"] + seconds)):
            session.get("https://api-seller.ozon.ru/read")
            session.get("https://api-seller.ozon.ru/read")
        self.assertGreaterEqual(clock["now"], 0.5)
        session.close()

    def test_adapter_does_not_hide_429_from_pacing_layer(self):
        with monthly.create_session() as session:
            retry = session.get_adapter("https://api-seller.ozon.ru/v3/posting/fbo/list").max_retries
            self.assertNotIn(429, retry.status_forcelist)
            self.assertFalse(retry.respect_retry_after_header)

    def test_empty_statuses_do_not_trigger_speculative_pages(self):
        with patch.object(monthly, "_fetch_fbo_page", return_value={"postings": [], "has_next": False}) as fetch:
            self.assertEqual(monthly.get_fbo_orders("a", "b", session=Mock()), [])
        self.assertEqual(fetch.call_count, 1)
        self.assertEqual(fetch.call_args.args[-1], "")

    def test_lock_excludes_other_process_and_releases_after_exception(self):
        target = self.directory / "costs.xlsx"
        code = """import sys
from scripts.file_lock import exclusive_file, FileBusyError
try:
    with exclusive_file(sys.argv[1]):
        print('acquired')
except FileBusyError:
    print('busy')
"""
        def child():
            return subprocess.run([sys.executable, "-c", code, str(target)], cwd=ROOT,
                                  capture_output=True, text=True, timeout=10, check=True).stdout.strip()
        with self.assertRaises(ValueError):
            with exclusive_file(target):
                with exclusive_file(target):
                    self.assertEqual(child(), "busy")
                raise ValueError("operation failed")
        self.assertEqual(child(), "acquired")

    def test_process_exit_does_not_leave_a_stale_lock(self):
        target = self.directory / "costs.xlsx"
        code = "import os,sys; from scripts.file_lock import exclusive_file; lock=exclusive_file(sys.argv[1]); lock.__enter__(); os._exit(0)"
        subprocess.run([sys.executable, "-c", code, str(target)], cwd=ROOT, timeout=10, check=True)
        with exclusive_file(target):
            pass

    def test_lock_registry_is_shared_by_both_import_styles(self):
        self.assertIs(importlib.import_module("file_lock"), importlib.import_module("scripts.file_lock"))

    def test_stale_editor_cannot_overwrite_a_newer_file(self):
        target = self.directory / "costs.xlsx"
        target.write_bytes(b"original")
        previous = file_signature(target)
        target.write_bytes(b"new content from another instance")
        writer = Mock()
        function = ui_function("write_costs_main_rows", {
            "exclusive_file": exclusive_file, "file_signature": file_signature,
            "_write_costs_main_rows": writer,
        })
        self.assertIn("изменился", function(target, [], expected_signature=previous))
        writer.assert_not_called()
        self.assertEqual(target.read_bytes(), b"new content from another instance")

    def test_automatic_refresh_keeps_unsaved_editors(self):
        function = ui_function("refresh_pricing_costs_preview", {"pricing_costs_dirty": {"value": True}})
        self.assertIsNone(function())  # No read, clearing, or control replacement is attempted.

    def test_manual_refresh_requires_explicit_discard(self):
        dialog = Mock(open=False)
        refresh = Mock()
        function = ui_function("handle_pricing_costs_refresh", {"pricing_costs_dirty": {"value": True},
            "discard_costs_dialog": dialog, "page": Mock(), "refresh_pricing_costs_preview": refresh})
        function()
        self.assertTrue(dialog.open)
        refresh.assert_not_called()

    def test_log_preview_is_bounded_and_full_history_survives_clear(self):
        log = SessionLog(self.directory / "session.log", max_lines=5)
        for number in range(20):
            log.append(f"line-{number}")
        self.assertEqual(list(log.lines), [f"line-{i}" for i in range(15, 20)])
        log.reset("cleared")
        self.assertEqual(list(log.lines), ["cleared"])
        self.assertIn("line-0\n", log.full_text())
        self.assertIn("line-19\n", log.full_text())
        before = log.full_text()
        log.reset("display only", persist=False)
        self.assertEqual(log.full_text(), before)

    def test_ui_log_does_not_retheme_history_per_line(self):
        store = SessionLog(self.directory / "session.log")
        log = Mock(controls=[])
        selectable = Mock(value="")
        theme = Mock()
        function = ui_function("append_log", {"log_store": store, "log": log,
            "log_selectable": selectable, "_log_line_control": lambda text: text, "_sync_log_theme": theme})
        for number in range(600):
            function(f"line-{number}")
        self.assertEqual(len(log.controls), 500)
        self.assertEqual(len(store.lines), 500)
        theme.assert_not_called()
        self.assertIn("line-0\n", store.full_text())

    def test_report_readers_close_books_on_early_return(self):
        book = Mock(sheetnames=[])
        target = self.directory / "report.xlsx"
        target.touch()
        with patch.object(ai_chat, "load_workbook", return_value=book):
            self.assertIsNone(ai_chat.load_report_summary(target))
        book.close.assert_called_once()
        book.reset_mock()
        with patch.object(recommended, "load_workbook", return_value=book):
            with self.assertRaises(ValueError):
                recommended.load_rates_from_report(target)
        book.close.assert_called_once()

    def test_abc_closes_source_when_parsing_fails(self):
        (self.directory / "Август 2026.xlsx").touch()
        source = Mock(sheet_names=["Заказы"])
        source.parse.side_effect = ValueError("broken sheet")
        with patch.object(pd, "ExcelFile", return_value=source):
            abc.merge_folder(str(self.directory), output_path=str(self.directory / "result.xlsx"))
        source.close.assert_called_once()

    def test_balance_writer_keeps_previous_excel_on_failure(self):
        folder = self.directory / "balance reports"
        folder.mkdir()
        target = folder / "2026-08-01_to_2026-08-30.xlsx"
        target.write_bytes(b"previous")
        with patch.object(pd.DataFrame, "to_excel", side_effect=ValueError("write failed")):
            balance.save_report({}, "2026-08-01", "2026-08-30", self.directory)
        self.assertEqual(target.read_bytes(), b"previous")

    def test_balance_writer_keeps_previous_json_on_failure(self):
        folder = self.directory / "balance reports"
        folder.mkdir()
        target = folder / "2026-08-01_to_2026-08-30.json"
        target.write_text('{"previous":true}', encoding="utf-8")
        with patch.object(json, "dump", side_effect=ValueError("write failed")):
            with self.assertRaises(ValueError):
                balance.save_report({}, "2026-08-01", "2026-08-30", self.directory)
        self.assertEqual(json.loads(target.read_text()), {"previous": True})

    def test_abc_writer_keeps_previous_report_on_failure(self):
        source = self.directory / "reports"
        source.mkdir()
        pd.DataFrame({"Статус": ["delivered"], "Артикул": [100], "Цена продажи": [100],
            "Количество шт.": [1], "Прибыль": [20], "Дата отгрузки": ["2026-08-01"]}).to_excel(source / "Август 2026.xlsx", sheet_name="Заказы", index=False)
        target = self.directory / "result.xlsx"
        target.write_bytes(b"previous")
        with patch.object(pd.DataFrame, "to_excel", side_effect=ValueError("write failed")):
            with self.assertRaises(Exception):
                abc.merge_folder(str(source), output_path=str(target))
        self.assertEqual(target.read_bytes(), b"previous")

    def test_fbo_writer_keeps_previous_report_on_failure(self):
        abc_path = self.directory / "abc.xlsx"
        abc_path.touch()
        folder = self.directory / fbo.STOCKS_REPORTS_DIR
        folder.mkdir()
        today = date.today()
        target = folder / f"Расчёт поставок {fbo.MONTHS_RU[today.month - 1]} {today.year}.xlsx"
        target.write_bytes(b"previous")
        summary = pd.DataFrame({"Артикул": [100], "Оценка по ABC": ["A"], "Оценка по XYZ": ["X"]})
        with patch.object(fbo, "_ensure_abc_xyz_report", return_value=abc_path), \
                patch.object(fbo, "_load_abc_xyz_itog_and_orders", return_value=(summary, pd.DataFrame())), \
                patch.object(fbo, "_daily_sales_90_from_orders", return_value={}), \
                patch.object(recommended, "get_fbo_stocks_by_offer_ids", return_value={"100": 2}), \
                patch.object(pd.DataFrame, "to_excel", side_effect=ValueError("write failed")):
            with self.assertRaises(ValueError):
                fbo.run(self.directory)
        self.assertEqual(target.read_bytes(), b"previous")


if __name__ == "__main__":
    unittest.main()
