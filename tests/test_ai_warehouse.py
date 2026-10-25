"""Offline checks for the AI assistant's analytics warehouse: loaders, incremental refresh, read-only access, catalog."""
import json
import os
import shutil
import sqlite3
import sys
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import ai_warehouse as wh  # noqa: E402
import search_db  # noqa: E402

TMP = ROOT / ".cache" / "ai-warehouse-tests"

ORDER_HEADER = ["Статус", "Номер заказа", "Название товара", "Артикул", "Количество шт.", "Цена продажи",
                "Комиссия за продажу Ozon", "Логистика (Включает операционные ошибки продавца)", "Сумма начисления",
                "Себестоимость", "Прибыль", "Дата отгрузки", "Схема"]


def save_book(path: Path, sheets):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        ws = book.create_sheet(title)
        for row in rows:
            ws.append(row)
    book.save(path)


def ozon_report(path: Path, orders, metrics):
    rows = [ORDER_HEADER + [None, None]]
    for order in orders:
        rows.append(list(order) + [None, None])
    # Блок показателей — в колонках P/Q (16/17), как в настоящем отчёте.
    for i, (name, value) in enumerate(metrics):
        while len(rows) <= i:
            rows.append([None] * 15)
        rows[i] = (rows[i] + [None] * 15)[:15] + [name, value]
    save_book(path, {
        "Заказы": rows,
        "Расходы вне заказов": [["Дата", "Вид расхода", "Код Ozon", "Номер (поставки/кампании)", "Сумма", "В чистой прибыли"],
                                ["2026-09-01", "Кросс-докинг", 46, "—", -100.5, "да"]],
        "Кампании": [["ID кампании", "Название кампании", "Расход за период (руб.)", "Клики", "ДРР (%)"],
                     ["42", "Браслеты", 1348.63, 331, 18.2]],
    })


def touch_later(path: Path):
    stat = path.stat()
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 2_000_000_000))


class WarehouseTestCase(unittest.TestCase):
    def setUp(self):
        shutil.rmtree(TMP, ignore_errors=True)
        TMP.mkdir(parents=True)
        self.addCleanup(shutil.rmtree, TMP, True)
        self.root = TMP
        ozon_report(self.root / "reports" / "Сентябрь 2026.xlsx", [
            ["delivered", "A-1", "Браслет", 7012, 1, 960, -460.8, -104.41, 394.79, -150, 244.79, "2026-09-19", "FBS"],
            ["cancelled", "A-2", "Доска", "1101", 2, 1375, 0, 0, 0, 0, 0, "2026-09-20", "FBO"],
            [None] * 13,
        ], [("Общая выручка", 103989), ("Чистая прибыль", 9305.14), ("Справка", "текст справки")])
        save_book(self.root / "reports" / "WB Сентябрь 2026.xlsx", {"Заказы": [
            ["Статус", "Номер заказа", "Название товара", "Артикул", "Количество шт.", "Цена продажи", "Сумма начисления",
             "Себестоимость", "Прибыль", "Дата", "Причина возврата (WB)", "Статус возврата (WB)", "Общая выручка", 22262],
            ["продажа", "eAQ", "Икона", "5004", 1, 1015, 1054.25, -250, 804.25, "2026-09-01", None, None],
        ]})
        save_book(self.root / "costs.xlsx", {
            "Основной": [["Артикул", "Себестоимость", "Минимальная цена продажи", "Текущая цена на Ozon"], [1101, 250, 1253, 1375]],
            "Акции": [["Артикул", "Минимальная цена продажи", "Бустинг 11.11", "Бустинг 11.11 — бустинг, %", "Пустая акция"],
                      [1101, 1253, 1375, 44.6, None]],
        })

    def refresh(self):
        return wh.refresh(self.root)

    def query(self, sql, *args):
        with closing(wh.connect(self.root, refresh_first=False)) as conn:
            return [tuple(row) for row in conn.execute(sql, args)]


class LoaderTests(WarehouseTestCase):
    def test_ozon_report_tables(self):
        self.refresh()
        self.assertEqual(self.query("SELECT period, artikul, quantity, profit, ship_date, scheme FROM ozon_orders ORDER BY order_number"),
                         [("2026-09", "7012", 1.0, 244.79, "2026-09-19", "FBS"), ("2026-09", "1101", 2.0, 0.0, "2026-09-20", "FBO")])
        self.assertEqual(self.query("SELECT metric, value, text FROM ozon_month_metrics ORDER BY rowid"),
                         [("Общая выручка", 103989.0, None), ("Чистая прибыль", 9305.14, None), ("Справка", None, "текст справки")])
        self.assertEqual(self.query("SELECT period, kind, amount FROM ozon_other_expenses"), [("2026-09", "Кросс-докинг", -100.5)])
        self.assertEqual(self.query("SELECT campaign_id, spend, clicks, drr_pct FROM ozon_ad_campaigns"), [("42", 1348.63, 331.0, 18.2)])

    def test_wildberries_reports_are_separate(self):
        self.refresh()
        self.assertEqual(self.query("SELECT COUNT(*) FROM ozon_orders WHERE artikul = '5004'"), [(0,)])
        self.assertEqual(self.query("SELECT period, artikul, profit FROM wb_orders"), [("2026-09", "5004", 804.25)])
        self.assertEqual(self.query("SELECT metric, value FROM wb_month_metrics"), [("Общая выручка", 22262.0)])

    def test_costs_and_promotions(self):
        self.refresh()
        self.assertEqual(self.query("SELECT artikul, cost, min_price, current_price, desired_price FROM products"),
                         [("1101", 250.0, 1253.0, 1375.0, None)])
        self.assertEqual(self.query("SELECT artikul, promotion, promo_price, boost_pct, min_price FROM product_promotions"),
                         [("1101", "Бустинг 11.11", 1375.0, 44.6, 1253.0)])

    def test_json_sources(self):
        analytics = self.root / "analytics history"
        (analytics / "snapshots").mkdir(parents=True)
        (analytics / "organic_last.json").write_text(json.dumps({
            "created": "2026-10-22T20:58:21", "periods": {"current": ["2026-09-24", "2026-10-22"]},
            "store": {"orders": 77.0, "premium": False, "ad_funnel": {"views": 13689.0}},
            "products": [{"offer_id": "1104", "sku": 18, "orders": 6.0, "locality": {"local": 1, "total": 4},
                          "actions": ["A", "B"], "price_index": {"color": "GREEN"},
                          "steps": [{"name": "Цена", "status": "ok", "text": "норма"}], "advice": ["Пополните остаток"],
                          "main_problem": "Наличие"}],
        }, ensure_ascii=False), encoding="utf-8")
        (analytics / "snapshots" / "2026-10-24.json").write_text(json.dumps({"1101": {"price": 1375.0, "stock": 3, "actions": 2}}), encoding="utf-8")
        supply = self.root / "stocks reports" / "Поставка FBO 23.10.2026.json"
        supply.parent.mkdir(parents=True)
        supply.write_text(json.dumps({
            "arrival": "2026-11-06",
            "supply": [{"article": "1308", "by_zone": {"Москва": 32, "СПб": 9}, "total": 41, "per_day": 1.05}],
            "fbs": {"rows": [{"article": "1308", "need": 24, "stock": 2, "shortage": 22}]},
        }, ensure_ascii=False), encoding="utf-8")
        log = self.root / "actions history" / "auto_actions_log.jsonl"
        log.parent.mkdir(parents=True)
        log.write_text('{"at": "2026-10-22T17:10:43", "action": "timer_update", "offer_id": "1101", "confirmed": true, "expires_at": "x"}\nне json\n', encoding="utf-8")
        self.refresh()
        self.assertEqual(self.query("SELECT artikul, orders, local_share, promotions_count, promotions, price_index_color, diagnosis, advice FROM organic_products"),
                         [("1104", 6.0, 25.0, 2, "A; B", "GREEN", "Цена [ok]: норма", "Пополните остаток")])
        self.assertEqual(self.query("SELECT metric, value FROM organic_store ORDER BY metric"),
                         [("ad_funnel_views", 13689.0), ("orders", 77.0)])
        self.assertEqual(self.query("SELECT date, artikul, price, stock, promotions_count FROM product_snapshots"),
                         [("2026-10-24", "1101", 1375.0, 3.0, 2)])
        self.assertEqual(self.query("SELECT plan_date, arrival, artikul, zone, quantity, total FROM fbo_supply_plan ORDER BY zone"),
                         [("2026-10-23", "2026-11-06", "1308", "Москва", 32.0, 41.0), ("2026-10-23", "2026-11-06", "1308", "СПб", 9.0, 41.0)])
        self.assertEqual(self.query("SELECT artikul, need, shortage FROM fbs_needs"), [("1308", 24.0, 22.0)])
        self.assertEqual(self.query("SELECT action, artikul, confirmed, details FROM auto_actions_log"),
                         [("timer_update", "1101", 1, '{"expires_at": "x"}')])

    def test_search_database_is_copied_read_only(self):
        path = self.root / "search data" / "search.db"
        with closing(search_db.connect(path)) as conn:
            search_db.upsert_products(conn, [{"sku": 18, "product_id": 5, "offer_id": "1104", "name": "Доска", "category": "Кухня"}])
            conn.commit()
        before = path.stat().st_mtime_ns
        self.refresh()
        self.assertEqual(self.query("SELECT sku, artikul, name FROM search_products"), [(18, "1104", "Доска")])
        self.assertEqual(path.stat().st_mtime_ns, before)


class RefreshTests(WarehouseTestCase):
    def test_only_changed_sources_are_rebuilt(self):
        self.assertEqual(self.refresh()["ozon_reports"], "updated")
        status = self.refresh()
        self.assertTrue(all(value == "fresh" for value in status.values()), status)
        touch_later(self.root / "costs.xlsx")
        status = self.refresh()
        self.assertEqual(status["costs"], "updated")
        self.assertEqual(status["ozon_reports"], "fresh")

    def test_deleted_file_removes_its_rows(self):
        self.refresh()
        (self.root / "reports" / "Сентябрь 2026.xlsx").unlink()
        self.refresh()
        self.assertEqual(self.query("SELECT COUNT(*) FROM ozon_orders"), [(0,)])
        self.assertEqual(self.query("SELECT COUNT(*) FROM wb_orders"), [(1,)])

    def test_broken_file_does_not_stop_other_sources_and_is_retried(self):
        self.refresh()
        (self.root / "costs.xlsx").write_bytes(b"not an xlsx")
        status = self.refresh()
        self.assertTrue(status["costs"].startswith("error:"))
        self.assertEqual(status["ozon_reports"], "fresh")
        self.assertEqual(self.query("SELECT COUNT(*) FROM products"), [(0,)])
        with closing(wh.connect(self.root, refresh_first=False)) as conn:
            catalog = {item["table"]: item for item in wh.describe(conn)}
        self.assertIn("error", catalog["products"])
        self.assertTrue(self.refresh()["costs"].startswith("error:"))  # подпись не сохранена — пробует снова

    def test_schema_version_change_rebuilds(self):
        self.refresh()
        with closing(sqlite3.connect(wh.warehouse_path(self.root))) as conn:
            conn.execute("UPDATE _meta SET value = '0' WHERE key = 'schema_version'")
            conn.commit()
        self.assertEqual(self.refresh()["ozon_reports"], "updated")

    def test_latest_abc_prefers_long_recent_report(self):
        for name, artikul in [("Сентябрь 2025-Август 2026", "1101"), ("Август 2026-Сентябрь 2026", "1102"), ("Январь 2025-Декабрь 2025", "1103")]:
            save_book(self.root / "ABC&XYZ reports" / f"{name}.xlsx", {"Итог": [
                ["Артикул", "Оценка по ABC", "Оценка по XYZ", "Общая оценка ABCXYZ"], [artikul, "A", "X", "AX"]]})
        self.refresh()
        self.assertEqual(self.query("SELECT report_period, artikul FROM abc_xyz_latest"), [("Сентябрь 2025-Август 2026", "1101")])

    def test_connection_is_read_only(self):
        with closing(wh.connect(self.root)) as conn:
            with self.assertRaises(sqlite3.OperationalError):
                conn.execute("DELETE FROM ozon_orders")

    def test_catalog_describes_every_table(self):
        with closing(wh.connect(self.root)) as conn:
            catalog = {item["table"]: item for item in wh.describe(conn)}
        self.assertEqual(set(catalog), set(wh.CATALOG) | set(wh.VIEWS))
        self.assertEqual(catalog["product_snapshot_latest"]["view_of"], "product_snapshots")
        self.assertEqual(catalog["ozon_orders"]["rows"], 2)
        self.assertEqual(catalog["ozon_orders"]["range"], ["2026-09", "2026-09"])
        self.assertIn("profit", catalog["ozon_orders"]["columns"])
        self.assertEqual(catalog["search_products"]["rows"], 0)
        sourced = {table for source in wh.SOURCES for table in source.tables}
        self.assertEqual(sourced, set(wh.CATALOG))


class HelperTests(unittest.TestCase):
    def test_period_from_name(self):
        self.assertEqual(wh.period_from_name("Сентябрь 2026"), "2026-09")
        self.assertEqual(wh.period_from_name("WB январь 2025"), "2025-01")
        self.assertIsNone(wh.period_from_name("Поставка FBO 23.10.2026"))

    def test_numbers(self):
        self.assertEqual(wh._num("1 234,5"), 1234.5)
        self.assertIsNone(wh._num("—"))
        self.assertIsNone(wh._num(True))
        self.assertEqual(wh._text(7012.0), "7012")


if __name__ == "__main__":
    unittest.main()
