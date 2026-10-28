"""WB monthly report in the Ozon report style: «Сводка» and «Заказы» on top, the old layout hidden as «Данные»."""
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from scripts import wb_finance as finance  # noqa: E402
from scripts import wb_monthly_report, wb_report_design as design  # noqa: E402


def report(**updates):
    row = {"reportId": 1, "currency": "RUB", "dateFrom": "2026-10-05", "dateTo": "2026-10-05",
           "retailAmountSum": "1004", "bankPaymentSum": "735.64", "deliveryServiceSum": "117",
           "paidStorageSum": "0", "paidAcceptanceSum": "0", "deductionSum": "0", "penaltySum": "0",
           "additionalPaymentSum": "0"}
    return dict(row, **updates)


def sale(**updates):
    # Реальная продажа 1205: покупатель заплатил 1004 при цене продавца 1440, WB перечисляет 852,64.
    row = {"reportId": 1, "rrdId": 2, "currency": "RUB", "sellerOperName": "Продажа", "srid": "s1", "nmId": 1205,
           "vendorCode": "1205", "title": "Доска", "quantity": 1, "retailAmount": "1004", "retailPriceWithDisc": "1440",
           "forPay": "852.64", "acquiringFee": "40.16", "deliveryService": "0", "saleDt": "2026-10-05T10:00:00",
           "officeName": "Склад поставщика - везу на склад WB"}
    return dict(row, **updates)


def delivery(**updates):
    row = {"reportId": 1, "rrdId": 1, "currency": "RUB", "sellerOperName": "Логистика", "srid": "s1", "nmId": 1205,
           "vendorCode": "1205", "quantity": 0, "retailAmount": "0", "forPay": "0", "deliveryService": "117"}
    return dict(row, **updates)


def data(details, reports=None):
    return {"reports": reports or [report()], "details": details, "version": 1, "year": 2026, "month": 10,
            "fetched_at": "2026-10-27", "account": "x"}


class OrdersTests(unittest.TestCase):
    def test_order_money_from_wb_details(self):
        rows = design.orders(data([delivery(), sale()]), {"1205": Decimal(400)})
        self.assertEqual(len(rows), 1)
        o = rows[0]
        self.assertEqual((o["status"], o["scheme"], o["qty"], o["article"]), ("Продажа", "FBS", 1, "1205"))
        self.assertEqual(o["price"], 1004)
        # Цена продавца 1440: WB удержал по ставке 1440 − 852,64 − 40,16 = 547,20 ₽, из них 436 ₽ — скидка покупателю.
        self.assertEqual((o["seller_price"], o["wb_discount"], o["commission"]), (1440, 436, -547.2))
        self.assertAlmostEqual(o["price"] + o["wb_discount"] + o["commission"] + o["acquiring"], 852.64, places=2)
        self.assertAlmostEqual(design.commission_rate([o]), 0.38, places=2)
        self.assertEqual(design.commission_label([o]), "Комиссия WB по ставке 38,0%")
        self.assertEqual((o["acquiring"], o["logistics"]), (-40.16, -117))
        self.assertEqual(o["accrued"], 735.64)                               # 852,64 − логистика 117
        self.assertEqual((o["cost"], o["profit"]), (-400, 335.64))

    def test_return_of_earlier_sale_gives_cost_back(self):
        back = sale(sellerOperName="Возврат", forPay="852.64", srid="old")
        o = design.orders(data([back]), {"1205": Decimal(400)})[0]
        self.assertEqual((o["status"], o["qty"], o["cost"]), ("Возврат", -1, 400))
        self.assertAlmostEqual(o["profit"], o["accrued"] + 400, places=2)

    def test_missing_cost_leaves_profit_empty(self):
        o = design.orders(data([sale()]), {})[0]
        self.assertIsNone(o["profit"])
        self.assertEqual(o["note"], "нет себестоимости в wb_costs.xlsx")

    def test_logistics_of_other_months_is_not_an_order(self):
        self.assertEqual(design.orders(data([delivery(srid="old")]), {}), [])

    def test_return_reason_goes_to_comment(self):
        o = design.orders(data([sale()]), {"1205": 400}, {"s1": {"reason": "Брак"}})[0]
        self.assertEqual(o["note"], "Причина возврата: Брак")

    def test_scheme_by_office(self):
        self.assertEqual(design.scheme("Склад WB"), "FBW")
        self.assertEqual(design.scheme("Коледино"), "FBW")
        self.assertEqual(design.scheme("Уральский федеральный округ МП"), "FBS")
        self.assertEqual(design.scheme(""), "")

    def test_weekly_payouts(self):
        reports = [report(dateFrom="2026-10-05", dateTo="2026-10-05", bankPaymentSum="100"),
                   report(dateFrom="2026-10-11", dateTo="2026-10-11", bankPaymentSum="50"),
                   report(dateFrom="2026-10-12", dateTo="2026-10-12", bankPaymentSum="-1")]
        weeks = design.weekly_payouts(data([], reports))
        self.assertEqual([(a.day, b.day, s) for a, b, s in weeks], [(5, 11, 150), (12, 12, -1)])


class ChargesTests(unittest.TestCase):
    def test_orders_and_charges_add_up_to_payout(self):
        payload = data([delivery(), sale(), delivery(srid="old", deliveryService="50", rrdId=3, vendorCode="1101"),
                        {"reportId": 1, "rrdId": 4, "sellerOperName": "Хранение", "srid": "st", "paidStorage": "0.11"},
                        {"reportId": 1, "rrdId": 5, "sellerOperName": "Возмещение издержек по перевозке/по складским "
                         "операциям с товаром", "srid": "x", "rebillLogisticCost": "12"}],
                       [report(bankPaymentSum="685.53")])
        rows, items = design.orders(payload, {}), design.charges(payload)
        self.assertAlmostEqual(sum(o["accrued"] for o in rows) + sum(i["amount"] for i in items), 685.53, places=2)
        names = {i["name"]: i["amount"] for i in items}
        self.assertEqual(names["Логистика по заказам других месяцев"], -50)
        self.assertEqual(names["Хранение на складах WB"], -0.11)

    def test_fast_withdrawal_fee_and_cancelled_orders_logistics(self):
        fee = {"reportId": 1, "rrdId": 6, "sellerOperName": "Разовое изменение срока перечисления денежных средств",
               "srid": "fast_withdraw_commission_1", "paymentSchedule": "387",
               "bonusTypeName": "Разовое изменение срока перечисления денежных средств"}
        cancel = delivery(srid="c1", rrdId=7, deliveryService="80.17", bonusTypeName="К клиенту при отмене")
        payload = data([delivery(), sale(), fee, cancel], [report(bankPaymentSum=str(735.64 - 387 - 80.17),
                                                                  paymentSchedule="387")])
        items = {i["name"]: i for i in design.charges(payload)}
        self.assertEqual(items["Досрочный вывод денег"]["amount"], -387)
        self.assertEqual(items["Досрочный вывод денег"]["srid"], "")
        self.assertEqual(items["Логистика отменённых заказов (туда и обратно)"]["amount"], -80.17)
        rows = design.orders(payload, {})
        self.assertAlmostEqual(sum(o["accrued"] for o in rows) + sum(i["amount"] for i in items.values()),
                               735.64 - 387 - 80.17, places=2)
        _rows, residual = finance.order_rows(payload, {})
        self.assertEqual(residual, Decimal("-467.17"))     # логистика отмены + плата за вывод — вне заказов

    def test_summary_shows_fast_withdrawal_line(self):
        fee = {"reportId": 1, "rrdId": 6, "sellerOperName": "Разовое изменение срока перечисления денежных средств",
               "srid": "fast", "paymentSchedule": "387"}
        payload = data([delivery(), sale(), fee], [report(bankPaymentSum=str(735.64 - 387), paymentSchedule="387")])
        with tempfile.TemporaryDirectory() as tmp:
            path = wb_monthly_report.build(Path(tmp), payload, {"1205": Decimal(400)})
            ws = load_workbook(path)["Сводка"]
        lines = {ws[f"B{r}"].value: ws[f"C{r}"].value for r in range(11, 25)}
        self.assertEqual(lines["Досрочный вывод денег"], -387)
        self.assertNotIn("Прочие операции WB", lines)

    def test_sheet_groups_paid_and_mentions_free_operations(self):
        from openpyxl import Workbook
        ws = Workbook().active
        design.build_charges(ws, [
            {"name": "Хранение на складах WB", "operation": "Хранение", "amount": -0.11, "date": None, "article": "",
             "srid": "st", "comment": None},
            {"name": "Возмещение издержек", "operation": "Возмещение издержек", "amount": 0, "date": None,
             "article": "", "srid": "x", "comment": None}])
        self.assertEqual((ws["A5"].value, ws["B5"].value, ws["A6"].value), ("Хранение на складах WB", -0.11, "Итого"))
        self.assertEqual(ws["A7"].value, "Без суммы, справочно: Возмещение издержек — 1")


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def build(self, costs):
        payload = data([delivery(), sale()])
        return wb_monthly_report.build(self.root, payload, costs), finance.summarize(payload, costs)

    def test_sheets_and_location(self):
        path, _summary = self.build({"1205": Decimal(400)})
        self.assertEqual(path, self.root / "reports" / "WB" / "2026" / "WB Октябрь 2026.xlsx")
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Сводка", "Заказы", "Расходы вне заказов", "Данные"])
        self.assertEqual(wb["Данные"].sheet_state, "hidden")
        self.assertEqual(wb["Данные"]["P1"].value, "Общая выручка")       # прежняя раскладка для AI-ассистента
        self.assertEqual(wb.active.title, "Сводка")

    def test_summary_lines_add_up_to_profit(self):
        path, summary = self.build({"1205": Decimal(400)})
        ws = load_workbook(path)["Сводка"]
        self.assertEqual(ws["B2"].value, "Wildberries · Октябрь 2026")
        lines = {}
        row = 11
        while ws[f"B{row}"].value != "Чистая прибыль":
            lines[ws[f"B{row}"].value] = ws[f"C{row}"].value
            row += 1
        self.assertEqual(ws[f"C{row}"].value, f"=SUM(C11:C{row - 1})")
        self.assertAlmostEqual(sum(lines.values()), float(summary["profit"]), places=2)
        self.assertEqual(lines["Логистика"], -117)
        self.assertEqual(lines["Скидка WB покупателям (оплатил WB)"], 436)
        self.assertEqual(lines["Комиссия WB по ставке 38,0%"], -547.2)
        self.assertEqual(lines["Себестоимость товаров"], -400)
        self.assertEqual(ws["C6"].value, float(summary["profit"]))

    def test_unknown_profit_is_a_dash_not_zero(self):
        path, summary = self.build({})
        self.assertIsNone(summary["profit"])
        ws = load_workbook(path)["Сводка"]
        self.assertEqual(ws["C6"].value, "—")
        texts = [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)]
        self.assertTrue(any(t.startswith("Заполните себестоимость") for t in texts))

    def test_orders_sheet(self):
        path, _summary = self.build({"1205": Decimal(400)})
        ws = load_workbook(path)["Заказы"]
        self.assertEqual(ws["B1"].value, "Номер заказа (srid)")
        self.assertEqual((ws["B2"].value, ws["I1"].value, ws["I2"].value, ws["O2"].value), ("s1", "Скидка WB (оплатил WB)", 436, 335.64))
        self.assertTrue(str(ws["O3"].value).startswith("=SUBTOTAL(109,O2:O2)"))

    def test_extra_sheets_are_styled_once(self):
        payload = data([delivery(), sale()])
        storage = {"total": Decimal(5), "by_nm": {"1205": {"total": Decimal(5), "subject": "Доски", "vendor_code": "1205"}}}
        path = wb_monthly_report.build(self.root, payload, {"1205": Decimal(400)}, storage=storage)
        wb_monthly_report.refresh_saved(self.root)   # перерисовка не добавляет вторую таблицу и вторые итоги
        book = load_workbook(path)
        ws = book["Хранение"]
        self.assertEqual(book.sheetnames, ["Сводка", "Заказы", "Расходы вне заказов", "Хранение", "Данные"])
        self.assertEqual(len(ws.tables), 1)
        self.assertEqual((ws["A3"].value, ws["C3"].value), ("Итого", "=SUBTOTAL(109,C2:C2)"))
        self.assertEqual(ws["C2"].number_format, design.MONEY)

    def test_rebuild_keeps_one_hidden_data_sheet(self):
        path, _summary = self.build({"1205": Decimal(400)})
        self.build({"1205": Decimal(400)})
        self.assertEqual(load_workbook(path).sheetnames.count("Данные"), 1)


if __name__ == "__main__":
    unittest.main()


class AiWarehouseTests(unittest.TestCase):
    def test_ai_reads_orders_and_metrics_from_hidden_data_sheet(self):
        from contextlib import closing
        import ai_warehouse as wh
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            wb_monthly_report.build(root, data([delivery(), sale()]), {"1205": Decimal(400)})
            wh.refresh(root)
            with closing(wh.connect(root, refresh_first=False)) as conn:
                orders = conn.execute("SELECT order_number, period FROM wb_orders").fetchall()
                metrics = conn.execute("SELECT COUNT(*) FROM wb_month_metrics").fetchone()[0]
        self.assertEqual([tuple(o) for o in orders], [("s1", "2026-10")])
        self.assertGreater(metrics, 0)
