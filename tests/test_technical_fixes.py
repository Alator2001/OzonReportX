"""Offline regressions for file integrity, failed requests and UI job lifecycle.

Run: python -m unittest discover -s tests -v
"""
import ast
import contextlib
import io
import os
from pathlib import Path
import queue
import subprocess
import sys
import tempfile
import threading
import types
import unittest
from collections import deque
from unittest.mock import Mock, patch

import pandas as pd
import requests
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
with patch.dict(os.environ, {"OZON_CLIENT_ID": "offline-test", "OZON_API_KEY": "offline-test"}):
    import Monthly_sales_report as monthly
    import balance_report as balance
    import first_run_setup as setup
    import price_management as pricing
    import recommended_prices as recommended
    import _auto_update as updater
    import file_io
    import ai_chat
    import update_prices


def ui_function(name, namespace):
    """Load a function without launching Flet or touching the user's files."""
    namespace.setdefault("deque", deque)
    tree = ast.parse((ROOT / "app_flet.py").read_text(encoding="utf-8-sig"))
    node = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == name)
    module = ast.Module(body=[
        ast.ImportFrom(module="__future__", names=[ast.alias(name="annotations")], level=0), node,
    ], type_ignores=[])
    ast.fix_missing_locations(module)
    exec(compile(module, "app_flet.py", "exec"), namespace)
    return namespace[name]


class OfflineTests(unittest.TestCase):
    def setUp(self):
        # Even a missing mock must never send requests to a real account.
        self.network = patch.object(requests.sessions.Session, "request", side_effect=AssertionError("Network forbidden"))
        self.network.start()
        self.addCleanup(self.network.stop)
        self.output = contextlib.redirect_stdout(io.StringIO())
        self.output.__enter__()
        self.addCleanup(self.output.__exit__, None, None, None)
        scratch = ROOT / ".cache" / "technical-fix-tests"
        scratch.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=scratch)
        self.directory = Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)

    def test_page_failure_does_not_return_partial_orders(self):
        for function, fetch_name in [(monthly.get_orders, "_fetch_fbs_page")]:
            with self.subTest(function=function.__name__):
                def fetch(_session, _start, _end, status, limit, offset):
                    if offset == 100:
                        raise requests.ConnectionError("offline")
                    return [{"posting_number": str(offset)}] if offset <= 200 else []
                with patch.object(monthly, fetch_name, side_effect=fetch):
                    with self.assertRaises(RuntimeError):
                        function("start", "end", session=Mock())

    def test_successful_pages_preserve_order(self):
        def fetch(_session, _start, _end, status, limit, offset):
            if status == "delivered" and offset < 200:
                return [{"posting_number": str(offset)}]
            return []
        for function, fetch_name in [(monthly.get_orders, "_fetch_fbs_page")]:
            with patch.object(monthly, fetch_name, side_effect=fetch):
                self.assertEqual(function("start", "end", session=Mock()), [
                    {"posting_number": "0"}, {"posting_number": "100"},
                ])

    def test_failed_transaction_page_does_not_return_partial_finances(self):
        session = Mock()
        good = Mock(status_code=200)
        good.json.return_value = {"accruals": [{"unit_number": "order", "total_amount": {"amount": "1", "currency": "RUB"}}], "last_id": "next"}
        bad = Mock(status_code=503, text="unavailable")
        bad.raise_for_status.side_effect = requests.HTTPError("HTTP 503")
        session.post.side_effect = [good, bad]
        with self.assertRaises(RuntimeError):
            monthly.get_transactions("order", "2026-09-01", "2026-09-01", session=session)
        self.assertEqual(session.post.call_count, 2)
        for call in session.post.call_args_list:
            self.assertEqual(call.kwargs["timeout"], (10, 60))

    def test_posting_requests_have_timeouts(self):
        session = Mock()
        session.post.return_value.json.return_value = {"result": {"postings": []}}
        monthly._fetch_fbs_page(session, "start", "end", "delivered", 100, 0)
        self.assertEqual(session.post.call_args.kwargs["timeout"], (10, 60))
        session.post.return_value.json.return_value = {"postings": [], "has_next": False, "cursor": ""}
        monthly._fetch_fbo_page(session, "start", "end", 100, "")
        self.assertEqual(session.post.call_args.kwargs["timeout"], (10, 60))

    def test_balance_reuses_segments_without_changing_amounts(self):
        def response(start, end):
            amount = -10 if start.endswith("-01") else -2
            return {"cashflows": {
                "star_products": {"amount": {"value": amount}},
                "product_placement_in_ozon_warehouses": {"amount": {"value": amount * 3}},
            }}
        with patch.object(balance, "get_balance_report", side_effect=response) as api:
            reports = balance.get_monthly_balance_reports(8, 2026)
            self.assertEqual(balance.get_star_products_for_month(8, 2026, reports), -12)
            self.assertEqual(balance.get_product_placement_in_ozon_warehouses_for_month(8, 2026, reports), -36)
            self.assertEqual(api.call_count, 2)
            self.assertEqual(api.call_args_list[1].args, ("2026-08-31", "2026-08-31"))

    def test_summarize_month_reconciles_payments_against_balance_and_finds_early_payment_fee(self):
        def response(start, end):
            first = start.endswith("-01")
            return {
                "total": {
                    "opening_balance": {"value": 5603.9 if first else 2208.34, "currency_code": "RUB"},
                    "closing_balance": {"value": 2208.34 if first else 3645.5, "currency_code": "RUB"},
                    "accrued": {"value": 25171.61 if first else 1437.16, "currency_code": "RUB"},
                    "payments": [{"value": -28567.17 if first else 0, "currency_code": "RUB"}],
                },
                "cashflows": {
                    "sales": {"amount": {"value": 95101}, "fee": {"value": -46782.55 if first else -2063.31}},
                    "returns": {"amount": {"value": -7313}, "fee": {"value": 3617.37 if first else 0}},
                    "services": [
                        {"name": "early_payment", "amount": {"value": -862.43 if first else 0}},
                        {"name": "acquiring", "amount": {"value": -949 if first else -10.56}},
                        {"name": "star_products", "amount": {"value": -1500.44 if first else -62.34}},
                    ],
                },
            }
        with patch.object(balance, "get_balance_report", side_effect=response):
            reports = balance.get_monthly_balance_reports(8, 2026)
            summary = balance.summarize_month(8, 2026, reports)
        self.assertEqual(summary["opening_balance"], 5603.9)
        self.assertEqual(summary["closing_balance"], 3645.5)
        self.assertAlmostEqual(summary["accrued"], 25171.61 + 1437.16)
        self.assertAlmostEqual(summary["payments"], -28567.17)
        self.assertAlmostEqual(summary["early_payment_fee"], -862.43)
        self.assertAlmostEqual(summary["services"]["acquiring"], -949 - 10.56)
        # Opening + accrued + payments should reconcile to closing, matching Ozon's own ledger.
        self.assertAlmostEqual(summary["opening_balance"] + summary["accrued"] + summary["payments"],
                               summary["closing_balance"], places=2)

    def test_failed_balance_segment_is_not_zero(self):
        with patch.object(balance, "get_balance_report", side_effect=[{}, RuntimeError("failed")]):
            with self.assertRaises(RuntimeError):
                balance.get_star_products_for_month(8, 2026)

    def test_atomic_write_keeps_original_on_writer_or_replace_failure(self):
        target = self.directory / "report.xlsx"
        target.write_bytes(b"previous report")
        with self.assertRaises(ValueError):
            with file_io.atomic_output_path(target) as temporary:
                temporary.write_bytes(b"incomplete")
                raise ValueError("writer failed")
        self.assertEqual(target.read_bytes(), b"previous report")
        with patch.object(file_io.os, "replace", side_effect=OSError("replace failed")):
            with self.assertRaises(OSError):
                monthly._safe_save_excel(pd.DataFrame({"value": [2]}), str(target))
        self.assertEqual(target.read_bytes(), b"previous report")
        self.assertEqual(list(self.directory.glob("~tmp_*")), [])

    def test_atomic_paths_are_unique(self):
        target = self.directory / "report.xlsx"
        with file_io.atomic_output_path(target) as first:
            with file_io.atomic_output_path(target) as second:
                self.assertNotEqual(first, second)
                second.write_bytes(b"second")
            first.write_bytes(b"first")
        self.assertEqual(target.read_bytes(), b"first")

    def test_costs_update_keeps_extra_sheets_and_migrates_legacy_sheet(self):
        target = self.directory / "costs.xlsx"
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb.active.append(["Артикул", "Себестоимость"])
        wb.active.append(["001", 10])
        notes = wb.create_sheet("Notes")
        notes["A1"] = "=1+2"
        notes["A1"].number_format = "0.00"
        wb.create_sheet("Акции")["A1"] = "keep"
        wb.save(target)
        wb.close()
        file_io.write_costs_dataframe(pd.DataFrame({"Артикул": ["001"], "Себестоимость": [20]}), target)
        wb = load_workbook(target)
        try:
            self.assertEqual(wb.sheetnames, ["Основной", "Notes", "Акции"])
            self.assertEqual(wb["Основной"]["B2"].value, 20)
            self.assertEqual(wb["Notes"]["A1"].value, "=1+2")
            self.assertEqual(wb["Notes"]["A1"].number_format, "0.00")
            self.assertEqual(wb["Акции"]["A1"].value, "keep")
        finally:
            wb.close()

    def test_costs_loader_prefers_main_sheet(self):
        target = self.directory / "costs.xlsx"
        with pd.ExcelWriter(target) as writer:
            pd.DataFrame({"Артикул": [1], "Себестоимость": [10]}).to_excel(writer, sheet_name="Sheet1", index=False)
            pd.DataFrame({"Артикул": [1], "Себестоимость": [20]}).to_excel(writer, sheet_name="Основной", index=False)
        df, _, cost_col = recommended.load_costs_df(target)
        self.assertEqual(df[cost_col].iloc[0], 20)

    def test_ai_costs_cache_refreshes_after_file_change(self):
        target = self.directory / "costs.xlsx"
        file_io.write_costs_dataframe(pd.DataFrame({"Артикул": [1], "Себестоимость": [10]}), target)
        tools = ai_chat.Tools(self.directory)
        self.assertEqual(tools._get_costs_df()["Себестоимость"].iloc[0], 10)
        previous_time = target.stat().st_mtime_ns
        file_io.write_costs_dataframe(pd.DataFrame({"Артикул": [1], "Себестоимость": [20]}), target)
        os.utime(target, ns=(previous_time + 2_000_000_000, previous_time + 2_000_000_000))
        self.assertEqual(tools._get_costs_df()["Себестоимость"].iloc[0], 20)
        target.unlink()
        self.assertIsNone(tools._get_costs_df())

    def test_temporary_reports_are_hidden_from_lists(self):
        folder = self.directory / "reports"
        folder.mkdir()
        for name in ["Август 2026.xlsx", "~tmp_failed.xlsx", "~$Август 2026.xlsx"]:
            (folder / name).write_bytes(b"fixture")
        expected = [folder / "Август 2026.xlsx"]
        self.assertEqual(ai_chat.list_available_reports(self.directory), expected)
        self.assertEqual(setup.list_report_files(self.directory, "reports"), expected)
        files = ui_function("files", {})
        self.assertEqual(files(folder), expected)

    def test_template_creation_handles_apostrophe_in_path(self):
        target = self.directory / "seller's folder"
        target.mkdir()
        with patch.object(setup, "prompt_yes_no", return_value=False):
            setup.ensure_costs(Path(sys.executable), target)
        wb = load_workbook(target / "costs.xlsx")
        try:
            self.assertEqual(list(next(wb.active.values)), ["артикул", "себестоимость"])
        finally:
            wb.close()

    def test_dependencies_reinstall_only_when_requirements_change(self):
        repo = self.directory
        (repo / "config").mkdir()
        requirements = repo / "config" / "requirements.txt"
        requirements.write_text("requests\nflet\n", encoding="utf-8")
        python = repo / "env" / "Scripts" / "python.exe"
        python.parent.mkdir(parents=True)
        with patch.object(setup, "run") as run:
            setup.ensure_deps(python, repo)
            self.assertEqual(run.call_args.args[0][-2:], ["-r", str(requirements)])
            run.reset_mock()
            setup.ensure_deps(python, repo)
            run.assert_not_called()
            requirements.write_text("requests\nflet\nplotly\n", encoding="utf-8")
            setup.ensure_deps(python, repo)
            self.assertEqual(run.call_count, 2)

    def test_setup_only_does_not_enter_menu_or_updater(self):
        with patch.object(sys, "argv", ["setup", "--setup-only"]), \
                patch.object(setup, "ensure_venv", return_value=(Path(sys.executable), False)), \
                patch.object(setup, "ensure_deps"), patch.object(setup, "check_for_updates") as updates, \
                patch.object(setup, "select_main_menu_option") as menu:
            setup.main()
            updates.assert_not_called()
            menu.assert_not_called()

    def test_discount_result_requires_confirmed_success_for_every_task(self):
        plan = {"approve_tasks": [{"id": 1}, {"id": 2}]}
        for response, expected in [
            ({"success_count": 0, "fail_count": 2}, False),
            ({"success_count": 1, "fail_count": 1}, False),
            ({"success_count": 0, "fail_count": 0}, False),
            ({"success_count": 2, "fail_count": 0}, True),
        ]:
            with patch.object(pricing, "approve_discount_requests", return_value=response):
                self.assertEqual(pricing.process_discount_request_plan(plan)["ok"], expected)

    def test_price_update_failure_returns_error_to_launcher(self):
        frame = pd.DataFrame({"Артикул": [1], update_prices.COL_MIN_PRICE: [100]})
        with patch.object(update_prices, "load_costs_df", return_value=(frame, "Артикул")), \
                patch.object(update_prices, "get_current_prices_from_ozon", return_value={"1": 200}), \
                patch.object(update_prices, "update_min_prices_on_ozon", return_value={"result": [
                    {"offer_id": "1", "updated": False, "errors": [{"message": "rejected"}]},
                ]}):
            with self.assertRaises(RuntimeError):
                update_prices.run(self.directory)

    def test_monthly_failure_does_not_publish_partial_workbook(self):
        (self.directory / "reports").mkdir()
        target = self.directory / "reports" / "Август 2026.xlsx"
        target.write_bytes(b"previous report")
        def write_orders(*args, **kwargs):
            Path(kwargs["output_file"]).write_bytes(b"partial report")
        with patch.object(monthly, "__file__", str(self.directory / "scripts" / "Monthly_sales_report.py")), \
                patch.object(monthly, "create_session", return_value=contextlib.nullcontext(Mock())), \
                patch.object(monthly, "get_orders", return_value=[]), \
                patch.object(monthly, "get_fbo_orders", return_value=[]), \
                patch.object(monthly, "to_excel", side_effect=write_orders), \
                patch.object(monthly, "calc_business_indicators", side_effect=RuntimeError("balance failed")):
            with self.assertRaises(RuntimeError):
                monthly.main(["--month", "8", "--year", "2026"])
        self.assertEqual(target.read_bytes(), b"previous report")

    def test_excel_metric_reader_preserves_fallback_values(self):
        target = self.directory / "metrics.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Заказы"
        ws.append(["status", None, None, None, None, None, "commission", "logistics"])
        ws.append(["returned", None, None, None, None, None, -2, -3])
        ws.append(["delivering", None, None, None, None, None, "-", 4])
        ws["P1"], ws["Q1"] = "Общая выручка", 100
        wb.save(target)
        wb.close()
        from datetime import datetime
        fn = ui_function("read_business_metrics", {"load_workbook": load_workbook, "datetime": datetime})
        metrics, _ = fn(target)
        self.assertEqual(metrics, {"Общая выручка": 100, "Количество возвращённых заказов": 1,
            "Количество заказов в доставке": 1, "Комиссии Ozon сумма": 2.0, "Логистика сумма": 7.0})

    def test_ui_rejects_second_job_and_reports_process_start_failure(self):
        page = Mock()
        events = queue.Queue()
        state = {"job_lock": threading.Lock(), "job_cancel_event": threading.Event(),
            "cancel_state": {}, "active_proc": {"proc": None}, "ui_queue": events,
            "set_status": Mock(), "set_log": Mock(), "toast": Mock(), "page": page,
            "ACCENT": "", "stream_process": Mock(side_effect=OSError("start failed"))}
        job = ui_function("job", state)
        job("first", ["python"])
        job("second", ["python"])
        page.run_thread.assert_called_once()
        state["toast"].assert_called_once()
        page.run_thread.call_args.args[0]()
        event = events.get_nowait()
        self.assertEqual(event[:3], ("finish", "first", 1))

    def test_ui_finish_releases_lock_even_when_refresh_fails(self):
        lock = threading.Lock()
        lock.acquire()
        page = Mock(controls=[])
        namespace = {"job_lock": lock, "cancel_state": {"requested": False}, "active_proc": {},
            "set_log": Mock(), "refresh_dashboard": Mock(side_effect=ValueError("bad workbook")),
            "refresh_pricing_costs_preview": Mock(), "append_log": Mock(), "set_status": Mock(),
            "PRIMARY": "green", "DANGER": "red", "chart_pending_action": {"type": None},
            "dashboard_chart_state": {"built": False}, "refresh_files": Mock(), "page": page}
        finish = ui_function("finish", namespace)
        finish("report", 1, "failed")
        self.assertFalse(lock.locked())
        self.assertIn("Ошибка", namespace["set_status"].call_args.args[0])

    def test_failed_discount_batch_retains_items_and_requires_refresh(self):
        plan = {"items": [{"id": 1}, {"id": 2}]}
        state = {"plan": plan, "busy": False}
        page = Mock(controls=[])
        namespace = {"discount_dialog_state": state, "page": page,
            "build_discount_auto_plan": Mock(return_value=plan), "render_discount_dialog": Mock(),
            "process_discount_request_plan": Mock(return_value={"ok": False, "message": "partial failure"}),
            "toast": Mock()}
        handle = ui_function("handle_discount_auto", namespace)
        handle()
        page.run_thread.call_args.args[0]()
        self.assertIs(state["plan"], plan)
        self.assertTrue(state["requires_refresh"])
        self.assertFalse(state["busy"])

    def test_cancelled_process_is_not_started(self):
        event = threading.Event()
        event.set()
        fn = ui_function("stream_process", {})
        self.assertEqual(fn(["must-not-start"], Mock(), cancel_event=event)[0], 130)

    def test_process_callback_failure_cleans_up_child(self):
        proc = Mock()
        proc.stdout = io.StringIO("line\n")
        proc.stdin = io.StringIO()
        proc.poll.return_value = None
        holder = {}
        fn = ui_function("stream_process", {"os": os, "subprocess": subprocess, "ROOT": ROOT})
        with patch.object(subprocess, "Popen", return_value=proc):
            with self.assertRaises(ValueError):
                fn(["mock"], Mock(side_effect=ValueError("UI failed")), proc_holder=holder)
        proc.kill.assert_called_once()
        self.assertIsNone(holder["proc"])
        self.assertTrue(proc.stdout.closed)

    def test_update_copy_failure_keeps_installed_files(self):
        repo = self.directory / "repo"
        source = self.directory / "release"
        repo.mkdir(); source.mkdir()
        (repo / "app.py").write_text("old", encoding="utf-8")
        (source / "app.py").write_text("new", encoding="utf-8")
        with patch.object(updater.shutil, "copy2", side_effect=OSError("copy failed")):
            with self.assertRaises(OSError):
                updater.apply_update(source, repo)
        self.assertEqual((repo / "app.py").read_text(), "old")

    def test_update_install_failure_rolls_back_and_preserves_user_data(self):
        repo = self.directory / "repo"
        source = self.directory / "release"
        repo.mkdir(); source.mkdir()
        for name in ["a.py", "b.py", "costs.xlsx", "margin_settings.json"]:
            (repo / name).write_text("old", encoding="utf-8")
            (source / name).write_text("new", encoding="utf-8")
        real_replace = os.replace
        def replace(src, dst):
            if Path(src).parent.name == "pending" and Path(src).name == "b.py":
                raise OSError("install failed")
            return real_replace(src, dst)
        with patch.object(updater.os, "replace", side_effect=replace):
            with self.assertRaises(OSError):
                updater.apply_update(source, repo)
        for name in ["a.py", "b.py", "costs.xlsx", "margin_settings.json"]:
            self.assertEqual((repo / name).read_text(), "old")

    def test_failed_rollback_keeps_recovery_files(self):
        repo = self.directory / "repo"
        source = self.directory / "release"
        repo.mkdir(); source.mkdir()
        (repo / "app.py").write_text("old", encoding="utf-8")
        (source / "app.py").write_text("new", encoding="utf-8")
        real_replace = os.replace
        def replace(src, dst):
            if Path(src).parent.name in ("pending", "previous"):
                raise OSError("locked")
            return real_replace(src, dst)
        with patch.object(updater.os, "replace", side_effect=replace):
            with self.assertRaises(RuntimeError):
                updater.apply_update(source, repo)
        recovery = list(self.directory.glob("ozon_update_*/previous/app.py"))
        self.assertEqual(len(recovery), 1)
        self.assertEqual(recovery[0].read_text(), "old")


if __name__ == "__main__":
    unittest.main()
