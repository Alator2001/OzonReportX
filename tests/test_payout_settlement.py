"""Offline checks for payout settlement: dating payouts, FIFO allocation, what lands in the report."""
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import payout_settlement as ps  # noqa: E402


def balance_api(payouts: dict[str, float], calls: list | None = None):
    """Fake v1/finance/balance: sums the payouts dated inside the period, like Ozon does."""
    def fetch(day_from, day_to):
        if calls is not None:
            calls.append((day_from, day_to))
        value = -sum(v for d, v in payouts.items() if day_from <= d <= day_to)
        return {"total": {"payments": [{"value": value}]}}
    return fetch


def report(path: Path, cost: float, profit: float) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
    ws["P1"], ws["Q1"] = "Общая выручка", 1.0
    ws["P2"], ws["Q2"] = "Чистая прибыль", profit
    ws["P3"], ws["Q3"] = "Итоговая себестоимость", -cost
    wb.save(path)
    return path


class PayoutDatesTests(unittest.TestCase):
    def test_two_payouts_in_one_month_are_found_on_their_days(self):
        fetch = balance_api({"2025-03-04": 100.0, "2025-03-31": 250.5})
        found = ps._month_payouts(2025, 3, date(2026, 1, 1), fetch)
        self.assertEqual(found, [(date(2025, 3, 4), 10000), (date(2025, 3, 31), 25050)])

    def test_closed_months_come_from_cache_and_current_month_is_refetched(self):
        with tempfile.TemporaryDirectory() as tmp:
            cache = Path(tmp) / "payouts.json"
            fetch = balance_api({"2025-03-10": 50.0, "2025-04-02": 7.0})
            ps.load_payouts(date(2025, 3, 1), date(2025, 4, 5), cache, fetch)
            calls = []
            fetch = balance_api({"2025-03-10": 50.0, "2025-04-02": 7.0, "2025-04-05": 1.0}, calls)
            found = ps.load_payouts(date(2025, 3, 1), date(2025, 4, 5), cache, fetch)
            self.assertTrue(all(c[0] >= "2025-04-01" for c in calls))
            self.assertEqual([k for _, k in found], [5000, 700, 100])


class AllocationTests(unittest.TestCase):
    def test_payouts_close_oldest_month_cost_first_then_profit_and_split(self):
        feb, mar = date(2025, 2, 1), date(2025, 3, 1)
        obligations = [(feb, {"Себестоимость": 1000, "Прибыль": 500}),
                       (mar, {"Себестоимость": 800, "Прибыль": 300})]
        parts, closed, rest = ps.allocate(obligations, [(date(2025, 3, 5), 1200), (date(2025, 4, 5), 2000)])
        self.assertEqual(parts[feb], [(date(2025, 3, 5), 1000, "Себестоимость", 1200),
                                      (date(2025, 3, 5), 200, "Прибыль", 1200),
                                      (date(2025, 4, 5), 300, "Прибыль", 2000)])
        self.assertEqual(closed[mar], {"Себестоимость": 800, "Прибыль": 300})
        self.assertEqual(rest, 600)

    def test_loss_month_owes_no_profit(self):
        jun = date(2026, 6, 1)
        parts, closed, rest = ps.allocate([(jun, {"Себестоимость": 1000, "Прибыль": -200})], [(date(2026, 7, 1), 1000)])
        self.assertEqual(closed[jun], {"Себестоимость": 1000, "Прибыль": 0})
        self.assertEqual(rest, 0)


class ReportWritingTests(unittest.TestCase):
    def test_reports_get_payout_rows_and_closing_status(self):
        with tempfile.TemporaryDirectory() as tmp:
            reports = Path(tmp) / "reports"
            reports.mkdir()
            report(reports / "Февраль 2025.xlsx", cost=100, profit=50)
            report(reports / "Март 2025.xlsx", cost=80, profit=20)
            report(reports / "WB Март 2025.xlsx", cost=999, profit=999)
            fetch = balance_api({"2025-03-14": 120.0, "2025-04-10": 60.0})
            # Отчёты лежат по-старому прямо в reports/ — сверка сама раскладывает их по reports/Ozon|WB/<год>/.
            summary = ps.update_reports(Path(tmp), today=date(2025, 4, 20), fetch=fetch)
            wb_reports, reports = reports / "WB" / "2025", reports / "Ozon" / "2025"
            self.assertEqual(sorted(p.name for p in reports.glob("*.xlsx")), ["Март 2025.xlsx", "Февраль 2025.xlsx"])

            self.assertEqual(summary["paid"], 180.0)
            self.assertEqual(summary["first_open_month"], date(2025, 3, 1))
            feb = load_workbook(reports / "Февраль 2025.xlsx")["Заказы"]
            self.assertEqual([feb.cell(r, 20).value for r in (2, 3, 4)], [100.0, 20.0, 30.0])
            self.assertEqual(feb["U3"].value, "Прибыль (часть выплаты 120.00 ₽)")
            self.assertEqual((feb["Y2"].value, feb["Y5"].value, feb["X7"].value), (0.0, 0.0, "Да"))
            mar = load_workbook(reports / "Март 2025.xlsx")["Заказы"]
            self.assertEqual((mar["X2"].value, mar["Y2"].value, mar["Y5"].value, mar["X7"].value), (30.0, 50.0, 20.0, "Нет"))
            self.assertIsNone(load_workbook(wb_reports / "WB Март 2025.xlsx")["Заказы"]["S1"].value)


if __name__ == "__main__":
    unittest.main()
