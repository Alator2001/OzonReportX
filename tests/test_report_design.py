"""Month report of the new look: hidden data sheets + visual sheets; readers accept old and new reports alike."""
import shutil
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import payout_settlement as ps  # noqa: E402
import recommended_prices as rp  # noqa: E402
import report_book as rb  # noqa: E402
import report_design as rd  # noqa: E402

HEAD = ["Статус", "Номер заказа", "Название товара", "Артикул", "Количество шт.", "Цена продажи",
        "Комиссия за продажу Ozon", "Логистика (Включает операционные ошибки продавца)", "Сумма начисления",
        "Себестоимость", "Прибыль", "Дата отгрузки", "Схема"]
PENDING = "Ожидает расчёта Ozon (трансграничный заказ)"


def legacy_report(path: Path) -> Path:
    """Отчёт старого вида, как его пишет Monthly_sales_report: три заказа, показатели P:Q, кампании, расходы."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
    ws.append(HEAD)
    ws.append(["delivered", "1-1", "Доска", 1302, 1, 1000, -500, -100, 400, -175, 225, "2026-09-02", "FBS"])
    ws.append(["cancelled", "1-2", "Доска", 1302, 1, 1000, "-", "-", -20, 0, -20, "2026-09-03", "FBS"])
    ws.append(["ожидает расчёта", "1-3", "Икона", 5003, 1, 1700, PENDING, PENDING, -150.63, 0, PENDING, "2026-09-04", "FBO"])
    for row, (label, value) in enumerate([
        ("Общая выручка", 1000), ("Чистая прибыль", 155), ("Итоговая себестоимость", -175),
        ("Продвижение Ozon", 30), ("Внешний маркетинг", 0), ("Средний чек", 1000), ("Расход хранения FBO", 10),
        ("Комиссии Ozon %", 50), ("Расходы Ozon вне заказов", 10),
    ], start=1):
        ws.cell(row=row, column=16, value=label)
        ws.cell(row=row, column=17, value=value)
    ws["Q15"] = 50
    charges = wb.create_sheet("Расходы вне заказов")
    charges.append(["Дата", "Вид расхода", "Код Ozon", "Номер (поставки/кампании)", "Сумма", "В чистой прибыли"])
    charges.append(["2026-09-07", "Кросс-докинг", 12, "2000064870634", -10, "вычтено отдельной строкой"])
    campaigns = wb.create_sheet("Кампании")
    campaigns.append(["ID кампании", "Название кампании", "Состояние", "Тип оплаты", "Тип объекта", "Бюджет (руб.)",
                      "Дневной бюджет (руб.)", "Недельный бюджет (руб.)", "Расход за период (руб.)", "Показы", "Клики",
                      "CTR (%)", "Средняя цена клика (руб.)", "Заказы (шт.)", "Заказы (руб.)", "ДРР (%)"])
    campaigns.append([1, "Браслеты", "inactive", None, "SKU", 0, 0, 0, 30, 1000, 10, 1, 3, 1, 900, 3.3])
    campaigns.append([2, "Архив", "archived", None, "SKU", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0])
    wb.save(path)
    return path


def no_payouts(_day_from, _day_to):
    return {"total": {"payments": [{"value": -100}] if _day_from <= "2026-10-08" <= _day_to else []}}


class ReportDesignTests(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.reports = self.dir / "reports" / "Ozon" / "2026"
        self.reports.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def designed(self, name="Сентябрь 2026.xlsx") -> Path:
        path = legacy_report(self.reports / name)
        rd.apply_design_file(path)
        return path

    def test_new_look_hides_data_unchanged_and_puts_summary_first(self):
        path = self.designed()
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames[:4], ["Сводка", "Заказы", "Реклама", "Расходы вне заказов"])
        self.assertEqual(wb.active.title, "Сводка")
        for name in rb.DATA_SHEETS.values():
            self.assertEqual(wb[name].sheet_state, "hidden")
        data = wb["Данные"]
        self.assertEqual([c.value for c in data[1]][:13], HEAD)
        self.assertEqual((data["P2"].value, data["Q2"].value, data["Q15"].value), ("Чистая прибыль", 155, 50))
        self.assertEqual(wb["Заказы"]["F2"].value, "Доставлен")

    def test_profit_breakdown_adds_up_to_net_profit_without_pending_orders(self):
        wb = load_workbook(self.designed())
        summary = wb["Сводка"]
        self.assertEqual(summary["B20"].value, "Чистая прибыль")
        self.assertAlmostEqual(sum(summary[f"C{r}"].value for r in range(11, 20)), 155)

    def test_empty_month_gets_a_valid_new_look(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Заказы"
        ws.append(HEAD)
        for row, (label, value) in enumerate([("Общая выручка", 0), ("Чистая прибыль", 0), ("Итоговая себестоимость", 0)], 1):
            ws.cell(row=row, column=16, value=label)
            ws.cell(row=row, column=17, value=value)
        wb.create_sheet("Расходы вне заказов").append(["Дата", "Вид расхода", "Код Ozon", "Номер", "Сумма", "В прибыли"])
        path = self.reports / "Январь 2025.xlsx"
        wb.save(path)
        rd.apply_design_file(path)
        result = load_workbook(path)
        self.assertEqual(result["Заказы"].tables, {})
        self.assertEqual(result["Заказы"]["A2"].value, "Заказов в этом месяце нет")
        self.assertEqual(ps.read_obligations(path), {"Себестоимость": 0, "Прибыль": 0})

    def test_redesign_is_repeatable(self):
        path = self.designed()
        rd.apply_design_file(path)
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames.count("Заказы"), 1)
        self.assertEqual(len(wb.sheetnames), 7)

    def test_readers_get_the_same_numbers_from_old_and_new_reports(self):
        old = legacy_report(self.reports / "Август 2026.xlsx")
        new = self.designed()
        self.assertEqual(ps.read_obligations(old), ps.read_obligations(new))
        old_costs, new_costs = rp.load_cost_structure(old), rp.load_cost_structure(new)
        self.assertEqual((old_costs.commission_rate, old_costs.logistics_avg, old_costs.logistics_by_article),
                         (new_costs.commission_rate, new_costs.logistics_avg, new_costs.logistics_by_article))
        self.assertEqual(old_costs.commission_rate, 0.5)

    def test_payout_settlement_writes_to_data_and_redraws_summary(self):
        old = legacy_report(self.reports / "Август 2026.xlsx")
        new = self.designed()
        ps.update_reports(self.dir, today=date(2026, 10, 20), fetch=no_payouts)
        self.assertEqual(load_workbook(old)["Заказы"]["S1"].value, "Дата выплаты Ozon")
        wb = load_workbook(new)
        self.assertEqual(wb["Данные"]["S1"].value, "Дата выплаты Ozon")
        texts = [c.value for row in wb["Сводка"].iter_rows() for c in row if isinstance(c.value, str)]
        self.assertTrue(any(t.startswith("⏳  Ozon ещё не рассчитался") for t in texts))
        self.assertEqual(wb.sheetnames[0], "Сводка")


if __name__ == "__main__":
    unittest.main()
