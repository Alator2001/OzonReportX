"""Logistics is taken in rubles per unit, commission as a share of price."""
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import recommended_prices as rp  # noqa: E402

HEADER = ["Статус", "Номер заказа", "Название товара", "Артикул", "Количество шт.", "Цена продажи",
          "Комиссия за продажу Ozon", rp.LOGISTICS_COLUMN]


def make_report(path: Path, rows, commission_pct=50.0):
    wb = Workbook()
    ws = wb.active
    ws.title = rp.ORDER_SHEET
    ws.append(HEADER)
    for article, qty, logistics in rows:
        ws.append(["delivered", "1", "Товар", article, qty, 1000, -500, logistics])
    ws["P15"] = "Комиссии Ozon %"
    ws["Q15"] = commission_pct
    wb.save(path)


class FormulaTests(unittest.TestCase):
    def test_prices_use_logistics_in_rubles(self):
        # (250 + 150) / (1 − 0,5 − 0,2) = 1333; (250 + 150) / (1 − 0,5 − 0,3) = 2000
        self.assertEqual(rp.compute_prices(250, 0.5, 150, 0.2, 0.3), (1333, 2000))

    def test_margin_does_not_scale_logistics_with_price(self):
        cheap = rp.compute_current_margin(1000, 250, 0.5, 150)
        dear = rp.compute_current_margin(2000, 250, 0.5, 150)
        self.assertAlmostEqual(cheap, 0.1)    # (1000 − 500 − 150 − 250) / 1000
        self.assertAlmostEqual(dear, 0.3)     # (2000 − 1000 − 150 − 250) / 2000 — логистика та же, 150 ₽
        self.assertIsNone(rp.compute_current_margin(0, 250, 0.5, 150))

    def test_impossible_margin_gives_no_price(self):
        self.assertEqual(rp.compute_prices(250, 0.8, 150, 0.25, 0.3), (None, None))


class StructureTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.path = Path(self.dir.name) / "Сентябрь 2026.xlsx"

    def tearDown(self):
        self.dir.cleanup()

    def test_article_average_and_store_fallback(self):
        make_report(self.path, [("1101", 1, 170), ("1101", 2, -340), ("7001", 1, 100), ("7002", 1, "-")])
        structure = rp.load_cost_structure(self.path)
        self.assertAlmostEqual(structure.commission_rate, 0.5)
        self.assertAlmostEqual(structure.logistics_avg, 610 / 4)        # все рубли / все штуки с логистикой
        self.assertAlmostEqual(structure.logistics_for("1101"), 170.0)  # 510 ₽ / 3 шт., две доставки
        self.assertAlmostEqual(structure.logistics_for("7001"), 152.5)  # одна доставка — среднее по магазину
        self.assertAlmostEqual(structure.logistics_for("9999"), 152.5)  # не продавался — среднее по магазину
        self.assertAlmostEqual(structure.logistics_for(None), 152.5)

    def test_series_average_when_article_has_few_shipments(self):
        # Серия 12xx: 1202 — две доставки по 210 ₽, 1203 — одна за 190 ₽; маленькие 13xx — по 110 ₽.
        make_report(self.path, [("1202", 1, 210), ("1202", 1, 210), ("1203", 1, 190),
                                ("1301", 1, 110), ("1302", 1, 110), ("1303", 1, 110), ("1304", 1, 110)])
        structure = rp.load_cost_structure(self.path)
        self.assertEqual(structure.logistics_source("1202"), (210.0, "артикул"))
        value, source = structure.logistics_source("1216")              # не продавался в прошлом месяце
        self.assertEqual(source, "серия")
        self.assertAlmostEqual(value, (210 * 2 + 190) / 3)
        self.assertEqual(structure.logistics_source("1301")[1], "серия")  # одна доставка, в серии 13xx — четыре
        self.assertEqual(structure.logistics_source("7001")[1], "магазин")
        self.assertIn("12xx — 203 ₽", structure.summary())

    def test_report_without_logistics_is_an_error(self):
        make_report(self.path, [("1101", 1, "-")])
        with self.assertRaises(ValueError):
            rp.load_cost_structure(self.path)


if __name__ == "__main__":
    unittest.main()
