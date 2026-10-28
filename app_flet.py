from __future__ import annotations

import os
import asyncio
import json
import math
import queue
import signal
import re
import requests
import subprocess
import sys
import threading
from types import SimpleNamespace
import time
from dataclasses import dataclass
from collections import deque
from uuid import uuid4
from datetime import date, datetime, timedelta
from pathlib import Path

import flet as ft
from openpyxl import load_workbook
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scripts import balance_report
from scripts.file_io import save_workbook_atomic
from scripts.report_paths import migrate as migrate_reports, ozon_report_path, ozon_reports_dir, wb_reports_dir
from scripts.report_book import CAMPAIGNS as REPORT_CAMPAIGNS, data_sheet as report_data_sheet, data_sheet_name as report_data_sheet_name
from scripts.file_lock import exclusive_file, file_signature
from scripts.ui_log import SessionLog
from scripts.wb_dashboard import HELP as WB_DASHBOARD_HELP, WbDashboard
from scripts.wb_prices_view import WbPricesView
from scripts.wb_promotions_view import WbPromotionsView
try:
    from scripts import auto_actions
except Exception:  # модуль автоакций необязателен для запуска интерфейса
    auto_actions = None
try:
    from scripts import limit_price
    LIMIT_IMPORT_ERROR = None
except Exception as exc:  # ползунок предельной цены необязателен для запуска интерфейса
    limit_price = None
    LIMIT_IMPORT_ERROR = str(exc)
try:
    from scripts import search_db, search_view
    SEARCH_IMPORT_ERROR = None
except Exception as exc:  # раздел «Поиск» необязателен для запуска интерфейса
    search_db = search_view = None
    SEARCH_IMPORT_ERROR = str(exc)
try:
    from scripts import stock_overview
    STOCK_IMPORT_ERROR = None
except Exception as exc:  # раздел «Остатки» необязателен для запуска интерфейса
    stock_overview = None
    STOCK_IMPORT_ERROR = str(exc)
from scripts.abc_view import load_abc_report as abc_load_report, matrix_counts as abc_matrix_counts, report_period_title as abc_period_title, xyz_group as abc_xyz_group
from scripts.marketplace_settings import FIELDS, HELP, read_settings, save_settings, needs_setup

try:
    import scripts.ai_chat as ai_chat_module
    from scripts.ai_chat import OLLAMA_HOST, OLLAMA_MODEL, check_ollama_model
    from scripts.ai_agent import Agent as AIAgent
    from scripts.ai_tools import AnalyticsTools
    AI_IMPORT_ERROR = None
except Exception as exc:
    ai_chat_module = None
    OLLAMA_HOST = "http://localhost:11434"
    OLLAMA_MODEL = "qwen3:4b"
    check_ollama_model = None
    AIAgent = None
    AnalyticsTools = None
    AI_IMPORT_ERROR = str(exc)

try:
    from scripts.price_management import (
        build_discount_request_plan,
        load_actions_calendar,
        process_discount_request_item,
        process_discount_request_plan,
        BOOST_COLUMN_SUFFIX,
    )
    PRICING_IMPORT_ERROR = None
except Exception as exc:
    BOOST_COLUMN_SUFFIX = " — бустинг, %"
    build_discount_request_plan = None
    load_actions_calendar = None
    process_discount_request_item = None
    process_discount_request_plan = None
    PRICING_IMPORT_ERROR = str(exc)


try:
    from scripts.ads_cpc import (
        build_ads_plan,
        execute_ads_plan,
        list_cpc_campaigns,
        load_settings as ads_load_settings,
        save_settings as ads_save_settings,
    )
    ADS_IMPORT_ERROR = None
except Exception as exc:
    build_ads_plan = None
    execute_ads_plan = None
    list_cpc_campaigns = None
    ads_load_settings = None
    ads_save_settings = None
    ADS_IMPORT_ERROR = str(exc)

from scripts.app_paths import (
    CODE_ROOT, DATA_ROOT, DATA_ROOT_SOURCE, DEFAULT_DATA_ROOT, FROZEN, SCRIPTS_DIR, python_executable, save_data_folder,
)

try:
    from scripts.organic_analysis import load_last_result as organic_load_last, run_analysis as organic_run_analysis
    ORGANIC_IMPORT_ERROR = None
except Exception as exc:
    organic_load_last = None
    organic_run_analysis = None
    ORGANIC_IMPORT_ERROR = str(exc)

# ROOT — данные пользователя (отчёты, costs.xlsx, .env), CODE_ROOT — код и картинки; в сборке это разные папки.
ROOT = DATA_ROOT
SCRIPTS = SCRIPTS_DIR
PYTHON = python_executable()
MONO_FONT = "Consolas" if os.name == "nt" else "Menlo" if sys.platform == "darwin" else "monospace"

BG = "#F5F5F7"
SURFACE = "#FFFFFF"
PRIMARY = "#0E6B5C"
PRIMARY_SOFT = "#E7F2EF"
ACCENT = "#A56B36"
TEXT = "#1D1D1F"
MUTED = "#6E6E73"
DANGER = "#A24438"
BORDER = "#D2D2D7"
COSTS_MAIN_SHEET = "Основной"
COSTS_MAIN_COLUMNS = [
    "Артикул",
    "Себестоимость",
    "Минимальная цена продажи",
    "Желательная цена продажи",
    "Текущая цена на Ozon",
    "Цена с учётом акций и скидок",
    "Текущая ожидаемая рентабельность",
    "Ожидаемая чистая прибыль",
]
# Колонки, которых может не быть в старом costs.xlsx до первой синхронизации.
COSTS_OPTIONAL_COLUMNS = {"Ожидаемая чистая прибыль"}
COSTS_MAIN_COLUMN_EXPANDS = {
    "Артикул": 12,
    "Себестоимость": 11,
    "Минимальная цена продажи": 14,
    "Желательная цена продажи": 14,
    "Текущая цена на Ozon": 14,
    "Цена с учётом акций и скидок": 16,
    "Текущая ожидаемая рентабельность": 14,
    "Ожидаемая чистая прибыль": 13,
    "Бустинг в акциях": 12,
}
# Колонка только для отображения: берётся с листа «Акции», в «Основной» не пишется.
PRICING_BOOST_COLUMN = "Бустинг в акциях"
COSTS_TABLE_COLUMNS = [*COSTS_MAIN_COLUMNS, PRICING_BOOST_COLUMN]
COSTS_ACTIONS_SHEET = "Акции"
# Название задачи синхронизации цен: по нему таблица показывает загрузку и блокируется.
PRICING_SYNC_TITLE = "Текущие цены и акции"
DASHBOARD_REFRESH_TITLE = "Обновление бизнес-сводки"
ABC_JOB_TITLE = "ABC/XYZ"
WB_SUPPLY_TITLE = "План запасов WB"
PRICING_CALCULATED_COLUMNS = ("Минимальная цена продажи", "Желательная цена продажи")
PRICING_SYNC_COLUMNS = (
    "Текущая цена на Ozon",
    "Цена с учётом акций и скидок",
    "Текущая ожидаемая рентабельность",
    "Ожидаемая чистая прибыль",
    PRICING_BOOST_COLUMN,
)
EDITABLE_COSTS_COLUMNS = {"Артикул", "Себестоимость"}
MONTHS = [
    "Январь",
    "Февраль",
    "Март",
    "Апрель",
    "Май",
    "Июнь",
    "Июль",
    "Август",
    "Сентябрь",
    "Октябрь",
    "Ноябрь",
    "Декабрь",
]
MONTH_INDEX = {name.lower(): idx for idx, name in enumerate(MONTHS, start=1)}
PRIMARY_TEXT_COLORS = {"#1D1D1F", "#F5F5F7", "#1D2A2E", "#F3F4F6", "#202123", "#ECECF1", "#1F1F22"}
MUTED_TEXT_COLORS = {"#6E6E73", "#A1A1A6", "#6F7A76", "#B6C0BC"}


@dataclass
class FolderSpec:
    title: str
    path: Path
    pattern: str = "*.xlsx"


def open_path(path: Path) -> None:
    if not path.exists():
        return
    if os.name == "nt":
        os.startfile(str(path))
        return
    subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", str(path)])


def pick_folder(prompt: str, initial: Path) -> Path | None:
    """Системный диалог выбора папки. Не ft.FilePicker: клиенту Flet на macOS не хватает разрешения
    на доступ к файлам (ENTITLEMENT_NOT_FOUND), и диалог ронял окно программы."""
    if sys.platform == "darwin":
        # Строки AppleScript экранируются как в JSON, но без \uXXXX: кириллицу AppleScript так не читает.
        quote = lambda text: json.dumps(text, ensure_ascii=False)  # noqa: E731
        script = f"POSIX path of (choose folder with prompt {quote(prompt)} default location (POSIX file {quote(str(initial))}))"
        result = subprocess.run(["osascript", "-e", script], capture_output=True, text=True)
    elif os.name == "nt":
        script = (
            "[Console]::OutputEncoding = [Text.Encoding]::UTF8; Add-Type -AssemblyName System.Windows.Forms; "
            "$d = New-Object System.Windows.Forms.FolderBrowserDialog; $d.Description = $env:PICK_PROMPT; "
            "$d.SelectedPath = $env:PICK_INITIAL; if ($d.ShowDialog() -eq 'OK') { $d.SelectedPath }"
        )
        result = subprocess.run(
            ["powershell", "-NoProfile", "-STA", "-Command", script], capture_output=True, text=True, encoding="utf-8",
            env={**os.environ, "PICK_PROMPT": prompt, "PICK_INITIAL": str(initial)},
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    else:
        result = subprocess.run(["zenity", "--file-selection", "--directory", f"--title={prompt}"], capture_output=True, text=True)
    path = result.stdout.strip() if result.returncode == 0 else ""
    return Path(path) if path else None


def kill_process_tree(proc: subprocess.Popen, force: bool = False) -> None:
    """Stop the child process together with the scripts it started."""
    if os.name == "nt":
        subprocess.run(
            ["taskkill", "/PID", str(proc.pid), "/T", "/F"],
            capture_output=True,
            text=True,
            timeout=10,
        )
        return
    # stream_process starts children in their own session, so the group id equals the child pid.
    try:
        os.killpg(proc.pid, signal.SIGKILL if force else signal.SIGTERM)
    except ProcessLookupError:
        pass


def stream_process(
    cmd: list[str],
    on_line,
    on_prompt=None,
    extra_env: dict[str, str] | None = None,
    proc_holder: dict[str, subprocess.Popen | None] | None = None,
    cancel_event: threading.Event | None = None,
) -> tuple[int, str]:
    if cancel_event is not None and cancel_event.is_set():
        return 130, "Остановлено пользователем."
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUNBUFFERED"] = "1"
    if extra_env:
        env.update(extra_env)
    proc = subprocess.Popen(
        cmd,
        cwd=ROOT,
        env=env,
        stdout=subprocess.PIPE,
        stdin=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
        bufsize=1,
        start_new_session=os.name != "nt",
    )
    if proc_holder is not None:
        proc_holder["proc"] = proc
    try:
        if cancel_event is not None and cancel_event.is_set():
            kill_process_tree(proc, force=True)
        lines = deque(maxlen=500)
        assert proc.stdout is not None
        for line in proc.stdout:
            lines.append(line)
            on_line(line)
            if on_prompt and proc.stdin is not None:
                lower = line.lower()
                label = None
                if "введите сумму затрат на продвижение ozon" in lower:
                    label = "Продвижение Ozon"
                elif "введите сумму затрат на внешний маркетинг" in lower:
                    label = "Внешний маркетинг"
                if label:
                    answer = on_prompt(label, line.strip())
                    try:
                        proc.stdin.write(f"{answer}\n")
                        proc.stdin.flush()
                    except (BrokenPipeError, OSError):
                        pass
        proc.wait()
        return proc.returncode, "".join(lines).strip()
    finally:
        if proc.poll() is None:
            try:
                kill_process_tree(proc, force=True)
            except Exception:
                proc.kill()
            proc.wait()
        if proc.stdout is not None:
            proc.stdout.close()
        if proc.stdin is not None:
            proc.stdin.close()
        if proc_holder is not None and proc_holder.get("proc") is proc:
            proc_holder["proc"] = None


def files(folder: Path, pattern: str = "*.xlsx") -> list[Path]:
    return sorted((p for p in folder.glob(pattern) if p.is_file() and not p.name.startswith(("~$", "~tmp_"))), key=lambda p: p.stat().st_mtime, reverse=True) if folder.exists() else []


def _period_sort_key(path: Path) -> tuple[int, int, int]:
    stem = path.stem.lower()
    matches = re.findall(r"(январь|февраль|март|апрель|май|июнь|июль|август|сентябрь|октябрь|ноябрь|декабрь)\s+(\d{4})", stem)
    if matches:
        month_name, year_text = matches[-1]
        return (int(year_text), MONTH_INDEX.get(month_name, 0), 1)
    try:
        return (datetime.fromtimestamp(path.stat().st_mtime).year, datetime.fromtimestamp(path.stat().st_mtime).month, 0)
    except OSError:
        return (0, 0, 0)


def files_by_relevance(folder: Path, pattern: str = "*.xlsx") -> list[Path]:
    if not folder.exists():
        return []
    batch = [p for p in folder.glob(pattern) if p.is_file() and not p.name.startswith(("~$", "~tmp_"))]
    return sorted(batch, key=lambda p: (_period_sort_key(p), p.stat().st_mtime if p.exists() else 0), reverse=True)


def report_path_for_period(month: int, year: int) -> Path:
    return ozon_report_path(ROOT, year, month)


def iter_periods(month: int, year: int, count: int) -> list[tuple[int, int]]:
    periods: list[tuple[int, int]] = []
    cursor_month = month
    cursor_year = year
    for _ in range(count):
        periods.append((cursor_month, cursor_year))
        cursor_month -= 1
        if cursor_month == 0:
            cursor_month = 12
            cursor_year -= 1
    periods.reverse()
    return periods


def _format_number(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, float):
        if value.is_integer():
            return f"{int(value):,}".replace(",", " ")
        return f"{value:,.2f}".replace(",", " ").replace(".", ",")
    if isinstance(value, int):
        return f"{value:,}".replace(",", " ")
    return str(value)


def _format_currency(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, (int, float)):
        return f"{_format_number(float(value))} ₽"
    return str(value)


def _format_percent(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, (int, float)):
        return f"{float(value):.2f}%".replace(".", ",")
    return str(value)


def _format_costs_preview_value(column: str, value) -> str:
    if value is None:
        return "—"
    if column == "Артикул":
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        text = str(value).strip()
        return text[:-2] if text.endswith(".0") else (text or "—")
    numeric = as_float(value)
    if numeric is not None:
        if "рентабельность" in column.lower():
            percent_value = numeric * 100 if abs(numeric) <= 1 else numeric
            return _format_percent(percent_value)
        return _format_currency(numeric)
    text = str(value).strip()
    return text or "—"


def _format_costs_editor_value(column: str, value) -> str:
    if value is None:
        return ""
    if column == "Артикул":
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        text = str(value).strip()
        return text[:-2] if text.endswith(".0") else text
    numeric = as_float(value)
    if numeric is not None:
        return str(int(numeric)) if float(numeric).is_integer() else str(numeric).replace(".", ",")
    return str(value).strip()


def _excel_color_to_hex(color) -> str | None:
    if color is None:
        return None
    color_type = getattr(color, "type", None)
    rgb = getattr(color, "rgb", None)
    if color_type == "rgb" and isinstance(rgb, str):
        hex_value = rgb[-6:].upper()
        if len(hex_value) == 6:
            return f"#{hex_value}"
    return None


def _excel_fill_color(cell) -> str | None:
    fill = getattr(cell, "fill", None)
    if fill is None or getattr(fill, "patternType", None) != "solid":
        return None
    return _excel_color_to_hex(getattr(fill, "fgColor", None)) or _excel_color_to_hex(getattr(fill, "startColor", None))


def _excel_font_color(cell) -> str | None:
    font = getattr(cell, "font", None)
    if font is None:
        return None
    return _excel_color_to_hex(getattr(font, "color", None))


def _hex_to_rgb(color: str | None) -> tuple[int, int, int] | None:
    if not color or not isinstance(color, str):
        return None
    value = color.lstrip("#")
    if len(value) != 6:
        return None
    try:
        return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))
    except ValueError:
        return None


def _excel_text_color(background_color: str | None, fallback_dark: bool) -> str:
    rgb = _hex_to_rgb(background_color)
    if rgb is None:
        return "#F3F4F6" if fallback_dark else TEXT
    red, green, blue = rgb
    luminance = (0.299 * red + 0.587 * green + 0.114 * blue) / 255
    return "#1D1D1F" if luminance >= 0.62 else "#F5F5F7"


def _build_costs_cell(column: str, cell_payload: dict[str, object] | None, dark: bool, expand: int) -> ft.Control:
    payload = cell_payload or {}
    value = payload.get("value")
    fill_color = payload.get("fill_color") if isinstance(payload.get("fill_color"), str) else None
    font_color = payload.get("font_color") if isinstance(payload.get("font_color"), str) else None
    resolved_text_color = font_color or _excel_text_color(fill_color, dark)
    return ft.Container(
        expand=expand,
        padding=ft.Padding.symmetric(horizontal=10, vertical=6),
        border_radius=10,
        bgcolor=fill_color,
        tooltip=payload.get("tooltip") if isinstance(payload.get("tooltip"), str) else None,
        data="preserve_excel_fill" if fill_color else None,
        content=ft.Text(
            _format_costs_preview_value(column, value),
            size=12,
            color=resolved_text_color,
            weight=ft.FontWeight.W_700 if column == "Артикул" else ft.FontWeight.W_500,
            selectable=True,
            data="preserve_excel_text" if (fill_color or font_color) else None,
        ),
    )


def _pricing_conditional_fill(
    column: str,
    row: dict[str, dict[str, object]],
    min_margin_pct: float = 23.0,
    desired_margin_pct: float = 30.0,
) -> str | None:
    min_price_payload = row.get("Минимальная цена продажи") or {}
    min_price = as_float(min_price_payload.get("value"))
    value_payload = row.get(column) or {}
    value = as_float(value_payload.get("value"))
    if value is None:
        return None

    if column in {"Текущая цена на Ozon", "Цена с учётом акций и скидок"}:
        if value <= 0 or min_price is None:
            return None
        return "#FFC7CE" if value < min_price else "#C6EFCE"

    if column == "Текущая ожидаемая рентабельность":
        profitability = value * 100 if abs(value) <= 1 else value
        if profitability < min_margin_pct:
            return "#FFC7CE"
        if min_margin_pct <= profitability <= desired_margin_pct:
            return "#C6EFCE"

    if column == "Ожидаемая чистая прибыль" and value < 0:
        return "#FFC7CE"
    return None


def _apply_costs_conditional_formats(
    rows: list[dict[str, object]],
    min_margin_pct: float = 23.0,
    desired_margin_pct: float = 30.0,
) -> None:
    for row in rows:
        typed_row = row if isinstance(row, dict) else {}
        for column in COSTS_MAIN_COLUMNS:
            payload = typed_row.get(column)
            if not isinstance(payload, dict):
                continue
            if payload.get("fill_color"):
                continue
            conditional_fill = _pricing_conditional_fill(column, typed_row, min_margin_pct=min_margin_pct, desired_margin_pct=desired_margin_pct)
            if conditional_fill:
                payload["fill_color"] = conditional_fill


def _read_action_boosts(wb) -> dict[str, list[tuple[str, float]]] | None:
    """
    {артикул: [(акция, бустинг %)]} с листа «Акции» (колонки «<акция> — бустинг, %»).
    None, если бустинг ещё не загружался (нет листа или колонок бустинга).
    """
    if COSTS_ACTIONS_SHEET not in wb.sheetnames:
        return None
    ws = wb[COSTS_ACTIONS_SHEET]
    header = [str(cell.value or "") for cell in next(ws.iter_rows(min_row=1, max_row=1), [])]
    boost_columns = [
        (index, name[: -len(BOOST_COLUMN_SUFFIX)])
        for index, name in enumerate(header)
        if name.endswith(BOOST_COLUMN_SUFFIX)
    ]
    if not boost_columns:
        return None
    result: dict[str, list[tuple[str, float]]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        article = _format_costs_preview_value("Артикул", row[0] if row else None)
        for index, action_name in boost_columns:
            boost = as_float(row[index]) if index < len(row) else None
            if boost is not None:
                result.setdefault(article, []).append((action_name, boost))
    return result


def _boost_cell_payload(boosts: list[tuple[str, float]] | None) -> dict[str, object]:
    """Максимальный бустинг товара; во всплывающей подсказке — все акции."""
    if boosts is None:
        return {"value": None, "fill_color": None, "font_color": None, "tooltip": "Нажмите «Синхронизировать», чтобы загрузить бустинг"}
    if not boosts:
        return {"value": None, "fill_color": None, "font_color": None, "tooltip": "Товар не участвует в акциях"}
    ordered = sorted(boosts, key=lambda pair: pair[1], reverse=True)

    def percent(boost: float) -> str:
        return f"{boost:g}%".replace(".", ",")

    best = percent(ordered[0][1])
    value = best if len(ordered) == 1 else f"{best} (+{len(ordered) - 1})"
    tooltip = "\n".join(f"{name}: {percent(boost)}" for name, boost in ordered)
    return {"value": value, "fill_color": None, "font_color": None, "tooltip": tooltip}


def read_costs_main_rows(
    costs_path: Path,
    min_margin_pct: float = 23.0,
    desired_margin_pct: float = 30.0,
) -> tuple[list[dict[str, object]], str | None]:
    if not costs_path.exists():
        return [], "Файл costs.xlsx не найден."

    try:
        wb = load_workbook(costs_path, data_only=True)
    except Exception as exc:
        return [], f"Не удалось открыть costs.xlsx: {exc}"

    try:
        if COSTS_MAIN_SHEET not in wb.sheetnames:
            return [], "Лист «Основной» не найден в costs.xlsx."
        ws = wb[COSTS_MAIN_SHEET]
        header_cells = next(ws.iter_rows(min_row=1, max_row=1), None)
        if not header_cells:
            return [], "Лист «Основной» пуст."

        header_map = {
            str(cell.value).strip().lower(): idx
            for idx, cell in enumerate(header_cells)
            if cell.value is not None and str(cell.value).strip()
        }
        missing = [column for column in COSTS_MAIN_COLUMNS if column.lower() not in header_map and column not in COSTS_OPTIONAL_COLUMNS]
        if missing:
            return [], f"В листе «Основной» нет колонок: {', '.join(missing)}."

        rows: list[dict[str, object]] = []
        for row in ws.iter_rows(min_row=2):
            item: dict[str, object] = {}
            item["__row_number"] = row[0].row if row else None
            has_payload = False
            for column in COSTS_MAIN_COLUMNS:
                cell_index = header_map.get(column.lower())
                cell = row[cell_index] if cell_index is not None and cell_index < len(row) else None
                value = cell.value if cell is not None else None
                if value not in (None, ""):
                    has_payload = True
                item[column] = {
                    "value": value,
                    "fill_color": _excel_fill_color(cell) if cell is not None else None,
                    "font_color": _excel_font_color(cell) if cell is not None else None,
                }
            if has_payload:
                rows.append(item)
        _apply_costs_conditional_formats(rows, min_margin_pct=min_margin_pct, desired_margin_pct=desired_margin_pct)
        boosts = _read_action_boosts(wb)
        for item in rows:
            article = _format_costs_preview_value("Артикул", item["Артикул"]["value"])
            item[PRICING_BOOST_COLUMN] = _boost_cell_payload(None if boosts is None else boosts.get(article, []))
        return rows, None
    except Exception as exc:
        return [], f"Не удалось прочитать лист «Основной»: {exc}"
    finally:
        wb.close()


def write_costs_main_rows(costs_path: Path, updates: list[dict[str, object]], expected_signature=None, saved_signature=None) -> str | None:
    try:
        with exclusive_file(costs_path):
            if expected_signature is not None and file_signature(costs_path) != expected_signature:
                return "Файл costs.xlsx изменился после открытия таблицы. Правки сохранены в редакторе; обновите данные перед повторным сохранением."
            error = _write_costs_main_rows(costs_path, updates)
            if error is None and saved_signature is not None:
                saved_signature["signature"] = file_signature(costs_path)
            return error
    except Exception as exc:
        return f"Не удалось сохранить costs.xlsx: {exc}"


def _write_costs_main_rows(costs_path: Path, updates: list[dict[str, object]]) -> str | None:
    if not costs_path.exists():
        return "Файл costs.xlsx не найден."

    try:
        wb = load_workbook(costs_path)
    except Exception as exc:
        return f"Не удалось открыть costs.xlsx для сохранения: {exc}"

    try:
        if COSTS_MAIN_SHEET not in wb.sheetnames:
            return "Лист «Основной» не найден в costs.xlsx."
        ws = wb[COSTS_MAIN_SHEET]
        header_cells = next(ws.iter_rows(min_row=1, max_row=1), None)
        if not header_cells:
            return "Лист «Основной» пуст."

        header_map = {
            str(cell.value).strip().lower(): idx + 1
            for idx, cell in enumerate(header_cells)
            if cell.value is not None and str(cell.value).strip()
        }
        missing = [column for column in EDITABLE_COSTS_COLUMNS if column.lower() not in header_map]
        if missing:
            return f"В листе «Основной» нет колонок для сохранения: {', '.join(sorted(missing))}."

        for update in updates:
            row_number = update.get("row_number")
            values = update.get("values")
            if not isinstance(row_number, int) or row_number < 2 or not isinstance(values, dict):
                continue
            for column in EDITABLE_COSTS_COLUMNS:
                if column not in values:
                    continue
                ws.cell(row=row_number, column=header_map[column.lower()]).value = values[column]

        save_workbook_atomic(wb, costs_path)
        return None
    except Exception as exc:
        return f"Не удалось сохранить costs.xlsx: {exc}"
    finally:
        wb.close()


def read_business_metrics(report_path: Path) -> tuple[dict[str, object], str | None]:
    if not report_path.exists():
        return {}, None
    wb = load_workbook(report_path, data_only=True, read_only=True)
    try:
        ws = report_data_sheet(wb)
        metrics: dict[str, object] = {}
        for label, value in ws.iter_rows(min_row=1, max_row=70, min_col=16, max_col=17, values_only=True):
            if label:
                metrics[str(label).strip()] = value
        needs_fallback = any(
            key not in metrics
            for key in (
                "Количество возвращённых заказов",
                "Количество заказов в доставке",
                "Комиссии Ozon сумма",
                "Логистика сумма",
            )
        )
        if needs_fallback:
            returned_count = 0
            delivering_count = 0
            commission_total = 0.0
            logistics_total = 0.0
            for row in ws.iter_rows(min_row=2, max_col=8, values_only=True):
                status_value = row[0]
                commission_value = row[6]
                logistics_value = row[7]
                if status_value is not None and str(status_value).strip().lower() == "returned":
                    returned_count += 1
                if status_value is not None and str(status_value).strip().lower() == "delivering":
                    delivering_count += 1
                try:
                    if commission_value is not None and str(commission_value).strip() not in ("", "-"):
                        commission_total += abs(float(commission_value))
                except (TypeError, ValueError):
                    pass
                try:
                    if logistics_value is not None and str(logistics_value).strip() not in ("", "-"):
                        logistics_total += abs(float(logistics_value))
                except (TypeError, ValueError):
                    pass
            metrics.setdefault("Количество возвращённых заказов", returned_count)
            metrics.setdefault("Количество заказов в доставке", delivering_count)
            metrics.setdefault("Комиссии Ozon сумма", commission_total)
            metrics.setdefault("Логистика сумма", logistics_total)
        built_at = datetime.fromtimestamp(report_path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
        return metrics, built_at
    finally:
        wb.close()


def read_campaign_metrics(report_path: Path) -> tuple[dict[str, float], list[dict[str, object]]]:
    if not report_path.exists():
        return {}, []
    wb = load_workbook(report_path, data_only=True, read_only=True)
    try:
        campaigns_sheet = report_data_sheet_name(wb.sheetnames, REPORT_CAMPAIGNS)
        if not campaigns_sheet:
            return {}, []
        ws = wb[campaigns_sheet]
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, ())
        header_names = [str(header).strip() if header else f"column_{idx}" for idx, header in enumerate(headers, start=1)]

        campaigns: list[dict[str, object]] = []
        totals = {
            "spend": 0.0,
            "impressions": 0.0,
            "clicks": 0.0,
            "orders": 0.0,
            "revenue": 0.0,
        }

        def safe_metric(row_map: dict[str, object], key: str) -> float:
            return as_float(row_map.get(key)) or 0.0

        for row in rows:
            row_map = dict(zip(header_names, row))
            spend = safe_metric(row_map, "Расход за период (руб.)")
            if spend <= 0:
                continue

            impressions = safe_metric(row_map, "Показы")
            clicks = safe_metric(row_map, "Клики")
            orders = safe_metric(row_map, "Заказы (шт.)")
            revenue = safe_metric(row_map, "Заказы (руб.)")
            ctr = safe_metric(row_map, "CTR (%)")
            cpc = safe_metric(row_map, "Средняя цена клика (руб.)")
            drr = safe_metric(row_map, "ДРР (%)")
            budget = safe_metric(row_map, "Бюджет (руб.)")
            daily_budget = safe_metric(row_map, "Дневной бюджет (руб.)")
            weekly_budget = safe_metric(row_map, "Недельный бюджет (руб.)")

            calc_cpm = (spend / impressions * 1000.0) if impressions > 0 else 0.0
            calc_cpa = (spend / orders) if orders > 0 else None
            calc_roas = (revenue / spend) if spend > 0 else None
            calc_cr = (orders / clicks * 100.0) if clicks > 0 else 0.0
            calc_avg_order_value = (revenue / orders) if orders > 0 else None
            budget_burn = (spend / budget * 100.0) if budget > 0 else None
            daily_budget_burn = (spend / daily_budget * 100.0) if daily_budget > 0 else None
            weekly_budget_burn = (spend / weekly_budget * 100.0) if weekly_budget > 0 else None

            if clicks >= 15 and orders <= 0:
                efficiency_label = "Трафик без заказов"
                efficiency_tone = "#F6DFDB"
                efficiency_reason = "Есть заметный трафик, но клики не конвертируются в заказ: кликов >= 15, заказов = 0."
            elif (
                calc_roas is not None
                and calc_avg_order_value is not None
                and calc_cpa is not None
                and orders >= 3
                and calc_roas >= 5
                and drr <= 20
                and calc_cr >= 5
                and calc_cpa <= calc_avg_order_value * 0.25
            ):
                efficiency_label = "Можно масштабировать"
                efficiency_tone = "#DCEFE5"
                efficiency_reason = "Сильная экономика кампании: ROAS >= 5, ДРР <= 20%, CR >= 5%, заказов >= 3 и CPA не выше 25% среднего чека."
            elif (
                calc_roas is not None
                and orders >= 1
                and calc_roas >= 3
                and drr <= 35
                and calc_cr >= 2
            ):
                efficiency_label = "Рабочая кампания"
                efficiency_tone = "#F4EDDB"
                efficiency_reason = "Кампания окупается и приносит заказы, но ещё не дотягивает до уровня масштабирования."
            elif (
                (calc_roas is not None and calc_roas < 2)
                or drr > 45
                or calc_cr < 1.5
                or (calc_avg_order_value is not None and calc_cpa is not None and calc_cpa > calc_avg_order_value * 0.45)
            ):
                efficiency_label = "Нужна оптимизация"
                efficiency_tone = "#F8E9DD"
                efficiency_reason = "Слабая эффективность по одной или нескольким метрикам: низкий ROAS, высокий ДРР, слабый CR или слишком дорогой заказ."
            else:
                efficiency_label = "Под наблюдением"
                efficiency_tone = "#EEF2F8"
                efficiency_reason = "Есть сигналы эффективности, но данных или запаса по метрикам пока недостаточно для уверенного решения."

            campaign = dict(row_map)
            campaign.update(
                {
                    "__spend": spend,
                    "__impressions": impressions,
                    "__clicks": clicks,
                    "__orders": orders,
                    "__revenue": revenue,
                    "__ctr": ctr,
                    "__cpc": cpc,
                    "__drr": drr,
                    "__cpm": calc_cpm,
                    "__cpa": calc_cpa,
                    "__roas": calc_roas,
                    "__cr": calc_cr,
                    "__avg_order_value": calc_avg_order_value,
                    "__budget_burn": budget_burn,
                    "__daily_budget_burn": daily_budget_burn,
                    "__weekly_budget_burn": weekly_budget_burn,
                    "__efficiency_label": efficiency_label,
                    "__efficiency_tone": efficiency_tone,
                    "__efficiency_reason": efficiency_reason,
                }
            )
            campaigns.append(campaign)

            totals["spend"] += spend
            totals["impressions"] += impressions
            totals["clicks"] += clicks
            totals["orders"] += orders
            totals["revenue"] += revenue

        if not campaigns:
            return {}, []

        campaigns.sort(key=lambda item: (item.get("__spend") or 0.0), reverse=True)

        summary = {
            "campaigns_count": float(len(campaigns)),
            "spend": totals["spend"],
            "impressions": totals["impressions"],
            "clicks": totals["clicks"],
            "orders": totals["orders"],
            "revenue": totals["revenue"],
            "ctr": (totals["clicks"] / totals["impressions"] * 100.0) if totals["impressions"] > 0 else 0.0,
            "cpc": (totals["spend"] / totals["clicks"]) if totals["clicks"] > 0 else 0.0,
            "cpm": (totals["spend"] / totals["impressions"] * 1000.0) if totals["impressions"] > 0 else 0.0,
            "cr": (totals["orders"] / totals["clicks"] * 100.0) if totals["clicks"] > 0 else 0.0,
            "cpa": (totals["spend"] / totals["orders"]) if totals["orders"] > 0 else 0.0,
            "roas": (totals["revenue"] / totals["spend"]) if totals["spend"] > 0 else 0.0,
            "drr": (totals["spend"] / totals["revenue"] * 100.0) if totals["revenue"] > 0 else 0.0,
            "avg_order_value": (totals["revenue"] / totals["orders"]) if totals["orders"] > 0 else 0.0,
        }
        return summary, campaigns
    finally:
        wb.close()


def margin_defaults() -> tuple[str, str]:
    path = ROOT / "margin_settings.json"
    if not path.exists():
        return "0.25", "0.30"
    try:
        import json

        data = json.loads(path.read_text(encoding="utf-8"))
        return f"{float(data.get('min_margin', 0.25)):.2f}", f"{float(data.get('desired_margin', 0.30)):.2f}"
    except Exception:
        return "0.25", "0.30"


def supply_defaults() -> tuple[str, str, str]:
    """Срок поставки, цикл и уровень сервиса из fbo_supply_settings.json."""
    lead, cover, level = "14", "30", "0.9"
    path = ROOT / "fbo_supply_settings.json"
    if path.exists():
        try:
            import json

            data = json.loads(path.read_text(encoding="utf-8"))
            lead = str(int(data.get("lead_days", 14)))
            cover = str(int(data.get("cover_days", 30)))
            level = f"{float(data.get('service_level', 0.9)):g}"
        except Exception:
            pass
    return lead, cover, level


def valid_supply(lead: str, cover: str, level: str) -> str | None:
    try:
        lead_days, cover_days = int(lead), int(cover)
        service = float(level.replace(",", "."))
    except ValueError:
        return "Срок и цикл — целые дни, уровень сервиса — число вроде 0.9."
    if not 1 <= lead_days <= 120 or not 7 <= cover_days <= 180:
        return "Срок поставки — от 1 до 120 дней, цикл — от 7 до 180 дней."
    if not 0.5 <= service <= 0.99:
        return "Уровень сервиса должен быть от 0.5 до 0.99."
    return None


def valid_month_year(month: str, year: str) -> str | None:
    try:
        m, y = int(month), int(year)
    except ValueError:
        return "Месяц и год должны быть числами."
    if not 1 <= m <= 12:
        return "Месяц должен быть в диапазоне 1-12."
    if not 2000 <= y <= 2100:
        return "Год должен быть в диапазоне 2000-2100."
    return None


def valid_abc(a_m: str, a_y: str, b_m: str, b_y: str) -> str | None:
    error = valid_month_year(a_m, a_y) or valid_month_year(b_m, b_y)
    if error:
        return error
    if (int(a_y), int(a_m)) > (int(b_y), int(b_m)):
        return "Начало периода должно быть не позже конца."
    return None


def valid_margin(a: str, b: str) -> str | None:
    try:
        ma, mb = float(a.replace(",", ".")), float(b.replace(",", "."))
    except ValueError:
        return "Маржа должна быть числом в формате 0.25."
    if not (0 < ma < 1 and 0 < mb < 1):
        return "Значения маржи должны быть в диапазоне (0, 1)."
    if ma >= mb:
        return "Минимальная маржа должна быть меньше желаемой."
    return None


def valid_money(value: str, label: str) -> str | None:
    try:
        amount = float((value or "0").replace(",", "."))
    except ValueError:
        return f"{label} должен быть числом."
    if amount < 0:
        return f"{label} не может быть отрицательным."
    return None


def card(
    title: str,
    subtitle: str,
    body: list[ft.Control],
    tone: str = SURFACE,
    width: int | None = None,
    expand: bool = False,
    col: int | float | dict[str, int] = 12,
) -> ft.Control:
    header_controls: list[ft.Control] = []
    if title.strip():
        header_controls.append(ft.Text(title, size=18, weight=ft.FontWeight.W_700, color=TEXT))
    if subtitle.strip():
        header_controls.append(ft.Text(subtitle, size=12, color=MUTED))
    return ft.Container(
        width=width,
        expand=expand,
        col=col,
        padding=22,
        bgcolor=tone,
        border_radius=24,
        border=ft.Border.all(1, BORDER),
        shadow=ft.BoxShadow(blur_radius=22, spread_radius=0, color="#0F000000", offset=ft.Offset(0, 8)),
        content=ft.Column(
            [*header_controls, *body],
            spacing=10,
        ),
    )


def metric(
    label: str,
    value: str,
    caption: str,
    tone: str,
    width: int = 220,
    col: int | float = 3,
    tooltip: str | None = None,
) -> ft.Control:
    caption_controls: list[ft.Control] = []
    if caption.strip():
        caption_controls.append(ft.Text(caption, size=12, color=MUTED))
    return ft.Container(
        width=width,
        col=col,
        padding=18,
        bgcolor=tone,
        border_radius=22,
        border=ft.Border.all(1, "#E5E5EA"),
        tooltip=tooltip,
        content=ft.Column(
            [
                ft.Text(label, size=11, color=MUTED, weight=ft.FontWeight.W_600),
                ft.Text(value, size=22, weight=ft.FontWeight.W_700, color=TEXT),
                *caption_controls,
            ],
            spacing=3,
        ),
    )


def hero_metric(
    label: str,
    value: str,
    caption: str,
    tone: str,
    accent: str,
    width: int = 320,
    col: int | float = 4,
    tooltip: str | None = None,
) -> ft.Control:
    accent_controls: list[ft.Control] = []
    if caption.strip():
        accent_controls.append(
            ft.Row(
                [
                    ft.Container(width=10, height=10, bgcolor=accent, border_radius=10),
                    ft.Text(caption, size=12, color=MUTED),
                ],
                spacing=6,
            )
        )
    return ft.Container(
        width=width,
        col=col,
        padding=24,
        bgcolor=tone,
        border_radius=26,
        border=ft.Border.all(1, BORDER),
        tooltip=tooltip,
        content=ft.Column(
            [
                ft.Text(label, size=11, color=MUTED, weight=ft.FontWeight.W_700),
                ft.Text(value, size=32, weight=ft.FontWeight.W_800, color=TEXT),
                *accent_controls,
            ],
            spacing=5,
        ),
    )


def kv_row(label: str, value: str) -> ft.Control:
    return ft.Row(
        [
            ft.Text(label, expand=True, color=MUTED, size=12),
            ft.Text(value, color=TEXT, size=12, weight=ft.FontWeight.W_700),
        ],
        alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
    )


def kv_row_compact(label: str, value: str, tooltip: str | None = None) -> ft.Control:
    return ft.Column(
        [
            ft.Text(label, color=MUTED, size=9, weight=ft.FontWeight.W_600),
            ft.Text(value, color=TEXT, size=13, weight=ft.FontWeight.W_700),
        ],
        spacing=0,
        tooltip=tooltip,
    )


def fact_chip(
    label: str,
    value: str,
    tone: str = "#FCFAF6",
    col: int | float = 6,
    tooltip: str | None = None,
    variant: str = "default",
) -> ft.Control:
    is_secondary = variant == "secondary"
    return ft.Container(
        col=col,
        padding=8 if is_secondary else 10,
        bgcolor=tone,
        border_radius=14 if is_secondary else 16,
        border=ft.Border.all(1, "#E5E5EA"),
        tooltip=tooltip,
        content=ft.Column(
            [
                ft.Text(label, color=MUTED, size=8 if is_secondary else 9, weight=ft.FontWeight.W_600),
                ft.Text(value, color=TEXT, size=12 if is_secondary else 13, weight=ft.FontWeight.W_700),
            ],
            spacing=1 if is_secondary else 2,
        ),
    )


def report_list_item(name: str, timestamp: str, on_open) -> ft.Control:
    return ft.Container(
        padding=ft.Padding.only(left=14, right=10, top=12, bottom=12),
        border_radius=18,
        bgcolor="#F7F7FA",
        border=ft.Border.all(1, "#E5E5EA"),
        content=ft.Row(
            [
                ft.Column(
                    [
                        ft.Text(name, size=13, color=TEXT, weight=ft.FontWeight.W_600),
                        ft.Text(timestamp, size=11, color=MUTED),
                    ],
                    spacing=2,
                    expand=True,
                ),
                ft.TextButton("Открыть", on_click=on_open),
            ],
            alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
        ),
    )


def as_float(value) -> float | None:
    try:
        if value is None:
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def choose_metric_tone(
    value: float | None,
    *,
    good_at_least: float | None = None,
    warn_at_least: float | None = None,
    good_at_most: float | None = None,
    warn_at_most: float | None = None,
    dark: bool = False,
    neutral_light: str = "#EEF2F8",
    neutral_dark: str = "#293240",
) -> str:
    if value is None:
        return neutral_dark if dark else neutral_light
    if good_at_least is not None and value >= good_at_least:
        return "#214039" if dark else "#EEF7F4"
    if warn_at_least is not None and value >= warn_at_least:
        return "#454033" if dark else "#F8F5EC"
    if good_at_most is not None and value <= good_at_most:
        return "#214039" if dark else "#EEF7F4"
    if warn_at_most is not None and value <= warn_at_most:
        return "#454033" if dark else "#F8F5EC"
    return "#4A3236" if dark else "#FAECEA"


def semantic_surface(
    kind: str,
    *,
    dark: bool = False,
) -> str:
    palettes = {
        "neutral": ("#FBFCFF", "#20262B"),
        "neutral_soft": ("#F7F8FC", "#24292E"),
        "neutral_alt": ("#F3F5FA", "#272C31"),
        "positive": ("#EEF7F4", "#214039"),
        "warning": ("#F8F5EC", "#454033"),
        "danger": ("#FAECEA", "#4A3236"),
        "info": ("#F2F5FC", "#2C3744"),
    }
    light, dark_value = palettes.get(kind, palettes["neutral"])
    return dark_value if dark else light


def chart_definition(chart_kind: str) -> tuple[str, str, str, str]:
    if chart_kind == "revenue":
        return "Оборот по месяцам", "Выручка", "оборот", "#4FB59E"
    if chart_kind == "margin":
        return "Рентабельность по месяцам", "Net Margin", "рентабельность", "#D9B44A"
    return "Чистая прибыль по месяцам", "Чистая прибыль", "прибыль", "#E58A45"


def iter_period_range(from_month: int, from_year: int, to_month: int, to_year: int) -> list[tuple[int, int]]:
    periods: list[tuple[int, int]] = []
    cursor_month = from_month
    cursor_year = from_year
    while (cursor_year, cursor_month) <= (to_year, to_month):
        periods.append((cursor_month, cursor_year))
        cursor_month += 1
        if cursor_month == 13:
            cursor_month = 1
            cursor_year += 1
    return periods


def themed_color(color: str | None, dark: bool) -> str | None:
    if color is None:
        return None
    light_to_dark = {
        "#1D1D1F": "#F5F5F7",
        "#6E6E73": "#A1A1A6",
        "#1D2A2E": "#F3F4F6",
        "#6F7A76": "#B6C0BC",
        "#FFFDF8": "#232A2F",
        "#F4F1EA": "#1A1F23",
        "#FCFAF6": "#20262B",
        "#F7F7FA": "#262A2F",
        "#EAF3F0": "#243934",
        "#EAF0F5": "#27323D",
        "#F5EDE4": "#3A312B",
        "#D9EEE8": "#27443E",
        "#E8F1EE": "#24343A",
        "#EEF3E9": "#2A342D",
        "#F7F3EC": "#2A2F33",
        "#F5E3D8": "#4A302E",
        "#F5DDD8": "#4A302E",
        "#F6E7DA": "#43342C",
        "#E8EEF8": "#2B3440",
        "#EFE6F5": "#352D3C",
        "#E6F1EE": "#263632",
        "#F2EBDD": "#433A2E",
        "#F4E4E1": "#453131",
        "#F3F7F1": "#24342C",
        "#F8F1EA": "#3A312A",
        "#EEF2F8": "#293240",
        "#F1ECE3": "#38332C",
        "#F3F1EC": "#2C302B",
        "#F3F0E8": "#32302C",
        "#F2EDE7": "#332E2B",
        "#F1EEF8": "#2E3040",
        "#F0F4EA": "#2A342D",
        "#F8E9DD": "#433129",
        "#E8F3EE": "#1E3A32",
        "#F6E3E0": "#4A2626",
        "#DCEFE5": "#24443B",
        "#F4EDDB": "#4A422D",
        "#F6DFDB": "#4A302E",
        "#FBFCFF": "#20262B",
        "#F7F8FC": "#24292E",
        "#F3F5FA": "#272C31",
        "#EEF7F4": "#214039",
        "#F8F5EC": "#454033",
        "#FAECEA": "#4A3236",
        "#F2F5FC": "#2C3744",
        "#E6DED1": "#3A3F43",
        "#D2D2D7": "#3A3A3C",
        "#E5E5EA": "#3A3A3C",
        "#ECECF1": "#2C2C2E",
        "#F5F5F7": "#23262A",
        "#E7F2EF": "#233933",
        "#D6D0C4": "#3D474B",
        "#F7F7F8": "#212121",
        "#E9EAED": "#2F2F31",
        "#202123": "#ECECF1",
        "#1F1F22": "#ECECF1",
        "#D9D9DE": "#3A3A3C",
        "#FFFFFF": "#2B2B2D",
    }
    dark_to_light = {v: k for k, v in light_to_dark.items()}
    return light_to_dark.get(color, color) if dark else dark_to_light.get(color, color)


UI_STATE_PATH = ROOT / ".cache" / "ui_state.json"


def load_ui_state(path: Path = UI_STATE_PATH) -> dict:
    """Настройки внешнего вида между запусками; повреждённый или пустой файл — значения по умолчанию."""
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def save_ui_state(updates: dict, path: Path = UI_STATE_PATH) -> None:
    state = load_ui_state(path)
    state.update(updates)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass


def system_prefers_dark(brightness: ft.Brightness | None) -> bool:
    """Тема системы (macOS/Windows); если Flet её не сообщил — тёмная, как раньше."""
    return brightness != ft.Brightness.LIGHT


def main(page: ft.Page) -> None:
    page.title = "OzonReportX"
    try:
        migrate_reports(ROOT)  # отчёты из общей кучи reports/ — в reports/Ozon|WB/<год>/
    except OSError as exc:
        print(f"⚠️ Не удалось разложить отчёты по папкам: {exc}")
    start_dark = system_prefers_dark(page.platform_brightness)
    page.theme_mode = ft.ThemeMode.DARK if start_dark else ft.ThemeMode.LIGHT
    page.bgcolor = "#1A1F23" if start_dark else BG
    page.padding = 16
    try:
        page.window.maximized = True
        page.window.resizable = True
    except Exception:
        page.window_width = 1600
        page.window_height = 960
    try:
        page.window_resizable = True
    except Exception:
        pass
    page.window_min_width = 1100
    page.window_min_height = 640
    # Страница не прокручивается: прокрутка внутри разделов и меню.
    page.scroll = None
    current_dark = {"value": start_dark}
    # Пока пользователь не выбрал тему кнопкой, приложение повторяет тему системы.
    follow_system_theme = {"value": True}
    # Свёрнутое боковое меню: только иконки, чтобы на маленьком экране осталось место для содержимого.
    sidebar_collapsed = {"value": bool(load_ui_state().get("sidebar_collapsed", False))}

    today = date.today()
    prev = today.replace(day=1) - timedelta(days=1)
    years = [str(y) for y in range(today.year - 2, today.year + 2)]
    month_opts = [ft.dropdown.Option(str(i), f"{i}. {m}") for i, m in enumerate(MONTHS, start=1)]
    m_min, m_des = margin_defaults()

    monthly_m = ft.Dropdown(label="Месяц", options=month_opts, value=str(today.month), width=230)
    monthly_y = ft.Dropdown(label="Год", options=[ft.dropdown.Option(y) for y in years], value=str(today.year), width=170)
    dashboard_m = ft.Dropdown(label="Месяц сводки", options=month_opts, value=str(today.month), width=230)
    dashboard_y = ft.Dropdown(label="Год сводки", options=[ft.dropdown.Option(y) for y in years], value=str(today.year), width=170)
    dashboard_chart_type = ft.Dropdown(
        label="График",
        options=[
            ft.dropdown.Option("profit", "Прибыль / Месяц"),
            ft.dropdown.Option("revenue", "Оборот / Месяц"),
            ft.dropdown.Option("margin", "Рентабельность / Месяц"),
        ],
        value="profit",
        width=230,
    )
    chart_from_month_value = prev.month - 4 if prev.month > 4 else prev.month + 8
    chart_from_year_value = prev.year if prev.month > 4 else prev.year - 1
    dashboard_chart_from_m = ft.Dropdown(label="С месяца", options=month_opts, value=str(chart_from_month_value), width=180)
    dashboard_chart_from_y = ft.Dropdown(label="Год", options=[ft.dropdown.Option(y) for y in years], value=str(chart_from_year_value), width=120)
    dashboard_chart_to_m = ft.Dropdown(label="По месяц", options=month_opts, value=str(prev.month), width=180)
    dashboard_chart_to_y = ft.Dropdown(label="Год", options=[ft.dropdown.Option(y) for y in years], value=str(prev.year), width=120)
    for chart_control in [
        dashboard_chart_type,
        dashboard_chart_from_m,
        dashboard_chart_from_y,
        dashboard_chart_to_m,
        dashboard_chart_to_y,
    ]:
        chart_control.color = TEXT
        chart_control.bgcolor = "#FBFCFF"
        chart_control.border_color = "#D9D9DE"
        chart_control.focused_border_color = PRIMARY
        chart_control.label_style = ft.TextStyle(size=11, color=MUTED)
    abc_fm = ft.Dropdown(label="Месяц от", options=month_opts, value=str(prev.month), width=190)
    abc_fy = ft.Dropdown(label="Год от", options=[ft.dropdown.Option(y) for y in years], value=str(prev.year), width=150)
    abc_tm = ft.Dropdown(label="Месяц до", options=month_opts, value=str(today.month), width=190)
    abc_ty = ft.Dropdown(label="Год до", options=[ft.dropdown.Option(y) for y in years], value=str(today.year), width=150)
    margin_min = ft.TextField(label="Минимальная маржа", value=m_min, width=210)
    supply_lead_default, supply_cover_default, supply_level_default = supply_defaults()
    supply_lead = ft.TextField(label="Срок поставки, дней", value=supply_lead_default, width=190,
                               tooltip="От решения о поставке до продажи на складе Ozon: сборка, слот, дорога, приёмка")
    supply_cover = ft.TextField(label="Цикл, дней", value=supply_cover_default, width=150,
                                tooltip="На сколько дней продаж должна хватать поставка")
    supply_level = ft.TextField(label="Уровень сервиса", value=supply_level_default, width=170,
                                tooltip="Вероятность не уйти в ноль на FBO за цикл: 0.8 — меньше товара, 0.95 — больше")
    margin_des = ft.TextField(label="Желаемая маржа", value=m_des, width=210)
    state = ft.Text("Готово", color=PRIMARY, weight=ft.FontWeight.W_700, size=12, expand=True, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)
    status_detail = ft.Text("Интерфейс готов к работе.", color=MUTED, size=11, max_lines=3, overflow=ft.TextOverflow.ELLIPSIS)
    status_indicator = ft.Container(width=8, height=8, border_radius=999, bgcolor=PRIMARY)
    spin = ft.ProgressRing(visible=False, width=18, height=18, stroke_width=2.4)
    cancel_button = ft.Button("Отмена", visible=False, bgcolor=DANGER, color="white", height=30)
    log = ft.ListView(spacing=6, auto_scroll=True, expand=True)
    log_selectable = ft.TextField(
        multiline=True,
        read_only=True,
        min_lines=30,
        max_lines=30,
        value="Интерфейс готов к работе.",
        border_radius=16,
        content_padding=ft.Padding.symmetric(horizontal=16, vertical=16),
    )
    log_store = SessionLog(ROOT / ".cache" / "logs" / f"ui-{datetime.now():%Y%m%d-%H%M%S}-{uuid4().hex}.log")
    log_store.reset("Интерфейс готов к работе.")
    log_state = {"lines": log_store.lines}
    dashboard_freshness = ft.Text("—", size=13, color=MUTED)
    dashboard_chart = ft.Image(src="", visible=False, border_radius=18)
    dashboard_chart_note = ft.Text("Выберите метрику и период, затем постройте график.", size=11, color=MUTED, visible=False)
    dashboard_chart_state_indicator = ft.Container(width=40, height=40, border_radius=14, alignment=ft.Alignment(0, 0))
    dashboard_chart_state_title = ft.Text("Постройте график", size=15, weight=ft.FontWeight.W_700, color=TEXT)
    dashboard_chart_state_caption = ft.Text("Выберите метрику и период, затем постройте график.", size=12, color=MUTED)
    dashboard_chart_state_card = ft.Container(
        padding=18,
        border_radius=22,
        bgcolor=semantic_surface("neutral_soft", dark=False),
        border=ft.Border.all(1, BORDER),
        content=ft.Row(
            [
                dashboard_chart_state_indicator,
                ft.Column(
                    [dashboard_chart_state_title, dashboard_chart_state_caption],
                    spacing=4,
                    expand=True,
                ),
            ],
            spacing=14,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
        ),
    )
    dashboard_chart_meta_title = ft.Text("", size=14, weight=ft.FontWeight.W_700, color=TEXT)
    dashboard_chart_meta_caption = ft.Text("", size=11, color=MUTED)
    dashboard_chart_meta_card = ft.Container(
        visible=False,
        padding=14,
        border_radius=20,
        bgcolor=semantic_surface("neutral_soft", dark=False),
        border=ft.Border.all(1, BORDER),
        content=ft.Row(
            [
                ft.Container(width=10, height=10, border_radius=999, bgcolor=PRIMARY),
                ft.Column([dashboard_chart_meta_title, dashboard_chart_meta_caption], spacing=3, expand=True),
            ],
            spacing=10,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
        ),
    )
    dashboard_chart_panel = ft.Column(
        [dashboard_chart_meta_card, dashboard_chart_state_card, dashboard_chart, dashboard_chart_note],
        spacing=10,
    )
    dashboard_chart_state = {"built": False, "mode": "empty"}
    chart_pending_action = {"type": None}
    active_proc: dict[str, subprocess.Popen | None] = {"proc": None}
    job_lock = threading.Lock()
    job_cancel_event = threading.Event()
    cancel_state = {"requested": False, "title": ""}
    status_ui_ready = {"value": False}
    abc_state: dict[str, object] = {"files": [], "selected": None, "data": None, "error": None, "filter": "all", "cell": None, "creating": False}
    abc_report_pick = ft.Dropdown(label="Отчёт", width=380, options=[])
    abc_subtitle = ft.Text("", size=13, color=MUTED)
    abc_body = ft.Column(spacing=16)
    abc_sync_label = ft.Text("Строю ABC/XYZ-анализ по месячным отчётам…", size=13, weight=ft.FontWeight.W_600, text_align=ft.TextAlign.CENTER)
    abc_sync_overlay = ft.Container(
        left=0, top=0, right=0, bottom=0, visible=False, border_radius=18, alignment=ft.Alignment(0, -1),
        padding=ft.Padding.only(top=150), data="preserve_excel_fill",
        content=ft.Container(
            padding=ft.Padding.symmetric(horizontal=28, vertical=20), border_radius=18, data="preserve_excel_fill",
            content=ft.Column([ft.ProgressRing(width=36, height=36, stroke_width=3), abc_sync_label],
                              spacing=14, tight=True, horizontal_alignment=ft.CrossAxisAlignment.CENTER),
        ),
    )
    files_col = ft.Column(spacing=14)
    pricing_costs_status = ft.Text("Читаю лист «Основной»...", size=13, color=MUTED, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)
    pricing_costs_status_shell = ft.Container(content=pricing_costs_status)
    pricing_stats_area = ft.Container()
    # Без своей прокрутки: таблица растёт по высоте, прокручивается только страница.
    pricing_costs_list = ft.Column(spacing=0)
    pricing_costs_save_button = ft.Button("Сохранить изменения", bgcolor=PRIMARY, color="white", disabled=True)
    pricing_costs_editors: dict[int, dict[str, ft.TextField]] = {}
    pricing_costs_dirty = {"value": False, "revision": 0}
    pricing_costs_snapshot = {"signature": None}
    ai_diag_title = ft.Text("Локальная модель готова к работе", size=18, weight=ft.FontWeight.W_800, color=TEXT)
    ai_diag_subtitle = ft.Text("Проверка Ollama и модели выполняется локально на устройстве пользователя.", size=12, color=MUTED)
    ai_status_dot = ft.Container(width=12, height=12, border_radius=999, bgcolor="#D9B44A")
    ai_status_label = ft.Text("Проверка статуса AI...", size=12, weight=ft.FontWeight.W_700, color=TEXT)
    ai_runtime_host = ft.Text(f"Ollama host: {OLLAMA_HOST}", size=11, color=MUTED)
    ai_runtime_model = ft.Text(f"Модель: {OLLAMA_MODEL}", size=11, color=MUTED)
    ai_runtime_note = ft.Text("AI-ассистент работает локально и не отправляет отчёты во внешний API.", size=12, color=MUTED)
    ai_model_select = ft.Dropdown(label="Модель", width=180, value=OLLAMA_MODEL, options=[])
    ai_phase_label = ft.Text("Текущий этап: idle", size=12, weight=ft.FontWeight.W_800, color=TEXT)
    ai_phase_detail = ft.Text("Модель ожидает новый запрос.", size=12, color=MUTED)
    ai_trace = ft.Column(spacing=8, scroll=ft.ScrollMode.AUTO, height=210)
    ai_messages = ft.ListView(spacing=18, auto_scroll=True, expand=True)
    ai_empty_state = ft.Container(
        alignment=ft.Alignment(0, 0),
        expand=True,
        content=ft.Container(
            width=760,
            padding=28,
            border_radius=32,
            bgcolor="#FFFFFF",
            border=ft.Border.all(1, "#E5E5EA"),
            content=ft.Column(
                [
                    ft.Text("Чем помочь по Ozon?", size=30, weight=ft.FontWeight.W_900, color=TEXT, text_align=ft.TextAlign.CENTER),
                    ft.Text(
                        "Спросите про прибыль, расходы, риски или конкретный месяц. Ассистент сам подтянет нужные данные из локальных отчётов.",
                        size=13,
                        color=MUTED,
                        text_align=ft.TextAlign.CENTER,
                        width=560,
                    ),
                ],
                spacing=10,
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            ),
        ),
    )
    ai_messages.controls.append(ai_empty_state)
    ai_input = ft.TextField(
        hint_text="Спросите про прибыль, расходы, риски...",
        multiline=True,
        min_lines=1,
        max_lines=6,
        shift_enter=True,
        border_radius=24,
        content_padding=ft.Padding.symmetric(horizontal=18, vertical=16),
        expand=True,
    )
    ai_send_button = ft.IconButton(icon=ft.Icons.ARROW_UPWARD_ROUNDED, tooltip="Отправить")
    ai_reset_button = ft.TextButton("Очистить контекст")
    ai_refresh_button = ft.TextButton("Проверить Ollama")
    ai_tools_button = ft.PopupMenuButton(content=ft.Text("Инструменты", size=12, weight=ft.FontWeight.W_700, color=TEXT))
    ai_new_chat_button = ft.TextButton("Новый диалог")
    ai_top_title = ft.Text("AI-ассистент", size=18, weight=ft.FontWeight.W_800, color=TEXT)
    ai_top_subtitle = ft.Text("Локальный анализ на Ollama", size=12, color=MUTED)
    ai_quick_month_button = ft.TextButton("Как дела в магазине")
    ai_quick_pressure_button = ft.TextButton("Что давит на прибыль")
    ai_quick_brief_button = ft.TextButton("Что пополнить")
    ai_quick_risk_button = ft.TextButton("Проблемные товары")
    ai_busy_label = ft.Text("Ассистент готов.", size=12, color=MUTED)
    ai_busy_ring = ft.ProgressRing(width=16, height=16, stroke_width=2.4, visible=False)
    ai_view: ft.Container | None = None
    ai_header_card: ft.Container | None = None
    ai_composer_card: ft.Container | None = None
    ai_state = {
        "busy": False,
        "conversation_history": [],
        "available": False,
        "status_message": "",
        "selected_model": OLLAMA_MODEL,
        "available_models": [],
        "request_seq": 0,
        "active_request_id": None,
        "cancelled_request_ids": set(),
        "phase": "idle",
        "trace": [],
        "debug_enabled": False,
        "stream_controls": {},
        "request_started_at": {},
    }
    # Один набор инструментов на сессию: живые данные Ozon кешируются между вопросами.
    ai_tools = AnalyticsTools(ROOT) if AnalyticsTools else None

    folder_specs = [
        FolderSpec("Отчёты Ozon", ROOT / "reports" / "Ozon", "**/*.xlsx"),
        FolderSpec("Отчёты WB", ROOT / "reports" / "WB", "**/*.xlsx"),
        FolderSpec("ABC/XYZ аналитика", ROOT / "ABC&XYZ reports"),
        FolderSpec("Поставки", ROOT / "stocks reports"),
        FolderSpec("Баланс", ROOT / "balance reports"),
    ]

    def toast(msg: str, error: bool = False) -> None:
        set_status("Ошибка" if error else "Готово", msg, DANGER if error else PRIMARY, busy=False)
        # Уведомление показывается через show_dialog: сразу и без перерисовки всего приложения. Раньше каждое
        # оставалось в overlay навсегда, а page.update() перерисовывал всё. О скрытии по таймеру Flet не сообщает,
        # поэтому прежние уведомления убираем из списка окон сами — иначе новое не покажется, а список будет расти.
        for old in [dialog for dialog in page._dialogs.controls if isinstance(dialog, ft.SnackBar)]:
            page._dialogs.controls.remove(old)
        page.show_dialog(ft.SnackBar(ft.Text(msg), bgcolor=DANGER if error else PRIMARY))

    def _log_line_control(line: str) -> ft.Text:
        return ft.Text(
            line if line else " ",
            size=12,
            color="#F3F4F6" if current_dark["value"] else TEXT,
            selectable=True,
            font_family=MONO_FONT,
        )

    def _sync_log_theme() -> None:
        line_color = "#F3F4F6" if current_dark["value"] else TEXT
        for control in log.controls:
            if isinstance(control, ft.Text):
                control.color = line_color
        log_selectable.color = line_color
        log_selectable.bgcolor = "#232A2F" if current_dark["value"] else "#FFFDF8"
        log_selectable.border_color = "#3D474B" if current_dark["value"] else "#D6D0C4"

    def _render_log_lines(lines: list[str]) -> None:
        payload = lines or ["Команда завершилась без вывода."]
        log.controls = [_log_line_control(line) for line in payload]
        log_selectable.value = "\n".join(payload)
        _sync_log_theme()

    def set_log(text: str, persist: bool = True) -> None:
        normalized = (text or "Команда завершилась без вывода.").replace("\r\n", "\n").replace("\r", "\n")
        log_store.reset(normalized, persist=persist)
        _render_log_lines(list(log_store.lines))

    def append_log(text: str) -> None:
        normalized = (text or "").replace("\r\n", "\n").replace("\r", "\n")
        extra_lines = normalized.split("\n") if normalized else [""]
        log_store.append(normalized)
        log.controls.extend(_log_line_control(line) for line in extra_lines[-500:])
        log.controls = log.controls[-500:]
        log_selectable.value = "\n".join(log_store.lines)

    def get_full_log_text() -> str:
        return log_store.full_text().strip() or "Лог пока пуст."

    # Тёмные тона статуса нечитаемы на тёмной подложке — в тёмной теме берём светлые аналоги.
    status_title_dark = {PRIMARY: "#70D4BC", ACCENT: "#E0A15F", DANGER: "#FF9C8F", MUTED: "#A1A1A6"}
    status_tone = {"value": PRIMARY}

    def set_status(title: str, detail: str = "", color: str | None = None, busy: bool = False, cancellable: bool = True) -> None:
        tone = color or PRIMARY
        status_tone["value"] = tone
        state.value = title
        state.color = status_title_dark.get(tone, tone) if current_dark["value"] else tone
        detail_text = (detail or "").replace("\n", " ").strip()
        if len(detail_text) > 140:
            detail_text = detail_text[:137].rstrip() + "..."
        status_detail.value = detail_text
        status_indicator.bgcolor = tone
        if tone == DANGER:
            status_shell.bgcolor = semantic_surface("danger", dark=current_dark["value"])
            status_shell.border = ft.Border.all(1, "#CFA8A1" if not current_dark["value"] else "#5A3B40")
        elif busy:
            status_shell.bgcolor = semantic_surface("info", dark=current_dark["value"])
            status_shell.border = ft.Border.all(1, "#C8D4E5" if not current_dark["value"] else "#344455")
        elif tone == PRIMARY:
            status_shell.bgcolor = semantic_surface("positive", dark=current_dark["value"])
            status_shell.border = ft.Border.all(1, "#C8DED6" if not current_dark["value"] else "#2C5047")
        else:
            status_shell.bgcolor = "#1F252A" if current_dark["value"] else "#FBFCFF"
            status_shell.border = ft.Border.all(1, "#353D42" if current_dark["value"] else "#E1E5EE")
        spin.visible = busy
        status_shell.tooltip = f"{title}. {detail_text}" if sidebar_collapsed["value"] and detail_text else (title if sidebar_collapsed["value"] else None)
        cancel_button.visible = (busy and cancellable) or active_proc["proc"] is not None
        cancel_row.visible = cancel_button.visible
        if status_ui_ready["value"] and page.controls:
            page.update(status_shell, status_indicator, state, status_detail, spin, cancel_row)

    def ai_surface_tokens() -> dict[str, str]:
        dark = current_dark["value"]
        return {
            "canvas": "#1D1D1F" if dark else "#F5F5F7",
            "composer": "#26282C" if dark else "#FFFFFF",
            "composer_border": "#383C41" if dark else "#D9D9DE",
            "text": "#F5F5F7" if dark else "#1D1D1F",
            "muted": "#A1A1A6" if dark else "#6E6E73",
            "assistant": "#F5F5F7" if dark else "#1D1D1F",
            "assistant_bg": "#26282C" if dark else "#FFFFFF",
            "assistant_border": "#383C41" if dark else "#E5E5EA",
            "user_bg": "#303236" if dark else "#ECEEF3",
            "user_text": "#F5F5F7" if dark else "#1D1D1F",
            "toolbar": "#26282C" if dark else "#FFFFFF",
            "toolbar_border": "#383C41" if dark else "#E5E7EB",
        }

    def ai_centered(control: ft.Control, width: int = 1040) -> ft.Control:
        return ft.Row(
            [ft.Container(content=control, width=width)],
            alignment=ft.MainAxisAlignment.CENTER,
        )

    def ai_markdown(text: str) -> ft.Control:
        dark = current_dark["value"]
        return ft.Markdown(
            value=text or "",
            selectable=True,
            extension_set=ft.MarkdownExtensionSet.GITHUB_FLAVORED,
            code_theme=ft.MarkdownCodeTheme.A11Y_DARK if dark else ft.MarkdownCodeTheme.A11Y_LIGHT,
            soft_line_break=True,
            shrink_wrap=True,
            fit_content=True,
        )

    def ai_message_bubble(role: str, text: str, note: str | None = None) -> ft.Control:
        tokens = ai_surface_tokens()
        is_user = role == "user"
        if is_user:
            return ai_centered(
                ft.Row(
                    [
                        ft.Container(
                            padding=ft.Padding.symmetric(horizontal=18, vertical=14),
                            border_radius=22,
                            bgcolor=tokens["user_bg"],
                            data="ai_user_bubble",
                            content=ft.Text(text, size=14, color=tokens["user_text"], selectable=True, data="ai_user_text"),
                        )
                    ],
                    alignment=ft.MainAxisAlignment.END,
                )
            )
        return ai_centered(
            ft.Container(
                padding=ft.Padding.symmetric(horizontal=22, vertical=18),
                border_radius=24,
                bgcolor=tokens["assistant_bg"],
                border=ft.Border.all(1, tokens["assistant_border"]),
                data="ai_assistant_bubble",
                content=ft.Column(
                    [ai_markdown(text)]
                    + ([ft.Text(note, size=11, color=tokens["muted"], data="ai_note_text")] if note else []),
                    spacing=10,
                    horizontal_alignment=ft.CrossAxisAlignment.START,
                ),
            )
        )

    def ai_append_message(role: str, text: str, note: str | None = None) -> None:
        if ai_messages.controls and ai_messages.controls[0] is ai_empty_state:
            ai_messages.controls.clear()
        ai_messages.controls.append(ai_message_bubble(role, text, note))

    def ai_format_elapsed(seconds: float | None) -> str | None:
        if seconds is None:
            return None
        if seconds < 1:
            return f"Ответ за {seconds * 1000:.0f} мс"
        if seconds < 10:
            return f"Ответ за {seconds:.1f} с"
        return f"Ответ за {seconds:.0f} с"

    def ai_append_streaming_placeholder(note: str | None = None) -> dict[str, ft.Control]:
        if ai_messages.controls and ai_messages.controls[0] is ai_empty_state:
            ai_messages.controls.clear()
        tokens = ai_surface_tokens()
        text_control = ft.Text("", size=15, color=tokens["assistant"], selectable=True)
        content_host = ft.Container(
            padding=ft.Padding.symmetric(horizontal=22, vertical=18),
            border_radius=24,
            bgcolor=tokens["assistant_bg"],
            border=ft.Border.all(1, tokens["assistant_border"]),
            data="ai_assistant_bubble",
            content=text_control,
        )
        note_control = ft.Text(note or "", size=11, color=tokens["muted"], data="ai_note_text")
        thinking_current = ft.Text(
            "",
            size=12,
            color=tokens["muted"],
            selectable=True,
        )
        thinking_anchor = ft.Text("", size=1, color=tokens["muted"])
        thinking_list = ft.ListView(
            controls=[thinking_current, thinking_anchor],
            spacing=3,
            auto_scroll=True,
            height=66,
            width=956,
        )
        thinking_toggle = ft.TextButton("Развернуть", visible=False)
        thinking_title = ft.Text("Ход рассуждения", size=11, weight=ft.FontWeight.W_700, color=tokens["muted"])
        thinking_wrap = ft.Container(
            visible=False,
            width=1040,
            padding=ft.Padding.symmetric(horizontal=14, vertical=12),
            border_radius=20,
            bgcolor=tokens["assistant_bg"],
            border=ft.Border.all(1, tokens["assistant_border"]),
            data="ai_thinking_wrap",
            content=ft.Column(
                [
                    ft.Row(
                        [thinking_title, thinking_toggle],
                        alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                        vertical_alignment=ft.CrossAxisAlignment.CENTER,
                    ),
                    thinking_list,
                ],
                spacing=6,
            ),
        )

        def toggle_thinking(_e=None) -> None:
            expanded = bool(getattr(thinking_toggle, "data", False))
            thinking_toggle.data = not expanded
            thinking_toggle.text = "Свернуть" if not expanded else "Развернуть"
            thinking_list.height = 220 if not expanded else 66
            if status_ui_ready["value"] and page.controls:
                page.update(thinking_wrap)

        thinking_toggle.on_click = toggle_thinking
        assistant_control = ai_centered(
            ft.Column(
                [thinking_wrap, content_host, note_control],
                spacing=8,
                horizontal_alignment=ft.CrossAxisAlignment.START,
            )
        )
        ai_messages.controls.append(assistant_control)
        return {
            "text": text_control,
            "content_host": content_host,
            "note": note_control,
            "thinking_list": thinking_list,
            "thinking_current": thinking_current,
            "thinking_anchor": thinking_anchor,
            "thinking_toggle": thinking_toggle,
            "thinking_wrap": thinking_wrap,
        }

    def ai_has_conversation() -> bool:
        return not (len(ai_messages.controls) == 1 and ai_messages.controls[0] is ai_empty_state)

    def ai_apply_chat_theme() -> None:
        tokens = ai_surface_tokens()
        if ai_view is not None:
            ai_view.bgcolor = tokens["canvas"]
        if ai_composer_card is not None:
            ai_composer_card.bgcolor = tokens["composer"]
            ai_composer_card.border = ft.Border.all(1, tokens["composer_border"])
        if ai_header_card is not None:
            ai_header_card.bgcolor = tokens["assistant_bg"]
            ai_header_card.border = ft.Border.all(1, tokens["assistant_border"])
        ai_send_button.icon_color = "white"
        ai_send_button.style = ft.ButtonStyle(
            bgcolor=PRIMARY,
            shape=ft.CircleBorder(),
            padding=ft.Padding.all(12),
        )

        for ctrl in [ai_refresh_button, ai_reset_button, ai_new_chat_button]:
            ctrl.style = ft.ButtonStyle(
                color=tokens["text"],
                bgcolor="transparent",
                shape=ft.RoundedRectangleBorder(radius=999),
                padding=ft.Padding.symmetric(horizontal=14, vertical=9),
            )

        ai_input.bgcolor = tokens["composer"]
        ai_input.border_color = tokens["composer_border"]
        ai_input.focused_border_color = PRIMARY
        ai_input.color = tokens["text"]
        ai_input.cursor_color = PRIMARY
        ai_input.hint_style = ft.TextStyle(color=tokens["muted"], size=14)
        ai_model_select.color = tokens["text"]
        ai_model_select.border_color = "transparent"
        ai_model_select.focused_border_color = PRIMARY
        ai_model_select.label_style = ft.TextStyle(color=tokens["muted"], size=11)
        ai_model_select.bgcolor = "transparent"
        ai_tools_button.content = (
            ft.Icon(ft.Icons.HANDYMAN_OUTLINED, size=20, color=tokens["text"])
            if sidebar_collapsed["value"]
            else ft.Text("Инструменты", size=12, weight=ft.FontWeight.W_700, color=tokens["text"])
        )
        ai_tools_button.tooltip = "Инструменты"
        ai_tools_button.menu_bgcolor = "#22272B" if current_dark["value"] else "#FBFCFF"
        ai_top_title.color = tokens["text"]
        ai_top_subtitle.color = tokens["muted"]
        ai_busy_label.color = tokens["muted"]
        ai_diag_subtitle.color = tokens["muted"]
        ai_empty_state.content.bgcolor = tokens["assistant_bg"]
        ai_empty_state.content.border = ft.Border.all(1, tokens["assistant_border"])
        if isinstance(ai_empty_state.content, ft.Container) and isinstance(ai_empty_state.content.content, ft.Column):
            for child in ai_empty_state.content.content.controls:
                if isinstance(child, ft.Text):
                    child.color = tokens["text"] if child.size and child.size >= 20 else tokens["muted"]

        def refresh_ai_control_theme(control) -> None:
            if control is None:
                return
            if isinstance(control, ft.Container):
                if control.data == "ai_user_bubble":
                    control.bgcolor = tokens["user_bg"]
                    control.border = None
                elif control.data in {"ai_assistant_bubble", "ai_thinking_wrap"}:
                    control.bgcolor = tokens["assistant_bg"]
                    if control.border is not None:
                        control.border = ft.Border.all(1, tokens["assistant_border"])
            if isinstance(control, ft.Text):
                if control.data == "ai_user_text":
                    control.color = tokens["user_text"]
                elif control.data == "ai_note_text":
                    control.color = tokens["muted"]
            for attr in ("content", "title", "subtitle", "leading", "trailing"):
                if hasattr(control, attr):
                    refresh_ai_control_theme(getattr(control, attr))
            for attr in ("controls", "actions"):
                if hasattr(control, attr):
                    for child in getattr(control, attr) or []:
                        refresh_ai_control_theme(child)

        refresh_ai_control_theme(ai_messages)
        for payload in ai_state["stream_controls"].values():
            if isinstance(payload, dict):
                thinking_list = payload.get("thinking_list")
                if isinstance(thinking_list, ft.ListView):
                    for control in thinking_list.controls:
                        if isinstance(control, ft.Text):
                            control.color = tokens["muted"]

    def ai_refresh_tools_menu() -> None:
        debug_label = "Debug режим ✓" if ai_state["debug_enabled"] else "Debug режим"
        menu_text_color = "#F5F5F7" if current_dark["value"] else "#1D1D1F"
        ai_tools_button.items = [
            ft.PopupMenuItem(
                content=ft.Row(
                    [
                        ft.Icon(ft.Icons.CHECK if ai_state["debug_enabled"] else ft.Icons.BUG_REPORT_OUTLINED, size=18, color=menu_text_color),
                        ft.Text(debug_label, size=14, color=menu_text_color, weight=ft.FontWeight.W_600),
                    ],
                    spacing=12,
                ),
                on_click=lambda _e: ai_toggle_debug(),
            ),
            ft.PopupMenuItem(
                content=ft.Row(
                    [
                        ft.Icon(ft.Icons.SYNC, size=18, color=menu_text_color),
                        ft.Text("Проверить Ollama", size=14, color=menu_text_color, weight=ft.FontWeight.W_600),
                    ],
                    spacing=12,
                ),
                on_click=lambda _e: ai_refresh_status(),
            ),
            ft.PopupMenuItem(
                content=ft.Row(
                    [
                        ft.Icon(ft.Icons.DELETE_OUTLINE, size=18, color=menu_text_color),
                        ft.Text("Очистить контекст", size=14, color=menu_text_color, weight=ft.FontWeight.W_600),
                    ],
                    spacing=12,
                ),
                on_click=ai_reset,
            ),
        ]
        if status_ui_ready["value"] and page.controls:
            page.update(ai_tools_button)

    def ai_toggle_debug() -> None:
        ai_state["debug_enabled"] = not ai_state["debug_enabled"]
        ai_refresh_tools_menu()
        if ai_state["debug_enabled"]:
            append_log("[AI DEBUG] Debug режим включён.")
            toast("Debug режим включён. Лог доступен по F2.")
        else:
            toast("Debug режим выключен.")

    def ai_trace_row(title: str, detail: str) -> ft.Control:
        return ft.Container(
            padding=12,
            border_radius=16,
            bgcolor="#FFFDF8" if not current_dark["value"] else "#232A2F",
            border=ft.Border.all(1, "#D6D0C4" if not current_dark["value"] else "#3D474B"),
            content=ft.Column(
                [
                    ft.Text(title, size=12, weight=ft.FontWeight.W_800, color=TEXT),
                    ft.Text(detail, size=11, color=MUTED, max_lines=4, overflow=ft.TextOverflow.ELLIPSIS),
                ],
                spacing=4,
            ),
        )

    def ai_set_phase(phase: str, detail: str) -> None:
        ai_state["phase"] = phase
        ai_phase_label.value = f"Текущий этап: {phase}"
        ai_phase_detail.value = detail
        if status_ui_ready["value"] and page.controls:
            page.update(ai_phase_label, ai_phase_detail)

    def ai_add_trace(title: str, detail: str) -> None:
        clean_detail = (detail or "").strip() or "—"
        ai_state["trace"].append((title, clean_detail))
        if len(ai_state["trace"]) > 10:
            ai_state["trace"] = ai_state["trace"][-10:]
        ai_trace.controls = [ai_trace_row(item_title, item_detail) for item_title, item_detail in ai_state["trace"]]
        if status_ui_ready["value"] and page.controls:
            page.update(ai_trace)

    def ai_apply_selected_model(model_name: str | None) -> None:
        model = (model_name or "").strip() or ai_state["selected_model"] or OLLAMA_MODEL
        ai_state["selected_model"] = model
        ai_runtime_model.value = f"Модель: {model}"
        ai_model_select.value = model
        if ai_chat_module is not None:
            ai_chat_module.OLLAMA_MODEL = model

    def ai_debug_log(title: str, payload: object = "") -> None:
        if not ai_state["debug_enabled"]:
            return
        if isinstance(payload, (dict, list)):
            body = json.dumps(payload, ensure_ascii=False, indent=2)
        elif payload is None:
            body = ""
        else:
            body = str(payload)
        block = f"[AI DEBUG] {title}"
        if body:
            block += f"\n{body}"
        summary = body.splitlines()[0][:180] if body else "Событие без деталей"
        title_map = {
            "Новый запрос": ("prepare", "Собираю payload и фиксирую активную модель."),
            "Контекст сессии": ("context", "Собираю контекст диалога и выбранный период."),
            "Этап A: запрос к модели": ("plan", "Модель решает, нужны ли данные из инструментов."),
            "DATA_REQUEST": ("tools", "AI запросил данные из локальных инструментов."),
            "Результаты инструментов": ("tools-ready", "Инструменты вернули данные для анализа."),
            "Сформированный data block": ("compose", "Подготавливаю данные для финального ответа модели."),
            "Этап B: финальный запрос к модели": ("answer", "Модель строит итоговый аналитический ответ."),
            "Финальный ответ для UI": ("done", "Ответ готов и будет показан в чате."),
            "Пользователь отменил AI-запрос": ("cancelled", "Текущий AI-запрос отменён пользователем."),
        }
        phase_payload = title_map.get(title)
        ui_queue.put(("ai_step", title, summary, phase_payload))
        ui_queue.put(("log_append", block))

    def ai_load_installed_models() -> tuple[list[str], str | None]:
        try:
            response = requests.get(f"{OLLAMA_HOST}/api/tags", timeout=5)
            response.raise_for_status()
            data = response.json()
        except requests.exceptions.RequestException as exc:
            return [], f"Ollama недоступен по адресу {OLLAMA_HOST}: {exc}"

        models = []
        for item in data.get("models", []) or []:
            if isinstance(item, dict) and item.get("name"):
                models.append(str(item["name"]))
        models = sorted(set(models))
        return models, None

    def ai_refresh_status(push_update: bool = True) -> None:
        selected_from_ui = (ai_model_select.value or "").strip()
        if selected_from_ui:
            ai_apply_selected_model(selected_from_ui)
        if AI_IMPORT_ERROR:
            ai_state["available"] = False
            ai_state["status_message"] = f"Не удалось загрузить AI-модуль: {AI_IMPORT_ERROR}"
            ai_model_select.options = []
        elif not check_ollama_model:
            ai_state["available"] = False
            ai_state["status_message"] = "AI-модуль недоступен."
        else:
            models, error = ai_load_installed_models()
            ai_state["available_models"] = models
            ai_model_select.options = [ft.dropdown.Option(model, model) for model in models]
            if error:
                ai_state["available"] = False
                ai_state["status_message"] = error
            else:
                current_model = ai_state["selected_model"]
                if models:
                    if current_model not in models:
                        ai_apply_selected_model(models[0])
                    else:
                        ai_apply_selected_model(current_model)
                ok, message = check_ollama_model()
                ai_state["available"] = ok
                ai_state["status_message"] = message if message else "Ollama запущен, модель установлена, локальный анализ доступен."

        ai_model_select.disabled = ai_state["busy"] or not bool(ai_state["available_models"])
        if ai_state["available_models"]:
            ai_runtime_note.value = f"Установлено моделей: {len(ai_state['available_models'])}. Можно переключать без перезапуска приложения."
        else:
            ai_runtime_note.value = "AI-ассистент работает локально и не отправляет отчёты во внешний API."

        ai_diag_title.value = "Локальная модель готова к работе" if ai_state["available"] else "Требуется внимание к AI-окружению"
        ai_diag_subtitle.value = ai_state["status_message"]
        ai_status_dot.bgcolor = "#4FB59E" if ai_state["available"] else DANGER
        ai_status_label.value = "Статус: готово" if ai_state["available"] else "Статус: требуется настройка"
        if push_update and status_ui_ready["value"] and page.controls:
            page.update(ai_diag_title, ai_diag_subtitle, ai_status_dot, ai_status_label, ai_runtime_host, ai_runtime_model, ai_runtime_note, ai_model_select)

    def ai_set_busy(busy: bool, detail: str) -> None:
        ai_state["busy"] = busy
        ai_busy_ring.visible = busy
        ai_busy_label.value = detail
        ai_send_button.disabled = busy
        ai_input.disabled = busy
        ai_reset_button.disabled = busy
        ai_refresh_button.disabled = busy
        ai_model_select.disabled = busy or not bool(ai_state["available_models"])
        if not busy and ai_state["phase"] not in ("done", "cancelled"):
            ai_set_phase("idle", "Модель ожидает новый запрос.")
        if status_ui_ready["value"] and page.controls:
            page.update(ai_busy_ring, ai_busy_label, ai_send_button, ai_input, ai_reset_button, ai_refresh_button, ai_model_select)

    def ai_reset(_e=None) -> None:
        if ai_state["active_request_id"] is not None:
            ai_state["cancelled_request_ids"].add(ai_state["active_request_id"])
            ai_state["request_started_at"].pop(ai_state["active_request_id"], None)
        ai_state["active_request_id"] = None
        ai_state["conversation_history"] = []
        ai_state["trace"] = []
        ai_trace.controls = []
        ai_messages.controls = [ai_empty_state]
        ai_sync_layout()
        ai_set_phase("idle", "Контекст и диагностическая лента очищены.")
        ai_busy_label.value = "Контекст очищен. Ассистент готов к новому диалогу."
        if status_ui_ready["value"] and page.controls:
            page.update(ai_body, ai_messages, ai_messages_shell, ai_welcome_shell, ai_composer_shell, ai_busy_label, ai_trace, ai_phase_label, ai_phase_detail)

    def ai_model_changed(e: ft.ControlEvent) -> None:
        selected = (e.control.value or "").strip()
        if not selected:
            return
        ai_apply_selected_model(selected)
        # Заново проверяем модель: без поддержки инструментов ассистент работать не может.
        ai_refresh_status()
        ai_busy_label.value = f"Активная модель: {selected}" if ai_state["available"] else ai_state["status_message"]
        if status_ui_ready["value"] and page.controls:
            page.update(ai_runtime_model, ai_busy_label, ai_model_select)

    def ai_submit(prompt_text: str) -> None:
        prompt = (prompt_text or "").strip()
        if not prompt or ai_state["busy"]:
            return
        ai_apply_selected_model(ai_model_select.value or ai_state["selected_model"])
        ai_state["request_seq"] += 1
        request_id = ai_state["request_seq"]
        ai_state["active_request_id"] = request_id
        ai_state["request_started_at"][request_id] = time.perf_counter()
        runtime_model = ai_chat_module.OLLAMA_MODEL if ai_chat_module is not None else ai_state["selected_model"]
        ai_debug_log("Новый запрос", {"request_id": request_id, "prompt": prompt, "model": ai_state["selected_model"], "runtime_model": runtime_model, "host": OLLAMA_HOST})
        ai_append_message("user", prompt)
        ai_input.value = ""
        ai_sync_layout()
        ai_set_busy(True, "Ассистент читает вопрос и строит план анализа.")
        set_status("AI анализирует запрос", prompt, ACCENT, busy=True)
        page.update(ai_body, ai_messages, ai_messages_shell, ai_welcome_shell, ai_composer_shell, ai_input, state, status_detail, spin, cancel_row)

        def worker() -> None:
            def is_cancelled() -> bool:
                return request_id in ai_state["cancelled_request_ids"] or ai_state["active_request_id"] != request_id

            if not ai_state["available"] or not AIAgent or not ai_tools:
                ai_debug_log("AI недоступен", ai_state["status_message"] or "AI недоступен.")
                ui_queue.put(("ai_finish", request_id, prompt, None, ai_state["status_message"] or "AI недоступен."))
                return

            def on_event(kind: str, payload: object) -> None:
                if is_cancelled():
                    return
                if kind == "status":
                    ui_queue.put(("ai_stream_start", request_id, str(payload)))
                elif kind == "thinking":
                    ui_queue.put(("ai_thinking_chunk", request_id, payload))
                elif kind == "answer_chunk":
                    ui_queue.put(("ai_stream_chunk", request_id, payload))
                elif kind == "answer_reset":
                    ui_queue.put(("ai_stream_reset", request_id))
                elif kind == "tool":
                    ui_queue.put(("ai_step", f"Данные: {payload['name']}", f"{payload['seconds']} с", ("tools", "Ассистент получил данные из программы.")))
                    ai_debug_log(f"Инструмент {payload['name']}", payload)
                elif kind == "model":
                    ai_debug_log(f"Ответ модели, шаг {payload['step']}", payload)
                    if payload.get("context_overflow"):
                        ai_debug_log("Окно контекста почти заполнено", "Ollama мог обрезать начало запроса.")

            try:
                ai_apply_selected_model(ai_model_select.value or ai_state["selected_model"])
                if is_cancelled():
                    return
                ui_queue.put(("ai_stream_start", request_id, "Думаю над вопросом…"))
                result = AIAgent(ROOT, tools=ai_tools).run(prompt, ai_state["conversation_history"], on_event=on_event, should_cancel=is_cancelled)
                if result.cancelled or is_cancelled():
                    ai_debug_log("AI запрос отменён", {"request_id": request_id})
                    return
                ai_debug_log("Итог запроса", {"steps": result.steps, "tool_calls": result.tool_calls, "seconds": result.seconds, "model": result.model})
                if result.answer:
                    ui_queue.put(("ai_finish", request_id, prompt, result.answer, result.model or ai_state["selected_model"]))
                else:
                    ui_queue.put(("ai_finish", request_id, prompt, None, result.error or "Не удалось получить ответ от модели."))
            except Exception as exc:
                ai_debug_log("Исключение AI-пайплайна", repr(exc))
                ui_queue.put(("ai_finish", request_id, prompt, None, f"Ошибка AI-анализа: {exc}"))

        page.run_thread(worker)

    def reveal_path(path: Path, label: str | None = None) -> None:
        target_label = label or path.name or str(path)
        set_status("Открытие файла", target_label, PRIMARY, busy=False)
        open_path(path)

    def cancel_current_action(_e=None) -> None:
        proc = active_proc["proc"]
        if ai_state["busy"] and ai_state["active_request_id"] is not None:
            request_id = ai_state["active_request_id"]
            ai_state["cancelled_request_ids"].add(request_id)
            ai_state["request_started_at"].pop(request_id, None)
            ai_state["active_request_id"] = None
            ai_set_busy(False, "AI-запрос отменён. Можно отправить новый вопрос.")
            ai_append_message("assistant", "Текущий AI-запрос остановлен пользователем.", "Запрос отменён")
            ai_debug_log("Пользователь отменил AI-запрос", {"request_id": request_id})
            set_status("AI остановлен", "Текущий запрос отменён пользователем.", DANGER, busy=False)
            if status_ui_ready["value"] and page.controls:
                page.update(ai_messages, ai_busy_label, state, status_detail, spin, cancel_row)
            return
        if proc is None and not job_lock.locked():
            set_status("Нет задачи", "Нечего отменять.", MUTED, busy=False)
            return
        cancel_state["requested"] = True
        job_cancel_event.set()
        set_status("Отмена...", cancel_state["title"] or "Останавливаем задачу.", DANGER, busy=True)
        if prompt_dialog is not None and prompt_dialog.open:
            prompt_dialog.open = False
            holder = prompt_state["holder"]
            event = prompt_state["event"]
            if isinstance(holder, dict):
                holder["value"] = "0"
            if isinstance(event, threading.Event):
                event.set()
            if page.controls:
                page.update(prompt_dialog)
        if proc is None:
            return
        try:
            kill_process_tree(proc)
        except Exception:
            try:
                proc.kill()
            except Exception:
                pass

    prompt_label = ft.Text("", color=TEXT, size=14, weight=ft.FontWeight.W_700)
    prompt_hint = ft.Text("", color=MUTED, size=12)
    prompt_input = ft.TextField(label="Сумма маркетинга", value="0", autofocus=True)
    prompt_context = ft.Text("", color=MUTED, size=12)
    prompt_state: dict[str, object | None] = {"holder": None, "event": None}
    prompt_dialog: ft.AlertDialog | None = None
    ui_queue: queue.Queue = queue.Queue()

    def set_dashboard_chart_feedback(kind: str, title: str, caption: str = "", meta_caption: str = "") -> None:
        dashboard_chart_state["mode"] = kind
        icon_specs = {
            "empty": (ft.Icons.AUTO_GRAPH_OUTLINED, semantic_surface("info", dark=False), "#4A6FA5"),
            "loading": (ft.Icons.HOURGLASS_TOP_ROUNDED, semantic_surface("neutral_alt", dark=False), PRIMARY),
            "error": (ft.Icons.ERROR_OUTLINE_ROUNDED, semantic_surface("danger", dark=False), DANGER),
            "ready": (ft.Icons.CHECK_CIRCLE_OUTLINE_ROUNDED, semantic_surface("positive", dark=False), PRIMARY),
        }
        icon_name, card_tone, accent_color = icon_specs.get(kind, icon_specs["empty"])
        dashboard_chart_state_indicator.bgcolor = card_tone
        dashboard_chart_state_indicator.content = (
            ft.ProgressRing(width=18, height=18, stroke_width=2.2, color=accent_color)
            if kind == "loading"
            else ft.Icon(icon_name, size=20, color=accent_color)
        )
        dashboard_chart_state_card.bgcolor = semantic_surface("neutral_soft", dark=False)
        dashboard_chart_state_card.border = ft.Border.all(1, BORDER)
        dashboard_chart_state_title.value = title
        dashboard_chart_state_caption.value = caption
        dashboard_chart_state_card.visible = kind != "ready"
        dashboard_chart_meta_card.visible = kind == "ready"
        if kind == "ready":
            dashboard_chart_meta_title.value = title
            dashboard_chart_meta_caption.value = meta_caption or caption
        if current_dark["value"] and "apply_control_theme" in locals():
            apply_control_theme(dashboard_chart_panel, True)

    set_dashboard_chart_feedback("empty", "Постройте график", "Выберите метрику и период, затем постройте график.")

    def infer_prompt_period(source_line: str) -> str | None:
        line = source_line or ""
        month_match = re.search(r"\b(0?[1-9]|1[0-2])\b", line)
        year_match = re.search(r"\b(20\d{2})\b", line)
        if month_match and year_match:
            month_num = int(month_match.group(1))
            if 1 <= month_num <= 12:
                return f"{MONTHS[month_num - 1]} {year_match.group(1)}"

        named_month_pattern = "|".join(MONTHS)
        named_match = re.search(rf"({named_month_pattern})\s+(20\d{{2}})", line, flags=re.IGNORECASE)
        if named_match:
            month_name = named_match.group(1)
            year = named_match.group(2)
            normalized = next((m for m in MONTHS if m.lower() == month_name.lower()), month_name)
            return f"{normalized} {year}"

        if cancel_state["title"] == "Месячный отчёт":
            month_value = int(monthly_m.value or today.month)
            year_value = int(monthly_y.value or today.year)
            return f"{MONTHS[month_value - 1]} {year_value}"

        if cancel_state["title"] == DASHBOARD_REFRESH_TITLE:
            month_value = int(dashboard_m.value or today.month)
            year_value = int(dashboard_y.value or today.year)
            return f"{MONTHS[month_value - 1]} {year_value}"

        return None

    def apply_prompt_dialog_theme() -> None:
        dark = current_dark["value"]
        prompt_input.color = "#F3F4F6" if dark else TEXT
        prompt_input.bgcolor = "#232A2F" if dark else "#FFFDF8"
        prompt_input.border_color = "#3D474B" if dark else "#D6D0C4"
        prompt_input.focused_border_color = PRIMARY
        prompt_input.label_style = ft.TextStyle(color="#B6C0BC" if dark else MUTED, size=12)
        prompt_label.color = "#F3F4F6" if dark else TEXT
        prompt_hint.color = "#B6C0BC" if dark else MUTED
        prompt_context.color = "#B6C0BC" if dark else MUTED
        if prompt_dialog is not None:
            prompt_dialog.bgcolor = "#232A2F" if dark else "#FFFDF8"
            prompt_dialog.title = ft.Text(
                "Нужны расходы на маркетинг",
                color="#F3F4F6" if dark else TEXT,
                weight=ft.FontWeight.W_800,
            )

    def apply_discount_dialog_theme() -> None:
        dark = current_dark["value"]
        divider_color = "#3D474B" if dark else "#D6D0C4"
        if isinstance(discount_dialog_state.get("plan"), dict):
            render_discount_dialog()
        discount_dialog_badge.color = "#B6C0BC" if dark else MUTED
        discount_dialog_status.color = "#B6C0BC" if dark else MUTED
        if discount_dialog is not None:
            discount_dialog.bgcolor = "#232A2F" if dark else "#FFFDF8"
            discount_dialog.title = ft.Text(
                "Заявки на скидку",
                color="#F3F4F6" if dark else TEXT,
                weight=ft.FontWeight.W_800,
            )
            content = discount_dialog.content
            if isinstance(content, ft.Container) and isinstance(content.content, ft.Column):
                for child in content.content.controls:
                    if isinstance(child, ft.Divider):
                        child.color = divider_color

    def request_prompt(label: str, source_line: str) -> str:
        holder: dict[str, str] = {"value": "0"}
        event = threading.Event()
        ui_queue.put(("prompt", label, source_line, holder, event))
        while not event.wait(timeout=0.1):
            if job_cancel_event.is_set():
                return "0"
        return holder["value"]

    def collect_dashboard_chart_data(
        from_month: int,
        from_year: int,
        to_month: int,
        to_year: int,
    ) -> tuple[list[str], list[float], list[float], list[float], float]:
        periods = iter_period_range(from_month, from_year, to_month, to_year)
        if is_wb_store():
            return wb_dash.chart_data(periods)
        labels: list[str] = []
        revenue_values: list[float] = []
        profit_values: list[float] = []
        margin_values: list[float] = []
        latest_source_mtime = 0.0
        for month_value, year_value in periods:
            period_report_path = report_path_for_period(month_value, year_value)
            metrics, _built = read_business_metrics(period_report_path)
            if not metrics:
                continue
            try:
                latest_source_mtime = max(latest_source_mtime, period_report_path.stat().st_mtime)
            except OSError:
                pass
            revenue = as_float(metrics.get("Общая выручка"))
            profit = as_float(metrics.get("Чистая прибыль"))
            margin = as_float(metrics.get("Рентабельность по чистой прибыли (Net Profit Margin) %"))
            if revenue is None and profit is None and margin is None:
                continue
            labels.append(f"{MONTHS[month_value - 1][:3]} {str(year_value)[-2:]}")
            revenue_values.append(revenue or 0.0)
            profit_values.append(profit or 0.0)
            margin_values.append(margin or 0.0)
        return labels, revenue_values, profit_values, margin_values, latest_source_mtime

    def render_dashboard_chart(from_month: int, from_year: int, to_month: int, to_year: int) -> None:
        labels, revenue_values, profit_values, margin_values, latest_source_mtime = collect_dashboard_chart_data(from_month, from_year, to_month, to_year)

        if not labels:
            dashboard_chart.visible = False
            dashboard_chart_note.value = "В выбранном диапазоне нет отчётов с бизнес-показателями."
            dashboard_chart_state["built"] = False
            set_dashboard_chart_feedback("empty", "Нет данных для графика", "В выбранном диапазоне пока нет готовых отчётов с бизнес-показателями.")
            return

        fig_bg = "#20262B"
        ax_bg = "#232A2F"
        text_color = "#F3F4F6"
        muted_color = "#B6C0BC"
        grid_color = "#3D474B"
        chart_kind = dashboard_chart_type.value or "profit"
        chart_dir = ROOT / ".cache"
        chart_dir.mkdir(exist_ok=True)

        plt.close("all")
        fig, ax = plt.subplots(figsize=(9.6, 4.4), dpi=140)
        fig.patch.set_facecolor(fig_bg)
        ax.set_facecolor(ax_bg)
        title, legend_label, note_label, line_color = chart_definition(chart_kind)
        if chart_kind == "revenue":
            series = revenue_values
        elif chart_kind == "margin":
            series = margin_values
        else:
            series = profit_values
        ax.plot(labels, series, color=line_color, linewidth=2.8, marker="o", label=legend_label)
        ax.grid(axis="y", color=grid_color, alpha=0.35, linewidth=0.8)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.spines["left"].set_color(grid_color)
        ax.spines["bottom"].set_color(grid_color)
        ax.tick_params(axis="x", colors=muted_color, labelsize=9)
        ax.tick_params(axis="y", colors=muted_color, labelsize=9)
        ax.set_title(title, loc="left", color=text_color, fontsize=13, pad=12)
        legend = ax.legend(
            frameon=False,
            loc="upper center",
            bbox_to_anchor=(0.5, -0.18),
            ncol=1,
            borderaxespad=0.0,
        )
        for text in legend.get_texts():
            text.set_color(text_color)
        fig.subplots_adjust(bottom=0.26, top=0.86, left=0.07, right=0.98)
        cache_stamp = int(latest_source_mtime) if latest_source_mtime else int(datetime.now().timestamp())
        chart_render_path = chart_dir / f"dashboard_trend_{active_store['value']}_{chart_kind}_{cache_stamp}.png"
        fig.savefig(chart_render_path, facecolor=fig_bg, bbox_inches="tight")
        plt.close(fig)

        dashboard_chart.src = str(chart_render_path.resolve())
        dashboard_chart.visible = True
        dashboard_chart.width = 980
        dashboard_chart.height = 420
        dashboard_chart_note.value = f"Период: {MONTHS[from_month - 1]} {from_year} - {MONTHS[to_month - 1]} {to_year}. Метрика: {note_label}."
        dashboard_chart_state["built"] = True
        set_dashboard_chart_feedback(
            "ready",
            title,
            f"{MONTHS[from_month - 1]} {from_year} - {MONTHS[to_month - 1]} {to_year}",
            f"Метрика: {note_label}. Период: {MONTHS[from_month - 1]} {from_year} - {MONTHS[to_month - 1]} {to_year}.",
        )

    def open_interactive_dashboard_chart(_e=None) -> None:
        chart_error = valid_dashboard_chart_period()
        if chart_error:
            toast(chart_error, True)
            return
        missing_periods = get_missing_chart_reports()
        if missing_periods:
            build_missing_chart_reports(missing_periods, "interactive")
            return
        from_month = int(dashboard_chart_from_m.value or prev.month)
        from_year = int(dashboard_chart_from_y.value or prev.year)
        to_month = int(dashboard_chart_to_m.value or prev.month)
        to_year = int(dashboard_chart_to_y.value or prev.year)
        labels, revenue_values, profit_values, margin_values, latest_source_mtime = collect_dashboard_chart_data(from_month, from_year, to_month, to_year)
        if not labels:
            set_dashboard_chart_feedback("empty", "Нет данных для интерактивного графика", "Сначала подготовьте отчёты за выбранный диапазон.")
            toast("Для интерактивного графика пока нет данных.", True)
            return
        try:
            import plotly.graph_objects as go
        except Exception:
            toast("Для интерактивного графика нужен пакет plotly. Установите зависимости.", True)
            return

        chart_kind = dashboard_chart_type.value or "profit"
        title, legend_label, note_label, line_color = chart_definition(chart_kind)
        if chart_kind == "revenue":
            series = revenue_values
            hover_template = "%{x}<br>Выручка: %{y:,.0f} ₽<extra></extra>"
        elif chart_kind == "margin":
            series = margin_values
            hover_template = "%{x}<br>Рентабельность: %{y:.2f}%<extra></extra>"
        else:
            series = profit_values
            hover_template = "%{x}<br>Чистая прибыль: %{y:,.0f} ₽<extra></extra>"

        paper_bg = "#1A1F23"
        plot_bg = "#232A2F"
        text_color = "#F3F4F6"
        grid_color = "#3D474B"

        fig = go.Figure()
        fig.add_trace(
            go.Scatter(
                x=labels,
                y=series,
                mode="lines+markers",
                name=legend_label,
                line={"color": line_color, "width": 3},
                marker={"size": 10, "color": line_color},
                hovertemplate=hover_template,
            )
        )
        fig.update_layout(
            title={"text": title, "x": 0.01, "xanchor": "left"},
            paper_bgcolor=paper_bg,
            plot_bgcolor=plot_bg,
            font={"color": text_color},
            hovermode="x",
            margin={"l": 50, "r": 24, "t": 60, "b": 80},
            legend={"orientation": "h", "yanchor": "top", "y": -0.18, "xanchor": "center", "x": 0.5},
        )
        fig.update_xaxes(showgrid=False, color=text_color)
        fig.update_yaxes(gridcolor=grid_color, zerolinecolor=grid_color, color=text_color)

        chart_dir = ROOT / ".cache"
        chart_dir.mkdir(exist_ok=True)
        cache_stamp = int(latest_source_mtime) if latest_source_mtime else int(datetime.now().timestamp())
        html_path = chart_dir / f"dashboard_interactive_{chart_kind}_{cache_stamp}.html"
        fig.write_html(str(html_path), include_plotlyjs="cdn", full_html=True)
        open_path(html_path)
        dashboard_chart_note.value = f"Открыт интерактивный график: {note_label}."
        set_dashboard_chart_feedback(
            "ready",
            title,
            f"{MONTHS[from_month - 1]} {from_year} - {MONTHS[to_month - 1]} {to_year}",
            f"Интерактивный график открыт. Метрика: {note_label}.",
        )
        set_status("График открыт", f"{MONTHS[from_month - 1]} {from_year} - {MONTHS[to_month - 1]} {to_year}", PRIMARY, busy=False)
        if page.controls:
            page.update(dashboard_chart_panel)

    # ---------------------------------------------------------------- Бизнес-сводка
    dashboard_state: dict[str, object] = {"metrics": None, "promo": {}, "campaigns": [], "built_at": None, "period": "",
                                          "store": "ozon"}
    DASHBOARD_HELP = {
        "Выручка": "Общая выручка за период по месячному отчёту. Основной показатель оборота.",
        "Чистая прибыль": "Прибыль после вычета себестоимости, комиссий, логистики, маркетинга и прочих расходов.",
        "Себестоимость": "Итоговая товарная себестоимость реализованных заказов за период.",
        "Маржа": "Рентабельность по чистой прибыли: чистая прибыль / выручка × 100%.",
        "Средний чек": "Выручка / количество заказов.",
        "Всего заказов": "Общее количество заказов за период.",
        "В пути": "Заказы в статусе доставки на момент формирования отчёта.",
        "Доставлено": "Успешно доставленные заказы.",
        "Отменено": "Отменённые заказы. Рост показателя требует проверки причин отмен.",
        "Возвраты": "Возвращённые заказы.",
        "Продвижение Ozon": "Расходы на внутреннее продвижение на Ozon за период.",
        "Внешний маркетинг": "Затраты на внешний маркетинг вне Ozon, введённые при сборке отчёта.",
        "Хранение FBO": "Затраты на хранение товаров на складах Ozon по схеме FBO.",
        "Комиссии Ozon": "Комиссии маркетплейса: доля в выручке и сумма.",
        "Логистика": "Логистические расходы Ozon: доля в выручке и сумма.",
        "Валовая прибыль": "Выручка минус себестоимость товара, до прочих операционных расходов.",
        "Операционные расходы": "Операционные расходы периода, влияющие на чистую прибыль.",
        "Вне заказов": "Расходы Ozon за месяц, не привязанные к заказам: кросс-докинг, обработка поставки, ранний вывод средств и т.п. "
                       "Вычитаются из чистой прибыли; подробности — на листе «Расходы вне заказов» в отчёте.",
        "Выплачено": "Реальный перевод от Ozon за период — из отчёта о балансе (Финансы → Баланс), а не расчётная величина.",
        "Начислено Ozon": "Начисления Ozon за период по календарным дням (отчёт о балансе). Сравните с нашим расчётом — методика группировки разная, небольшое расхождение нормально.",
        "Ранний вывод": "Комиссия Ozon за досрочный вывод средств.",
        "Рекламный расход": "Расход по кампаниям с ненулевым расходом за период.",
        "Выручка с рекламы": "Выручка, атрибутированная рекламным кампаниям.",
        "ROAS": "Выручка с рекламы / рекламный расход.",
        "ДРР": "Рекламный расход / выручка с рекламы × 100%. Чем ниже, тем эффективнее.",
        "CTR": "Клики / показы × 100%: насколько объявление привлекает клик.",
        "CR": "Заказы / клики × 100%: как трафик превращается в покупку.",
        "CPC": "Рекламный расход / клики.",
        "CPA": "Стоимость заказа: рекламный расход / заказы.",
        "CPM": "Стоимость 1000 показов.",
    }

    def dash_pal() -> dict:
        dark = current_dark["value"]
        return {
            "dark": dark,
            "ink": "#F3F4F6" if dark else TEXT,
            "muted": "#9AA5A1" if dark else MUTED,
            "surface": semantic_surface("neutral", dark=dark),
            "border": "#343C42" if dark else "#E6E8EE",
            "separator": "#2C3338" if dark else "#EEF0F4",
            "good": "#5FD0B1" if dark else "#0E7A62",
            "warn": "#E5B567" if dark else "#9A6420",
            "bad": "#FF8F82" if dark else "#C2412F",
            "accent": "#7FB2FF" if dark else "#2A6FD6",
        }

    def dtext(value: str, size: int = 13, color: str | None = None, weight=None, **kwargs) -> ft.Text:
        return ft.Text(value, size=size, color=color or dash_pal()["ink"], weight=weight, data="preserve_excel_text", **kwargs)

    def dpanel(content: ft.Control, padding: int = 20, **kwargs) -> ft.Container:
        pal = dash_pal()
        return ft.Container(content=content, padding=padding, border_radius=18, bgcolor=pal["surface"],
                            border=ft.Border.all(1, pal["border"]), data="preserve_excel_fill", **kwargs)

    def dash_section_title(title: str, caption: str = "") -> ft.Control:
        pal = dash_pal()
        items: list[ft.Control] = [dtext(title, 20, weight=ft.FontWeight.W_700)]
        if caption:
            items.append(dtext(caption, 13, pal["muted"]))
        return ft.Container(padding=ft.Padding.only(top=10, left=4), content=ft.Column(items, spacing=2))

    def dash_stat(label: str, value: str, caption: str = "", color: str | None = None, help_key: str | None = None) -> ft.Control:
        pal = dash_pal()
        controls: list[ft.Control] = [
            dtext(label, 13, pal["muted"], ft.FontWeight.W_600),
            dtext(value, 30, color, ft.FontWeight.W_700, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS),
        ]
        if caption:
            controls.append(dtext(caption, 12, pal["muted"], max_lines=1, overflow=ft.TextOverflow.ELLIPSIS))
        return ft.Container(col={"xs": 6, "md": 3}, padding=ft.Padding.symmetric(vertical=4, horizontal=4),
                            tooltip=DASHBOARD_HELP.get(help_key or label), content=ft.Column(controls, spacing=2))

    def dash_stats_panel(stats: list[ft.Control]) -> ft.Control:
        return dpanel(ft.ResponsiveRow(stats, spacing=24, run_spacing=18), padding=24)

    DASH_ROW_HEIGHT = 40

    def dash_list(title: str, rows: list[tuple], rows_total: int, col=None) -> ft.Control:
        """Сгруппированный список: подпись слева, значение справа, тонкие разделители. rows: (label, value, note, color, help_key)."""
        pal = dash_pal()
        lines: list[ft.Control] = []
        for index, (label, value, note, color, help_key) in enumerate(rows):
            right: list[ft.Control] = []
            if note:
                right.append(dtext(note, 12, pal["muted"]))
            right.append(dtext(value, 14, color, ft.FontWeight.W_600))
            lines.append(ft.Container(
                height=DASH_ROW_HEIGHT, data="preserve_excel_fill",
                border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])) if index else None,
                tooltip=DASHBOARD_HELP.get(help_key or label),
                content=ft.Row([dtext(label, 13, pal["muted"], expand=True, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS),
                                ft.Row(right, spacing=8, tight=True)],
                               vertical_alignment=ft.CrossAxisAlignment.CENTER),
            ))
        body_height = rows_total * DASH_ROW_HEIGHT + 2
        return dpanel(
            ft.Column([dtext(title, 15, weight=ft.FontWeight.W_700),
                       ft.Container(height=body_height, content=ft.Column(lines, spacing=0))], spacing=8),
            padding=ft.Padding.only(left=20, right=20, top=16, bottom=10),
            col=col or {"xs": 12, "md": 6, "xl": 3},
        )

    def dash_money(value) -> str:
        return _format_currency(value)

    def is_wb_store() -> bool:
        return active_store["value"] == "wb"

    def render_dashboard(push_update: bool = True) -> None:
        if is_wb_store():
            # Тот же экран, что у Ozon; показатели и блоки WB — из scripts/wb_dashboard.py.
            dashboard_freshness.value = wb_dash.freshness()
            dashboard_metrics_area.controls = wb_dash.controls(launch_dashboard_report)
            apply_control_theme(dashboard_metrics_area, current_dark["value"])
            if push_update and page.controls:
                page.update(dashboard_metrics_area, dashboard_freshness)
            return
        pal = dash_pal()
        metrics = dashboard_state["metrics"]
        period = str(dashboard_state["period"] or "")
        if not metrics:
            dashboard_freshness.value = f"{period} · отчёта ещё нет" if period else "Отчёта ещё нет"
            dashboard_metrics_area.controls = [dpanel(
                ft.Column([
                    ft.Icon(ft.Icons.INSERT_CHART_OUTLINED, size=44, color=pal["muted"]),
                    dtext(f"Отчёта за {period.lower()} ещё нет" if period else "Отчёта ещё нет", 18, weight=ft.FontWeight.W_700),
                    dtext("Сформируйте отчёт — сводка заполнится автоматически.", 13, pal["muted"]),
                    ft.Container(height=6),
                    ft.Button("Сформировать отчёт", icon=ft.Icons.REFRESH, bgcolor=PRIMARY, color="white", on_click=lambda _e: launch_dashboard_report()),
                ], spacing=6, horizontal_alignment=ft.CrossAxisAlignment.CENTER, tight=True),
                padding=48, alignment=ft.Alignment(0, 0),
            )]
        else:
            built_at = str(dashboard_state["built_at"] or "")
            fresh = ""
            if len(built_at) >= 16:
                fresh = f" · обновлено {built_at[8:10]}.{built_at[5:7]} в {built_at[11:16]}"
            dashboard_freshness.value = period + fresh
            revenue = as_float(metrics.get("Общая выручка")) or 0.0
            profit = as_float(metrics.get("Чистая прибыль")) or 0.0
            margin = as_float(metrics.get("Рентабельность по чистой прибыли (Net Profit Margin) %"))
            orders = metrics.get("Общее количество заказов")
            margin_color = pal["good"] if (margin or 0) >= 20 else pal["warn"] if (margin or 0) >= 5 else pal["bad"]

            def share(amount) -> str:
                value = as_float(amount)
                return _format_percent(abs(value) / revenue * 100) if value and revenue else ""

            cancelled = as_float(metrics.get("Количество отменённых заказов")) or 0
            returns = as_float(metrics.get("Количество возвращённых заказов")) or 0
            gap = as_float(metrics.get("Расхождение: баланс Ozon минус наш расчёт"))
            groups = [
                ("Заказы", [
                    ("Всего", _format_number(orders), "", None, "Всего заказов"),
                    ("В пути", _format_number(metrics.get("Количество заказов в доставке")), "", None, "В пути"),
                    ("Доставлено", _format_number(metrics.get("Количество доставленных заказов")), "", None, "Доставлено"),
                    ("Отменено", _format_number(cancelled), "", pal["bad"] if cancelled else None, "Отменено"),
                    ("Возвраты", _format_number(returns), "", pal["bad"] if returns else None, "Возвраты"),
                ]),
                ("Экономика", [
                    ("Себестоимость", dash_money(abs(as_float(metrics.get("Итоговая себестоимость")) or 0.0)), share(metrics.get("Итоговая себестоимость")), None, "Себестоимость"),
                    ("Валовая прибыль", dash_money(metrics.get("Валовая прибыль", metrics.get("COGS (валовая прибыль)"))), "", None, "Валовая прибыль"),
                    ("Операционные расходы", dash_money(metrics.get("Операционные расходы")), "", None, "Операционные расходы"),
                    ("Чистая прибыль", dash_money(profit), "", pal["good"] if profit >= 0 else pal["bad"], "Чистая прибыль"),
                ]),
                ("Расходы Ozon", [
                    ("Комиссии", dash_money(metrics.get("Комиссии Ozon сумма")), _format_percent(metrics.get("Комиссии Ozon %")), None, "Комиссии Ozon"),
                    ("Логистика", dash_money(metrics.get("Логистика сумма")), _format_percent(metrics.get("Логистика %")), None, "Логистика"),
                    ("Продвижение", dash_money(metrics.get("Продвижение Ozon")), share(metrics.get("Продвижение Ozon")), None, "Продвижение Ozon"),
                    ("Хранение FBO", dash_money(metrics.get("Расход хранения FBO")), share(metrics.get("Расход хранения FBO")), None, "Хранение FBO"),
                    ("Внешний маркетинг", dash_money(metrics.get("Внешний маркетинг")), share(metrics.get("Внешний маркетинг")), None, "Внешний маркетинг"),
                    ("Вне заказов", dash_money(metrics.get("Расходы Ozon вне заказов")) if metrics.get("Расходы Ozon вне заказов") is not None else "—",
                     share(metrics.get("Расходы Ozon вне заказов")), None, "Вне заказов"),
                ]),
                ("Выплаты Ozon", [
                    ("Выплачено", dash_money(metrics.get("Выплачено Ozon за период (реальный перевод)")), "", None, "Выплачено"),
                    ("Начислено Ozon", dash_money(metrics.get("Начислено Ozon за период (по балансу, календарные дни)")), "", None, "Начислено Ozon"),
                    ("Наш расчёт", dash_money(metrics.get("Сумма начисления по нашему расчёту (по датам отгрузки)")), "", None, "Начислено Ozon"),
                    ("Расхождение", dash_money(gap), "", pal["warn"] if gap and abs(gap) >= 1 else None, "Начислено Ozon"),
                    ("Баланс: начало → конец", f"{dash_money(metrics.get('Входящий баланс на начало периода'))} → {dash_money(metrics.get('Исходящий баланс на конец периода'))}", "", None, "Выплачено"),
                ]),
            ]
            rows_total = max(len(rows) for _title, rows in groups)
            controls: list[ft.Control] = [
                dash_stats_panel([
                    dash_stat("Выручка", dash_money(revenue), f"{_format_number(orders)} заказов"),
                    dash_stat("Чистая прибыль", dash_money(profit), "после всех расходов", pal["good"] if profit >= 0 else pal["bad"]),
                    dash_stat("Маржа", _format_percent(margin), "чистая прибыль / выручка", margin_color),
                    dash_stat("Средний чек", dash_money(metrics.get("Средний чек")), "выручка / заказы"),
                ]),
                ft.ResponsiveRow([dash_list(title, rows, rows_total) for title, rows in groups], spacing=16, run_spacing=16),
            ]
            controls += render_dashboard_promotion(pal)
            dashboard_metrics_area.controls = controls
        if push_update and page.controls:
            page.update(dashboard_metrics_area, dashboard_freshness)

    def render_dashboard_promotion(pal: dict) -> list[ft.Control]:
        summary = dashboard_state["promo"] or {}
        campaigns = list(dashboard_state["campaigns"] or [])
        if not campaigns:
            return [dash_section_title("Продвижение"),
                    dpanel(dtext("В отчёте нет рекламных кампаний с расходом за этот месяц.", 13, pal["muted"]))]
        roas = as_float(summary.get("roas")) or 0.0
        drr = as_float(summary.get("drr"))
        roas_color = pal["good"] if roas >= 5 else pal["warn"] if roas >= 3 else pal["bad"]
        drr_color = None if drr is None else pal["good"] if drr <= 20 else pal["warn"] if drr <= 35 else pal["bad"]
        roas_text = f"{roas:.2f}×".replace(".", ",")
        total_revenue = as_float((dashboard_state["metrics"] or {}).get("Общая выручка")) or 0.0
        ad_revenue = as_float(summary.get("revenue")) or 0.0
        funnel = [
            ("Показы", _format_number(summary.get("impressions")), "", None, None),
            ("Клики", _format_number(summary.get("clicks")), f"CTR {_format_percent(summary.get('ctr'))}", None, "CTR"),
            ("Заказы", _format_number(summary.get("orders")), f"CR {_format_percent(summary.get('cr'))}", None, "CR"),
        ]
        costs = [
            ("Клик", dash_money(summary.get("cpc")), "CPC", None, "CPC"),
            ("Заказ", dash_money(summary.get("cpa")), "CPA", None, "CPA"),
            ("1000 показов", dash_money(summary.get("cpm")), "CPM", None, "CPM"),
        ]
        outcome = [
            ("Средний чек", dash_money(summary.get("avg_order_value")), "", None, "Средний чек"),
            ("Доля в общей выручке", _format_percent(ad_revenue / total_revenue * 100) if total_revenue else "—", "", None, None),
            ("Кампаний с расходом", _format_number(summary.get("campaigns_count")), "", None, None),
        ]
        best_roas = max(campaigns, key=lambda item: item.get("__roas") or 0.0)
        top_revenue = max(campaigns, key=lambda item: item.get("__revenue") or 0.0)
        label_colors = {"Можно масштабировать": pal["good"], "Рабочая кампания": pal["accent"], "Нужна оптимизация": pal["warn"],
                        "Трафик без заказов": pal["bad"]}
        columns = [("Кампания", 4), ("Оценка", 3), ("Расход", 2), ("Выручка", 2), ("ROAS", 1), ("ДРР", 1), ("Заказы", 1), ("CPA", 2)]

        def cell(content: ft.Control, weight: int, right: bool = False) -> ft.Control:
            return ft.Container(expand=weight, alignment=ft.Alignment(1 if right else -1, 0), content=content)

        table: list[ft.Control] = [ft.Row(
            [cell(dtext(name, 12, pal["muted"], ft.FontWeight.W_600), weight, index >= 2) for index, (name, weight) in enumerate(columns)],
            spacing=12,
        )]
        for campaign in sorted(campaigns, key=lambda item: -(item.get("__spend") or 0.0)):
            name = str(campaign.get("Название кампании") or f"Кампания {campaign.get('ID кампании') or '—'}")
            badges = []
            if campaign is best_roas and len(campaigns) > 1:
                badges.append("лучший ROAS")
            if campaign is top_revenue and len(campaigns) > 1:
                badges.append("лидер по выручке")
            caption = f"ID {campaign.get('ID кампании') or '—'}" + (" · " + ", ".join(badges) if badges else "")
            label = str(campaign.get("__efficiency_label") or "Без оценки")
            table.append(ft.Container(
                padding=ft.Padding.symmetric(vertical=10), data="preserve_excel_fill",
                border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])),
                content=ft.Row([
                    cell(ft.Column([dtext(name, 13, weight=ft.FontWeight.W_600, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS),
                                    dtext(caption, 11, pal["muted"], max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)], spacing=1), 4),
                    cell(ft.Container(tooltip=str(campaign.get("__efficiency_reason") or ""),
                                      content=ft.Row([ft.Container(width=8, height=8, border_radius=4, bgcolor=label_colors.get(label, pal["muted"]),
                                                                   data="preserve_excel_fill"),
                                                      dtext(label, 12, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)], spacing=6, tight=True)), 3),
                    cell(dtext(dash_money(campaign.get("__spend")), 13), 2, True),
                    cell(dtext(dash_money(campaign.get("__revenue")), 13), 2, True),
                    cell(dtext(f"{(campaign.get('__roas') or 0.0):.2f}×".replace(".", ","), 13), 1, True),
                    cell(dtext(_format_percent(campaign.get("__drr")), 13), 1, True),
                    cell(dtext(_format_number(campaign.get("__orders")), 13), 1, True),
                    cell(dtext(dash_money(campaign.get("__cpa")) if campaign.get("__cpa") is not None else "—", 13), 2, True),
                ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.CENTER),
            ))
        return [
            dash_section_title("Продвижение", "Кампании с фактическим расходом за месяц"),
            dash_stats_panel([
                dash_stat("Рекламный расход", dash_money(summary.get("spend")), f"{_format_number(summary.get('campaigns_count'))} кампаний"),
                dash_stat("Выручка с рекламы", dash_money(summary.get("revenue")), f"{_format_number(summary.get('orders'))} заказов"),
                dash_stat("ROAS", roas_text, "выручка / расход", roas_color),
                dash_stat("ДРР", _format_percent(drr), "доля рекламы в выручке", drr_color),
            ]),
            ft.ResponsiveRow([
                dash_list("Воронка", funnel, 3, col={"xs": 12, "md": 4}),
                dash_list("Стоимость", costs, 3, col={"xs": 12, "md": 4}),
                dash_list("Итог рекламы", outcome, 3, col={"xs": 12, "md": 4}),
            ], spacing=16, run_spacing=16),
            dpanel(ft.Column([dtext("Кампании", 15, weight=ft.FontWeight.W_700), ft.Column(table, spacing=0)], spacing=10),
                   padding=ft.Padding.only(left=20, right=20, top=16, bottom=8)),
        ]

    def refresh_dashboard(push_update: bool = True) -> None:
        selected_month = int(dashboard_m.value or today.month)
        selected_year = int(dashboard_y.value or today.year)
        if is_wb_store():
            wb_dash.load(selected_year, selected_month)
            period_title = f"{MONTHS[selected_month - 1]} {selected_year}"
            set_status("Данные загружены" if wb_dash.data else "Нет отчёта", f"Wildberries · {period_title}",
                       PRIMARY if wb_dash.data else DANGER, busy=False)
            render_dashboard(push_update)
            return
        report_path = report_path_for_period(selected_month, selected_year)
        set_status("Чтение отчёта", f"{report_path.name}", ACCENT, busy=True)
        try:
            metrics, built_at = read_business_metrics(report_path)
            promotion_summary, promotion_campaigns = read_campaign_metrics(report_path)
        except Exception as exc:
            set_status("Ошибка чтения отчёта", str(exc), DANGER, busy=False)
            toast(f"Не удалось прочитать {report_path.name}: {exc}", True)
            return
        period_title = f"{MONTHS[selected_month - 1]} {selected_year}"
        dashboard_state.update(metrics=metrics or None, promo=promotion_summary or {}, campaigns=promotion_campaigns or [],
                               built_at=built_at, period=period_title)
        if metrics:
            set_status("Данные загружены", period_title, PRIMARY, busy=False)
        else:
            set_status("Нет отчёта", period_title, DANGER, busy=False)
        render_dashboard(push_update)
        if push_update and page.controls:
            page.update(dashboard_chart_panel)

    def valid_dashboard_chart_period() -> str | None:
        error = valid_month_year(dashboard_chart_from_m.value or "", dashboard_chart_from_y.value or "") or valid_month_year(
            dashboard_chart_to_m.value or "", dashboard_chart_to_y.value or ""
        )
        if error:
            return error
        if not (dashboard_chart_type.value or ""):
            return "Сначала выберите тип графика."
        if (int(dashboard_chart_from_y.value or 0), int(dashboard_chart_from_m.value or 0)) > (
            int(dashboard_chart_to_y.value or 0),
            int(dashboard_chart_to_m.value or 0),
        ):
            return "Начало периода графика должно быть не позже конца."
        return None

    def get_missing_chart_reports() -> list[tuple[int, int]]:
        if is_wb_store():
            return []  # месяц WB грузится минутами — график строим по уже загруженным месяцам
        from_month = int(dashboard_chart_from_m.value or prev.month)
        from_year = int(dashboard_chart_from_y.value or prev.year)
        to_month = int(dashboard_chart_to_m.value or prev.month)
        to_year = int(dashboard_chart_to_y.value or prev.year)
        missing: list[tuple[int, int]] = []
        for month_value, year_value in iter_period_range(from_month, from_year, to_month, to_year):
            if not report_path_for_period(month_value, year_value).exists():
                missing.append((month_value, year_value))
        return missing

    def get_missing_abc_reports() -> list[tuple[int, int]]:
        from_month = int(abc_fm.value or prev.month)
        from_year = int(abc_fy.value or prev.year)
        to_month = int(abc_tm.value or today.month)
        to_year = int(abc_ty.value or today.year)
        missing: list[tuple[int, int]] = []
        for month_value, year_value in iter_period_range(from_month, from_year, to_month, to_year):
            if not report_path_for_period(month_value, year_value).exists():
                missing.append((month_value, year_value))
        return missing

    def build_missing_chart_reports(missing_periods: list[tuple[int, int]], action_type: str) -> None:
        if not missing_periods:
            return
        if not job_lock.acquire(blocking=False):
            toast("Дождитесь завершения текущей операции.", True)
            return
        job_cancel_event.clear()
        chart_pending_action["type"] = action_type
        period_labels = ", ".join(f"{MONTHS[m - 1]} {y}" for m, y in missing_periods)
        if action_type == "abc":
            title = "Подготовка отчётов для ABC/XYZ"
            lead = "Готовим monthly-отчёты для ABC/XYZ"
        else:
            title = "Подготовка данных для графика"
            lead = "Подготавливаем недостающие отчёты для графика"
        cancel_state["requested"] = False
        cancel_state["title"] = title
        set_status("Готовим отчёты", period_labels, ACCENT, busy=True)
        if action_type != "abc":
            set_dashboard_chart_feedback("loading", "Подготавливаем данные", f"Собираем недостающие отчёты: {period_labels}.")
        set_log(f"{lead}: {period_labels}")
        page.update()

        def worker() -> None:
            log_lines = deque([f"{lead}: {period_labels}\n"], maxlen=500)

            def on_line(line: str) -> None:
                log_lines.append(line)
                ui_queue.put(("log_append", line))

            last_code = 0
            for month_value, year_value in missing_periods:
                if job_cancel_event.is_set():
                    ui_queue.put(("finish", title, 130, "Остановлено пользователем."))
                    return
                cmd = [
                    str(PYTHON),
                    str(SCRIPTS / "Monthly_sales_report.py"),
                    "--month",
                    str(month_value),
                    "--year",
                    str(year_value),
                ]
                log_lines.append(f"\n$ {' '.join(cmd)}\n")
                ui_queue.put(("log_append", f"\n$ {' '.join(cmd)}\n"))
                try:
                    last_code, output = stream_process(
                        cmd,
                        on_line,
                        on_prompt=request_prompt,
                        extra_env={"OZONREPORTX_UI_PROMPTS": "1"},
                        proc_holder=active_proc,
                        cancel_event=job_cancel_event,
                    )
                except Exception as exc:
                    last_code, output = 1, f"Не удалось выполнить задачу: {exc}"
                if last_code != 0:
                    ui_queue.put(("finish", title, last_code, output))
                    return
            ui_queue.put(("finish", title, 0, "".join(log_lines).rstrip()))

        try:
            page.run_thread(worker)
        except Exception as exc:
            ui_queue.put(("finish", title, 1, f"Не удалось запустить задачу: {exc}"))

    def build_dashboard_chart(_e=None) -> None:
        chart_error = valid_dashboard_chart_period()
        if chart_error:
            toast(chart_error, True)
            return
        missing_periods = get_missing_chart_reports()
        if missing_periods:
            build_missing_chart_reports(missing_periods, "build")
            return
        from_month = int(dashboard_chart_from_m.value or prev.month)
        from_year = int(dashboard_chart_from_y.value or prev.year)
        to_month = int(dashboard_chart_to_m.value or prev.month)
        to_year = int(dashboard_chart_to_y.value or prev.year)
        set_status("Строим график", f"{MONTHS[from_month - 1]} {from_year} - {MONTHS[to_month - 1]} {to_year}", ACCENT, busy=True)
        set_dashboard_chart_feedback("loading", "Строим график", f"{MONTHS[from_month - 1]} {from_year} - {MONTHS[to_month - 1]} {to_year}")
        if page.controls:
            page.update(dashboard_chart_panel)
        render_dashboard_chart(from_month, from_year, to_month, to_year)
        if dashboard_chart_state["built"]:
            set_status("График построен", dashboard_chart_note.value, PRIMARY, busy=False)
        else:
            set_status("Нет данных", "Для графика нет отчётов.", DANGER, busy=False)
        if page.controls:
            page.update(dashboard_chart_panel)

    def launch_abc_analysis() -> None:
        abc_error = valid_abc(abc_fm.value or "", abc_fy.value or "", abc_tm.value or "", abc_ty.value or "")
        if abc_error:
            toast(abc_error, True)
            return
        missing_periods = get_missing_abc_reports()
        if missing_periods:
            build_missing_chart_reports(missing_periods, "abc")
            return
        job(
            ABC_JOB_TITLE,
            [
                str(PYTHON),
                str(SCRIPTS / "ABC_XYZ_analytics_report.py"),
                "-i",
                str(ozon_reports_dir(ROOT)),
                "--output_dir",
                "ABC&XYZ reports",
                "--from-month",
                abc_fm.value or "",
                "--from-year",
                abc_fy.value or "",
                "--to-month",
                abc_tm.value or "",
                "--to-year",
                abc_ty.value or "",
            ],
        )

    def dashboard_period_changed(_e=None) -> None:
        selected_month = int(dashboard_m.value or today.month)
        selected_year = int(dashboard_y.value or today.year)
        if is_wb_store():
            refresh_dashboard()
            return
        report_path = report_path_for_period(selected_month, selected_year)
        set_log(f"Открываем отчёт за {MONTHS[selected_month - 1]} {selected_year}: {report_path.name}")
        set_status("Открытие периода", f"{MONTHS[selected_month - 1]} {selected_year}", ACCENT, busy=True)
        refresh_dashboard()
        if not report_path.exists():
            set_log(f"Нет данных за {MONTHS[selected_month - 1]} {selected_year}. Нажмите «Сформировать отчёт» для получения данных.")
        elif page.controls:
            page.update(log)

    def apply_control_theme(control, dark: bool) -> None:
        if control is None:
            return
        if isinstance(control, ft.Text):
            if control.data == "preserve_excel_text":
                return
            current = (control.color or "").upper() if isinstance(control.color, str) else control.color
            if current in PRIMARY_TEXT_COLORS:
                control.color = "#F5F5F7" if dark else "#1D1D1F"
            elif current in MUTED_TEXT_COLORS:
                control.color = "#A1A1A6" if dark else "#6E6E73"
            else:
                control.color = themed_color(control.color, dark)
        if isinstance(control, ft.Dropdown):
            control.color = "#F5F5F7" if dark else "#1D1D1F"
            control.bgcolor = themed_color(control.bgcolor, dark)
            control.border_color = themed_color(control.border_color, dark)
            if control.focused_border_color:
                control.focused_border_color = PRIMARY
            if getattr(control, "label_style", None):
                control.label_style = ft.TextStyle(
                    size=getattr(control.label_style, "size", 11) or 11,
                    color="#A1A1A6" if dark else "#6E6E73",
                )
        if isinstance(control, ft.Container):
            if control.data != "preserve_excel_fill":
                control.bgcolor = themed_color(control.bgcolor, dark)
                if control.border is not None and hasattr(control.border, "top") and control.border.top is not None:
                    border_color = themed_color(control.border.top.color, dark)
                    control.border = ft.Border.all(control.border.top.width, border_color)
        for attr in ("content", "title", "subtitle", "leading", "trailing"):
            if hasattr(control, attr):
                child = getattr(control, attr)
                apply_control_theme(child, dark)
        for attr in ("controls", "actions", "destinations"):
            if hasattr(control, attr):
                children = getattr(control, attr) or []
                for child in children:
                    apply_control_theme(child, dark)

    def refresh_files() -> None:
        files_col.controls = []
        for spec in folder_specs:
            rows: list[ft.Control] = [ft.TextButton("Открыть папку", on_click=lambda _e, x=spec.path: reveal_path(x, x.name))]
            batch = files_by_relevance(spec.path, spec.pattern)[:12]
            if not batch:
                rows.append(ft.Text("Файлы пока не найдены.", color=MUTED))
            else:
                for p in batch:
                    rows.append(
                        ft.Row(
                            [
                                ft.Text(p.name, expand=True),
                                ft.Text(datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M"), color=MUTED),
                                ft.IconButton(ft.Icons.OPEN_IN_NEW, on_click=lambda _e, x=p: reveal_path(x, x.name)),
                            ],
                            alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                        )
                    )
            files_col.controls.append(card(spec.title, f"Папка: {spec.path.name}", rows))


        if current_dark["value"]:
            apply_control_theme(files_col, True)
        abc_files_changed()

    # ---------------------------------------------------------------- ABC/XYZ-анализ
    ABC_COLUMN_INFO = [
        ("X", "стабильный"), ("Y1", "слабая сезонность"), ("Y2", "сезонный"), ("Y3", "сильная сезонность"),
        ("Y", "умеренные колебания"), ("Z", "нерегулярный"), ("—", "мало продаж"),
    ]
    ABC_ROW_INFO = [("A", "≈80% прибыли"), ("B", "следующие ≈15%"), ("C", "последние ≈5%")]
    ABC_FILTERS = [("all", "Все"), ("A", "A"), ("B", "B"), ("C", "C"), ("X", "X"), ("Y", "Y"), ("Z", "Z"), ("—", "Мало данных")]

    def abc_cell_tone(abc: str, xyz: str) -> str:
        pal = dash_pal()
        dark = pal["dark"]
        tones = {
            "good": "#1E3A33" if dark else "#E6F4EE",
            "mid": "#3A3426" if dark else "#FBF2E2",
            "bad": "#42292C" if dark else "#FBE8E5",
            "none": "#252B30" if dark else "#F4F5F8",
        }
        if xyz == "—":
            return tones["none"]
        score = {"A": 0, "B": 1, "C": 2}.get(abc, 1) + {"X": 0, "Y1": 1, "Y": 1.5, "Y2": 2, "Y3": 3, "Z": 4}.get(xyz, 2)
        return tones["good"] if score <= 1 else tones["mid"] if score <= 3 else tones["bad"]

    def abc_files_changed() -> None:
        files = sorted(files_by_relevance(ROOT / "ABC&XYZ reports"), key=lambda p: p.stat().st_mtime, reverse=True)
        abc_state["files"] = files
        abc_report_pick.options = [ft.dropdown.Option(str(p), abc_period_title(p)) for p in files]
        selected = abc_state.get("selected")
        if not files:
            abc_state.update(selected=None, data=None)
        elif selected not in [str(p) for p in files]:
            # Анализ за один месяц ничего не говорит о стабильности спроса — по умолчанию берём многомесячный.
            multi = [p for p in files if len(set(re.findall(r"\d{4}", p.stem))) > 1 or p.stem.split("-")[0] != p.stem.split("-")[-1]]
            abc_select(str((multi or files)[0]), push_update=False)
            return
        abc_report_pick.value = abc_state.get("selected")
        render_abc_view(push_update=False)

    def abc_select(path: str | None, push_update: bool = True) -> None:
        if not path or abc_load_report is None:
            return
        abc_state.update(selected=path, filter="all", cell=None, error=None)
        abc_report_pick.value = path
        try:
            abc_state["data"] = abc_load_report(Path(path))
        except Exception as exc:
            abc_state.update(data=None, error=f"Не удалось прочитать отчёт: {exc}")
        render_abc_view(push_update=push_update)

    def abc_toggle_create(show: bool | None = None) -> None:
        abc_state["creating"] = (not abc_state["creating"]) if show is None else show
        abc_create_panel.visible = abc_state["creating"]
        if page.controls:
            page.update(abc_create_panel)

    def abc_set_filter(value: str, cell=None) -> None:
        abc_state["filter"] = value
        abc_state["cell"] = cell
        render_abc_view()

    def abc_item_matches(item: dict) -> bool:
        cell = abc_state.get("cell")
        if cell:
            return (item["abc"], abc_xyz_group(item["xyz"])) == cell
        value = abc_state.get("filter") or "all"
        if value == "all":
            return True
        if value in ("A", "B", "C"):
            return item["abc"] == value
        group = abc_xyz_group(item["xyz"])
        if value == "Y":
            return group.startswith("Y")
        return group == value

    def abc_sparkline(values: list[float], months: list[str]) -> ft.Control:
        pal = dash_pal()
        peak = max(values) if values and max(values) > 0 else 1
        bars = []
        for value, month in zip(values, months):
            bars.append(ft.Container(
                expand=1, height=26, alignment=ft.Alignment(0, 1), tooltip=f"{month}: {_format_number(value)} шт.",
                content=ft.Container(height=max(2, 24 * value / peak), border_radius=2, data="preserve_excel_fill",
                                     bgcolor=pal["accent"] if value else pal["separator"]),
            ))
        return ft.Row(bars, spacing=2)

    def abc_badge(code: str, abc: str, xyz: str) -> ft.Control:
        return ft.Container(padding=ft.Padding.symmetric(horizontal=8, vertical=3), border_radius=8, bgcolor=abc_cell_tone(abc, abc_xyz_group(xyz)),
                            data="preserve_excel_fill", content=dtext(code, 12, weight=ft.FontWeight.W_700))

    def render_abc_view(push_update: bool = True) -> None:
        pal = dash_pal()
        data = abc_state.get("data")
        files = abc_state.get("files") or []
        if data:
            built = data["built_at"]
            abc_subtitle.value = f"{data['period']} · построен {built:%d.%m.%Y в %H:%M} · {len(data['items'])} товаров"
        else:
            abc_subtitle.value = "Какие товары приносят прибыль и насколько стабилен их спрос"
        controls: list[ft.Control] = []
        if abc_state.get("error"):
            controls.append(dpanel(dtext(abc_state["error"], 13, pal["bad"])))
        if not files:
            controls.append(dpanel(
                ft.Column([
                    ft.Icon(ft.Icons.GRID_VIEW, size=44, color=pal["muted"]),
                    dtext("Анализов пока нет", 18, weight=ft.FontWeight.W_700),
                    dtext("Выберите период — программа соберёт ABC/XYZ-анализ по месячным отчётам.", 13, pal["muted"]),
                    ft.Container(height=6),
                    ft.Button("Новый анализ", icon=ft.Icons.ADD, bgcolor=PRIMARY, color="white", on_click=lambda _e: abc_toggle_create(True)),
                ], spacing=6, horizontal_alignment=ft.CrossAxisAlignment.CENTER, tight=True),
                padding=48, alignment=ft.Alignment(0, 0),
            ))
        if data:
            items = data["items"]
            total_profit = data["total_profit"] or 0.0
            group_a = [i for i in items if i["abc"] == "A"]
            share_a = sum(i["profit"] or 0 for i in group_a) / total_profit * 100 if total_profit else 0
            stable = [i for i in items if abc_xyz_group(i["xyz"]) in ("X", "Y1")]
            risky = [i for i in items if abc_xyz_group(i["xyz"]) in ("Z", "Y3")]
            controls.append(dash_stats_panel([
                dash_stat("Прибыль за период", _format_currency(round(total_profit)), f"{len(items)} товаров с продажами", help_key="Прибыль за период"),
                dash_stat("Группа A", f"{len(group_a)} тов.", f"дают {share_a:.0f}% прибыли", pal["good"]),
                dash_stat("Стабильный спрос", f"{len(stable)} тов.", "X и Y1 — легко планировать запасы"),
                dash_stat("Требуют внимания", f"{len(risky)} тов.", "нерегулярный или сильно сезонный спрос", pal["warn"] if risky else None),
            ]))

            # Матрица ABC × XYZ
            cells = abc_matrix_counts(items)
            columns = [(key, caption) for key, caption in ABC_COLUMN_INFO
                       if key not in ("Y", "—") or any(col == key for (_row, col) in cells)]
            header = [ft.Container(width=120)] + [
                ft.Container(expand=1, content=ft.Column([dtext(key if key != "—" else "?", 14, weight=ft.FontWeight.W_700),
                                                          dtext(caption, 11, pal["muted"], max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)],
                                                         spacing=0, horizontal_alignment=ft.CrossAxisAlignment.CENTER))
                for key, caption in columns
            ]
            grid: list[ft.Control] = [ft.Row(header, spacing=8)]
            for row_key, row_caption in ABC_ROW_INFO:
                row_cells: list[ft.Control] = [ft.Container(width=120, content=ft.Column([
                    dtext(row_key, 18, weight=ft.FontWeight.W_800), dtext(row_caption, 11, pal["muted"])], spacing=0))]
                for col_key, _caption in columns:
                    articles = cells.get((row_key, col_key), [])
                    selected = abc_state.get("cell") == (row_key, col_key)
                    content = ft.Column([
                        dtext(str(len(articles)) if articles else "—", 20, None if articles else pal["muted"], ft.FontWeight.W_700),
                        dtext(", ".join(articles), 11, pal["muted"], max_lines=2, overflow=ft.TextOverflow.ELLIPSIS),
                    ], spacing=2, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
                    row_cells.append(ft.Container(
                        expand=1, height=86, border_radius=12, padding=8, alignment=ft.Alignment(0, 0), data="preserve_excel_fill",
                        bgcolor=abc_cell_tone(row_key, col_key) if articles else None,
                        border=ft.Border.all(2, pal["accent"]) if selected else ft.Border.all(1, pal["separator"]),
                        tooltip=", ".join(articles) if articles else None,
                        on_click=(lambda _e, cell=(row_key, col_key): abc_set_filter("all", None if abc_state.get("cell") == cell else cell)) if articles else None,
                        content=content,
                    ))
                grid.append(ft.Row(row_cells, spacing=8))
            controls.append(dpanel(ft.Column([
                ft.Row([dtext("Матрица ABC × XYZ", 15, weight=ft.FontWeight.W_700),
                        dtext("ABC — вклад в прибыль, XYZ — стабильность спроса. Нажмите на ячейку, чтобы отфильтровать таблицу.", 12, pal["muted"], expand=True)],
                       spacing=14, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ft.Column(grid, spacing=8),
            ], spacing=14)))

            # Таблица товаров
            active = abc_state.get("filter") or "all"
            chips: list[ft.Control] = []
            for key, label in ABC_FILTERS:
                is_active = active == key and not abc_state.get("cell")
                chips.append(ft.Container(
                    padding=ft.Padding.symmetric(horizontal=12, vertical=6), border_radius=999, data="preserve_excel_fill",
                    bgcolor=pal["accent"] if is_active else None, border=None if is_active else ft.Border.all(1, pal["border"]),
                    on_click=lambda _e, value=key: abc_set_filter(value),
                    content=dtext(label, 12, "#FFFFFF" if is_active else pal["ink"], ft.FontWeight.W_600),
                ))
            if abc_state.get("cell"):
                code = "".join(abc_state["cell"])
                chips.append(ft.Container(
                    padding=ft.Padding.symmetric(horizontal=12, vertical=6), border_radius=999, bgcolor=pal["accent"], data="preserve_excel_fill",
                    on_click=lambda _e: abc_set_filter("all"),
                    content=ft.Row([dtext(f"Ячейка {code}", 12, "#FFFFFF", ft.FontWeight.W_600), ft.Icon(ft.Icons.CLOSE, size=14, color="#FFFFFF")],
                                   spacing=4, tight=True),
                ))
            shown = [i for i in items if abc_item_matches(i)]
            months = data["months"]
            table_columns = [("Артикул", 2, False), ("Класс", 2, False), ("Прибыль", 3, True), ("Доля", 2, True),
                             (f"Продажи: {months[0][:3].lower()} — {months[-1][:3].lower()}" if months else "Продажи", 5, False),
                             ("Шт.", 2, True), ("Колебания", 2, True), ("Что делать", 8, False)]

            def tcell(content: ft.Control, weight: int, right: bool = False) -> ft.Control:
                return ft.Container(expand=weight, alignment=ft.Alignment(1 if right else -1, 0), content=content)

            rows: list[ft.Control] = [ft.Row([tcell(dtext(name, 12, pal["muted"], ft.FontWeight.W_600), w, r) for name, w, r in table_columns], spacing=14)]
            for index, item in enumerate(shown):
                recs = item.get("recommendations") or []
                tip = item["title"] + ("\n• " + "\n• ".join(recs) if recs else "")
                cv = item.get("cv")
                rows.append(ft.Container(
                    padding=ft.Padding.symmetric(vertical=9), data="preserve_excel_fill",
                    border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])),
                    tooltip=tip,
                    content=ft.Row([
                        tcell(dtext(item["article"], 14, weight=ft.FontWeight.W_700), 2),
                        tcell(abc_badge(item["code"] or item["abc"], item["abc"], item["xyz"]), 2),
                        tcell(dtext(_format_currency(round(item["profit"])) if item["profit"] is not None else "—", 13), 3, True),
                        tcell(dtext(f"{item['share'] * 100:.1f}%".replace(".", ",") if item["share"] is not None else "—", 13, pal["muted"]), 2, True),
                        tcell(abc_sparkline(item["sales"], months), 5),
                        tcell(dtext(_format_number(item["total_qty"]) if item["total_qty"] is not None else "—", 13), 2, True),
                        tcell(dtext(f"{cv:.2f}".replace(".", ",") if cv else "—", 13, pal["muted"]), 2, True),
                        tcell(ft.Column([dtext(item["title"] or "—", 12, weight=ft.FontWeight.W_600, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)]
                                        + ([dtext(recs[0], 11, pal["muted"], max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)] if recs else []),
                                        spacing=1, tight=True), 8),
                    ], spacing=14, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ))
            if not shown:
                rows.append(ft.Container(padding=20, alignment=ft.Alignment(0, 0), content=dtext("В этой группе товаров нет.", 13, pal["muted"])))
            controls.append(dpanel(ft.Column([
                ft.Row([dtext("Товары", 15, weight=ft.FontWeight.W_700), dtext(f"{len(shown)} из {len(items)}", 12, pal["muted"])],
                       spacing=10, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ft.Row(chips, spacing=8, run_spacing=8, wrap=True),
                ft.Column(rows, spacing=0),
                dtext("Колебания — коэффициент вариации продаж по месяцам: до 0,2 стабильно, больше 0,5 — нерегулярно. "
                      "Наведите на строку, чтобы увидеть все рекомендации.", 11, pal["muted"]),
            ], spacing=12), padding=ft.Padding.only(left=20, right=20, top=16, bottom=14)))
        abc_body.controls = controls
        if push_update and page.controls:
            page.update(abc_body, abc_subtitle, abc_report_pick)

    def abc_open_selected(_e=None) -> None:
        path = abc_state.get("selected")
        if path:
            reveal_path(Path(path), Path(path).name)
        else:
            toast("Сначала выберите или создайте анализ.", True)

    def abc_after_job(code: int) -> None:
        """После построения анализа — сразу показываем новый отчёт."""
        set_abc_sync(False)
        if code != 0:
            return
        try:
            expected = ROOT / "ABC&XYZ reports" / f"{MONTHS[int(abc_fm.value) - 1]} {abc_fy.value}-{MONTHS[int(abc_tm.value) - 1]} {abc_ty.value}.xlsx"
        except (TypeError, ValueError):
            expected = None
        abc_files_changed()
        if expected is not None and expected.exists():
            abc_select(str(expected), push_update=False)
            abc_toggle_create(False)
        render_abc_view()

    def set_abc_sync(active: bool) -> None:
        dark = current_dark["value"]
        abc_sync_overlay.visible = active
        abc_sync_overlay.bgcolor = ("#CC1E2328" if dark else "#CCF5F7FB") if active else None
        abc_sync_overlay.content.bgcolor = "#2A3136" if dark else "#FFFFFF"
        abc_sync_overlay.content.border = ft.Border.all(1, "#465056" if dark else BORDER)
        abc_sync_label.color = "#F3F4F6" if dark else TEXT
        abc_body.disabled = active
        if page.controls:
            page.update(abc_sync_overlay, abc_body)

    def set_pricing_costs_dirty(flag: bool, push_update: bool = False) -> None:
        pricing_costs_dirty["value"] = flag
        pricing_costs_save_button.disabled = not flag
        pricing_costs_save_button.text = "Сохранить изменения *" if flag else "Сохранить изменения"
        if push_update and page.controls:
            page.update(pricing_costs_save_button)

    def set_pricing_sync(active: bool) -> None:
        """Пока идёт синхронизация, таблица закрыта заглушкой с загрузкой и не принимает клики."""
        dark = current_dark["value"]
        pricing_sync_overlay.visible = active
        pricing_sync_overlay.bgcolor = ("#CC1E2328" if dark else "#CCF5F7FB") if active else None
        pricing_sync_label.color = "#F3F4F6" if dark else TEXT
        pricing_costs_list.disabled = active
        if active:
            pricing_costs_save_button.disabled = True
        if page.controls:
            page.update(pricing_sync_overlay, pricing_costs_list, pricing_costs_save_button)

    def mark_pricing_costs_dirty(_e=None) -> None:
        pricing_costs_dirty["revision"] += 1
        if not pricing_costs_dirty["value"]:
            set_pricing_costs_dirty(True, push_update=True)

    def build_costs_edit_field(row_number: int, column: str, cell_payload: dict[str, object] | None, dark: bool, expand: int) -> ft.Control:
        payload = cell_payload or {}
        fill_color = payload.get("fill_color") if isinstance(payload.get("fill_color"), str) else None
        font_color = payload.get("font_color") if isinstance(payload.get("font_color"), str) else None
        resolved_text_color = font_color or _excel_text_color(fill_color, dark)
        field = ft.TextField(
            value=_format_costs_editor_value(column, payload.get("value")),
            expand=True,
            dense=True,
            text_size=12,
            color=resolved_text_color,
            bgcolor="transparent" if fill_color else ("#232A2F" if dark else "#FFFFFF"),
            border_color="transparent" if fill_color else ("#D6D0C4" if not dark else "#465056"),
            focused_border_color=PRIMARY,
            border_radius=10,
            content_padding=ft.Padding.symmetric(horizontal=10, vertical=6),
            on_change=mark_pricing_costs_dirty,
        )
        pricing_costs_editors.setdefault(row_number, {})[column] = field
        return ft.Container(
            expand=expand,
            bgcolor=fill_color,
            border_radius=10,
            padding=2 if not fill_color else 0,
            data="preserve_excel_fill" if fill_color else None,
            content=field,
        )

    def pricing_stats_panel(rows: list[dict[str, object]], min_margin_pct: float, desired_margin_pct: float) -> ft.Control:
        """Ключевые показатели таблицы цен — в той же строке показателей, что и бизнес-сводка."""
        pal = dash_pal()

        def number(row: dict[str, object], column: str) -> float | None:
            payload = row.get(column)
            return as_float(payload.get("value")) if isinstance(payload, dict) else None

        below_min = 0
        margins: list[float] = []
        profits: list[float] = []
        no_cost = 0
        for row in rows:
            price, floor = number(row, "Текущая цена на Ozon"), number(row, "Минимальная цена продажи")
            if price and floor and price < floor:
                below_min += 1
            margin = number(row, "Текущая ожидаемая рентабельность")
            if margin is not None:
                margins.append(margin * 100 if abs(margin) <= 1 else margin)
            profit = number(row, "Ожидаемая чистая прибыль")
            if profit is not None:
                profits.append(profit)
            if not number(row, "Себестоимость"):
                no_cost += 1
        avg_margin = sum(margins) / len(margins) if margins else None
        margin_color = None
        if avg_margin is not None:
            margin_color = pal["bad"] if avg_margin < min_margin_pct else pal["good"] if avg_margin <= desired_margin_pct else pal["accent"]
        pct = lambda value: f"{value:.1f}%".replace(".", ",")
        stats = [
            dash_stat("Товаров", str(len(rows)), f"без себестоимости: {no_cost}" if no_cost else "все с себестоимостью"),
            dash_stat("Ниже минимальной цены", str(below_min), "текущая цена Ozon < минимальной",
                      pal["bad"] if below_min else pal["good"]),
            dash_stat("Средняя рентабельность", pct(avg_margin) if avg_margin is not None else "—",
                      f"цель {pct(min_margin_pct)} – {pct(desired_margin_pct)}", margin_color),
            dash_stat("Прибыль с единицы", _format_currency(sum(profits) / len(profits)) if profits else "—",
                      "в среднем по товарам", pal["bad"] if profits and sum(profits) < 0 else None),
        ]
        return dpanel(ft.ResponsiveRow(stats, spacing=24, run_spacing=12), padding=ft.Padding.symmetric(horizontal=24, vertical=16))

    def refresh_pricing_costs_preview(push_update: bool = False) -> str | None:
        if pricing_costs_dirty["value"]:
            return None
        saved_min_margin, saved_desired_margin = margin_defaults()
        active_min_margin = margin_min.value or saved_min_margin
        active_desired_margin = margin_des.value or saved_desired_margin
        if valid_margin(active_min_margin, active_desired_margin):
            active_min_margin, active_desired_margin = saved_min_margin, saved_desired_margin
        min_margin_pct = float(active_min_margin.replace(",", ".")) * 100
        desired_margin_pct = float(active_desired_margin.replace(",", ".")) * 100

        try:
            signature = file_signature(ROOT / "costs.xlsx")
        except OSError:
            signature = None
        rows, error = read_costs_main_rows(
            ROOT / "costs.xlsx",
            min_margin_pct=min_margin_pct,
            desired_margin_pct=desired_margin_pct,
        )
        if signature is not None:
            try:
                if file_signature(ROOT / "costs.xlsx") != signature:
                    return "Файл изменился во время чтения. Обновите таблицу ещё раз."
            except OSError as exc:
                return f"Не удалось проверить costs.xlsx: {exc}"
        # An edit event may have arrived while Excel was being read.
        if pricing_costs_dirty["value"]:
            return None
        pricing_costs_snapshot["signature"] = signature
        dark = current_dark["value"]
        pricing_costs_editors.clear()
        set_pricing_costs_dirty(False)

        if error:
            pricing_costs_status.value = error
            pricing_costs_status.color = "#E6BBB5" if dark else DANGER
            pricing_stats_area.content = None
            pricing_costs_list.controls = [
                ft.Container(
                    padding=18,
                    border_radius=18,
                    bgcolor=semantic_surface("danger", dark=dark),
                    border=ft.Border.all(1, "#D9A49D" if not dark else "#6E464C"),
                    content=ft.Text(error, size=12, color="#F3F4F6" if dark else DANGER),
                )
            ]
        else:
            refreshed_at = datetime.fromtimestamp((ROOT / "costs.xlsx").stat().st_mtime)
            pricing_costs_status.value = f"{len(rows)} товаров · лист «Основной» · обновлено {refreshed_at:%d.%m в %H:%M}"
            pricing_costs_status.color = "#B6C0BC" if dark else MUTED
            pricing_stats_area.content = pricing_stats_panel(rows, min_margin_pct, desired_margin_pct)

            pal = dash_pal()
            header = ft.Container(
                padding=ft.Padding.only(left=8, right=8, bottom=8),
                content=ft.Row(
                    [ft.Text(column, expand=COSTS_MAIN_COLUMN_EXPANDS[column], size=11, weight=ft.FontWeight.W_600, color=pal["muted"],
                             max_lines=2, overflow=ft.TextOverflow.ELLIPSIS) for column in COSTS_TABLE_COLUMNS],
                    spacing=12,
                    vertical_alignment=ft.CrossAxisAlignment.END,
                ),
            )
            row_controls: list[ft.Control] = [header]
            for row in rows:
                row_number = int(row.get("__row_number") or 0)
                row_controls.append(
                    ft.Container(
                        padding=ft.Padding.symmetric(horizontal=8, vertical=6),
                        border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])),
                        data="preserve_excel_fill",
                        content=ft.Row(
                            [
                                build_costs_edit_field(row_number, column, row.get(column), dark, COSTS_MAIN_COLUMN_EXPANDS[column])
                                if column in EDITABLE_COSTS_COLUMNS and row_number >= 2
                                else _build_costs_cell(column, row.get(column), dark, COSTS_MAIN_COLUMN_EXPANDS[column])
                                for column in COSTS_TABLE_COLUMNS
                            ],
                            spacing=12,
                            vertical_alignment=ft.CrossAxisAlignment.CENTER,
                        ),
                    )
                )
            pricing_costs_list.controls = row_controls

        if push_update and page.controls:
            page.update(pricing_costs_status, pricing_stats_area, pricing_costs_list, pricing_costs_save_button)
        return error

    def _save_pricing_costs(_e=None) -> None:
        revision = pricing_costs_dirty["revision"]
        updates: list[dict[str, object]] = []
        for row_number, editors in pricing_costs_editors.items():
            article_value = (editors.get("Артикул").value or "").strip() if editors.get("Артикул") else ""
            cost_raw = (editors.get("Себестоимость").value or "").strip() if editors.get("Себестоимость") else ""
            if cost_raw:
                try:
                    cost_value = float(cost_raw.replace(" ", "").replace(",", "."))
                except ValueError:
                    set_status("Ошибка сохранения", f"Строка {row_number}: себестоимость должна быть числом.", DANGER, busy=False)
                    return
            else:
                cost_value = None
            updates.append(
                {
                    "row_number": row_number,
                    "values": {
                        "Артикул": article_value or None,
                        "Себестоимость": cost_value,
                    },
                }
            )

        saved_signature = {}
        error = write_costs_main_rows(ROOT / "costs.xlsx", updates, expected_signature=pricing_costs_snapshot["signature"], saved_signature=saved_signature)
        if error:
            set_status("Ошибка сохранения", error, DANGER, busy=False)
            return
        pricing_costs_snapshot["signature"] = saved_signature["signature"]
        if pricing_costs_dirty["revision"] != revision:
            set_status("Изменения сохранены", "В редакторе есть более новые несохранённые правки.", ACCENT, busy=False)
            return
        set_pricing_costs_dirty(False)
        refresh_error = refresh_pricing_costs_preview(push_update=True)
        if refresh_error:
            set_status("Изменения сохранены", "Значения записаны, но превью не удалось перечитать.", ACCENT, busy=False)
        else:
            set_status("Изменения сохранены", "Артикул и себестоимость обновлены в costs.xlsx.", PRIMARY, busy=False)

    def handle_pricing_costs_save(_e=None) -> None:
        if not job_lock.acquire(blocking=False):
            toast("Дождитесь завершения текущей операции перед сохранением.", True)
            return
        try:
            _save_pricing_costs(_e)
        finally:
            job_lock.release()

    def handle_pricing_costs_refresh(_e=None) -> None:
        if pricing_costs_dirty["value"]:
            discard_costs_dialog.open = True
            page.update(discard_costs_dialog)
            return
        error = refresh_pricing_costs_preview(push_update=True)
        if error:
            set_status("Ошибка таблицы", error, DANGER, busy=False)
        else:
            set_status("Таблица обновлена", "Лист «Основной» перечитан из costs.xlsx.", PRIMARY, busy=False)

    def close_discard_costs_dialog(_e=None) -> None:
        discard_costs_dialog.open = False
        page.update(discard_costs_dialog)

    def discard_costs_edits(_e=None) -> None:
        close_discard_costs_dialog()
        set_pricing_costs_dirty(False)
        handle_pricing_costs_refresh()

    discard_costs_dialog = ft.AlertDialog(
        modal=True,
        title=ft.Text("Есть несохранённые изменения"),
        content=ft.Text("Обновление таблицы удалит введённые правки. Оставьте их в редакторе или явно отмените."),
        actions=[ft.TextButton("Оставить правки", on_click=close_discard_costs_dialog),
                 ft.TextButton("Отменить правки и обновить", on_click=discard_costs_edits)],
    )
    page.overlay.append(discard_costs_dialog)
    pricing_costs_save_button.on_click = handle_pricing_costs_save

    def refresh_report_lists(_e=None) -> None:
        refresh_files()
        set_status("Списки обновлены", "Отчёты перечитаны.", PRIMARY, busy=False)
        if page.controls:
            page.update(files_col, abc_body, abc_report_pick)

    def finish(title: str, code: int, output: str) -> None:
        was_cancelled = cancel_state["requested"]
        active_proc["proc"] = None
        cancel_state["requested"] = False
        cancel_state["title"] = ""
        set_log(output, persist=False)
        append_log(f"\nЗадача «{title}» завершена, код {code}.")
        if code != 0:
            append_log(output)
        if title == PRICING_SYNC_TITLE:
            set_pricing_sync(False)
        if title == DASHBOARD_REFRESH_TITLE:
            set_dashboard_sync(False)
        try:
            refresh_dashboard()
            refresh_pricing_costs_preview()
            if title == "План поставки FBO":
                render_supply_results()
            if title == WB_SUPPLY_TITLE:
                render_wb_supply_results()
            if title == SEARCH_COLLECT_TITLE:
                search_load()
        except Exception as exc:
            append_log(f"\nНе удалось обновить интерфейс: {exc}")
        finally:
            if job_lock.locked():
                job_lock.release()
        if was_cancelled:
            set_status(f"Отменено: {title}", "Остановлено пользователем.", DANGER, busy=False)
        else:
            set_status(f"Готово: {title}" if code == 0 else f"Ошибка ({code}): {title}", "Задача завершена.", PRIMARY if code == 0 else DANGER, busy=False)
        pending_action = chart_pending_action["type"]
        chart_pending_action["type"] = None
        if was_cancelled:
            pending_action = None
        if code == 0 and pending_action == "build":
            build_dashboard_chart()
        elif code == 0 and pending_action == "interactive":
            open_interactive_dashboard_chart()
        elif code == 0 and pending_action == "abc":
            launch_abc_analysis()
        elif code == 0 and dashboard_chart_state["built"] and not was_cancelled:
            build_dashboard_chart()
        refresh_files()
        if title == ABC_JOB_TITLE:
            abc_after_job(code)
        if page.controls:
            page.update(state, spin, cancel_row, files_col, abc_body, abc_report_pick, pricing_costs_list, pricing_costs_status, pricing_stats_area, pricing_costs_save_button, log)

    def job(title: str, cmd: list[str], prompt_for_costs: bool = False) -> None:
        if not job_lock.acquire(blocking=False):
            toast("Дождитесь завершения текущей операции.", True)
            return
        job_cancel_event.clear()
        cancel_state["requested"] = False
        cancel_state["title"] = title
        if title == PRICING_SYNC_TITLE:
            set_pricing_sync(True)
        if title == DASHBOARD_REFRESH_TITLE:
            set_dashboard_sync(True)
        if title == ABC_JOB_TITLE:
            set_abc_sync(True)
        set_status(f"Выполняется: {title}", "Фоновый процесс запущен.", ACCENT, busy=True)
        set_log(f"$ {' '.join(cmd)}")
        page.update()

        def worker() -> None:
            def on_line(line: str) -> None:
                ui_queue.put(("log_append", line))

            try:
                code, output = stream_process(
                    cmd,
                    on_line,
                    on_prompt=request_prompt if prompt_for_costs else None,
                    extra_env={"OZONREPORTX_UI_PROMPTS": "1"} if prompt_for_costs else None,
                    proc_holder=active_proc,
                    cancel_event=job_cancel_event,
                )
            except Exception as exc:
                code, output = 1, f"Не удалось выполнить задачу: {exc}"
            ui_queue.put(("finish", title, code, output))

        try:
            page.run_thread(worker)
        except Exception as exc:
            ui_queue.put(("finish", title, 1, f"Не удалось запустить задачу: {exc}"))

    def launch_recommended_prices() -> None:
        margin_error = valid_margin(margin_min.value or "", margin_des.value or "")
        if margin_error:
            toast(margin_error, True)
            return
        job(
            "Рекомендованные цены",
            [
                str(PYTHON),
                str(SCRIPTS / "ui_actions.py"),
                "recommended-prices",
                "--min-margin",
                (margin_min.value or "").replace(",", "."),
                "--desired-margin",
                (margin_des.value or "").replace(",", "."),
            ],
        )

    def open_selected_dashboard_report() -> None:
        month_error = valid_month_year(dashboard_m.value or "", dashboard_y.value or "")
        if month_error:
            toast(month_error, True)
            return
        selected_month, selected_year = int(dashboard_m.value or today.month), int(dashboard_y.value or today.year)
        report_path = (wb_dash.report_path(selected_year, selected_month) if is_wb_store()
                       else report_path_for_period(selected_month, selected_year))
        if not report_path.exists():
            toast("Отчёт за выбранный период пока не найден.", True)
            return
        reveal_path(report_path, report_path.name)

    def launch_dashboard_report() -> None:
        month_error = valid_month_year(dashboard_m.value or "", dashboard_y.value or "")
        if month_error:
            toast(month_error, True)
            return
        if is_wb_store():
            page.run_task(wb_dash.download, int(dashboard_y.value or today.year), int(dashboard_m.value or today.month))
            return
        monthly_m.value = dashboard_m.value
        monthly_y.value = dashboard_y.value
        job(
            DASHBOARD_REFRESH_TITLE,
            [
                str(PYTHON),
                str(SCRIPTS / "Monthly_sales_report.py"),
                "--month",
                dashboard_m.value or "",
                "--year",
                dashboard_y.value or "",
            ],
            prompt_for_costs=True,
        )

    dashboard_m.on_select = dashboard_period_changed
    dashboard_y.on_select = dashboard_period_changed

    prompt_dialog = ft.AlertDialog(
        modal=True,
        title=ft.Text("Нужны расходы на маркетинг"),
        content=ft.Container(
            width=520,
            padding=ft.Padding.symmetric(horizontal=4, vertical=2),
            content=ft.Column([prompt_label, prompt_hint, prompt_context, prompt_input], tight=True, spacing=10),
        ),
        actions=[
            ft.TextButton("Использовать 0", on_click=lambda _e: None),
            ft.Button("Подтвердить", on_click=lambda _e: None),
        ],
        actions_alignment=ft.MainAxisAlignment.END,
        open=False,
    )

    def prompt_submit(_e=None) -> None:
        holder = prompt_state["holder"]
        event = prompt_state["event"]
        if isinstance(holder, dict):
            holder["value"] = (prompt_input.value or "0").replace(",", ".")
        prompt_dialog.open = False
        page.update(prompt_dialog)
        if isinstance(event, threading.Event):
            event.set()

    prompt_dialog.actions = [
        ft.TextButton("Использовать 0", on_click=prompt_submit),
        ft.Button("Подтвердить", on_click=prompt_submit),
    ]
    page.overlay.append(prompt_dialog)

    discount_dialog_title = ft.Text("Заявки на скидку", color=TEXT, size=18, weight=ft.FontWeight.W_800)
    discount_dialog_badge = ft.Text("", color=MUTED, size=12, weight=ft.FontWeight.W_600)
    discount_dialog_status = ft.Text("", color=MUTED, size=12)
    discount_dialog_badge_shell = ft.Container(
        padding=ft.Padding.symmetric(horizontal=12, vertical=8),
        border_radius=999,
        bgcolor=semantic_surface("neutral_alt", dark=False),
        border=ft.Border.all(1, "#E5E5EA"),
        content=discount_dialog_badge,
    )
    discount_overview_row = ft.ResponsiveRow(run_spacing=8, spacing=8)
    discount_dialog_list = ft.Column(
        controls=[ft.Text("Нажмите «Обновить заявки», чтобы загрузить список.", size=12, color=MUTED)],
        spacing=8,
    )
    discount_dialog_state: dict[str, object] = {"plan": None, "busy": False, "processing_ids": set()}
    discount_dialog: ft.AlertDialog | None = None
    discount_requests_refresh_button = ft.IconButton(
        icon=ft.Icons.SYNC,
        tooltip="Обновить заявки",
        icon_size=18,
        style=ft.ButtonStyle(shape=ft.CircleBorder(), padding=ft.Padding.all(10)),
    )
    discount_requests_auto_button = ft.Button("Обработать автоматически", bgcolor=ACCENT, color="white", disabled=True)

    def close_discount_dialog(_e=None) -> None:
        discount_dialog_state["plan"] = None
        discount_dialog_state["busy"] = False
        discount_dialog_state["processing_ids"] = set()
        render_discount_dialog()
        if page.controls:
            page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)

    def current_discount_items() -> list[dict[str, object]]:
        plan = discount_dialog_state.get("plan")
        if not isinstance(plan, dict):
            return []
        items = plan.get("items")
        if not isinstance(items, list):
            return []
        return [item for item in items if isinstance(item, dict)]

    def discount_row_tokens(item: dict[str, object]) -> dict[str, str]:
        dark = current_dark["value"]
        price_state = str(item.get("price_state") or "")
        price_delta = as_float(item.get("price_delta"))

        if price_state == "below_min":
            delta_label = "Ниже минимума"
            if price_delta is not None:
                delta_label = f"Ниже на {_format_currency(abs(price_delta))}"
            return {
                "row_bg": semantic_surface("danger", dark=dark),
                "row_border": "#D9A49D" if not dark else "#6E464C",
                "title_color": "#F3F4F6" if dark else TEXT,
                "value_color": "#FFD2CB" if dark else DANGER,
                "detail_color": "#E6BBB5" if dark else "#7D443E",
                "badge_bg": "#F6DFDB" if not dark else "#5A383D",
                "badge_border": "#D9A49D" if not dark else "#81555C",
                "badge_text": "#8F342B" if not dark else "#FFD2CB",
                "badge_label": "Ниже минимума",
                "delta_label": delta_label,
                "approve_color": ACCENT if not dark else "#E0A15F",
                "approve_tooltip": "Одобрить вручную, несмотря на цену ниже минимума",
            }

        if price_state == "ok" or item.get("recommended_action") == "approve":
            if price_delta is None:
                delta_label = "Можно одобрять"
            elif abs(price_delta) < 0.01:
                delta_label = "Ровно по минимуму"
            else:
                delta_label = f"Запас {_format_currency(price_delta)}"
            return {
                "row_bg": semantic_surface("positive", dark=dark),
                "row_border": "#9DC9BA" if not dark else "#2E5A50",
                "title_color": "#F3F4F6" if dark else TEXT,
                "value_color": "#D8FFF3" if dark else PRIMARY,
                "detail_color": "#C5E0D7" if dark else "#355D54",
                "badge_bg": "#DCEFE5" if not dark else "#295046",
                "badge_border": "#9DC9BA" if not dark else "#3A6E62",
                "badge_text": "#1F5E51" if not dark else "#D8FFF3",
                "badge_label": "Цена нормальная",
                "delta_label": delta_label,
                "approve_color": PRIMARY if not dark else "#70D4BC",
                "approve_tooltip": "Одобрить заявку",
            }

        return {
            "row_bg": semantic_surface("warning", dark=dark),
            "row_border": "#CFBE8A" if not dark else "#655C43",
            "title_color": "#F3F4F6" if dark else TEXT,
            "value_color": "#F6E3B0" if dark else ACCENT,
            "detail_color": "#D8C698" if dark else "#72552A",
            "badge_bg": "#F4EDDB" if not dark else "#5A5037",
            "badge_border": "#CFBE8A" if not dark else "#7B6F4E",
            "badge_text": "#7A5A20" if not dark else "#F6E3B0",
            "badge_label": "Проверить вручную",
            "delta_label": "Нет полной привязки к минимальной цене",
            "approve_color": ACCENT if not dark else "#E0A15F",
            "approve_tooltip": "Одобрить заявку вручную",
        }

    def build_discount_auto_plan() -> dict[str, object]:
        items = current_discount_items()
        approve_tasks = [item.get("approve_payload") for item in items if item.get("recommended_action") == "approve" and item.get("approve_payload")]
        decline_tasks = [item.get("decline_payload") for item in items if item.get("recommended_action") != "approve" and item.get("decline_payload")]
        return {
            "items": items,
            "approve_tasks": approve_tasks,
            "decline_tasks": decline_tasks,
        }

    def remove_discount_item(task_id: object) -> None:
        plan = discount_dialog_state.get("plan")
        if not isinstance(plan, dict):
            return
        items = current_discount_items()
        plan["items"] = [item for item in items if str(item.get("id")) != str(task_id)]

    def restore_discount_item(item: dict[str, object]) -> None:
        plan = discount_dialog_state.get("plan")
        if not isinstance(plan, dict):
            return
        items = current_discount_items()
        items.append(item)
        items.sort(key=lambda row: str(row.get("id") or ""))
        plan["items"] = items

    def render_discount_dialog() -> None:
        items = current_discount_items()
        processing_ids = discount_dialog_state.get("processing_ids")
        processing_ids = processing_ids if isinstance(processing_ids, set) else set()
        dark = current_dark["value"]
        busy = discount_dialog_state.get("busy") is True
        requires_refresh = discount_dialog_state.get("requires_refresh") is True
        neutral_border = "#D8DEE8" if not dark else "#3E474E"
        discount_overview_row.controls = []
        if not isinstance(discount_dialog_state.get("plan"), dict):
            discount_dialog_badge.value = "Собираю список" if busy else "Заявки не загружены"
            discount_dialog_status.value = (
                "Проверяю новые скидочные заявки и готовлю рекомендации."
                if busy
                else "Загрузите список, чтобы увидеть рекомендации и обработать заявки прямо на странице."
            )
            discount_dialog_badge_shell.bgcolor = semantic_surface("info" if busy else "neutral_alt", dark=dark)
            discount_dialog_badge_shell.border = ft.Border.all(
                1,
                ("#C7D7EF" if not dark else "#425264") if busy else neutral_border,
            )
            discount_dialog_list.controls = []
            discount_requests_surface.visible = False
            discount_requests_auto_button.disabled = True
            return
        total = len(items)
        discount_dialog_badge.value = f"Новых заявок: {total}"
        auto_ready_items = [
            item
            for item in items
            if item.get("recommended_action") == "approve" and item.get("approve_payload") is not None
        ]
        attention_items = [item for item in items if item not in auto_ready_items]
        discount_overview_row.controls = [
            fact_chip("Новых", str(total), tone=semantic_surface("neutral_alt", dark=dark), col=4, variant="secondary"),
            fact_chip(
                "Можно авто",
                str(len(auto_ready_items)),
                tone=semantic_surface("positive", dark=dark) if auto_ready_items else semantic_surface("neutral_alt", dark=dark),
                col=4,
                variant="secondary",
            ),
            fact_chip(
                "Требуют внимания",
                str(len(attention_items)),
                tone=semantic_surface("warning", dark=dark) if attention_items else semantic_surface("neutral_alt", dark=dark),
                col=4,
                variant="secondary",
            ),
        ]
        if not items:
            discount_dialog_status.value = "Новых заявок нет — когда появятся, здесь будет рекомендация по каждой."
            discount_dialog_badge_shell.bgcolor = semantic_surface("neutral_alt", dark=dark)
            discount_dialog_badge_shell.border = ft.Border.all(1, neutral_border)
            discount_dialog_list.controls = []
            discount_requests_surface.visible = False
            discount_requests_auto_button.disabled = True
            return

        discount_dialog_badge_shell.bgcolor = (
            semantic_surface("warning", dark=dark)
            if attention_items
            else semantic_surface("positive", dark=dark)
        )
        discount_dialog_badge_shell.border = ft.Border.all(
            1,
            "#CFBE8A" if (attention_items and not dark) else "#655C43" if attention_items else "#9DC9BA" if not dark else "#2E5A50",
        )
        discount_dialog_status.value = (
            "Идёт обработка заявок. По завершении список обновится."
            if busy
            else "Зелёные строки проходят автоматически. Рискованные решения остаются доступны для ручного подтверждения."
        )
        header = ft.Container(
            padding=10,
            bgcolor=semantic_surface("neutral_alt", dark=dark),
            border_radius=12,
            border=ft.Border.all(1, "#D8DEE8" if not dark else "#3E474E"),
            content=ft.Row(
                [
                    ft.Text("Артикул", expand=3, size=11, weight=ft.FontWeight.W_700, color="#A1A1A6" if dark else MUTED),
                    ft.Text("Мин. цена продажи", expand=2, size=11, weight=ft.FontWeight.W_700, color="#A1A1A6" if dark else MUTED),
                    ft.Text("Цена по скидке", expand=2, size=11, weight=ft.FontWeight.W_700, color="#A1A1A6" if dark else MUTED),
                    ft.Text("Действие", width=96, size=11, weight=ft.FontWeight.W_700, color="#A1A1A6" if dark else MUTED, text_align=ft.TextAlign.CENTER),
                ],
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
            ),
        )
        rows: list[ft.Control] = [header]
        for item in items:
            task_id = item.get("id")
            item_id = str(task_id or "")
            reason = str(item.get("recommendation_reason") or "")
            is_processing = item_id in processing_ids or busy or requires_refresh
            row_tokens = discount_row_tokens(item)
            approve_disabled = is_processing or item.get("approve_payload") is None
            approve_tooltip = row_tokens["approve_tooltip"] if not approve_disabled else (reason or "Одобрение недоступно для этой заявки")
            rows.append(
                ft.Container(
                    padding=10,
                    bgcolor=row_tokens["row_bg"],
                    border_radius=12,
                    border=ft.Border.all(1, row_tokens["row_border"]),
                    tooltip=reason,
                    content=ft.Row(
                        [
                            ft.Container(
                                expand=3,
                                content=ft.Column(
                                    [
                                        ft.Text(
                                            str(item.get("offer_id") or item.get("sku") or "—"),
                                            size=12,
                                            color=row_tokens["title_color"],
                                            weight=ft.FontWeight.W_700,
                                        ),
                                        ft.Container(
                                            padding=ft.Padding.symmetric(horizontal=8, vertical=3),
                                            bgcolor=row_tokens["badge_bg"],
                                            border=ft.Border.all(1, row_tokens["badge_border"]),
                                            border_radius=999,
                                            content=ft.Text(
                                                row_tokens["badge_label"],
                                                size=10,
                                                color=row_tokens["badge_text"],
                                                weight=ft.FontWeight.W_700,
                                            ),
                                        ),
                                    ],
                                    spacing=4,
                                    tight=True,
                                ),
                            ),
                            ft.Container(
                                expand=2,
                                content=ft.Column(
                                    [
                                        ft.Text(_format_currency(item.get("min_price")), size=12, color=row_tokens["title_color"], weight=ft.FontWeight.W_600),
                                        ft.Text("Порог Ozon", size=10, color=row_tokens["detail_color"]),
                                    ],
                                    spacing=3,
                                    tight=True,
                                ),
                            ),
                            ft.Container(
                                expand=2,
                                content=ft.Column(
                                    [
                                        ft.Text(
                                            _format_currency(item.get("requested_price")),
                                            size=12,
                                            color=row_tokens["value_color"],
                                            weight=ft.FontWeight.W_700,
                                        ),
                                        ft.Text(row_tokens["delta_label"], size=10, color=row_tokens["detail_color"]),
                                    ],
                                    spacing=3,
                                    tight=True,
                                ),
                            ),
                            ft.Row(
                                [
                                    ft.IconButton(
                                        icon=ft.Icons.CHECK_CIRCLE_OUTLINE,
                                        icon_color=row_tokens["approve_color"],
                                        tooltip=approve_tooltip,
                                        disabled=approve_disabled,
                                        on_click=lambda _e, x=item: submit_discount_decision(x, "approve"),
                                    ),
                                    ft.IconButton(
                                        icon=ft.Icons.CANCEL_OUTLINED,
                                        icon_color=DANGER,
                                        tooltip="Отклонить заявку",
                                        disabled=is_processing,
                                        on_click=lambda _e, x=item: submit_discount_decision(x, "decline"),
                                    ),
                                ],
                                width=96,
                                spacing=0,
                                alignment=ft.MainAxisAlignment.CENTER,
                            ),
                        ],
                        vertical_alignment=ft.CrossAxisAlignment.CENTER,
                    ),
                )
        )
        discount_dialog_list.controls = rows
        discount_requests_surface.visible = True
        discount_requests_auto_button.disabled = busy or requires_refresh

    def submit_discount_decision(item: dict[str, object], action: str) -> None:
        if discount_dialog_state.get("requires_refresh") or discount_dialog_state["busy"]:
            return
        if process_discount_request_item is None:
            toast(f"Модуль обработки скидок недоступен: {PRICING_IMPORT_ERROR or 'ошибка импорта'}", True)
            return
        processing_ids = discount_dialog_state.get("processing_ids")
        if not isinstance(processing_ids, set):
            processing_ids = set()
            discount_dialog_state["processing_ids"] = processing_ids
        item_id = str(item.get("id") or "")
        processing_ids.add(item_id)
        remove_discount_item(item.get("id"))
        render_discount_dialog()
        if page.controls:
            page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)

        def worker() -> None:
            result = process_discount_request_item(item, action)
            processing_ids.discard(item_id)
            if result.get("ok"):
                toast(str(result.get("message") or "Готово."))
                if not current_discount_items():
                    toast("Все заявки обработаны.")
                    render_discount_dialog()
                else:
                    render_discount_dialog()
                    if page.controls:
                        page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)
            else:
                restore_discount_item(item)
                render_discount_dialog()
                toast(str(result.get("message") or "Не удалось обработать заявку."), True)
                if page.controls:
                    page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)

        page.run_thread(worker)

    def handle_discount_auto(_e=None) -> None:
        if discount_dialog_state["busy"] or discount_dialog_state.get("requires_refresh"):
            return
        plan = build_discount_auto_plan()
        if not plan.get("items"):
            render_discount_dialog()
            return
        if process_discount_request_plan is None:
            toast(f"Модуль обработки скидок недоступен: {PRICING_IMPORT_ERROR or 'ошибка импорта'}", True)
            return
        discount_dialog_state["busy"] = True
        render_discount_dialog()
        if page.controls:
            page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)

        def worker() -> None:
            try:
                result = process_discount_request_plan(plan)
            except Exception as exc:
                result = {"ok": False, "message": f"Ошибка обработки заявок: {exc}"}
            discount_dialog_state["busy"] = False
            if result.get("ok"):
                toast(str(result.get("message") or "Автоматическая обработка завершена."))
                discount_dialog_state["plan"] = {"items": [], "ok": True, "message": str(result.get("message") or "")}
                render_discount_dialog()
                if page.controls:
                    page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)
            else:
                discount_dialog_state["requires_refresh"] = True
                render_discount_dialog()
                toast(str(result.get("message") or "Не удалось автоматически обработать заявки."), True)
                # Aggregate API counts do not identify successful items reliably.
                # Reload before allowing a retry to avoid sending successful items twice.
                if page.controls:
                    page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)

        page.run_thread(worker)

    def open_discount_requests_dialog(_e=None, quiet: bool = False, on_done=None) -> None:
        if discount_dialog_state["busy"]:
            return
        if build_discount_request_plan is None:
            toast(f"Модуль обработки скидок недоступен: {PRICING_IMPORT_ERROR or 'ошибка импорта'}", True)
            return
        set_status("Заявки на скидку", "Собираю рекомендации по заявкам...", ACCENT, busy=True)
        discount_dialog_state["busy"] = True
        discount_dialog_state["plan"] = None
        discount_dialog_state["processing_ids"] = set()
        render_discount_dialog()
        if page.controls:
            page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)

        def worker() -> None:
            try:
                load_discount_plan()
            finally:
                if on_done is not None:
                    # Продолжение (синхронизация цен) — в потоке интерфейса, после загрузки заявок.
                    ui_queue.put(("call", on_done))

        def load_discount_plan() -> None:
            try:
                plan = build_discount_request_plan(ROOT)
            except Exception as exc:
                plan = {"ok": False, "message": f"Не удалось загрузить заявки: {exc}", "items": []}
            set_status("Заявки на скидку", str(plan.get("message") or "Готово"), PRIMARY if plan.get("ok") else DANGER, busy=False)
            if not plan.get("ok"):
                discount_dialog_state["plan"] = {"items": [], "ok": False, "message": str(plan.get("message") or "")}
                discount_dialog_state["busy"] = False
                discount_dialog_state["processing_ids"] = set()
                render_discount_dialog()
                if page.controls:
                    page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)
                toast(str(plan.get("message") or "Не удалось подготовить заявки."), True)
                return
            discount_dialog_state["plan"] = plan
            discount_dialog_state["requires_refresh"] = False
            discount_dialog_state["busy"] = False
            discount_dialog_state["processing_ids"] = set()
            render_discount_dialog()
            apply_discount_dialog_theme()
            if page.controls:
                page.update(discount_dialog_badge, discount_dialog_badge_shell, discount_dialog_status, discount_overview_row, discount_dialog_list, discount_requests_surface, discount_requests_auto_button)
            if not plan.get("items") and not quiet:
                toast(str(plan.get("message") or "Новых заявок нет."))

        page.run_thread(worker)
    discount_requests_refresh_button.on_click = open_discount_requests_dialog
    discount_requests_auto_button.on_click = handle_discount_auto

    clipboard_service = ft.Clipboard()
    page.services.append(clipboard_service)

    async def copy_log_to_clipboard_async() -> None:
        try:
            await clipboard_service.set(get_full_log_text())
            toast("Полный лог скопирован в буфер обмена.")
        except Exception as exc:
            toast(f"Не удалось скопировать лог: {exc}", True)

    def copy_log_to_clipboard(_e=None) -> None:
        page.run_task(copy_log_to_clipboard_async)

    def clear_log(_e=None) -> None:
        set_log("Отображение журнала очищено. Полная история сохранена в файле.")
        if page.controls:
            page.update(log, log_selectable)

    log_dialog = ft.AlertDialog(
        modal=False,
        title=ft.Text("Операционный лог"),
        content=ft.Container(content=log_selectable, width=920, height=620),
        actions=[
            ft.TextButton("Открыть файл журнала", on_click=lambda _e: reveal_path(log_store.path)),
            ft.TextButton("Очистить лог", on_click=clear_log),
            ft.TextButton("Копировать лог", on_click=copy_log_to_clipboard),
            ft.TextButton("Закрыть", on_click=lambda _e: close_log_dialog()),
        ],
        actions_alignment=ft.MainAxisAlignment.END,
        open=False,
    )

    def open_log_dialog() -> None:
        log_dialog.open = True
        page.update(log_dialog, log_selectable)

    def close_log_dialog() -> None:
        log_dialog.open = False
        page.update(log_dialog)

    page.overlay.append(log_dialog)

    def handle_keyboard(e: ft.KeyboardEvent) -> None:
        if e.key == "F2":
            if not ai_state["debug_enabled"]:
                toast("Включите Debug режим через «Инструменты», чтобы открыть лог.", True)
                return
            if log_dialog.open:
                close_log_dialog()
            else:
                open_log_dialog()

    page.on_keyboard_event = handle_keyboard

    async def ui_event_loop() -> None:
        while True:
            processed = False
            while True:
                try:
                    event = ui_queue.get_nowait()
                except queue.Empty:
                    break
                processed = True
                kind = event[0]
                if kind == "log":
                    current = event[1]
                    set_log(current)
                    page.update(log, log_selectable, state, spin)
                elif kind == "log_append":
                    append_log(event[1])
                    page.update(log, log_selectable, state, spin)
                elif kind == "ai_step":
                    _, title, summary, phase_payload = event
                    if phase_payload:
                        phase, phase_detail = phase_payload
                        ai_set_phase(phase, phase_detail)
                    ai_add_trace(title, summary)
                    page.update(ai_trace, ai_phase_label, ai_phase_detail)
                elif kind == "ai_stream_start":
                    _, request_id, note = event
                    if request_id in ai_state["cancelled_request_ids"]:
                        continue
                    existing_payload = ai_state["stream_controls"].get(request_id)
                    if isinstance(existing_payload, dict):
                        note_control = existing_payload.get("note")
                        if note_control is not None:
                            note_control.value = note
                    else:
                        ai_state["stream_controls"][request_id] = ai_append_streaming_placeholder(note)
                    ai_sync_layout()
                    page.update(ai_body, ai_messages, ai_messages_shell, ai_welcome_shell, ai_composer_shell)
                elif kind == "ai_thinking_chunk":
                    _, request_id, chunk = event
                    if request_id in ai_state["cancelled_request_ids"]:
                        continue
                    stream_payload = ai_state["stream_controls"].get(request_id)
                    thinking_list = stream_payload.get("thinking_list") if isinstance(stream_payload, dict) else None
                    thinking_current = stream_payload.get("thinking_current") if isinstance(stream_payload, dict) else None
                    thinking_anchor = stream_payload.get("thinking_anchor") if isinstance(stream_payload, dict) else None
                    thinking_toggle = stream_payload.get("thinking_toggle") if isinstance(stream_payload, dict) else None
                    thinking_wrap = stream_payload.get("thinking_wrap") if isinstance(stream_payload, dict) else None
                    if thinking_list is not None and thinking_current is not None and thinking_wrap is not None and thinking_anchor is not None:
                        parts = chunk.split("\n")
                        for idx, part in enumerate(parts):
                            if idx == 0:
                                thinking_current.value = (thinking_current.value or "") + part
                            else:
                                thinking_current = ft.Text(part, size=12, color=ai_surface_tokens()["muted"], selectable=True)
                                thinking_list.controls.insert(max(len(thinking_list.controls) - 1, 0), thinking_current)
                        stream_payload["thinking_current"] = thinking_current
                        if thinking_anchor in thinking_list.controls:
                            thinking_list.controls.remove(thinking_anchor)
                        thinking_list.controls.append(thinking_anchor)
                        thinking_wrap.visible = True
                        if thinking_toggle is not None:
                            thinking_toggle.visible = True
                        page.update(ai_messages, thinking_wrap, thinking_list)
                elif kind == "ai_stream_reset":
                    _, request_id = event
                    stream_payload = ai_state["stream_controls"].get(request_id)
                    text_control = stream_payload.get("text") if isinstance(stream_payload, dict) else None
                    if text_control is not None and request_id not in ai_state["cancelled_request_ids"]:
                        text_control.value = ""
                        page.update(ai_messages)
                elif kind == "ai_stream_chunk":
                    _, request_id, chunk = event
                    if request_id in ai_state["cancelled_request_ids"]:
                        continue
                    stream_payload = ai_state["stream_controls"].get(request_id)
                    text_control = stream_payload.get("text") if isinstance(stream_payload, dict) else None
                    if text_control is not None:
                        text_control.value = (text_control.value or "") + chunk
                        page.update(ai_messages)
                elif kind == "prompt":
                    _, label, source_line, holder, waiter = event
                    prompt_state["holder"] = holder
                    prompt_state["event"] = waiter
                    period_label = infer_prompt_period(source_line)
                    prompt_label.value = (
                        f"Введите сумму расходов на маркетинг для отчёта за {period_label}."
                        if period_label
                        else "Введите сумму расходов на маркетинг для отчёта, который сейчас строится."
                    )
                    prompt_hint.value = "Укажите общую сумму в рублях. Если расходов не было, оставьте 0."
                    prompt_context.value = f"Источник запроса: {source_line}" if source_line else ""
                    prompt_input.label = f"{label} (руб.)" if label else "Сумма маркетинга (руб.)"
                    prompt_input.value = "0"
                    apply_prompt_dialog_theme()
                    set_status("Ожидается ввод", period_label or label, ACCENT, busy=False)
                    prompt_dialog.open = True
                    page.update(prompt_label, prompt_hint, prompt_context, prompt_input, prompt_dialog, log)
                elif kind == "finish":
                    _, title, code, output = event
                    finish(title, code, output)
                elif kind == "run_job":
                    _, title, cmd = event
                    job(title, cmd)
                elif kind == "ads_rebuild":
                    ads_build(keep_result=True)
                elif kind == "call":
                    event[1]()
                elif kind == "ai_finish":
                    _, request_id, prompt, answer, meta = event
                    if request_id in ai_state["cancelled_request_ids"]:
                        ai_state["stream_controls"].pop(request_id, None)
                        ai_state["request_started_at"].pop(request_id, None)
                        ai_state["cancelled_request_ids"].discard(request_id)
                        page.update(log, log_selectable)
                        continue
                    if ai_state["active_request_id"] not in (None, request_id):
                        continue
                    ai_state["active_request_id"] = None
                    elapsed = None
                    started_at = ai_state["request_started_at"].pop(request_id, None)
                    if started_at is not None:
                        elapsed = max(0.0, time.perf_counter() - started_at)
                    elapsed_note = ai_format_elapsed(elapsed)
                    streamed_payload = ai_state["stream_controls"].pop(request_id, None)
                    streamed_control = streamed_payload.get("text") if isinstance(streamed_payload, dict) else None
                    streamed_host = streamed_payload.get("content_host") if isinstance(streamed_payload, dict) else None
                    streamed_note = streamed_payload.get("note") if isinstance(streamed_payload, dict) else None
                    streamed_thinking_list = streamed_payload.get("thinking_list") if isinstance(streamed_payload, dict) else None
                    streamed_toggle = streamed_payload.get("thinking_toggle") if isinstance(streamed_payload, dict) else None
                    streamed_wrap = streamed_payload.get("thinking_wrap") if isinstance(streamed_payload, dict) else None
                    if answer:
                        if streamed_control is not None:
                            streamed_control.value = answer
                            if streamed_host is not None:
                                streamed_host.content = ai_markdown(answer)
                            if streamed_note is not None:
                                streamed_note.value = elapsed_note or ""
                            if streamed_thinking_list is not None and streamed_wrap is not None:
                                has_thinking = any(isinstance(ctrl, ft.Text) and (ctrl.value or "").strip() for ctrl in streamed_thinking_list.controls)
                                if not has_thinking:
                                    streamed_wrap.visible = False
                                elif streamed_toggle is not None:
                                    streamed_toggle.visible = True
                        else:
                            ai_append_message("assistant", answer, elapsed_note)
                        ai_state["conversation_history"].append({"role": "user", "content": prompt})
                        ai_state["conversation_history"].append({"role": "assistant", "content": answer})
                        if len(ai_state["conversation_history"]) > 6:
                            ai_state["conversation_history"] = ai_state["conversation_history"][-6:]
                        ai_set_busy(False, "Ответ готов. Можно продолжать разбор месяца.")
                        set_status("AI ответил", "Локальный анализ завершён.", PRIMARY, busy=False)
                    else:
                        ai_append_message("assistant", str(meta), elapsed_note or "Диагностика runtime")
                        ai_set_busy(False, str(meta))
                        set_status("AI недоступен", str(meta), DANGER, busy=False)
                    ai_sync_layout()
                    page.update(ai_body, ai_messages, ai_messages_shell, ai_welcome_shell, ai_composer_shell, ai_busy_label, state, status_detail, spin, cancel_row)
            await asyncio.sleep(0.01 if processed else 0.03)

    ai_refresh_status(push_update=False)
    ai_refresh_tools_menu()
    ai_apply_chat_theme()

    dashboard_sync_label = ft.Text("Обновляю отчёт: загружаю продажи, расходы и выплаты из Ozon…",
                                   size=13, weight=ft.FontWeight.W_600, text_align=ft.TextAlign.CENTER)
    dashboard_sync_cancel = ft.TextButton("Отмена", visible=False, on_click=lambda _e: wb_dash.cancel())
    dashboard_sync_overlay = ft.Container(
        left=0, top=0, right=0, bottom=0,
        visible=False,
        border_radius=26,
        alignment=ft.Alignment(0, -1),
        padding=ft.Padding.only(top=150),
        data="preserve_excel_fill",
        content=ft.Container(
            padding=ft.Padding.symmetric(horizontal=28, vertical=20),
            border_radius=18,
            data="preserve_excel_fill",
            content=ft.Column(
                [ft.ProgressRing(width=36, height=36, stroke_width=3), dashboard_sync_label, dashboard_sync_cancel],
                spacing=14, tight=True, horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            ),
        ),
    )
    dashboard_metrics_area = ft.Column(spacing=16)

    def set_dashboard_sync(active: bool, text: str | None = None) -> None:
        """Пока обновляется отчёт, показатели закрыты заглушкой с загрузкой и не принимают клики.
        text — подпись под индикатором (WB пишет сюда ход загрузки); у WB есть своя кнопка «Отмена»."""
        dark = current_dark["value"]
        if text:
            dashboard_sync_label.value = text
        elif active:
            dashboard_sync_label.value = "Обновляю отчёт: загружаю продажи, расходы и выплаты из Ozon…"
        dashboard_sync_cancel.visible = active and is_wb_store()
        dashboard_sync_overlay.visible = active
        dashboard_sync_overlay.bgcolor = ("#CC1E2328" if dark else "#CCF5F7FB") if active else None
        dashboard_sync_label.color = "#F3F4F6" if dark else TEXT
        dashboard_sync_overlay.content.bgcolor = "#2A3136" if dark else "#FFFFFF"
        dashboard_sync_overlay.content.border = ft.Border.all(1, "#465056" if dark else BORDER)
        dashboard_metrics_area.disabled = active
        if page.controls:
            page.update(dashboard_sync_overlay, dashboard_metrics_area)

    dashboard_m.label = "Месяц"
    dashboard_m.width = 190
    dashboard_y.label = "Год"
    dashboard_y.width = 120
    dashboard_more = ft.PopupMenuButton(
        icon=ft.Icons.MORE_HORIZ,
        tooltip="Ещё",
        items=[
            ft.PopupMenuItem(content=ft.Text("Открыть отчёт в Excel"), on_click=lambda _e: open_selected_dashboard_report()),
            ft.PopupMenuItem(content=ft.Text("Открыть папку с отчётами"), on_click=lambda _e: reveal_path(wb_reports_dir(ROOT), "reports/WB") if is_wb_store()
                             else reveal_path(ozon_reports_dir(ROOT), "reports/Ozon")),
            ft.PopupMenuItem(content=ft.Text("Себестоимость WB"), visible=False, on_click=lambda _e: wb_dash.open_costs()),
        ],
    )
    dash_header = ft.ResponsiveRow(
        [
            ft.Container(
                col={"xs": 12, "lg": 5},
                content=ft.Column([ft.Text("Бизнес-сводка", size=26, weight=ft.FontWeight.W_800, color=TEXT), dashboard_freshness], spacing=2),
            ),
            ft.Container(
                col={"xs": 12, "lg": 7},
                alignment=ft.Alignment(1, 0),
                content=ft.Row(
                    [
                        dashboard_m,
                        dashboard_y,
                        ft.Button("Обновить", icon=ft.Icons.REFRESH, bgcolor=PRIMARY, color="white", on_click=lambda _e: launch_dashboard_report()),
                        dashboard_more,
                    ],
                    spacing=10,
                    alignment=ft.MainAxisAlignment.END,
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                ),
            ),
        ],
        vertical_alignment=ft.CrossAxisAlignment.CENTER,
        run_spacing=12,
    )

    dash = ft.Column(
        [
            dash_header,
            ft.Stack([dashboard_metrics_area, dashboard_sync_overlay]),
            ft.Container(
                padding=ft.Padding.only(top=10, left=4),
                content=ft.Column([
                    ft.Text("Динамика", size=20, weight=ft.FontWeight.W_700, color=TEXT),
                    ft.Text("Прибыль, оборот или рентабельность по месяцам", size=13, color=MUTED),
                ], spacing=2),
            ),
            ft.Container(
                padding=20,
                border_radius=18,
                bgcolor=semantic_surface("neutral", dark=False),
                border=ft.Border.all(1, BORDER),
                content=ft.Column(
                    [
                        ft.Row(
                            [
                                dashboard_chart_type,
                                dashboard_chart_from_m,
                                dashboard_chart_from_y,
                                ft.Text("—", size=16, color=MUTED),
                                dashboard_chart_to_m,
                                dashboard_chart_to_y,
                                ft.Button("Построить", icon=ft.Icons.SHOW_CHART, bgcolor=PRIMARY, color="white", on_click=build_dashboard_chart),
                                ft.TextButton("Интерактивный", on_click=open_interactive_dashboard_chart),
                            ],
                            spacing=10,
                            run_spacing=10,
                            wrap=True,
                            vertical_alignment=ft.CrossAxisAlignment.CENTER,
                        ),
                        dashboard_chart_panel,
                    ],
                    spacing=16,
                ),
            ),
        ],
        spacing=14,
    )

    for control, width in ((abc_fm, 180), (abc_fy, 120), (abc_tm, 180), (abc_ty, 120)):
        control.width = width
    abc_fm.label, abc_fy.label, abc_tm.label, abc_ty.label = "С месяца", "Год", "По месяц", "Год"

    abc_report_pick.on_select = lambda e: abc_select(e.control.value)
    abc_more = ft.PopupMenuButton(
        icon=ft.Icons.MORE_HORIZ,
        tooltip="Ещё",
        items=[
            ft.PopupMenuItem(content=ft.Text("Открыть отчёт в Excel"), on_click=abc_open_selected),
            ft.PopupMenuItem(content=ft.Text("Папка с анализами"), on_click=lambda _e: reveal_path(ROOT / "ABC&XYZ reports", "ABC&XYZ reports")),
        ],
    )
    abc_create_panel = ft.Container(
        visible=False,
        padding=20,
        border_radius=18,
        bgcolor=semantic_surface("neutral", dark=False),
        border=ft.Border.all(1, BORDER),
        content=ft.Column(
            [
                ft.Text("Новый анализ", size=15, weight=ft.FontWeight.W_700, color=TEXT),
                ft.Text("Анализ строится по месячным отчётам. Недостающие месяцы программа сформирует сама.", size=12, color=MUTED),
                ft.Row(
                    [
                        abc_fm, abc_fy, ft.Text("—", size=16, color=MUTED), abc_tm, abc_ty,
                        ft.Button("Провести анализ", icon=ft.Icons.PLAY_ARROW_ROUNDED, bgcolor=PRIMARY, color="white", on_click=lambda _e: launch_abc_analysis()),
                        ft.TextButton("Отмена", on_click=lambda _e: abc_toggle_create(False)),
                    ],
                    spacing=10, run_spacing=10, wrap=True, vertical_alignment=ft.CrossAxisAlignment.CENTER,
                ),
            ],
            spacing=10,
        ),
    )
    reports = ft.Column(
        [
            ft.ResponsiveRow(
                [
                    ft.Container(
                        col={"xs": 12, "lg": 5},
                        content=ft.Column([ft.Text("ABC/XYZ-анализ", size=26, weight=ft.FontWeight.W_800, color=TEXT), abc_subtitle], spacing=2),
                    ),
                    ft.Container(
                        col={"xs": 12, "lg": 7},
                        alignment=ft.Alignment(1, 0),
                        content=ft.Row(
                            [
                                abc_report_pick,
                                ft.Button("Новый анализ", icon=ft.Icons.ADD, bgcolor=PRIMARY, color="white", on_click=lambda _e: abc_toggle_create()),
                                abc_more,
                            ],
                            spacing=10, alignment=ft.MainAxisAlignment.END, vertical_alignment=ft.CrossAxisAlignment.CENTER,
                        ),
                    ),
                ],
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
                run_spacing=12,
            ),
            abc_create_panel,
            ft.Stack([abc_body, abc_sync_overlay]),
        ],
        spacing=16,
        horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
    )

    def pricing_panel(content: ft.Control, padding=20, **kwargs) -> ft.Container:
        """Панель в стиле бизнес-сводки: светлая поверхность, тонкая рамка, скругление 18."""
        return ft.Container(content=content, padding=padding, border_radius=18, bgcolor=semantic_surface("neutral", dark=False),
                            border=ft.Border.all(1, BORDER), **kwargs)

    def save_margin(_e=None) -> None:
        error = valid_margin(margin_min.value or "", margin_des.value or "")
        if error:
            toast(error, True)
            return
        job("Сохранение маржи", [str(PYTHON), str(SCRIPTS / "ui_actions.py"), "save-margin",
                                 "--min-margin", (margin_min.value or "").replace(",", "."),
                                 "--desired-margin", (margin_des.value or "").replace(",", ".")])

    for field, label in ((margin_min, "Мин. маржа"), (margin_des, "Желаемая")):
        field.label = label
        field.width = 120
        field.dense = True
        field.tooltip = "Доля, например 0.25 = 25%. Минимальная маржа — порог «ниже нельзя», желаемая — цель расчёта."
        field.on_submit = save_margin

    pricing_more = ft.PopupMenuButton(
        icon=ft.Icons.MORE_HORIZ,
        tooltip="Ещё",
        items=[
            ft.PopupMenuItem(content=ft.Text("Обновить минимальные цены в Ozon"),
                             on_click=lambda _e: job("Обновление минимальных цен", [str(PYTHON), str(SCRIPTS / "ui_actions.py"), "update-min-prices"])),
            ft.PopupMenuItem(content=ft.Text("Перечитать таблицу из costs.xlsx"), on_click=handle_pricing_costs_refresh),
            ft.PopupMenuItem(content=ft.Text("Открыть costs.xlsx"), on_click=lambda _e: reveal_path(ROOT / "costs.xlsx", "costs.xlsx")),
            ft.PopupMenuItem(content=ft.Text("Открыть папку проекта"), on_click=lambda _e: reveal_path(ROOT, "корень проекта")),
        ],
    )
    pricing_header = ft.ResponsiveRow(
        [
            ft.Container(
                col={"xs": 12, "lg": 4},
                content=ft.Column([ft.Text("Цены и скидки", size=26, weight=ft.FontWeight.W_800, color=TEXT), pricing_costs_status_shell], spacing=2),
            ),
            ft.Container(
                col={"xs": 12, "lg": 8},
                alignment=ft.Alignment(1, 0),
                content=ft.Row(
                    [
                        margin_min,
                        margin_des,
                        ft.IconButton(icon=ft.Icons.SAVE_OUTLINED, tooltip="Сохранить маржу", on_click=save_margin),
                        ft.Button("Рассчитать цены", icon=ft.Icons.CALCULATE_OUTLINED, bgcolor=ACCENT, color="white",
                                  tooltip="Пересчитает минимальную и желательную цену продажи по марже и расходам прошлого месяца",
                                  on_click=lambda _e: launch_recommended_prices()),
                        ft.Button("Синхронизировать", icon=ft.Icons.SYNC, bgcolor=PRIMARY, color="white",
                                  tooltip="Обновит текущие цены, рентабельность, прибыль и бустинг из Ozon. Запускается сама при входе в раздел",
                                  on_click=lambda _e: job(PRICING_SYNC_TITLE, [str(PYTHON), str(SCRIPTS / "ui_actions.py"), "refresh-pricing"])),
                        pricing_more,
                    ],
                    spacing=10,
                    wrap=True,
                    run_spacing=8,
                    alignment=ft.MainAxisAlignment.END,
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                ),
            ),
        ],
        vertical_alignment=ft.CrossAxisAlignment.CENTER,
        run_spacing=12,
    )

    discount_requests_surface = ft.Container(
        visible=False,
        content=ft.Column([discount_overview_row, discount_dialog_list], spacing=10),
    )
    discount_panel = pricing_panel(
        ft.Column(
            [
                ft.Row(
                    [
                        ft.Text("Заявки на скидку", size=15, weight=ft.FontWeight.W_700, color=TEXT),
                        discount_dialog_badge_shell,
                        ft.Container(expand=True, content=discount_dialog_status),
                        discount_requests_auto_button,
                        discount_requests_refresh_button,
                    ],
                    spacing=10,
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                ),
                discount_requests_surface,
            ],
            spacing=10,
        ),
        padding=ft.Padding.only(left=20, right=12, top=10, bottom=10),
    )

    pricing_sync_label = ft.Text("Синхронизация с Ozon: обновляю цены, рентабельность, прибыль и бустинг…",
                                 size=13, weight=ft.FontWeight.W_600, text_align=ft.TextAlign.CENTER)
    pricing_sync_overlay = ft.Container(
        left=0, top=0, right=0, bottom=0,
        visible=False,
        border_radius=18,
        # Таблица длинная: подпись загрузки — у её верхнего края, а не в середине за экраном.
        alignment=ft.Alignment(0, -1),
        padding=ft.Padding.only(top=120),
        data="preserve_excel_fill",
        content=ft.Column(
            [ft.ProgressRing(width=36, height=36, stroke_width=3), pricing_sync_label],
            spacing=14, tight=True, horizontal_alignment=ft.CrossAxisAlignment.CENTER,
        ),
    )
    pricing_costs_surface = ft.Container(content=ft.Stack([pricing_costs_list, pricing_sync_overlay]))
    pricing_table_panel = pricing_panel(
        ft.Column(
            [
                ft.Row(
                    [
                        ft.Text("Товары", size=15, weight=ft.FontWeight.W_700, color=TEXT),
                        ft.Text("Артикул и себестоимость можно править прямо в таблице. Наведите на бустинг, чтобы увидеть акции.",
                                size=12, color=MUTED, expand=True, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS),
                        pricing_costs_save_button,
                    ],
                    spacing=10,
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                ),
                pricing_costs_surface,
            ],
            spacing=8,
        ),
        padding=ft.Padding.only(left=20, right=20, top=12, bottom=8),
    )

    pricing = ft.Column(
        [pricing_header, pricing_stats_area, discount_panel, pricing_table_panel],
        spacing=14,
    )

    # ---------------------------------------------------------------- Акции: предельная цена (ползунок по артикулу)
    # Один цвет на акцию во всём разделе (календарь, шкала ползунка, значки): слот закрепляется за акцией один раз,
    # в порядке id — так цвет не зависит от того, какой артикул или блок отрисован первым. Белый текст читается на всех.
    ACTION_PALETTE = {False: ["#2a78d6", "#d9622b", "#16936a", "#7b52e8", "#b07d10", "#cf3f74", "#138a9e", "#8d6440"],
                      True: ["#4f93ea", "#e9783f", "#25ad82", "#9877f5", "#c99a1e", "#e3588b", "#24a5ba", "#a47d58"]}
    action_color_slots: dict[str, int] = {}

    def register_action_colors(action_ids) -> None:
        for action_id in sorted({str(a) for a in action_ids}, key=lambda a: (len(a), a)):
            action_color_slots.setdefault(action_id, len(action_color_slots))

    def action_color(action_id) -> str:
        palette = ACTION_PALETTE[current_dark["value"]]
        slot = action_color_slots.setdefault(str(action_id), len(action_color_slots))
        return palette[slot % len(palette)]

    def short_action_name(name: str) -> str:
        """«Эластичный бустинг. Без ограничения срока действия» → «Эластичный бустинг»; «Максимальный …» → «Макс. …»."""
        short = name.split(". ")[0].rstrip(".").replace("Максимальный бустинг", "Макс. бустинг")
        return short.split(" с ")[0] if len(short) > 28 else short

    LIMIT_SLIDER_INSET = 12
    LIMIT_PAGE = 12  # карточек за раз: 50+ ползунков разом заметно тормозят открытие раздела и смену темы
    limit_state: dict[str, object] = {"data": None, "busy": False, "error": None, "started": False, "prices": {}, "filter": "all",
                                      "elasticity": None, "shown": LIMIT_PAGE, "collapsed": False, "action_filter": None}
    limit_panel = ft.Container(padding=18, border_radius=18, data="preserve_excel_fill")

    def limit_articles() -> list[dict]:
        data = limit_state.get("data") or {}
        return list(data.get("articles") or [])

    def limit_current(article: dict) -> int | None:
        return int(round(article["price"])) if article.get("price") else None

    def limit_chosen(article: dict) -> int | None:
        return limit_state["prices"].get(article["offer_id"], limit_current(article))

    def limit_changed() -> dict[str, int]:
        return {a["offer_id"]: limit_state["prices"][a["offer_id"]] for a in limit_articles()
                if a["offer_id"] in limit_state["prices"] and limit_state["prices"][a["offer_id"]] != limit_current(a)}

    def limit_load(_e=None) -> None:
        if limit_state["busy"] or limit_price is None:
            return
        limit_state.update(busy=True, error=None, started=True)
        render_limit_panel()

        def worker() -> None:
            try:
                data = limit_price.load_board(ROOT)
                if not data.get("ok"):
                    limit_state["error"] = data.get("message")
                limit_state["data"] = data
                limit_state["prices"] = {}
            except Exception as exc:
                limit_state["error"] = str(exc)
            finally:
                limit_state["busy"] = False
            ui_queue.put(("call", render_limit_panel))

        page.run_thread(worker)

    def limit_elasticity() -> float:
        if limit_state["elasticity"] is None:
            limit_state["elasticity"] = limit_price.load_elasticity(ROOT) if limit_price else 2.0
        return float(limit_state["elasticity"])

    def limit_set_elasticity(e: ft.ControlEvent) -> None:
        limit_state["elasticity"] = float(e.control.value)
        try:
            limit_price.save_elasticity(ROOT, limit_state["elasticity"])
        except OSError as exc:
            toast(f"Не удалось сохранить настройку: {exc}", True)
        render_limit_panel()

    def limit_set_all(mode: str) -> None:
        for article in limit_articles():
            price = limit_price.balance_price(article, limit_elasticity()) if mode == "auto" else limit_current(article)
            if price is not None:
                limit_state["prices"][article["offer_id"]] = price
        render_limit_panel()

    def limit_money(value) -> str:
        return _format_currency(value) if value else "—"

    def limit_range(article: dict) -> tuple[int, int]:
        values = [v for v in (article.get("min_price"), article.get("desired_price"), article.get("price"), article.get("ozon_min"))
                  if v] + [a["threshold"] for a in article["actions"]]
        low, high = math.floor(min(values) * 0.95), math.ceil(max(values) * 1.03)
        return low, max(high, low + 10)

    def limit_band(low: int, high: int, start: float, end: float, color: str, height: int, tooltip: str | None = None) -> ft.Control:
        """Отрезок [start, end] на шкале [low, high] — те же пропорции, что у ползунка под ним."""
        span = high - low
        left = max(0, int((min(start, end) - low) / span * 1000))
        mid = max(1, int((max(start, end) - min(start, end)) / span * 1000))
        right = max(0, 1000 - left - mid)
        parts: list[ft.Control] = []
        if left:
            parts.append(ft.Container(expand=left, height=height))
        parts.append(ft.Container(expand=mid, height=height, bgcolor=color, border_radius=height, tooltip=tooltip, data="preserve_excel_fill"))
        if right:
            parts.append(ft.Container(expand=right, height=height))
        return ft.Row(parts, spacing=0)

    def limit_status(article: dict, price: int) -> list[ft.Control]:
        """Компактные значки акций при этой цене: цвет акции, участие с бустингом, войдёт / выйдет."""
        pal = dash_pal()
        chips: list[ft.Control] = []
        for entry in limit_price.evaluate(article, price):
            color = action_color(entry["action_id"])
            boost = entry["boost"]
            future = entry["change"] == "future"
            when = f"с {entry['start'][8:10]}.{entry['start'][5:7]} " if entry.get("start") else ""
            if future and entry["participating"]:
                own = entry["action_price"]
                below = bool(own and article.get("min_price") and own < article["min_price"])
                text, text_color = f"{when}по {limit_money(own)}", (pal["bad"] if below else pal["ink"])
            elif future:
                text, text_color = f"{when}при ≤ {limit_money(entry['threshold'])}", pal["muted"]
            elif entry["in"]:
                text, text_color = (f"{boost:g}%" if boost is not None else "в акции"), pal["ink"]
            else:
                text, text_color = f"≤ {limit_money(entry['threshold'])}", pal["muted"]
            change = {"enter": ("↑ войдёт", pal["good"]), "exit": ("↓ выйдет", pal["bad"])}.get(entry["change"])
            active = entry["in"] and not future or (future and entry["participating"])
            tooltip = (f"{entry['name']}\nПорог участия: {limit_money(entry['threshold'])}"
                       + (f"\nПорог акции у Ozon — {limit_money(entry['ozon_threshold'])}, но товар сейчас в ней по "
                          f"{limit_money(entry['action_price'])}: считаем, что по этой цене он остаётся" if entry.get("above_limit") else "")
                       + ("\nАкция ещё не началась: участие «с даты» со своей ценой, ползунок его не меняет" if future else "")
                       + (f"\nБаллы Ozon за скидку: {entry['points']:g}% от цены продажи — в счёт комиссии и услуг, входят в прибыль"
                          if entry.get("points") else "")
                       + ("\nДобавлен Ozon автоматически" if entry["added_by_ozon"] else "")
                       + ("\nOzon: скидка не соответствует условиям — 3 рабочих дня на исправление цены" if entry["alert"] else ""))
            chips.append(ft.Container(
                padding=ft.Padding.only(left=4, right=8, top=3, bottom=3), border_radius=999, data="preserve_excel_fill",
                bgcolor=pal["surface"], border=ft.Border.all(1, color if active else pal["border"]), tooltip=tooltip,
                content=ft.Row([
                    ft.Container(width=14, height=14, border_radius=7, bgcolor=color if active else None,
                                 border=None if active else ft.Border.all(2, color), data="preserve_excel_fill"),
                    dtext(short_action_name(entry["name"]), 11, pal["muted"], max_lines=1, overflow=ft.TextOverflow.ELLIPSIS),
                    dtext(text, 11, text_color, ft.FontWeight.W_700),
                    *([dtext(f"распродажа +{limit_price.SALE_BONUS:g}%", 11, pal["warn"], ft.FontWeight.W_700)]
                      if entry.get("sale") and entry["in"] and not future else []),
                    *([dtext(change[0], 11, change[1], ft.FontWeight.W_700)] if change else []),
                    *([dtext("⚠", 11, pal["warn"])] if entry["alert"] else []),
                ], spacing=5, tight=True, vertical_alignment=ft.CrossAxisAlignment.CENTER),
            ))
        return chips

    def limit_warnings(article: dict, price: int) -> list[tuple[str, str]]:
        """Предупреждения (текст, цвет). Пусто — всё в порядке."""
        pal = dash_pal()
        notes = []
        if article.get("min_price") and price < article["min_price"]:
            notes.append((f"ниже вашей минимальной {limit_money(article['min_price'])} — продажа в минус", pal["bad"]))
        if article.get("ozon_min") and price < article["ozon_min"]:
            notes.append((f"ниже ограничения Ozon {limit_money(article['ozon_min'])} — Ozon опустит его до этой цены", pal["warn"]))
        for entry in article["actions"]:
            own = entry["action_price"]
            if entry["start"] and entry["participating"] and own and article.get("min_price") and own < article["min_price"]:
                notes.append((f"с {entry['start'][8:10]}.{entry['start'][5:7]} в «{short_action_name(entry['name'])}» по {limit_money(own)} — "
                              "ниже минимальной; уберите в кабинете Ozon (участие «с даты» ползунок не меняет)", pal["bad"]))
        if any(entry["alert"] for entry in article["actions"]):
            notes.append(("Ozon пометил скидку как несоответствующую условиям: 3 рабочих дня на исправление", pal["warn"]))
        return notes

    def limit_row(article: dict, live: dict) -> ft.Control:
        """Карточка артикула в одну строку: сведения · шкала с ползунком и значками акций · цена, прибыль, бустинг."""
        pal = dash_pal()
        low, high = limit_range(article)
        current = limit_current(article)
        chosen = limit_chosen(article) or current or low
        chosen = min(max(chosen, low), high)
        price_big, profit_big, boost_big = (dtext("", 20, weight=ft.FontWeight.W_800, max_lines=1, no_wrap=True) for _ in range(3))
        price_note, profit_note, boost_note = (dtext("", 11, pal["muted"]) for _ in range(3))
        status_row = ft.Row(spacing=6, run_spacing=6, wrap=True)
        warning_text = dtext("", 11, pal["muted"], visible=False)
        best = limit_price.balance_price(article, limit_elasticity())
        best_button = ft.Container(
            visible=bool(best), padding=ft.Padding.symmetric(horizontal=8, vertical=3), border_radius=999, data="preserve_excel_fill",
            border=ft.Border.all(1, pal["accent"]), tooltip="Цена с наибольшей оценкой «прибыль с продажи × продажи» между минимальной и желательной",
            content=ft.Row([ft.Icon(ft.Icons.AUTO_FIX_HIGH, size=12, color=pal["accent"]), dtext(f"баланс {limit_money(best)}", 11, pal["accent"], ft.FontWeight.W_700)],
                           spacing=4, tight=True),
        )

        def pct(value: float) -> str:
            return f"{value:.1f}%".replace(".", ",").replace(",0%", "%")

        def refresh(price: int, push: bool) -> None:
            profit = limit_price.unit_profit(article, price)
            boost = limit_price.total_boost(article, price)
            sale = limit_price.sale_bonus(article, price)
            now_profit = limit_price.unit_profit(article, current) if current else None
            now_effect = (limit_price.total_boost(article, current) + limit_price.sale_bonus(article, current)) if current else 0.0
            moved = bool(current and price != current)
            price_big.value = limit_money(price)
            price_big.color = pal["accent"] if moved else pal["ink"]
            price_note.value = (f"{'+' if price > current else '−'}{_format_currency(abs(price - current))} к {limit_money(current)}" if moved
                                else "текущая")
            profit_big.value = _format_currency(round(profit)) if profit is not None else "—"
            profit_big.color = pal["bad"] if profit is not None and profit < 0 else (pal["good"] if profit is not None else pal["muted"])
            profit_note.value = (f"сейчас {_format_currency(round(now_profit))}" if now_profit is not None and moved
                                 else ("с продажи" if profit is not None else "нет отчёта или себестоимости"))
            points = limit_price.points_percent(article, price)
            if points and profit is not None:
                profit_note.value += f" · в т.ч. баллы Ozon {pct(points)}: +{_format_currency(round(price * points / 100))}"
            effect = boost + sale
            boost_big.value = f"+{effect:.0f}%" if effect else "0%"
            boost_big.color = pal["ink"] if effect else pal["muted"]
            parts = ([f"бустинг {pct(boost)}"] if boost else []) + ([f"распродажа {pct(sale)}"] if sale else [])
            boost_note.value = (" + ".join(parts) or "нет акций") + (f" · сейчас +{pct(now_effect)}" if moved and now_effect else
                                                                     (" · сейчас 0%" if moved else ""))
            status_row.controls = limit_status(article, price)
            notes = limit_warnings(article, price)
            warning_text.visible = bool(notes)
            warning_text.value = "⚠ " + " · ".join(text for text, _ in notes) if notes else ""
            warning_text.color = notes[0][1] if notes else pal["muted"]
            if push and page.controls:
                # Счётчики в шапке панели — сразу, без перерисовки всех артикулов.
                count = len(limit_changed())
                live["apply"].content = f"Применить в Ozon ({count})"
                live["apply"].disabled = not count or bool(limit_state["busy"])
                live["reset"].disabled = not count or bool(limit_state["busy"])
                live["changed"].value = f"Изменённые ({count})"
                page.update(stats, status_row, warning_text, live["apply"], live["reset"], live["changed"])

        def on_change(e: ft.ControlEvent) -> None:
            price = int(round(float(e.control.value)))
            limit_state["prices"][article["offer_id"]] = price
            refresh(price, push=True)

        def stat(label: str, big: ft.Control, note: ft.Control) -> ft.Control:
            return ft.Column([dtext(label, 11, pal["muted"], ft.FontWeight.W_600), big, note], spacing=0, expand=True)

        stats = ft.Row([stat("Цена", price_big, price_note), stat("Прибыль", profit_big, profit_note), stat("Эффект акций", boost_big, boost_note)],
                       spacing=10, vertical_alignment=ft.CrossAxisAlignment.START)
        refresh(chosen, push=False)

        # Шкала: тонкая дорожка, зелёный отрезок мин–жел и засечки порогов акций цветом акции (левее засечки — товар в акции).
        span = high - low
        track_layers: list[ft.Control] = [
            ft.Container(left=0, right=0, top=7, height=4, bgcolor=pal["separator"], border_radius=2, data="preserve_excel_fill"),
            ft.Container(left=0, right=0, top=5, height=8, content=limit_band(
                low, high, article["min_price"], article["desired_price"] or article["min_price"], pal["good"], 8,
                f"Между минимальной {limit_money(article['min_price'])} и желательной {limit_money(article['desired_price'])}")),
        ]
        for entry in sorted(article["actions"], key=lambda e: e["threshold"]):
            tick = max(3, int(span * 0.006))
            track_layers.append(ft.Container(left=0, right=0, top=0, height=18, content=limit_band(
                low, high, entry["threshold"] - tick / 2, entry["threshold"] + tick / 2, action_color(entry["action_id"]), 18,
                f"{entry['name']}: участие при цене до {limit_money(entry['threshold'])}")))
        track = ft.Container(padding=ft.Padding.symmetric(horizontal=LIMIT_SLIDER_INSET), content=ft.Stack(track_layers, height=18))
        slider = ft.Slider(min=low, max=high, value=chosen, divisions=high - low, label="{value}", round=0,
                           active_color=pal["accent"], padding=ft.Padding.symmetric(horizontal=LIMIT_SLIDER_INSET), on_change=on_change)

        def go_best(_e=None) -> None:
            price = min(max(best, low), high)
            slider.value = price
            limit_state["prices"][article["offer_id"]] = price
            refresh(price, push=False)
            page.update(slider)
            refresh(price, push=True)

        best_button.on_click = go_best
        info = ft.Column([
            ft.Row([dtext(article["offer_id"], 16, weight=ft.FontWeight.W_800), best_button], spacing=8,
                   vertical_alignment=ft.CrossAxisAlignment.CENTER, wrap=True),
            dtext(f"мин. {limit_money(article['min_price'])} · жел. {limit_money(article['desired_price'])}", 11, pal["muted"]),
            *([dtext(f"ограничение Ozon {limit_money(article['ozon_min'])}", 11,
                     pal["warn"] if article.get("min_price") and article["ozon_min"] < article["min_price"] else pal["muted"],
                     tooltip="Минимальная цена в Ozon («ограничение для акций»). Если она ниже вашей минимальной, Ozon может сам "
                             "добавить товар в акцию и опустить цену. Обновить: «Цены и скидки» → «⋯» → «Обновить минимальные цены в Ozon».")]
              if article.get("ozon_min") else []),
        ], spacing=2)
        scale = ft.Column([
            track,
            slider,
            ft.Row([dtext(limit_money(low), 10, pal["muted"]), dtext(limit_money(high), 10, pal["muted"])],
                   alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
            status_row,
            warning_text,
        ], spacing=2)
        return ft.Container(
            padding=ft.Padding.only(left=16, right=16, top=10, bottom=10), border_radius=14, data="preserve_excel_fill",
            border=ft.Border.all(1, pal["border"]),
            content=ft.ResponsiveRow([
                ft.Container(col={"xs": 12, "md": 2}, content=info),
                ft.Container(col={"xs": 12, "md": 6}, content=scale),
                ft.Container(col={"xs": 12, "md": 4}, content=stats),
            ], spacing=16, run_spacing=8, vertical_alignment=ft.CrossAxisAlignment.START),
        )

    def limit_confirm(_e=None) -> None:
        changes = limit_changed()
        if not changes:
            return
        articles = [a for a in limit_articles() if a["offer_id"] in changes]
        lines: list[ft.Control] = []
        for article in articles:
            price = changes[article["offer_id"]]
            result = limit_price.evaluate(article, price)
            enter = [e["name"] for e in result if e["change"] == "enter"]
            leave = [e["name"] for e in result if e["change"] == "exit"]
            text = f"{article['offer_id']}: {limit_money(limit_current(article))} → {limit_money(price)}"
            if enter:
                text += " · войдёт: " + ", ".join(enter)
            if leave:
                text += " · выйдет: " + ", ".join(leave)
            if not limit_price.plan_changes([article], {article["offer_id"]: price}):
                text += " · вне акций — меняется цена товара"
            lines.append(ft.Text(text, size=13))
        dark = current_dark["value"]

        def run(_e=None) -> None:
            page.pop_dialog()
            if not job_lock.acquire(blocking=False):
                toast("Дождитесь завершения текущей операции.", True)
                return
            title = "Предельная цена"
            set_status(title, "Отправляю цены в Ozon и проверяю участие в акциях…", ACCENT, busy=True, cancellable=False)

            def worker() -> None:
                try:
                    result = limit_price.apply_prices(articles, changes)
                except Exception as exc:
                    result = {"ok": False, "done": 0, "message": f"Не удалось: {exc}", "failed": [], "errors": []}
                finally:
                    job_lock.release()
                message = str(result.get("message") or "")
                append_log(f"{title}: {message}")
                for item in result.get("items") or []:
                    append_log(f"  {item['offer_id']} · {item['action_name']}: {item['change']} по {item['price']} ₽")
                for item in result.get("failed") or []:
                    append_log(f"  ✖ {item['offer_id']} · {item['action_name']}: {item['reason']}")
                for error in result.get("errors") or []:
                    append_log(f"  ✖ {error}")
                for notice in result.get("notices") or []:
                    append_log(f"  ⚠ {notice['text']}")
                set_status(title, message, PRIMARY if result.get("ok") else DANGER, busy=False)
                ui_queue.put(("call", lambda r=result: limit_show_result(r)))
                ui_queue.put(("call", limit_load))
                if result.get("done"):
                    ui_queue.put(("run_job", "Обновление цен и акций", [str(PYTHON), str(SCRIPTS / "ui_actions.py"), "refresh-pricing"]))
                    if calendar_state["started"]:
                        ui_queue.put(("call", calendar_load))

            page.run_thread(worker)

        page.show_dialog(ft.AlertDialog(
            modal=True,
            title=ft.Text(f"Отправить цены в Ozon ({len(changes)})?", size=18, weight=ft.FontWeight.W_800, color="#F3F4F6" if dark else TEXT),
            bgcolor="#232A2F" if dark else "#FFFDF8",
            content=ft.Container(width=760, height=min(420, 60 + 26 * len(lines)), content=ft.Column([
                ft.Text("Цена товара одна на все акции: ниже — что изменится по каждому артикулу. После отправки участие "
                        "перечитывается из Ozon.", size=12, color="#B6C0BC" if dark else MUTED),
                *lines,
            ], spacing=6, scroll=ft.ScrollMode.AUTO)),
            actions=[ft.TextButton("Отмена", on_click=lambda _e: page.pop_dialog()),
                     ft.Button("Отправить в Ozon", bgcolor=PRIMARY, color="white", on_click=run)],
            actions_alignment=ft.MainAxisAlignment.END,
        ))

    def limit_show_result(result: dict) -> None:
        """Итог отправки: что применилось, что сообщил Ozon (сдвиг ограничения, выход из акций) и что не прошло."""
        dark = current_dark["value"]
        ink, muted = ("#F3F4F6", "#B6C0BC") if dark else (TEXT, MUTED)
        warn, bad, good = ("#E0A15F", "#FF8F82", "#5FD0B1") if dark else (ACCENT, DANGER, PRIMARY)
        notices = result.get("notices") or []
        lowered = [n for n in notices if n["kind"] == "limit"]
        rows: list[ft.Control] = [ft.Text(str(result.get("message") or ""), size=13, color=good if result.get("ok") else bad,
                                          weight=ft.FontWeight.W_600)]
        if lowered:
            rows.append(ft.Text("Ozon опустил «ограничение для акций»", size=14, weight=ft.FontWeight.W_700, color=warn))
            rows += [ft.Text("• " + n["text"], size=12, color=ink) for n in lowered]
            rows.append(ft.Text("Если цена ниже вашей минимальной, Ozon может сам добавить товар в акцию в минус. Вернуть ограничение: "
                                "«Цены и скидки» → «⋯» → «Обновить минимальные цены в Ozon».", size=12, color=muted))
        left = [n for n in notices if n["kind"] == "left"]
        if left:
            rows.append(ft.Text("Вышли из акций", size=14, weight=ft.FontWeight.W_700, color=ink))
            rows += [ft.Text("• " + n["text"], size=12, color=ink if n.get("expected") else warn) for n in left]
        other = [n for n in notices if n["kind"] == "other"]
        if other:
            rows.append(ft.Text("Ozon также сообщил", size=14, weight=ft.FontWeight.W_700, color=ink))
            rows += [ft.Text("• " + n["text"], size=12, color=muted) for n in other]
        problems = [f"{item['offer_id']} · {item['action_name']}: {item['reason']}" for item in result.get("failed") or []] + list(result.get("errors") or [])
        if problems:
            rows.append(ft.Text("Не применилось", size=14, weight=ft.FontWeight.W_700, color=bad))
            rows += [ft.Text("• " + text, size=12, color=ink) for text in problems]
        page.show_dialog(ft.AlertDialog(
            modal=True,
            title=ft.Text("Предельная цена: итог", size=18, weight=ft.FontWeight.W_800, color=ink),
            bgcolor="#232A2F" if dark else "#FFFDF8",
            content=ft.Container(width=700, content=ft.Column(rows, spacing=6, tight=True, scroll=ft.ScrollMode.AUTO)),
            actions=[ft.Button("Понятно", bgcolor=PRIMARY, color="white", on_click=lambda _e: page.pop_dialog())],
            actions_alignment=ft.MainAxisAlignment.END,
        ))

    def limit_set_filter(mode: str) -> None:
        limit_state["filter"] = mode
        limit_state["shown"] = LIMIT_PAGE
        render_limit_panel()

    def limit_filter_action(action_id) -> None:
        """Показать в панели только товары, которым доступна эта акция (None — все). Вызывается из календаря и легенды."""
        limit_state.update(action_filter=action_id, collapsed=False, shown=LIMIT_PAGE)
        render_limit_panel()
        render_actions_calendar()

    def limit_toggle_collapsed(_e=None) -> None:
        limit_state["collapsed"] = not limit_state["collapsed"]
        render_limit_panel()

    def limit_show_more(_e=None) -> None:
        limit_state["shown"] = int(limit_state["shown"]) + LIMIT_PAGE
        render_limit_panel()

    def render_limit_panel(push_update: bool = True) -> None:
        pal = dash_pal()
        data = limit_state.get("data") or {}
        articles = limit_articles()
        changed = limit_changed()
        busy = bool(limit_state["busy"])
        apply_button = ft.Button(f"Применить в Ozon ({len(changed)})", icon=ft.Icons.SEND, bgcolor=PRIMARY, color="white",
                                 disabled=not changed or busy, on_click=limit_confirm)
        reset_button = ft.TextButton("Сбросить", disabled=busy or not changed, on_click=lambda _e: limit_set_all("reset"))
        changed_text = dtext(f"Изменённые ({len(changed)})", 12, "#FFFFFF" if limit_state["filter"] == "changed" else pal["ink"], ft.FontWeight.W_600)
        live = {"apply": apply_button, "reset": reset_button, "changed": changed_text}
        loaded = str(data.get("loaded") or "")
        help_text = ("Двигайте ползунок — сразу видно цену, прибыль с продажи, общий бустинг и в каких акциях будет товар. "
                     "Зелёный отрезок — между минимальной и желательной ценой. Цветные засечки — пороги акций: пока цена левее "
                     "засечки, товар в этой акции. Цвет акции один во всём разделе.\n\n"
                     "Эффект акций = наибольший бустинг из акций (бустинги не складываются, как и в кабинете Ozon) + 30% за распродажу "
                     "(«Чёрная пятница», «11.11» и т. п.), пока она идёт. Надбавки Ozon за цель и индекс цен сюда не входят.")
        if data.get("economics_note"):
            help_text += "\n\n" + str(data["economics_note"]) + " Прибыль с продажи = цена × (1 − комиссия) − логистика − себестоимость; " \
                         "бустинг — наибольший из идущих акций, где товар будет при этой цене."
        collapsed = bool(limit_state["collapsed"])
        summary = f"{len(articles)} товаров" + (f" · изменено {len(changed)}" if changed else "") if articles else ""
        header = ft.Row([
            ft.IconButton(icon=ft.Icons.EXPAND_MORE if collapsed else ft.Icons.EXPAND_LESS, icon_size=20,
                          tooltip="Развернуть" if collapsed else "Свернуть", on_click=limit_toggle_collapsed),
            dtext("Предельная цена", 18, weight=ft.FontWeight.W_800),
            ft.Icon(ft.Icons.INFO_OUTLINE, size=16, color=pal["muted"], tooltip=help_text),
            dtext(summary, 12, pal["muted"]) if collapsed else ft.Container(),
        ], spacing=4, vertical_alignment=ft.CrossAxisAlignment.CENTER)
        toolbar = ft.Row([
            ft.Button("Подобрать баланс", icon=ft.Icons.AUTO_FIX_HIGH, bgcolor=ACCENT, color="white", disabled=busy or not articles,
                      tooltip="Каждому артикулу — цена между минимальной и желательной с наибольшей оценкой «прибыль с продажи × продажи»: "
                              "продажи растут от бустинга и от более низкой цены. В Ozon ничего не отправляется до «Применить».",
                      on_click=lambda _e: limit_set_all("auto")),
            ft.Dropdown(label="Реакция покупателей на цену", width=230, dense=True, value=str(limit_elasticity()) if limit_price else None,
                        options=[ft.dropdown.Option(str(value), f"{title} ({value:g})") for value, title in
                                 (limit_price.ELASTICITY_PRESETS.items() if limit_price else [])],
                        tooltip="Насколько продажи растут при снижении цены. Слабая (1): −10% цены ≈ +10% продаж; средняя (2) — +20%; "
                                "сильная (3) — +30%. Влияет только на «Подобрать баланс» и «Лучший баланс».",
                        on_select=limit_set_elasticity),
            reset_button,
            apply_button,
            ft.IconButton(icon=ft.Icons.REFRESH, tooltip="Обновить из Ozon", disabled=busy, on_click=limit_load),
        ], spacing=8, run_spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER, wrap=True)
        toolbar.alignment = ft.MainAxisAlignment.END
        if collapsed:
            limit_panel.content = ft.Row([header], vertical_alignment=ft.CrossAxisAlignment.CENTER)
            limit_panel.padding = ft.Padding.symmetric(horizontal=12, vertical=6)
            limit_panel.bgcolor = pal["surface"]
            limit_panel.border = ft.Border.all(1, pal["border"])
            if push_update and page.controls and host.controls and host.controls[0] is actions:
                page.update(limit_panel)
            return
        limit_panel.padding = 18
        controls: list[ft.Control] = [ft.ResponsiveRow([
            ft.Container(col={"xs": 12, "lg": 3}, content=header),
            ft.Container(col={"xs": 12, "lg": 9}, alignment=ft.Alignment(1, 0), content=toolbar),
        ], run_spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER)]
        legend: dict[str, dict] = {}
        for article in articles:
            for entry in article["actions"]:
                legend.setdefault(str(entry["action_id"]), entry)
        register_action_colors(legend)
        action_filter = limit_state.get("action_filter")
        if action_filter is not None and str(action_filter) not in legend:
            limit_state["action_filter"] = action_filter = None
        if legend:
            controls.append(ft.Row([
                ft.Container(
                    padding=ft.Padding.symmetric(horizontal=8, vertical=3), border_radius=999, data="preserve_excel_fill",
                    border=ft.Border.all(1.5, action_color(action_id)) if str(action_filter) == action_id else ft.Border.all(1, "#00000000"),
                    on_click=lambda _e, a=entry["action_id"]: limit_filter_action(None if limit_state.get("action_filter") == a else a),
                    content=ft.Row([ft.Container(width=10, height=10, border_radius=5, bgcolor=action_color(action_id), data="preserve_excel_fill"),
                        dtext(short_action_name(entry["name"]) + (f" · с {entry['start'][8:10]}.{entry['start'][5:7]}" if entry.get("start") else "")
                              + (" · распродажа" if entry.get("sale") else ""), 11, pal["ink"])],
                       spacing=5, tight=True),
                    tooltip=entry["name"] + "\nНажмите — показать только товары этой акции"
                            + (f"\nРаспродажа: пока акция идёт и товар в ней, к продажам +{limit_price.SALE_BONUS:g}% "
                               "сверх бустинга (подборки, баннеры, спрос дня распродажи)" if entry.get("sale") else ""))
                for action_id, entry in sorted(legend.items(), key=lambda item: action_color_slots.get(item[0], 0))
            ], spacing=6, run_spacing=6, wrap=True))
        if limit_price is None:
            controls.append(dtext(f"Модуль предельной цены не загрузился: {LIMIT_IMPORT_ERROR}", 13, pal["bad"]))
        if busy:
            controls.append(ft.Row([ft.ProgressRing(width=20, height=20, stroke_width=2.4),
                                    dtext("Загружаю акции, участников и кандидатов…", 13, pal["muted"])], spacing=10))
        if limit_state.get("error"):
            controls.append(dtext(f"Не удалось загрузить: {limit_state['error']}", 13, pal["bad"]))
        if articles:
            mode = limit_state["filter"]
            visible = [a for a in articles if mode == "all" or (mode == "changed" and a["offer_id"] in changed)
                       or (mode == "warn" and limit_warnings(a, limit_chosen(a) or 0))]
            if action_filter is not None:
                def in_action(article: dict) -> dict | None:
                    return next((e for e in article["actions"] if str(e["action_id"]) == str(action_filter)), None)
                # Сначала участники акции, затем те, кто может войти.
                visible = sorted((a for a in visible if in_action(a)), key=lambda a: not in_action(a)["participating"])

            def chip(label: str, key: str, text: ft.Control | None = None) -> ft.Control:
                active = mode == key
                return ft.Container(padding=ft.Padding.symmetric(horizontal=12, vertical=6), border_radius=999, data="preserve_excel_fill",
                                    bgcolor=pal["accent"] if active else None, border=None if active else ft.Border.all(1, pal["border"]),
                                    on_click=lambda _e, k=key: limit_set_filter(k),
                                    content=text or dtext(label, 12, "#FFFFFF" if active else pal["ink"], ft.FontWeight.W_600))

            filter_chip = []
            if action_filter is not None:
                entry = legend[str(action_filter)]
                members = sum(1 for a in visible if any(str(e["action_id"]) == str(action_filter) and e["participating"] for e in a["actions"]))
                filter_chip = [ft.Container(
                    padding=ft.Padding.symmetric(horizontal=10, vertical=5), border_radius=999, data="preserve_excel_fill",
                    bgcolor=action_color(action_filter), on_click=lambda _e: limit_filter_action(None), tooltip="Снять фильтр по акции",
                    content=dtext(f"{short_action_name(entry['name'])}: участвуют {members} из {len(visible)}  ✕", 12, "#FFFFFF", ft.FontWeight.W_700))]
            controls.append(ft.Row([*filter_chip, chip(f"Все ({len(articles)})", "all"), chip("", "changed", changed_text),
                                    chip("С предупреждениями", "warn"),
                                    dtext(f"обновлено {loaded[11:16]}" if loaded else "", 11, pal["muted"])],
                                   spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER))
            shown = int(limit_state["shown"])
            controls += [limit_row(article, live) for article in visible[:shown]]
            if len(visible) > shown:
                rest = len(visible) - shown
                controls.append(ft.Row([ft.TextButton(f"Показать ещё {min(rest, LIMIT_PAGE)} из {rest}", icon=ft.Icons.EXPAND_MORE,
                                                      on_click=limit_show_more)], alignment=ft.MainAxisAlignment.CENTER))
            if not visible:
                controls.append(dtext("Под этот фильтр артикулов нет.", 12, pal["muted"]))
        elif not busy and limit_state["started"] and not limit_state.get("error"):
            controls.append(dtext(str(data.get("message") or "Нет артикулов с доступными акциями."), 13, pal["muted"]))
        limit_panel.content = ft.Column(controls, spacing=10)
        limit_panel.bgcolor = pal["surface"]
        limit_panel.border = ft.Border.all(1, pal["border"])
        if push_update and page.controls and host.controls and host.controls[0] is actions:
            page.update(limit_panel)
    actions_calendar = ft.Container(padding=18, border_radius=18, data="preserve_excel_fill")
    auto_panel = ft.Container(padding=18, border_radius=18, data="preserve_excel_fill")
    actions = ft.Column(
        [
            ft.Text("Акции", size=26, weight=ft.FontWeight.W_800, color=TEXT),
            ft.Text("Предельная цена задаёт участие во всех акциях Ozon сразу: товар в акции, пока его цена не выше её порога.",
                    size=13, color=MUTED),
            actions_calendar,
            limit_panel,
            auto_panel,
        ],
        spacing=18,
        horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
    )

    # Палитра графиков (справочная палитра dataviz: первые три слота проверены попарно для обеих тем).
    supply_zone_palette = {False: ["#2a78d6", "#eb6834", "#1baf7a"], True: ["#3987e5", "#d95926", "#199e70"]}
    supply_season_color = {False: "#4a3aa7", True: "#9085e9"}
    supply_results = ft.Column(spacing=18, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)

    def latest_supply_summary() -> dict | None:
        folder = ROOT / "stocks reports"
        files = sorted(folder.glob("Поставка FBO *.json"), key=lambda path: path.stat().st_mtime) if folder.exists() else []
        if not files:
            return None
        try:
            return json.loads(files[-1].read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return None

    def supply_short_zone(name: str) -> str:
        return {"Москва, МО и Дальние регионы": "Москва", "Санкт-Петербург и СЗО": "СПб"}.get(name, name)

    def supply_date(value: str, with_year: bool = False) -> str:
        day = datetime.strptime(value, "%Y-%m-%d").date()
        return day.strftime("%d.%m.%Y" if with_year else "%d.%m")

    def render_supply_results(push_update: bool = True) -> None:
        dark = current_dark["value"]
        ink = "#F3F4F6" if dark else TEXT
        muted = "#B6C0BC" if dark else MUTED
        border = "#384046" if dark else "#E1E5EE"
        band = "#353F46" if dark else "#E6EBF3"
        summary = latest_supply_summary()

        def text(value: str, size: int = 12, color: str | None = None, weight=None, **kwargs) -> ft.Text:
            return ft.Text(value, size=size, color=color or ink, weight=weight, data="preserve_excel_text", **kwargs)

        def fill(color: str | None, **kwargs) -> ft.Container:
            return ft.Container(bgcolor=color, data="preserve_excel_fill", **kwargs)

        def panel(title: str, caption: str, content: ft.Control) -> ft.Control:
            return fill(
                semantic_surface("neutral", dark=dark),
                padding=18,
                border_radius=18,
                border=ft.Border.all(1, border),
                content=ft.Column(
                    [text(title, 15, weight=ft.FontWeight.W_700), text(caption, 11, muted), content],
                    spacing=10,
                ),
            )

        if not summary:
            supply_results.controls = [
                panel("Результатов пока нет", "Нажмите «Рассчитать поставку» — здесь появятся прогноз, сезонность и что везти.", ft.Container())
            ]
            if push_update and page.controls:
                page.update(supply_results)
            return

        zones = [zone["name"] for zone in summary["zones"]]
        zone_order = [*zones, summary["fbs_zone"]]
        palette = supply_zone_palette[dark]
        zone_color = {zone: palette[index % len(palette)] for index, zone in enumerate(zone_order)}
        settings = summary["settings"]
        forecast = summary["forecast"]
        window_start, window_end = summary["window"]
        last_day = (datetime.strptime(window_end, "%Y-%m-%d") - timedelta(days=1)).strftime("%Y-%m-%d")

        def swatch(color: str) -> ft.Control:
            return fill(color, width=10, height=10, border_radius=3)

        def legend(items: list[tuple[str, str]]) -> ft.Control:
            return ft.Row(
                [ft.Row([swatch(color), text(label, 11, muted)], spacing=6, tight=True) for label, color in items],
                spacing=16,
                wrap=True,
            )

        # 1. Ключевые цифры
        if forecast["method"] == "season":
            forecast_note = f"год назад {forecast['last_year_units']:.0f} шт. × рост {forecast['growth']:.2f}".replace(".", ",")
        else:
            forecast_note = "по средней за последние 30 дней"
        cost = summary.get("cost")  # нет в расчётах, сделанных до появления себестоимости
        fbs = summary.get("fbs")    # нет в расчётах, сделанных до появления FBS
        tile_col = 4 if fbs or not cost else 3
        tiles = [
            hero_metric("Прогноз магазина за цикл", f"{forecast['units']:.0f} шт.", forecast_note,
                        semantic_surface("neutral_soft", dark=dark), supply_season_color[dark], width=330, col=tile_col)
        ]
        for zone in summary["zones"]:
            safety = max(0, zone["ship"] - max(0.0, zone["demand"] - zone["on_arrival"]))
            zone_cost = f" · {_format_currency(zone['cost'])}" if "cost" in zone else ""
            tiles.append(hero_metric(
                f"К поставке · {supply_short_zone(zone['name'])}", f"{zone['ship']} шт.",
                f"спрос {zone['demand']:.0f} + страховка {safety:.0f} · {zone['articles']} арт.{zone_cost}",
                semantic_surface("neutral_soft", dark=dark), zone_color[zone["name"]], width=330, col=tile_col,
            ))
        if cost:
            missing = cost["missing_articles"]
            cost_note = (
                f"без себестоимости в costs.xlsx: {', '.join(missing[:5])}{'…' if len(missing) > 5 else ''} — не учтены"
                if missing else "количество × себестоимость из costs.xlsx"
            )
            tiles.append(hero_metric(
                "Себестоимость поставки", _format_currency(cost["total"]), cost_note,
                semantic_surface("neutral_soft", dark=dark), ACCENT, width=330, col=tile_col,
            ))
        if fbs:
            fbs_color = zone_color[summary["fbs_zone"]]
            fbs_until = supply_date(last_day)
            tiles.append(hero_metric(
                "Нужно на своём складе (FBS)", f"{fbs['need']} шт.",
                f"до {fbs_until} · спрос ~{fbs['demand']:.0f} шт. + страховка · есть {fbs['stock']} шт.",
                semantic_surface("neutral_soft", dark=dark), fbs_color, width=330, col=tile_col,
            ))
            missing_cost = fbs.get("missing_cost_articles") or []
            shortage_note = (
                f"на {_format_currency(fbs['shortage_cost'])} по себестоимости" + (" (не у всех есть себестоимость)" if missing_cost else "")
                if fbs["shortage"] else "остатка хватит на эти дни"
            )
            tiles.append(hero_metric(
                "Не хватает на FBS", f"{fbs['shortage']} шт.", shortage_note,
                semantic_surface("danger" if fbs["shortage"] else "positive", dark=dark), DANGER if fbs["shortage"] else PRIMARY,
                width=330, col=tile_col,
            ))
        header = text(
            f"Расчёт от {supply_date(summary['created'], True)} · товар придёт ~{supply_date(summary['arrival'])} · "
            f"продаётся {supply_date(window_start)}–{supply_date(last_day)} · срок {settings['lead_days']} дн., "
            f"цикл {settings['cover_days']} дн., уровень сервиса {settings['service_level']:.0%}",
            12, muted,
        )

        # 2. Уровень сервиса
        level_chips = []
        for option in summary["sensitivity"]:
            total = sum(option["by_zone"].values())
            selected = abs(option["level"] - float(settings["service_level"])) < 1e-6
            level_chips.append(fill(
                semantic_surface("positive" if selected else "neutral_soft", dark=dark),
                padding=ft.Padding.symmetric(horizontal=16, vertical=10),
                border_radius=12,
                border=ft.Border.all(2 if selected else 1, PRIMARY if selected else border),
                tooltip=" · ".join(f"{supply_short_zone(zone)}: {qty} шт." for zone, qty in option["by_zone"].items()),
                content=ft.Column(
                    [text(f"{option['level']:.0%}", 15, weight=ft.FontWeight.W_800), text(f"{total} шт.", 12, muted)],
                    spacing=2, tight=True, horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                ),
            ))
        sensitivity_panel = panel(
            "Сколько везти при другом уровне сервиса",
            "Уровень сервиса — вероятность не уйти в ноль на FBO за цикл. Если FBO опустеет, заказ уйдёт FBS, только медленнее. "
            "Чтобы получить список по артикулам, пересчитайте с нужным уровнем.",
            ft.Row(level_chips, spacing=10, wrap=True),
        )

        # 3. Сезонность
        season = summary["season"]
        while season and season[0]["units"] == 0:
            season = season[1:]
        base_from, base_to = (datetime.strptime(value, "%Y-%m-%d") for value in summary["base_window"])
        base_months = set()
        cursor = base_from.replace(day=1)
        while cursor < base_to:
            base_months.add(cursor.strftime("%Y-%m"))
            cursor = (cursor.replace(day=28) + timedelta(days=4)).replace(day=1)
        current_month = summary["created"][:7]
        by_month = {row["month"]: row["units"] for row in season}
        peak = max((row["units"] for row in season), default=0) or 1
        columns = []
        for row in season:
            month = row["month"]
            year, mon = month.split("-")
            last_year = by_month.get(f"{int(year) - 1}-{mon}")
            partial = month == current_month
            tooltip = f"{mon}.{year}: {row['units']} шт."
            if last_year is not None:
                tooltip += f" · год назад {last_year} шт."
            if partial:
                tooltip += " · месяц ещё идёт"
            labeled = month in base_months or row["units"] == peak
            columns.append(fill(
                band if month in base_months else None,
                expand=1,
                border_radius=8,
                padding=ft.Padding.only(left=2, right=2, top=4, bottom=2),
                tooltip=tooltip,
                content=ft.Column(
                    [
                        text(str(row["units"]) if labeled else "", 10, muted),
                        fill(
                            supply_season_color[dark],
                            height=max(2, row["units"] / peak * 150) if row["units"] else 0,
                            border_radius=ft.BorderRadius.only(top_left=4, top_right=4),
                            opacity=0.45 if partial else 1,
                        ),
                        text(f"{mon}.{year[2:]}", 9, muted),
                    ],
                    spacing=4,
                    alignment=ft.MainAxisAlignment.END,
                    horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
                ),
            ))
        season_panel = panel(
            "Продажи по месяцам",
            f"Подсвечены месяцы год назад: продажи за {supply_date(summary['base_window'][0])}–"
            f"{(base_to - timedelta(days=1)).strftime('%d.%m.%Y')} × рост дают прогноз на цикл. Наведите на столбец, чтобы сравнить с прошлым годом.",
            ft.Container(height=200, content=ft.Row(columns, spacing=2, vertical_alignment=ft.CrossAxisAlignment.END)),
        )

        # 4. Спрос по кластерам
        clusters = summary["clusters"]
        shown = [row for row in clusters if row["share"] >= 1.0]
        hidden = [row for row in clusters if row["share"] < 1.0]
        if hidden:
            shown.append({"cluster": f"Остальные ({len(hidden)})", "zone": None, "share": round(sum(r["share"] for r in hidden), 1)})
        top_share = max((row["share"] for row in shown), default=1) or 1
        cluster_rows = []
        for row in shown:
            color = zone_color.get(row["zone"], "#8E8E93")
            width = int(row["share"] / top_share * 1000)
            cluster_rows.append(ft.Row(
                [
                    ft.Container(width=190, content=text(row["cluster"], 11, ink, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)),
                    ft.Row(
                        [fill(color, expand=max(1, width), height=12, border_radius=ft.BorderRadius.only(top_right=4, bottom_right=4)),
                         ft.Container(expand=max(1, 1000 - width))],
                        expand=True, spacing=0,
                    ),
                    ft.Container(width=52, content=text(f"{row['share']:g}%".replace(".", ","), 11, muted, text_align=ft.TextAlign.RIGHT)),
                ],
                spacing=10,
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
            ))
        zone_totals = {zone: sum(row["share"] for row in clusters if row["zone"] == zone) for zone in zone_order}
        cluster_panel = panel(
            "Откуда покупатели — и какой склад их обслуживает",
            "Доля спроса по кластеру покупателя за год. Цвет — склад, за которым закреплён кластер: склад FBO продаёт не только своему региону.",
            ft.Column(
                [legend([(f"{supply_short_zone(zone)} — {zone_totals[zone]:.0f}%", zone_color[zone]) for zone in zone_order]), *cluster_rows],
                spacing=6,
            ),
        )

        # 5. Что везти
        supply_rows = summary["supply"]
        biggest = max((row["total"] for row in supply_rows), default=1) or 1
        header_row = ft.Row(
            [
                ft.Container(width=70, content=text("Артикул", 11, muted, weight=ft.FontWeight.W_700)),
                *[ft.Container(width=64, content=text(supply_short_zone(zone), 11, muted, weight=ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT)) for zone in zones],
                ft.Container(width=56, content=text("Итого", 11, muted, weight=ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT)),
                *([ft.Container(width=110, content=text("Себестоимость", 11, muted, weight=ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT))] if cost else []),
                ft.Container(width=86, content=text("Прогноз/день", 11, muted, weight=ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT)),
                ft.Container(expand=True),
            ],
            spacing=12,
        )
        table_rows: list[ft.Control] = [header_row]
        for row in supply_rows:
            segments = []
            for zone in zones:
                qty = row["by_zone"].get(zone, 0)
                if qty:
                    segments.append(fill(zone_color[zone], expand=qty, height=12, border_radius=4,
                                         tooltip=f"{supply_short_zone(zone)}: {qty} шт."))
            segments.append(ft.Container(expand=max(1, biggest - row["total"])) if row["total"] < biggest else ft.Container(width=0))
            table_rows.append(ft.Row(
                [
                    ft.Container(width=70, content=text(row["article"], 12, weight=ft.FontWeight.W_700)),
                    *[ft.Container(width=64, content=text(str(row["by_zone"].get(zone, 0) or "—"), 12, text_align=ft.TextAlign.RIGHT)) for zone in zones],
                    ft.Container(width=56, content=text(str(row["total"]), 12, weight=ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT)),
                    *([ft.Container(
                        width=110,
                        tooltip=f"{row['total']} шт. × {_format_currency(row['unit_cost'])}" if row.get("unit_cost") else "Нет себестоимости в costs.xlsx",
                        content=text(_format_currency(row.get("cost")), 12, text_align=ft.TextAlign.RIGHT),
                    )] if cost else []),
                    ft.Container(width=86, content=text(f"{row['per_day']:.2f}".replace(".", ","), 12, muted, text_align=ft.TextAlign.RIGHT)),
                    ft.Row(segments, expand=True, spacing=2),
                ],
                spacing=12,
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
            ))
        total_by_zone = {zone: sum(row["by_zone"].get(zone, 0) for row in supply_rows) for zone in zones}
        supply_panel = panel(
            f"Что везти — {len(supply_rows)} артикулов, {sum(total_by_zone.values())} шт."
            + (f" на {_format_currency(cost['total'])} по себестоимости" if cost else ""),
            "По убыванию количества. Полный расчёт по каждой строке — в Excel на листе «Расчёт».",
            ft.Column(
                [legend([(f"{supply_short_zone(zone)} — {total_by_zone[zone]} шт.", zone_color[zone]) for zone in zones]), *table_rows],
                spacing=8,
            ),
        )

        # 6. Свой склад (FBS) на те же дни
        fbs_panel = None
        if fbs:
            bad = "#FF9C8F" if dark else DANGER
            fbs_rows_data = fbs["rows"]
            fbs_peak = max((max(row["need"], row["stock"] or 0) for row in fbs_rows_data), default=1) or 1

            def fbs_cell(value: str, width: int, color: str | None = None, weight=None, tooltip: str | None = None) -> ft.Control:
                return ft.Container(width=width, tooltip=tooltip, content=text(value, 12, color, weight, text_align=ft.TextAlign.RIGHT))

            fbs_header = ft.Row(
                [
                    ft.Container(width=70, content=text("Артикул", 11, muted, weight=ft.FontWeight.W_700)),
                    *[ft.Container(width=w, content=text(label, 11, muted, weight=ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT))
                      for label, w in (("Регионы FBS", 90), ("С FBO", 64), ("Спрос", 64), ("Нужно", 60), ("Есть", 56), ("Не хватает", 84))],
                    *([ft.Container(width=110, content=text("Себестоимость", 11, muted, weight=ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT))] if cost else []),
                    ft.Container(expand=True),
                ],
                spacing=12,
            )
            fbs_table: list[ft.Control] = [fbs_header]
            for row in fbs_rows_data:
                stock_now = row["stock"] or 0
                covered = min(stock_now, row["need"])
                segments = []
                if covered:
                    segments.append(fill(fbs_color, expand=covered, height=12, border_radius=4, tooltip=f"Есть: {stock_now} шт."))
                if row["shortage"]:
                    segments.append(fill(bad, expand=row["shortage"], height=12, border_radius=4, tooltip=f"Не хватает: {row['shortage']} шт."))
                extra = max(0, stock_now - row["need"])
                if extra:
                    segments.append(fill(band, expand=extra, height=12, border_radius=4, tooltip=f"Сверх нужного: {extra} шт."))
                used = covered + row["shortage"] + extra
                segments.append(ft.Container(expand=max(1, fbs_peak - used)) if used < fbs_peak else ft.Container(width=0))
                own = row["own_lead"] + row["own_window"]
                fbs_table.append(ft.Row(
                    [
                        ft.Container(width=70, content=text(row["article"], 12, weight=ft.FontWeight.W_700)),
                        fbs_cell(f"{own:.1f}".replace(".", ","), 90, muted,
                                 tooltip=f"До прихода поставки {row['own_lead']:.1f} шт. + в окне {row['own_window']:.1f} шт.".replace(".", ",")),
                        fbs_cell(f"{row['spill']:.1f}".replace(".", ",") if row["spill"] else "—", 64, muted,
                                 tooltip="Что FBO не покроет до прихода поставки + мелкий спрос зон FBO, куда не везём"),
                        fbs_cell(f"{row['demand']:.1f}".replace(".", ","), 64),
                        fbs_cell(str(row["need"]), 60, weight=ft.FontWeight.W_700),
                        fbs_cell("—" if row["stock"] is None else str(row["stock"]), 56),
                        fbs_cell(str(row["shortage"]) if row["shortage"] else "—", 84, bad if row["shortage"] else muted, ft.FontWeight.W_700),
                        *([fbs_cell(_format_currency(row["cost"]) if row.get("cost") else "—", 110)] if cost else []),
                        ft.Row(segments, expand=True, spacing=2),
                    ],
                    spacing=12,
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                ))
            fbs_panel = panel(
                f"Свой склад (FBS) на те же дни — нужно {fbs['need']} шт., не хватает {fbs['shortage']} шт.",
                f"С {supply_date(fbs['horizon'][0])} по {supply_date(last_day)}: заказы регионов без склада FBO, то, что FBO не покроет до прихода поставки, "
                f"и мелкий спрос зон FBO, куда не везём. «Нужно» — запас, чтобы с вероятностью {settings['service_level']:.0%} не уйти в ноль.",
                ft.Column(
                    [legend([("есть на складе", fbs_color), ("не хватает", bad), ("сверх нужного", band)]), *fbs_table],
                    spacing=8,
                ),
            )

        supply_results.controls = [
            header,
            ft.ResponsiveRow(tiles, run_spacing=14, spacing=14),
            sensitivity_panel,
            supply_panel,
            *([fbs_panel] if fbs_panel else []),
            ft.ResponsiveRow(
                [ft.Container(col={"md": 12, "xl": 6}, content=season_panel), ft.Container(col={"md": 12, "xl": 6}, content=cluster_panel)],
                run_spacing=14, spacing=14,
            ),
        ]
        if push_update and page.controls:
            page.update(supply_results)

    # ---------------------------------------------------------------- Реклама: оплата за клик
    ads_defaults = ads_load_settings(ROOT) if ads_load_settings else {"profit_share": 0.4, "lookback_days": 21}
    ads_share = ft.TextField(label="Доля прибыли на рекламу, %", value=f"{float(ads_defaults['profit_share']) * 100:g}", width=230,
                             tooltip="Сколько прибыли с заказа можно отдать за клики, которые к нему привели")
    ads_days = ft.TextField(label="Период анализа, дней", value=str(int(ads_defaults["lookback_days"])), width=190)
    ads_results = ft.Column(spacing=18, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)
    ads_apply_button = ft.Button("Применить отмеченные", bgcolor=PRIMARY, color="white")
    ads_state: dict[str, object] = {"plan": None, "result": None, "busy": False}
    ads_action_labels = {"bid_down": "Снизить ставку", "bid_up": "Повысить ставку", "remove": "Убрать из рекламы", "keep": "Без изменений", "info": "Только информация"}

    def ads_selected_count() -> int:
        plan = ads_state.get("plan")
        if not isinstance(plan, dict):
            return 0
        return sum(1 for item in plan["items"] if item.get("selected") and item["action"] in ("bid_down", "bid_up", "remove"))

    def ads_rub(value, digits: int = 0) -> str:
        if value is None:
            return "—"
        return f"{value:,.{digits}f} ₽".replace(",", " ").replace(".", ",")

    def render_ads_results(push_update: bool = True) -> None:
        dark = current_dark["value"]
        ink = "#F3F4F6" if dark else TEXT
        muted = "#B6C0BC" if dark else MUTED
        border = "#384046" if dark else "#E1E5EE"
        good = "#70D4BC" if dark else PRIMARY
        bad = "#FF9C8F" if dark else DANGER
        warn = "#E0A15F" if dark else ACCENT

        def text(value: str, size: int = 12, color: str | None = None, weight=None, **kwargs) -> ft.Text:
            return ft.Text(value, size=size, color=color or ink, weight=weight, data="preserve_excel_text", **kwargs)

        def box(content: ft.Control, **kwargs) -> ft.Container:
            return ft.Container(
                content=content, padding=18, border_radius=18, data="preserve_excel_fill",
                bgcolor=semantic_surface("neutral", dark=dark), border=ft.Border.all(1, border), **kwargs,
            )

        status_icons = {
            "ok": (ft.Icons.CHECK_CIRCLE_OUTLINE, good),
            "warn": (ft.Icons.WARNING_AMBER_ROUNDED, warn),
            "bad": (ft.Icons.ERROR_OUTLINE, bad),
            "info": (ft.Icons.INFO_OUTLINE, muted),
        }

        def status_line(step: dict) -> ft.Control:
            icon, color = status_icons.get(step["status"], status_icons["info"])
            return ft.Row(
                [
                    ft.Icon(icon, size=16, color=color),
                    ft.Column([text(step["name"], 12, weight=ft.FontWeight.W_700), text(step["text"], 12)], spacing=1, expand=True),
                ],
                spacing=8,
                vertical_alignment=ft.CrossAxisAlignment.START,
            )

        def advice_block(lines: list) -> ft.Control:
            return ft.Column(
                [text("Рекомендации", 12, warn, ft.FontWeight.W_700), *[text(f"→ {line}", 12) for line in lines]],
                spacing=3,
            )

        plan, result = ads_state.get("plan"), ads_state.get("result")
        controls: list[ft.Control] = []
        if ads_state.get("loading"):
            controls.append(box(ft.Row([ft.ProgressRing(width=22, height=22, stroke_width=2.4),
                                        text("Подбираю ставки: статистика товаров, конкурентные ставки, остатки — около минуты…", 13, muted)], spacing=12)))
        if isinstance(result, dict):
            failures = result.get("failed") or []
            lines = [text(result["message"], 13, good if result.get("ok") else bad, ft.FontWeight.W_700)]
            lines += [text(f"✖ {f.get('offer_id') or f['sku']} — {f['campaign_title']}: {f['error']}", 12, bad) for f in failures]
            controls.append(box(ft.Column(lines, spacing=6)))
        if ads_state.get("loading"):
            pass
        elif not isinstance(plan, dict):
            controls.append(box(text("Нажмите «Подобрать ставки» — программа посчитает, сколько выгодно платить за клик по каждому товару. "
                                     "В Ozon ничего не отправляется, пока вы не нажмёте «Применить».", 12, muted)))
        elif not plan.get("campaigns"):
            controls.append(box(text(plan.get("message") or "Нет работающих кампаний.", 13, muted)))
        else:
            period = plan["period"]
            controls.append(text(
                f"Расчёт от {plan['created'][:16].replace('T', ' ')} · статистика {supply_date(period[0])}–{supply_date(period[1])} · "
                f"на рекламу до {float(plan['settings']['profit_share']):.0%} прибыли с заказа",
                12, muted,
            ))
            for campaign in plan["campaigns"]:
                items = [i for i in plan["items"] if i["campaign_id"] == campaign["id"]]
                after = campaign["profit_after_ads"]
                share = campaign["profit_share_spent"]
                budget = f"{ads_rub(campaign['weekly_budget'])} в неделю" if campaign["weekly_budget"] else (
                    f"{ads_rub(campaign['daily_budget'])} в день" if campaign["daily_budget"] else "без лимита")
                tiles = [
                    hero_metric("Расход на рекламу", ads_rub(campaign["spend"]), f"{campaign['clicks']:.0f} кликов",
                                semantic_surface("neutral_soft", dark=dark), ACCENT, width=260, col=3),
                    hero_metric("Заказов с рекламы", f"{campaign['orders']:.0f}",
                                f"конверсия {campaign['conversion']:.1%}".replace(".", ",") if campaign.get("conversion") else "",
                                semantic_surface("neutral_soft", dark=dark), PRIMARY, width=260, col=3),
                    hero_metric("Прибыль с них до рекламы", ads_rub(campaign["profit_before_ads"]), "по прибыли с продажи из costs.xlsx",
                                semantic_surface("neutral_soft", dark=dark), PRIMARY, width=260, col=3),
                    hero_metric("Осталось после рекламы", ads_rub(after),
                                f"реклама съела {share:.0%} прибыли" if share is not None else "",
                                semantic_surface("danger" if after < 0 else "positive", dark=dark), DANGER if after < 0 else PRIMARY, width=260, col=3),
                ]
                header = ft.Row(
                    [
                        ft.Container(width=40),
                        ft.Container(width=210, content=text("Товар", 11, muted, ft.FontWeight.W_700)),
                        *[ft.Container(width=w, content=text(t, 11, muted, ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT))
                          for t, w in (("Клики", 56), ("Заказы", 60), ("Расход", 80), ("₽ на заказ", 84), ("Безубыт.", 74), ("Ставка", 120), ("Ozon советует", 100))],
                        ft.Container(expand=True, content=text("Решение", 11, muted, ft.FontWeight.W_700)),
                    ],
                    spacing=10,
                )
                rows: list[ft.Control] = [header]
                for item in items:
                    actionable = item["action"] in ("bid_down", "bid_up", "remove")

                    def on_toggle(e: ft.ControlEvent, target=item) -> None:
                        target["selected"] = bool(e.control.value)
                        refresh_ads_apply()

                    cpo = item["spend"] / item["orders_total"] if item["orders_total"] else None
                    if item["action"] in ("bid_down", "bid_up"):
                        bid_text = f"{item['current_bid']:.2f} → {item['recommended_bid']:.2f}".replace(".", ",")
                    else:
                        bid_text = f"{item['current_bid']:.2f}".replace(".", ",")
                    action_color = {"bid_down": warn, "bid_up": good, "remove": bad}.get(item["action"], muted)
                    orders_text = f"{item['orders_total']:.0f}" + (f" ({item['model_orders']:.0f} вар.)" if item.get("model_orders") else "")
                    details = ft.Container(
                        visible=bool(item.get("_open")),
                        padding=ft.Padding.only(left=50, top=6, bottom=4, right=8),
                        content=ft.Column(
                            [*[status_line(step) for step in item.get("steps") or []],
                             *([advice_block(item["advice"])] if item.get("advice") else [])],
                            spacing=8,
                        ),
                    )
                    details_button = ft.TextButton("Скрыть шаги" if item.get("_open") else "Подробнее: 6 шагов")

                    def toggle_details(_e, target=item, panel=details, button=details_button) -> None:
                        target["_open"] = not target.get("_open")
                        panel.visible = target["_open"]
                        button.content = "Скрыть шаги" if target["_open"] else "Подробнее: 6 шагов"
                        page.update(panel, button)

                    details_button.on_click = toggle_details
                    stage_label = {"test": "Тест · ", "mature": ""}.get(item.get("stage") or "", "")
                    rows.append(ft.Container(
                        padding=ft.Padding.symmetric(vertical=6),
                        border=ft.Border.only(top=ft.BorderSide(1, border)),
                        data="preserve_excel_fill",
                        content=ft.Column(
                            [
                                ft.Row(
                                    [
                                        ft.Container(width=40, content=ft.Checkbox(value=bool(item.get("selected")), on_change=on_toggle) if actionable else None),
                                        ft.Container(width=210, content=ft.Column([
                                            text(str(item.get("offer_id") or item["sku"]), 12, weight=ft.FontWeight.W_700),
                                            text(item.get("title") or "", 10, muted, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS),
                                        ], spacing=0, tight=True)),
                                        *[ft.Container(width=w, content=text(v, 12, text_align=ft.TextAlign.RIGHT)) for v, w in (
                                            (f"{item['clicks']:.0f}", 56),
                                            (orders_text, 60),
                                            (ads_rub(item["spend"]), 80),
                                            (ads_rub(cpo), 84),
                                            (f"{item['breakeven_bid']:.2f}".replace(".", ",") if item.get("breakeven_bid") is not None else "—", 74),
                                            (bid_text, 120),
                                            (f"{item['competitive_bid']:.2f}".replace(".", ",") if item.get("competitive_bid") else "—", 100),
                                        )],
                                        ft.Container(expand=True, content=text(ads_action_labels[item["action"]], 12, action_color, ft.FontWeight.W_700)),
                                    ],
                                    spacing=10,
                                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                                ),
                                ft.Row(
                                    [
                                        ft.Container(padding=ft.Padding.only(left=50), expand=True, content=text(stage_label + item["reason"], 11, muted)),
                                        details_button,
                                    ],
                                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                                ),
                                details,
                            ],
                            spacing=2,
                        ),
                    ))
                controls.append(box(ft.Column(
                    [
                        text(campaign["title"], 16, weight=ft.FontWeight.W_800),
                        text(f"Стратегия «{campaign['strategy']}» · бюджет {budget}", 12, muted),
                        ft.ResponsiveRow(tiles, run_spacing=12, spacing=12),
                        *([ft.Container(
                            padding=12,
                            border_radius=12,
                            data="preserve_excel_fill",
                            bgcolor=semantic_surface("neutral_soft", dark=dark),
                            content=ft.Column([text("Что видно по кампании", 13, weight=ft.FontWeight.W_800),
                                               *[status_line(note) for note in campaign["advice"]]], spacing=8),
                        )] if campaign.get("advice") else []),
                        *rows,
                    ],
                    spacing=10,
                )))
            controls.append(text(
                "«Безубыт.» — ставка, при которой реклама съедает всю прибыль с заказа. «Ozon советует» — конкурентная ставка Ozon. "
                "Заказы учитывают покупки соседних вариантов товара после клика (вар.).", 11, muted,
            ))
        ads_results.controls = controls
        # Плитки сводки (hero_metric) окрашиваются темой; наши тексты и заливки помечены и не меняются.
        apply_control_theme(ads_results, dark)
        refresh_ads_apply(push_update=False)
        if push_update and page.controls:
            try:
                page.update(ads_results, ads_apply_button)
            except Exception:
                pass  # вкладку уже закрыли — покажем при следующем открытии

    def refresh_ads_apply(push_update: bool = True) -> None:
        count = ads_selected_count()
        ads_apply_button.content = f"Применить отмеченные ({count})"
        ads_apply_button.disabled = count == 0 or bool(ads_state["busy"])
        ads_apply_button.opacity = 0.45 if ads_apply_button.disabled else 1
        if push_update and page.controls:
            try:
                page.update(ads_apply_button)
            except Exception:
                pass

    def ads_read_settings() -> dict | None:
        try:
            share = float((ads_share.value or "").replace(",", ".")) / 100
            days = int(ads_days.value or "")
        except ValueError:
            toast("Доля прибыли — число, период — целое число дней.", True)
            return None
        if not 0.05 <= share <= 0.9 or not 7 <= days <= 60:
            toast("Доля прибыли — от 5 до 90%, период — от 7 до 60 дней.", True)
            return None
        settings = {**ads_load_settings(ROOT), "profit_share": share, "lookback_days": days}
        settings.pop("min_profit_to_cost", None)  # условие убрано, старое значение из файла не нужно
        ads_save_settings(ROOT, settings)
        return settings

    def ads_build(_e=None, keep_result: bool = False) -> None:
        if build_ads_plan is None:
            toast(f"Модуль рекламы недоступен: {ADS_IMPORT_ERROR or 'ошибка импорта'}", True)
            return
        settings = ads_read_settings()
        if settings is None:
            return
        if not job_lock.acquire(blocking=False):
            toast("Дождитесь завершения текущей операции.", True)
            return
        campaign_id = ads_nav["campaign_id"]
        ads_state["busy"] = True
        if not keep_result:
            ads_state["result"] = None
        set_status("Реклама", "Собираю статистику кампании из Ozon...", ACCENT, busy=True, cancellable=False)
        ads_state["plan"] = None
        ads_state["loading"] = True
        render_ads_results()

        def worker() -> None:
            try:
                ads_state["plan"] = build_ads_plan(ROOT, settings, campaign_ids=[campaign_id] if campaign_id else None)
                ads_state["plan_campaign"] = campaign_id
                set_status("Реклама", str(ads_state["plan"]["message"]), PRIMARY, busy=False)
            except Exception as exc:
                ads_state["plan"] = None
                set_status("Реклама", f"Не удалось собрать данные: {exc}", DANGER, busy=False)
                ads_state["result"] = {"ok": False, "message": f"Не удалось собрать данные: {exc}", "failed": []}
            finally:
                ads_state["busy"] = False
                ads_state["loading"] = False
                job_lock.release()
            render_ads_results()

        page.run_thread(worker)

    def ads_apply(_e=None) -> None:
        plan = ads_state.get("plan")
        if not isinstance(plan, dict) or ads_selected_count() == 0 or execute_ads_plan is None:
            return
        if not job_lock.acquire(blocking=False):
            toast("Дождитесь завершения текущей операции.", True)
            return
        ads_state["busy"] = True
        set_status("Реклама", "Отправляю ставки в Ozon и проверяю применение...", ACCENT, busy=True, cancellable=False)
        refresh_ads_apply()

        def worker() -> None:
            try:
                ads_state["result"] = execute_ads_plan(ROOT, plan["items"])
            except Exception as exc:
                ads_state["result"] = {"ok": False, "message": f"Ошибка отправки в Ozon: {exc}", "failed": []}
            finally:
                ads_state["busy"] = False
                job_lock.release()
            result = ads_state["result"]
            append_log(f"Реклама: {result['message']}")
            for failure in result.get("failed") or []:
                append_log(f"  ✖ {failure.get('offer_id') or failure['sku']}: {failure['error']}")
            set_status("Реклама", result["message"], PRIMARY if result.get("ok") else DANGER, busy=False)
            # Пересобираем рекомендации, чтобы показать фактические ставки после изменений.
            ui_queue.put(("ads_rebuild",))

        page.run_thread(worker)

    ads_apply_button.on_click = ads_apply

    def open_ads_reports() -> None:
        folder = ROOT / "ads reports"
        folder.mkdir(parents=True, exist_ok=True)
        reveal_path(folder, "ads reports")

    # ---------------------------------------------------------------- Реклама: навигация (тип → кампании → кампания)
    ADS_TYPES = [
        ("cpc", "Оплата за клик", "Поднимайте товары в поиске и рекомендациях Ozon и платите за клики. С помощью инструмента можно привлечь потенциальных покупателей и увеличить продажи"),
        ("cpo", "Оплата за заказ", "Поднимайте товары выше в поиске Ozon и платите только за заказы. С помощью инструмента можно усилить продвижение товаров в «Оплате за клик»"),
        ("media", "Медийная реклама", "Повышайте узнаваемость бренда, рассказывайте о новинках и акциях с помощью баннеров на главной странице, в каталоге и поиске Ozon"),
        ("collections", "Подборки товаров", "Создавайте страницу с вашими товарами и делитесь ей с покупателями в рассылках, соцсетях или за пределами Ozon"),
        ("external", "Внешнее продвижение", "Продвигайте товары в соцсетях с помощью блогеров, чтобы повысить охваты и привлечь внимание покупателей"),
        ("promo", "Акции", "Участвуйте в акциях от Ozon, чтобы повысить продажи. Акционным бустингом можно управлять: чем выше скидка, тем больше бустинга"),
        ("b2b", "Продажи юридическим лицам", "Юрлица чаще заказывают оптом и повышают средний чек"),
    ]
    ADS_PERIOD_PRESETS = (7, 14, 31, 90)
    ads_nav: dict[str, object] = {
        "screen": "types", "campaign_id": None, "preset": 31,
        "from": date.today() - timedelta(days=30), "to": date.today(),
        "list": None, "list_busy": False, "list_error": None, "opened": set(),
    }
    ads_view = ft.Column(spacing=18, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)
    ads_custom_from = ft.TextField(label="С", width=150, dense=True)
    ads_custom_to = ft.TextField(label="По", width=150, dense=True)

    def ads_palette() -> dict:
        dark = current_dark["value"]
        return {
            "dark": dark,
            "ink": "#F3F4F6" if dark else TEXT,
            "muted": "#B6C0BC" if dark else MUTED,
            "border": "#384046" if dark else "#E1E5EE",
            "surface": semantic_surface("neutral", dark=dark),
            "soft": semantic_surface("neutral_soft", dark=dark),
            "good": "#70D4BC" if dark else PRIMARY,
            "bad": "#FF9C8F" if dark else DANGER,
        }

    def ads_text(value: str, size: int = 12, color: str | None = None, weight=None, **kwargs) -> ft.Text:
        return ft.Text(value, size=size, color=color or ads_palette()["ink"], weight=weight, data="preserve_excel_text", **kwargs)

    def ads_box(content: ft.Control, **kwargs) -> ft.Container:
        pal = ads_palette()
        options = {"padding": 18, "border_radius": 18, "bgcolor": pal["surface"], "border": ft.Border.all(1, pal["border"])}
        options.update(kwargs)
        return ft.Container(content=content, data="preserve_excel_fill", **options)

    def ads_header(title: str, subtitle: str, back: tuple[str, object] | None = None) -> list[ft.Control]:
        controls: list[ft.Control] = []
        if back:
            controls.append(ft.Row([ft.TextButton(back[0], icon=ft.Icons.ARROW_BACK, on_click=lambda _e: back[1]())]))
        controls += [ft.Text(title, size=26, weight=ft.FontWeight.W_800, color=TEXT), ft.Text(subtitle, size=13, color=MUTED)]
        return controls

    def ads_go(screen: str, campaign_id: str | None = None) -> None:
        ads_nav["screen"] = screen
        if campaign_id is not None:
            ads_nav["campaign_id"] = campaign_id
        render_ads_view()
        if screen == "list" and ads_nav["list"] is None and not ads_nav["list_busy"]:
            ads_load_list()
        if screen == "campaign":
            campaign = ads_current_campaign()
            if ads_state.get("plan_campaign") != campaign_id:
                ads_state["plan"], ads_state["result"] = None, None
                render_ads_results()
            # При первом открытии кампании подбор ставок запускается сам — на настройках по умолчанию.
            if campaign and campaign["running"] and campaign_id not in ads_nav["opened"]:
                ads_nav["opened"].add(campaign_id)
                ads_build()

    def ads_current_campaign() -> dict | None:
        data = ads_nav.get("list")
        if not isinstance(data, dict):
            return None
        return next((c for c in data["campaigns"] if c["id"] == ads_nav["campaign_id"]), None)

    def ads_set_period(preset: int | None = None) -> None:
        if preset is not None:
            ads_nav["preset"] = preset
            ads_nav["to"] = date.today()
            ads_nav["from"] = date.today() - timedelta(days=preset - 1)
            ads_custom_from.value, ads_custom_to.value = f"{ads_nav['from']:%d.%m.%Y}", f"{ads_nav['to']:%d.%m.%Y}"
        else:
            try:
                start = datetime.strptime((ads_custom_from.value or "").strip(), "%d.%m.%Y").date()
                end = datetime.strptime((ads_custom_to.value or "").strip(), "%d.%m.%Y").date()
            except ValueError:
                toast("Даты в формате ДД.ММ.ГГГГ, например 01.09.2026.", True)
                return
            if start > end or end > date.today():
                toast("Начало периода должно быть раньше конца, а конец — не позже сегодняшнего дня.", True)
                return
            ads_nav["preset"], ads_nav["from"], ads_nav["to"] = None, start, end
        ads_load_list()

    def ads_load_list() -> None:
        if list_cpc_campaigns is None:
            toast(f"Модуль рекламы недоступен: {ADS_IMPORT_ERROR or 'ошибка импорта'}", True)
            return
        if ads_nav["list_busy"]:
            return
        ads_nav["list_busy"], ads_nav["list_error"] = True, None
        render_ads_view()
        start, end = ads_nav["from"], ads_nav["to"]
        set_status("Реклама", f"Загружаю кампании за {start:%d.%m}–{end:%d.%m}...", ACCENT, busy=True, cancellable=False)

        def worker() -> None:
            try:
                ads_nav["list"] = list_cpc_campaigns(start, end)
                set_status("Реклама", f"Кампаний за период: {len(ads_nav['list']['campaigns'])}.", PRIMARY, busy=False)
            except Exception as exc:
                ads_nav["list_error"] = str(exc)
                set_status("Реклама", f"Не удалось загрузить кампании: {exc}", DANGER, busy=False)
            finally:
                ads_nav["list_busy"] = False
            render_ads_view()

        page.run_thread(worker)

    def ads_metric(label: str, value: str) -> ft.Control:
        pal = ads_palette()
        return ft.Container(
            width=118,
            content=ft.Column([ads_text(label, 10, pal["muted"]), ads_text(value, 14, weight=ft.FontWeight.W_700)], spacing=1, tight=True),
        )

    def ads_campaign_metrics(c: dict) -> list[ft.Control]:
        pct = lambda v: "—" if v is None else f"{v * 100:.1f}%".replace(".", ",")
        rub2 = lambda v: "—" if v is None else f"{v:.2f} ₽".replace(".", ",")
        return [
            ads_metric("Расход", ads_rub(c["spend"])), ads_metric("Показы", _format_number(c["views"])),
            ads_metric("Клики", _format_number(c["clicks"])), ads_metric("CTR", pct(c["ctr"])),
            ads_metric("Цена клика", rub2(c["cpc"])), ads_metric("Заказы", _format_number(c["orders"])),
            ads_metric("Выручка", ads_rub(c["revenue"])), ads_metric("ДРР", pct(c["drr"])),
            ads_metric("₽ на заказ", ads_rub(c["cpo"])),
        ]

    def ads_status_badge(c: dict) -> ft.Control:
        pal = ads_palette()
        color = pal["good"] if c["running"] else pal["muted"]
        return ft.Container(
            padding=ft.Padding.symmetric(horizontal=10, vertical=4), border_radius=999, border=ft.Border.all(1, color),
            data="preserve_excel_fill", content=ads_text(c["state_name"], 11, color, ft.FontWeight.W_700),
        )

    def ads_budget_text(c: dict) -> str:
        if c["weekly_budget"]:
            return f"{ads_rub(c['weekly_budget'])} в неделю"
        return f"{ads_rub(c['daily_budget'])} в день" if c["daily_budget"] else "без лимита"

    def render_ads_types() -> list[ft.Control]:
        pal = ads_palette()
        cards = []
        for key, title, description in ADS_TYPES:
            if key == "cpc":
                footer = ft.Button("Открыть кампании", bgcolor=PRIMARY, color="white", on_click=lambda _e: ads_go("list"))
            elif key == "promo":
                footer = ft.TextButton("Открыть раздел «Акции»", on_click=lambda _e: open_section(actions))
            else:
                footer = ads_text("Скоро", 12, pal["muted"], ft.FontWeight.W_700)
            enabled = key in ("cpc", "promo")
            cards.append(ft.Container(
                col={"md": 6, "xl": 4},
                content=ads_box(
                    ft.Column([ads_text(title, 16, weight=ft.FontWeight.W_800), ads_text(description, 12, pal["muted"]), footer], spacing=10),
                    height=190, opacity=1 if enabled else 0.55,
                    on_click=(lambda _e: ads_go("list")) if key == "cpc" else None,
                    ink=key == "cpc",
                ),
            ))
        return [*ads_header("Реклама", "Продвижение товаров. Пока программа работает с «Оплатой за клик»."),
                ft.ResponsiveRow(cards, run_spacing=14, spacing=14)]

    def render_ads_period_bar() -> ft.Control:
        pal = ads_palette()
        chips = []
        for days in ADS_PERIOD_PRESETS:
            selected = ads_nav["preset"] == days
            chips.append(ft.Container(
                padding=ft.Padding.symmetric(horizontal=14, vertical=8), border_radius=999,
                bgcolor=semantic_surface("positive", dark=pal["dark"]) if selected else None,
                border=ft.Border.all(2 if selected else 1, PRIMARY if selected else pal["border"]),
                data="preserve_excel_fill", ink=True,
                on_click=lambda _e, d=days: ads_set_period(d),
                content=ads_text(f"{days} дней" if days != 31 else "31 день", 12, weight=ft.FontWeight.W_700),
            ))
        start, end = ads_nav["from"], ads_nav["to"]
        if not ads_custom_from.value:
            ads_custom_from.value, ads_custom_to.value = f"{start:%d.%m.%Y}", f"{end:%d.%m.%Y}"
        return ads_box(ft.Row(
            [
                ads_text("Период:", 12, pal["muted"], ft.FontWeight.W_700), *chips,
                ft.Container(width=12), ads_custom_from, ads_custom_to,
                ft.TextButton("Показать свои даты", on_click=lambda _e: ads_set_period(None)),
                ft.Container(width=12),
                ads_text(f"Показано: {start:%d.%m.%Y} — {end:%d.%m.%Y}", 12, pal["muted"]),
            ],
            wrap=True, spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER,
        ), padding=12)

    def render_ads_list() -> list[ft.Control]:
        pal = ads_palette()
        controls = [*ads_header("Оплата за клик", "Кампании, которые работают сейчас или работали в выбранном периоде. Нажмите на кампанию, чтобы подобрать ставки.",
                                back=("Типы рекламы", lambda: ads_go("types"))),
                    render_ads_period_bar()]
        if ads_nav["list_busy"]:
            controls.append(ads_box(ft.Row([ft.ProgressRing(width=22, height=22, stroke_width=2.4),
                                            ads_text("Загружаю кампании и показатели за период…", 13, pal["muted"])], spacing=12)))
        elif ads_nav["list_error"]:
            controls.append(ads_box(ads_text(f"Не удалось загрузить кампании: {ads_nav['list_error']}", 13, pal["bad"])))
        elif isinstance(ads_nav["list"], dict):
            campaigns = ads_nav["list"]["campaigns"]
            if not campaigns:
                controls.append(ads_box(ads_text("За этот период кампаний с показами не было.", 13, pal["muted"])))
            for c in campaigns:
                controls.append(ads_box(
                    ft.Column([
                        ft.Row([ads_text(c["title"], 16, weight=ft.FontWeight.W_800), ads_status_badge(c), ft.Container(expand=True),
                                ft.Icon(ft.Icons.CHEVRON_RIGHT, color=pal["muted"])],
                               spacing=10, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                        ads_text(f"Стратегия «{c['strategy']}» · бюджет {ads_budget_text(c)}" + (f" · с {c['from_date']}" if c["from_date"] else ""), 12, pal["muted"]),
                        ft.Row(ads_campaign_metrics(c), wrap=True, spacing=8, run_spacing=8),
                    ], spacing=8),
                    ink=True, on_click=lambda _e, cid=c["id"]: ads_go("campaign", cid),
                ))
        return controls

    def render_ads_campaign() -> list[ft.Control]:
        pal = ads_palette()
        campaign = ads_current_campaign()
        back = ("Все кампании", lambda: ads_go("list"))
        if campaign is None:
            return [*ads_header("Кампания", "Кампания не найдена в выбранном периоде.", back=back)]
        controls = [*ads_header(campaign["title"], f"Стратегия «{campaign['strategy']}» · бюджет {ads_budget_text(campaign)}", back=back),
                    ads_box(ft.Column([
                        ft.Row([ads_status_badge(campaign),
                                ads_text(f"Показатели за {ads_nav['from']:%d.%m.%Y} — {ads_nav['to']:%d.%m.%Y}", 12, pal["muted"])], spacing=10),
                        ft.Row(ads_campaign_metrics(campaign), wrap=True, spacing=8, run_spacing=8),
                    ], spacing=10))]
        if not campaign["running"]:
            controls.append(ads_box(ads_text("Кампания сейчас не работает — подбор ставок доступен только для активных кампаний.", 13, pal["muted"])))
            return controls
        controls += [
            card(
                "Подбор ставок",
                "Безубыточная ставка = конверсия клика в заказ × прибыль с продажи. Рекомендуем долю от неё и меняем ставку не больше чем на 25% за раз. "
                "Убыточные товары предлагаем убрать из рекламы. В Ozon уходит только то, что вы отметите.",
                [
                    ft.Row([ads_share, ads_days], wrap=True),
                    ft.Row([
                        ft.Button("Подобрать ставки", bgcolor=PRIMARY, color="white", on_click=ads_build),
                        ads_apply_button,
                        ft.TextButton("Журнал изменений", on_click=lambda _e: open_ads_reports()),
                    ], wrap=True),
                ],
                tone="#EAF0F5",
            ),
            ads_results,
        ]
        return controls

    def render_ads_view(push_update: bool = True) -> None:
        screen = ads_nav["screen"]
        if screen == "types":
            ads_view.controls = render_ads_types()
        elif screen == "list":
            ads_view.controls = render_ads_list()
        else:
            ads_view.controls = render_ads_campaign()
        apply_control_theme(ads_view, current_dark["value"])
        if push_update and page.controls and host.controls and host.controls[0] is ads_view:
            page.update(ads_view)

    # ---------------------------------------------------------------- Анализ органических продаж
    analysis_state: dict[str, object] = {
        "days": 28, "result": organic_load_last(ROOT) if organic_load_last else None,
        "busy": False, "progress": "", "error": None, "started": False, "open": set(),
    }
    analysis_view = ft.Column(spacing=18, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)
    ANALYSIS_STATUS = {"bad": "Проблема", "warn": "Внимание", "info": "Инфо", "ok": "В порядке"}

    def analysis_colors() -> dict:
        dark = current_dark["value"]
        return {
            "organic": "#3987e5" if dark else "#2a78d6",
            "ads": "#d95926" if dark else "#eb6834",
            "status": {"bad": "#FF9C8F" if dark else DANGER, "warn": "#E0A15F" if dark else ACCENT,
                       "ok": "#70D4BC" if dark else PRIMARY, "info": "#B6C0BC" if dark else MUTED},
        }

    def analysis_run(_e=None) -> None:
        if organic_run_analysis is None:
            toast(f"Модуль анализа недоступен: {ORGANIC_IMPORT_ERROR or 'ошибка импорта'}", True)
            return
        if analysis_state["busy"]:
            return
        # Анализ грузит заказы за 24 месяца — в общей очереди, чтобы не делить лимит Ozon с другой задачей.
        if not job_lock.acquire(blocking=False):
            toast("Дождитесь завершения текущей операции.", True)
            return
        analysis_state.update(busy=True, error=None, progress="Готовлю анализ…", started=True)
        set_status("Анализ", "Собираю данные по продажам, рекламе, ценам и карточкам…", ACCENT, busy=True, cancellable=False)
        render_analysis_view()
        days = int(analysis_state["days"])

        def progress(message: str) -> None:
            analysis_state["progress"] = message
            ui_queue.put(("call", render_analysis_view))

        def worker() -> None:
            try:
                analysis_state["result"] = organic_run_analysis(ROOT, days=days, progress=progress)
                products = analysis_state["result"]["products"]
                problems = sum(1 for p in products if p["main_status"] in ("bad", "warn"))
                set_status("Анализ", f"Готово: товаров {len(products)}, с проблемами {problems}.", PRIMARY, busy=False)
            except Exception as exc:
                analysis_state["error"] = str(exc)
                set_status("Анализ", f"Не удалось выполнить анализ: {exc}", DANGER, busy=False)
            finally:
                analysis_state["busy"] = False
                job_lock.release()
            ui_queue.put(("call", render_analysis_view))

        try:
            page.run_thread(worker)
        except Exception as exc:
            analysis_state.update(busy=False, error=f"Не удалось запустить анализ: {exc}")
            job_lock.release()
            render_analysis_view()

    def analysis_set_days(days: int) -> None:
        # Только выбор периода: запрос в Ozon уходит по кнопке «Начать анализ».
        if analysis_state["busy"] or analysis_state["days"] == days:
            return
        analysis_state["days"] = days
        render_analysis_view()

    def analysis_delta(current: float, previous: float) -> str:
        if not previous:
            return "новые продажи" if current else "без изменений"
        return f"{(current - previous) / previous * 100:+.0f}% к прошлому периоду".replace(".", ",")

    def analysis_funnel(title: str, stages: list[tuple[str, float]], color: str, note: str) -> ft.Control:
        """Воронка: ширина полосы — логарифм значения (иначе заказы не видны рядом с тысячами показов)."""
        import math
        pal = ads_palette()
        top = max((math.log10(v + 1) for _, v in stages), default=1) or 1
        rows: list[ft.Control] = [ads_text(title, 14, weight=ft.FontWeight.W_800), ads_text(note, 11, pal["muted"])]
        previous = None
        for label, value in stages:
            share = max(0.06, math.log10(value + 1) / top)
            conversion = f"{value / previous * 100:.1f}% от предыдущего шага".replace(".", ",") if previous else ""
            rows.append(ft.Row(
                [
                    ft.Container(width=150, content=ads_text(label, 12, pal["muted"])),
                    ft.Row([ft.Container(expand=int((1 - share) * 500)), ft.Container(expand=int(share * 1000), height=26, border_radius=6, bgcolor=color, data="preserve_excel_fill"),
                            ft.Container(expand=int((1 - share) * 500))], expand=True, spacing=0),
                    ft.Container(width=110, content=ads_text(_format_number(value), 13, weight=ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT)),
                    ft.Container(width=190, content=ads_text(conversion, 11, pal["muted"])),
                ],
                spacing=10, vertical_alignment=ft.CrossAxisAlignment.CENTER,
            ))
            previous = value
        return ads_box(ft.Column(rows, spacing=8))

    def analysis_weekly_chart(weeks: list[dict]) -> ft.Control:
        pal, colors = ads_palette(), analysis_colors()
        top = max((w["total"] for w in weeks), default=1) or 1
        bars = []
        for w in weeks:
            ads_part = min(w["ads"], w["total"])
            organic_part = max(0.0, w["total"] - ads_part)
            week = datetime.strptime(w["week"], "%Y-%m-%d").date()
            bars.append(ft.Container(
                expand=1, tooltip=f"Неделя с {week:%d.%m}: всего {w['total']:.0f}, органика {organic_part:.0f}, реклама {ads_part:.0f}",
                content=ft.Column(
                    [
                        ads_text(f"{w['total']:.0f}", 10, pal["muted"]),
                        ft.Container(height=ads_part / top * 150, bgcolor=colors["ads"], border_radius=ft.BorderRadius.only(top_left=4, top_right=4), data="preserve_excel_fill"),
                        ft.Container(height=organic_part / top * 150, bgcolor=colors["organic"], data="preserve_excel_fill"),
                        ads_text(f"{week:%d.%m}", 9, pal["muted"]),
                    ],
                    spacing=2, alignment=ft.MainAxisAlignment.END, horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
                ),
            ))
        legend = ft.Row([
            ft.Row([ft.Container(width=10, height=10, border_radius=3, bgcolor=colors["organic"], data="preserve_excel_fill"), ads_text("органика", 11, pal["muted"])], spacing=6, tight=True),
            ft.Row([ft.Container(width=10, height=10, border_radius=3, bgcolor=colors["ads"], data="preserve_excel_fill"), ads_text("реклама", 11, pal["muted"])], spacing=6, tight=True),
        ], spacing=16)
        return ads_box(ft.Column([
            ads_text("Заказы по неделям", 14, weight=ft.FontWeight.W_800),
            ads_text("Органика = все заказы − заказы, которые Ozon приписал рекламе. Наведите на столбец, чтобы увидеть цифры.", 11, pal["muted"]),
            legend,
            ft.Container(height=200, content=ft.Row(bars, spacing=4, vertical_alignment=ft.CrossAxisAlignment.END)),
        ], spacing=8))

    def analysis_product_row(p: dict) -> ft.Control:
        pal, colors = ads_palette(), analysis_colors()
        status_color = colors["status"][p["main_status"]]
        is_open = p["offer_id"] in analysis_state["open"]
        details = ft.Container(visible=is_open, padding=ft.Padding.only(left=12, top=8, right=8, bottom=4))
        step_icons = {"ok": ft.Icons.CHECK_CIRCLE_OUTLINE, "warn": ft.Icons.WARNING_AMBER_ROUNDED, "bad": ft.Icons.ERROR_OUTLINE, "info": ft.Icons.INFO_OUTLINE}

        def build_details() -> list[ft.Control]:
            lines: list[ft.Control] = []
            for step in p["steps"]:
                lines.append(ft.Row([
                    ft.Icon(step_icons[step["status"]], size=16, color=colors["status"][step["status"]]),
                    ft.Column([ads_text(step["name"], 12, weight=ft.FontWeight.W_700), ads_text(step["text"], 12)], spacing=1, expand=True),
                ], spacing=8, vertical_alignment=ft.CrossAxisAlignment.START))
            ad = p.get("ads") or {}
            if ad.get("views"):
                lines.append(analysis_funnel("Воронка товара по рекламному трафику",
                                             [("Показы", ad["views"]), ("Клики", ad["clicks"]), ("В корзину", ad["to_cart"]), ("Заказы", ad["orders"])],
                                             colors["ads"], "Те же покупатели приходят и из поиска: слабый шаг здесь — слабый и в органике."))
            if p.get("events"):
                lines.append(ft.Column([ads_text("Что менялось (по снимкам программы)", 12, weight=ft.FontWeight.W_700),
                                        *[ads_text(f"• {e}", 12) for e in p["events"]]], spacing=2))
            if p.get("advice"):
                lines.append(ft.Column([ads_text("Рекомендации", 12, colors["status"]["warn"], ft.FontWeight.W_700),
                                        *[ads_text(f"→ {a}", 12) for a in p["advice"]]], spacing=3))
            return lines

        if is_open:
            details.content = ft.Column(build_details(), spacing=8)
        button = ft.TextButton("Скрыть" if is_open else "Подробнее")

        def toggle(_e) -> None:
            if p["offer_id"] in analysis_state["open"]:
                analysis_state["open"].discard(p["offer_id"])
                details.visible = False
            else:
                analysis_state["open"].add(p["offer_id"])
                details.content = ft.Column(build_details(), spacing=8)
                apply_control_theme(details, current_dark["value"])
                details.visible = True
            button.content = "Скрыть" if details.visible else "Подробнее"
            page.update(details, button)

        button.on_click = toggle
        organic_text = f"{p['organic']:.0f} ({p['organic_prev']:.0f})"
        return ft.Container(
            padding=ft.Padding.symmetric(vertical=8), border=ft.Border.only(top=ft.BorderSide(1, pal["border"])), data="preserve_excel_fill",
            content=ft.Column([
                ft.Row([
                    ft.Container(width=90, content=ads_text(str(p["offer_id"]), 13, weight=ft.FontWeight.W_700)),
                    ft.Container(width=110, content=ads_text(f"{p['orders']:.0f} ({p['orders_prev']:.0f})", 12, text_align=ft.TextAlign.RIGHT)),
                    ft.Container(width=120, content=ads_text(organic_text, 12, text_align=ft.TextAlign.RIGHT)),
                    ft.Container(width=80, content=ads_text(f"{p['ad_orders']:.0f}", 12, text_align=ft.TextAlign.RIGHT)),
                    ft.Container(width=16),
                    ft.Container(
                        padding=ft.Padding.symmetric(horizontal=10, vertical=3), border_radius=999, border=ft.Border.all(1, status_color), data="preserve_excel_fill",
                        content=ads_text(f"{ANALYSIS_STATUS[p['main_status']]}: {p['main_problem']}" if p["main_problem"] else "В порядке", 11, status_color, ft.FontWeight.W_700),
                    ),
                    ft.Container(expand=True, content=ads_text((p["advice"] or [""])[0], 11, pal["muted"], max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)),
                    button,
                ], spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                details,
            ], spacing=2),
        )

    def render_analysis_view(push_update: bool = True) -> None:
        pal, colors = ads_palette(), analysis_colors()
        controls: list[ft.Control] = [
            ft.Text("Анализ органических продаж", size=26, weight=ft.FontWeight.W_800, color=TEXT),
            ft.Text("Почему товары продаются или не продаются без рекламы: наличие и доставка, цена, карточка, воронка, поиск.", size=13, color=MUTED),
        ]
        chips = []
        for days in (7, 14, 28, 56):
            selected = analysis_state["days"] == days
            chips.append(ft.Container(
                padding=ft.Padding.symmetric(horizontal=14, vertical=8), border_radius=999, ink=True, data="preserve_excel_fill",
                bgcolor=semantic_surface("positive", dark=pal["dark"]) if selected else None,
                border=ft.Border.all(2 if selected else 1, PRIMARY if selected else pal["border"]),
                on_click=lambda _e, d=days: analysis_set_days(d),
                content=ads_text(f"{days} дней", 12, weight=ft.FontWeight.W_700),
            ))
        controls.append(ads_box(ft.Row([
            ads_text("Период:", 12, pal["muted"], ft.FontWeight.W_700), *chips,
            ft.Button("Начать анализ", bgcolor=PRIMARY, color="white", on_click=analysis_run, disabled=bool(analysis_state["busy"])),
            ads_text("Сравниваем с таким же предыдущим периодом. 28 дней = 4 полные недели — так сравнивают аналитики маркетплейсов.", 11, pal["muted"]),
        ], wrap=True, spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER), padding=12))

        if analysis_state["busy"]:
            controls.append(ads_box(ft.Row([ft.ProgressRing(width=22, height=22, stroke_width=2.4),
                                            ads_text(str(analysis_state["progress"]), 13, pal["muted"])], spacing=12)))
        if analysis_state["error"]:
            controls.append(ads_box(ads_text(f"Не удалось выполнить анализ: {analysis_state['error']}", 13, colors["status"]["bad"])))

        result = analysis_state["result"]
        if not isinstance(result, dict):
            if not analysis_state["busy"]:
                controls.append(ads_box(ads_text("Нажмите «Начать анализ» — займёт 1–3 минуты.", 13, pal["muted"])))
            analysis_view.controls = controls
            apply_control_theme(analysis_view, current_dark["value"])
            if push_update and page.controls and host.controls and host.controls[0] is analysis_view:
                page.update(analysis_view)
            return

        if not analysis_state["busy"] and result.get("days") not in (None, analysis_state["days"]):
            controls.append(ads_box(ads_text(
                f"Показан анализ за {result.get('days')} дней. Нажмите «Начать анализ», чтобы посчитать за {analysis_state['days']} дней.",
                13, pal["muted"])))

        periods = result["periods"]
        fmt = lambda value: datetime.strptime(value, "%Y-%m-%d").strftime("%d.%m")
        end = (datetime.strptime(periods["current"][1], "%Y-%m-%d") - timedelta(days=1)).strftime("%d.%m")
        prev_end = (datetime.strptime(periods["previous"][1], "%Y-%m-%d") - timedelta(days=1)).strftime("%d.%m")
        controls.append(ads_text(
            f"Анализ от {result['created'][:16].replace('T', ' ')} · {fmt(periods['current'][0])}–{end} против {fmt(periods['previous'][0])}–{prev_end}", 12, pal["muted"]))
        source_chips = []
        for source in result["sources"]:
            color = colors["status"].get(source["status"], pal["muted"])
            source_chips.append(ft.Container(
                padding=ft.Padding.symmetric(horizontal=10, vertical=4), border_radius=999, border=ft.Border.all(1, color), data="preserve_excel_fill",
                tooltip=source["text"], content=ads_text(source["name"], 11, color, ft.FontWeight.W_700)))
        controls.append(ft.Row([ads_text("Источники:", 11, pal["muted"]), *source_chips], wrap=True, spacing=6))

        store = result["store"]
        season = store.get("season_change")
        season_text = f"по сезону год назад {season * 100:+.0f}%".replace(".", ",") if season is not None else "сезон не определён"
        share = store["organic"] / store["orders"] if store["orders"] else None
        controls.append(ft.ResponsiveRow([
            hero_metric("Заказов всего", f"{store['orders']:.0f}", analysis_delta(store["orders"], store["orders_prev"]),
                        semantic_surface("neutral_soft", dark=pal["dark"]), PRIMARY, width=260, col=3),
            hero_metric("Органические", f"{store['organic']:.0f}", f"{analysis_delta(store['organic'], store['organic_prev'])}; {season_text}",
                        semantic_surface("neutral_soft", dark=pal["dark"]), colors["organic"], width=260, col=3),
            hero_metric("С рекламы", f"{store['ad_orders']:.0f}", analysis_delta(store["ad_orders"], store["ad_orders_prev"]),
                        semantic_surface("neutral_soft", dark=pal["dark"]), colors["ads"], width=260, col=3),
            hero_metric("Доля органики", f"{share * 100:.0f}%" if share is not None else "—", "чем выше, тем меньше зависимость от рекламы",
                        semantic_surface("neutral_soft", dark=pal["dark"]), colors["organic"], width=260, col=3),
        ], run_spacing=14, spacing=14))

        funnels = []
        if store.get("search_funnel"):
            sf = store["search_funnel"]
            funnels.append(analysis_funnel("Поиск: от запроса до заказа",
                                           [("Искали", sf["searched"]), ("Увидели товар", sf["viewed"]), ("Заказали", sf["ordered"])],
                                           colors["organic"], "Данные Premium: сколько покупателей искали по запросам с вашими товарами, увидели их и заказали."))
        af = store["ad_funnel"]
        if af.get("views"):
            funnels.append(analysis_funnel("Воронка карточек по рекламному трафику",
                                           [("Показы", af["views"]), ("Клики", af["clicks"]), ("В корзину", af["to_cart"]), ("Заказы", af["orders"])],
                                           colors["ads"], "Поисковую воронку Ozon без Premium не отдаёт, поэтому карточки оцениваем по рекламным покупателям — карточка у них та же."))
        if not store.get("search_funnel"):
            funnels.append(ads_box(ads_text("Поисковая воронка (запросы → просмотры → заказы) и позиции появятся здесь автоматически, "
                                            "когда будет подключена подписка Premium.", 12, pal["muted"])))
        controls.extend(funnels)
        if result.get("weeks"):
            controls.append(analysis_weekly_chart(result["weeks"]))

        products = result["products"]
        top = [p for p in products if p["advice"] and p["main_status"] in ("bad", "warn")][:6]
        if top:
            controls.append(ads_box(ft.Column(
                [ads_text("Что сделать в первую очередь", 14, weight=ft.FontWeight.W_800),
                 *[ft.Row([ft.Container(width=70, content=ads_text(str(p["offer_id"]), 12, weight=ft.FontWeight.W_700)),
                           ads_text(p["advice"][0], 12, expand=True)], vertical_alignment=ft.CrossAxisAlignment.START) for p in top]],
                spacing=8)))
        header = ft.Row([
            ft.Container(width=90, content=ads_text("Артикул", 11, pal["muted"], ft.FontWeight.W_700)),
            ft.Container(width=110, content=ads_text("Заказы (было)", 11, pal["muted"], ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT)),
            ft.Container(width=120, content=ads_text("Органика (было)", 11, pal["muted"], ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT)),
            ft.Container(width=80, content=ads_text("Реклама", 11, pal["muted"], ft.FontWeight.W_700, text_align=ft.TextAlign.RIGHT)),
            ft.Container(width=16),
            ads_text("Главная проблема и первая рекомендация", 11, pal["muted"], ft.FontWeight.W_700),
        ], spacing=8)
        controls.append(ads_box(ft.Column(
            [ads_text(f"Товары — {len(products)}, по важности", 14, weight=ft.FontWeight.W_800), header, *[analysis_product_row(p) for p in products]],
            spacing=4)))
        analysis_view.controls = controls
        apply_control_theme(analysis_view, current_dark["value"])
        if push_update and page.controls and host.controls and host.controls[0] is analysis_view:
            page.update(analysis_view)

    # ---------------------------------------------------------------- Акции: календарь
    calendar_state: dict[str, object] = {"data": None, "busy": False, "error": None, "started": False, "open": set(), "offset": 0,
                                         "expanded": False}
    CALENDAR_PAST_DAYS, CALENDAR_DAYS = 3, 35
    MONTHS_RU = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
    WEEKDAYS_RU = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]

    def calendar_load(_e=None) -> None:
        if load_actions_calendar is None:
            toast(f"Модуль акций недоступен: {PRICING_IMPORT_ERROR or 'ошибка импорта'}", True)
            return
        if calendar_state["busy"]:
            return
        calendar_state.update(busy=True, error=None, started=True)
        render_actions_calendar()

        def worker() -> None:
            try:
                calendar_state["data"] = load_actions_calendar(ROOT)
            except Exception as exc:
                calendar_state["error"] = str(exc)
            finally:
                calendar_state["busy"] = False
            ui_queue.put(("call", render_actions_calendar))

        page.run_thread(worker)

    def calendar_boost_text(action: dict) -> str:
        boosts = [p["boost"] for p in action["participants"] if p.get("boost") is not None]
        if action.get("fixed_boost") is not None:
            return f"буст {action['fixed_boost']:g}%"
        if boosts:
            low, high = min(boosts), max(boosts)
            return f"буст {low:g}%" if low == high else f"буст {low:g}–{high:g}%"
        return "буст по скидке"

    def calendar_shift(days: int | None) -> None:
        """Листание календаря: ±неделя или возврат к сегодняшнему дню (None)."""
        calendar_state["offset"] = 0 if days is None else int(calendar_state.get("offset") or 0) + days
        render_actions_calendar()

    def render_actions_calendar(push_update: bool = True) -> None:
        """
        Календарь акций сверху раздела. Компактный (по умолчанию): строка на акцию — полоса на шкале по неделям, число
        товаров и бустинг, «Товары ›» фильтрует ползунки предельной цены по этой акции. Развёрнутый: ещё даты, сколько
        можно добавить, автодобавление Ozon и участники. Цвет акции — общий для раздела (action_color).
        """
        pal = ads_palette()
        dark = pal["dark"]
        today_color = "#FF9C8F" if dark else DANGER
        warn = "#E0A15F" if dark else ACCENT
        legend_color = "#4f93ea" if dark else "#2a78d6"  # образец полосы в легенде; сами полосы — цветом акции
        week_line = "#2F363B" if dark else "#ECEEF3"
        expanded = bool(calendar_state.get("expanded"))
        label_width = 230 if expanded else 200
        side_width = 150
        offset = int(calendar_state.get("offset") or 0)
        today = date.today()
        start = today - timedelta(days=CALENDAR_PAST_DAYS) + timedelta(days=offset)
        days = [start + timedelta(days=i) for i in range(CALENDAR_DAYS)]
        end = days[-1]
        data = calendar_state["data"]

        def segments(parts: list[tuple[int, ft.Control | None]]) -> ft.Control:
            """Строка из отрезков по дням окна: (сколько дней, содержимое)."""
            return ft.Row([ft.Container(expand=n, content=c) for n, c in parts if n > 0], spacing=0)

        def week_grid(height: int) -> ft.Control:
            # Тонкие линии только по понедельникам — вместо сетки по каждому дню.
            parts, run = [], 0
            for day in days:
                if day.weekday() == 0 and run:
                    parts.append(run)
                    run = 0
                run += 1
            parts.append(run)
            return ft.Row([ft.Container(expand=n, height=height, border=ft.Border.only(left=ft.BorderSide(1, week_line)) if i else None,
                                        data="preserve_excel_fill") for i, n in enumerate(parts)], spacing=0)

        def today_line(height: int) -> ft.Control | None:
            if not start <= today <= end:
                return None
            before = (today - start).days
            return segments([(before, None), (1, ft.Container(alignment=ft.Alignment(0, 0), content=ft.Container(
                width=2, height=height, bgcolor=today_color, data="preserve_excel_fill"))), (CALENDAR_DAYS - before - 1, None)])

        title_info = ""
        if isinstance(data, dict):
            running = [a for a in data["actions"] if a.get("start") and a.get("end") and date.fromisoformat(a["start"]) <= today <= date.fromisoformat(a["end"])]
            upcoming_count = sum(1 for a in data["actions"] if a.get("start") and date.fromisoformat(a["start"]) > today)
            title_info = f"идут {len(running)}" + (f" · скоро {upcoming_count}" if upcoming_count else "")
        controls: list[ft.Control] = [ft.Row([
            ads_text("Календарь акций", 15, weight=ft.FontWeight.W_800),
            ads_text(title_info, 12, pal["muted"]),
            ft.Container(expand=True),
            ft.IconButton(ft.Icons.CHEVRON_LEFT, tooltip="Неделя назад", icon_size=18, on_click=lambda _e: calendar_shift(-7)),
            ft.TextButton("Сегодня", on_click=lambda _e: calendar_shift(None), disabled=offset == 0),
            ft.IconButton(ft.Icons.CHEVRON_RIGHT, tooltip="Неделя вперёд", icon_size=18, on_click=lambda _e: calendar_shift(7)),
            ft.IconButton(ft.Icons.REFRESH, tooltip="Обновить акции", icon_size=18, on_click=calendar_load, disabled=bool(calendar_state["busy"])),
            ft.TextButton("Свернуть" if expanded else "Подробнее", icon=ft.Icons.EXPAND_LESS if expanded else ft.Icons.EXPAND_MORE,
                          on_click=lambda _e: calendar_toggle()),
        ], spacing=2, vertical_alignment=ft.CrossAxisAlignment.CENTER)]
        if calendar_state["busy"] and not data:
            controls.append(ft.Row([ft.ProgressRing(width=18, height=18, stroke_width=2.4), ads_text("Загружаю акции и участников…", 12, pal["muted"])], spacing=10))
        if calendar_state["error"]:
            controls.append(ads_text(f"Не удалось загрузить акции: {calendar_state['error']}", 12, pal["bad"]))

        if isinstance(data, dict):
            register_action_colors(action["id"] for action in data["actions"])
            # Шкала: месяц у первой недели и понедельники с числом.
            axis, run_start = [], 0
            for i, day in enumerate(days):
                if i and (day.weekday() == 0):
                    axis.append((i - run_start, days[run_start]))
                    run_start = i
            axis.append((CALENDAR_DAYS - run_start, days[run_start]))
            axis_row = ft.Row([ft.Container(expand=n, padding=ft.Padding.only(left=4), content=ads_text(
                (f"{MONTHS_RU[d.month - 1][:3].lower()} " if index == 0 or d.day <= 7 else "") + f"{d.day}", 10, pal["muted"], ft.FontWeight.W_600,
                max_lines=1, overflow=ft.TextOverflow.CLIP)) for index, (n, d) in enumerate(axis)], spacing=0)
            controls.append(ft.Row([ft.Container(width=label_width), ft.Container(expand=True, content=axis_row), ft.Container(width=side_width)], spacing=10))

            bar_height = 18 if expanded else 14
            row_height = bar_height + 8
            shown = 0
            active_filter = limit_state.get("action_filter")
            for action in data["actions"]:
                if not action.get("start") or not action.get("end"):
                    continue
                first, last = date.fromisoformat(action["start"]), date.fromisoformat(action["end"])
                if last < start or first > end:
                    continue
                shown += 1
                participating = action["participating"]
                color = action_color(action["id"])
                bar_from, bar_to = max(first, start), min(last, end)
                before = (bar_from - start).days
                length = (bar_to - bar_from).days + 1
                after = CALENDAR_DAYS - before - length
                label = f"{participating} тов. · {calendar_boost_text(action)}"
                inside = length >= 9
                bar = ft.Container(
                    height=bar_height, border_radius=bar_height // 2, padding=ft.Padding.symmetric(horizontal=8), alignment=ft.Alignment(-1, 0),
                    bgcolor=color if participating else pal["soft"], border=None if participating else ft.Border.all(1.5, color),
                    data="preserve_excel_fill",
                    tooltip=f"{action['name']}\n{first:%d.%m.%Y} — {last:%d.%m.%Y}\nУчаствует {participating}, можно добавить ещё {action['can_add']}",
                    content=ads_text(label, 10, "#FFFFFF" if participating else pal["muted"], ft.FontWeight.W_700,
                                     max_lines=1, overflow=ft.TextOverflow.ELLIPSIS) if inside else None,
                )
                outside = None if inside else ft.Container(padding=ft.Padding.only(left=6), alignment=ft.Alignment(-1, 0), content=ads_text(
                    label, 10, pal["ink"], ft.FontWeight.W_700, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS))
                auto_dates = {date.fromisoformat(d) for d in action.get("auto_add_dates") or []}
                markers = segments([(1, ft.Container(alignment=ft.Alignment(0, 1), content=ft.Container(
                    width=7, height=7, border_radius=4, bgcolor=warn, data="preserve_excel_fill", tooltip=f"{day:%d.%m}: Ozon добавит товары сам")
                    if day in auto_dates else None)) for day in days])
                layers = [ft.Container(left=0, right=0, top=0, height=row_height, content=week_grid(row_height)),
                          ft.Container(left=0, right=0, top=4, height=bar_height, content=segments([(before, None), (length, bar), (after, outside)]))]
                line = today_line(row_height)
                if line:
                    layers.append(ft.Container(left=0, right=0, top=0, height=row_height, content=line))
                layers.append(ft.Container(left=0, right=0, top=0, height=row_height, content=markers))
                timeline = ft.Stack(layers, height=row_height, expand=True)

                filtered = active_filter == action["id"]
                filter_button = ft.TextButton(
                    "Показаны ✕" if filtered else "Товары ›",
                    tooltip="Показать в «Предельной цене» только товары этой акции" if not filtered else "Снять фильтр по акции",
                    on_click=lambda _e, a=action["id"]: limit_filter_action(None if limit_state.get("action_filter") == a else a))
                name = ads_text(action["name"] if expanded else short_action_name(action["name"]), 12, pal["bad"] if action.get("error") else None,
                                ft.FontWeight.W_700, max_lines=2 if expanded else 1, overflow=ft.TextOverflow.ELLIPSIS, expand=True)
                label_parts: list[ft.Control] = [ft.Row([
                    ft.Container(width=10, height=10, border_radius=5, bgcolor=color, data="preserve_excel_fill"), name], spacing=6)]
                if expanded:
                    upcoming = sorted(d for d in auto_dates if d >= today)
                    label_parts.append(ads_text(f"{first:%d.%m} — {last:%d.%m.%Y} · можно ещё {action['can_add']}", 10, pal["muted"]))
                    if upcoming:
                        label_parts.append(ads_text("● Ozon добавит сам: " + ", ".join(f"{d:%d.%m}" for d in upcoming), 10, warn))
                row = ft.Container(
                    border_radius=8, bgcolor=pal["soft"] if filtered else None, data="preserve_excel_fill",
                    padding=ft.Padding.symmetric(vertical=2),
                    content=ft.Row([
                        ft.Container(width=label_width, content=ft.Column(label_parts, spacing=1, tight=True)),
                        timeline,
                        ft.Container(width=side_width, alignment=ft.Alignment(1, 0), content=filter_button),
                    ], spacing=10, vertical_alignment=ft.CrossAxisAlignment.CENTER))
                controls.append(row)
                if expanded and action["participants"]:
                    chips = [ft.Container(
                        padding=ft.Padding.symmetric(horizontal=8, vertical=3), border_radius=8, border=ft.Border.all(1, pal["border"]),
                        data="preserve_excel_fill",
                        content=ads_text(f"{item['offer_id']} · {item['action_price']:,.0f} ₽".replace(",", " ")
                                         + (f" · {item['boost']:g}%" if item.get("boost") is not None else ""), 10))
                        for item in action["participants"]]
                    controls.append(ft.Container(padding=ft.Padding.only(left=label_width + 10, bottom=4),
                                                 content=ft.Row(chips, wrap=True, spacing=4, run_spacing=4)))
            if not shown:
                controls.append(ads_text("В эти дни акций нет — пролистайте назад или вперёд.", 12, pal["muted"]))
            if expanded:
                controls.append(ft.Row([
                    ft.Container(width=14, height=8, border_radius=4, bgcolor=legend_color, data="preserve_excel_fill"), ads_text("ваши товары участвуют", 10, pal["muted"]),
                    ft.Container(width=14, height=8, border_radius=4, bgcolor=pal["soft"], border=ft.Border.all(1.5, legend_color), data="preserve_excel_fill"),
                    ads_text("не участвуют", 10, pal["muted"]),
                    ft.Container(width=7, height=7, border_radius=4, bgcolor=warn, data="preserve_excel_fill"), ads_text("автодобавление Ozon", 10, pal["muted"]),
                    ft.Container(width=2, height=10, bgcolor=today_color, data="preserve_excel_fill"), ads_text("сегодня", 10, pal["muted"]),
                    ads_text(f"· обновлено {data['loaded'][11:16]}", 10, pal["muted"]),
                ], spacing=6, wrap=True))
        actions_calendar.content = ft.Column(controls, spacing=4 if not expanded else 6)
        actions_calendar.padding = ft.Padding.symmetric(horizontal=16, vertical=10)
        actions_calendar.bgcolor = pal["surface"]
        actions_calendar.border = ft.Border.all(1, pal["border"])
        apply_control_theme(actions_calendar, current_dark["value"])
        if push_update and page.controls and host.controls and host.controls[0] is actions:
            page.update(actions_calendar)

    def calendar_toggle() -> None:
        calendar_state["expanded"] = not calendar_state.get("expanded")
        render_actions_calendar()

    # ---------------------------------------------------------------- Акции: автодобавление и учёт минимальной цены
    auto_state: dict[str, object] = {"data": None, "busy": False, "error": None, "started": False, "selected": set(), "filter": "all"}
    AUTO_SOON_HOURS = 72
    AUTO_ACTIONS = {
        "auto_add_off": ("Отключить автодобавление",
                         "Ozon перестанет сам добавлять эти товары в акции. Уже добавленные останутся — их убирайте в акциях. "
                         "Срока у отключения нет; программа будет сверять его с Ozon и предупредит, если Ozon включит обратно."),
        "auto_add_on": ("Включить автодобавление", "Ozon снова сможет сам добавлять эти товары в акции."),
        "min_price_on": ("Учитывать минимальную цену",
                         "Ozon не будет добавлять товары в акции ниже минимальной цены. Учёт работает до сброса таймера — "
                         "дату окончания программа покажет после применения. Нужна заданная минимальная цена."),
        "timer_update": ("Продлить таймер",
                         "Подтверждаем Ozon, что минимальная цена актуальна: учёт минимальной цены продлится ещё на период. "
                         "Цены не меняются, переносится только дата сброса."),
    }

    def auto_load(_e=None) -> None:
        if auto_actions is None:
            return
        if auto_state["busy"]:
            return
        auto_state.update(busy=True, error=None, started=True)
        render_auto_panel()

        def worker() -> None:
            try:
                auto_state["data"] = auto_actions.load_state(ROOT)
            except Exception as exc:
                auto_state["error"] = str(exc)
            finally:
                auto_state["busy"] = False
            ui_queue.put(("call", render_auto_panel))

        page.run_thread(worker)

    def auto_left_text(hours: float | None) -> str:
        if hours is None:
            return ""
        if hours <= 0:
            return "истёк"
        days, rest = divmod(int(hours), 24)
        return f"{days} д {rest} ч" if days else f"{rest} ч"

    def auto_rows_filtered() -> list[dict]:
        data = auto_state.get("data") or {}
        rows = list(data.get("rows") or [])
        mode = auto_state.get("filter")
        if mode == "auto_add_on":
            rows = [r for r in rows if r["auto_add"]]
        elif mode == "soon":
            rows = [r for r in rows if r["hours_left"] is not None and r["hours_left"] <= AUTO_SOON_HOURS]
        elif mode == "min_off":
            rows = [r for r in rows if not r["min_price_on"]]
        return rows

    def auto_set_filter(mode: str) -> None:
        auto_state["filter"] = mode
        render_auto_panel()

    def auto_select(ids: set | None) -> None:
        auto_state["selected"] = set(ids or ())
        render_auto_panel()

    def auto_confirm(action: str) -> None:
        data = auto_state.get("data") or {}
        chosen = [r for r in data.get("rows") or [] if r["product_id"] in auto_state["selected"]]
        if not chosen:
            toast("Отметьте товары в таблице.", True)
            return
        title, explanation = AUTO_ACTIONS[action]
        dark = current_dark["value"]
        articles = ", ".join(r["offer_id"] for r in chosen)
        body = ft.Column([
            ft.Text(explanation, size=13, color="#B6C0BC" if dark else MUTED),
            ft.Text(f"Товары ({len(chosen)}): {articles}", size=13, color="#F3F4F6" if dark else TEXT),
        ], spacing=10, tight=True)

        def run(_e=None) -> None:
            page.pop_dialog()
            if not job_lock.acquire(blocking=False):
                toast("Дождитесь завершения текущей операции.", True)
                return
            set_status(title, "Отправляю изменение в Ozon и проверяю результат…", ACCENT, busy=True, cancellable=False)

            def worker() -> None:
                try:
                    result = auto_actions.apply_change(ROOT, action, chosen)
                except Exception as exc:
                    result = {"ok": False, "message": f"Не удалось: {exc}", "failed": [], "state": None}
                finally:
                    job_lock.release()
                message = str(result.get("message") or "")
                append_log(f"{title}: {message}")
                for entry in result.get("failed") or []:
                    append_log(f"  ✖ {entry['offer_id']}: {entry['reason']}")
                if result.get("state"):
                    auto_state["data"] = {"loaded": datetime.now().isoformat(timespec="seconds"),
                                          "rows": sorted(result["state"], key=lambda r: (len(r["offer_id"]), r["offer_id"]))}
                auto_state["selected"] = set()
                set_status(title, message, PRIMARY if result.get("ok") else DANGER, busy=False)
                failed = result.get("failed") or []
                toast(message + (" Причины — в журнале задач." if failed else ""), bool(failed) or not result.get("ok"))
                ui_queue.put(("call", render_auto_panel))

            page.run_thread(worker)

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text(f"{title}?", size=18, weight=ft.FontWeight.W_800, color="#F3F4F6" if dark else TEXT),
            bgcolor="#232A2F" if dark else "#FFFDF8",
            content=ft.Container(width=620, content=body),
            actions=[ft.TextButton("Отмена", on_click=lambda _e: page.pop_dialog()),
                     ft.Button(f"Отправить в Ozon ({len(chosen)})", bgcolor=PRIMARY, color="white", on_click=run)],
            actions_alignment=ft.MainAxisAlignment.END,
        )
        # show_dialog сразу отправляет окно на экран и убирает его из страницы после закрытия.
        # Раньше окно добавлялось в overlay и обновлялось только само: клиент его не знал, и оно
        # появлялось лишь при следующем полном обновлении страницы (иногда через много секунд).
        page.show_dialog(dialog)

    def render_auto_panel(push_update: bool = True) -> None:
        pal = dash_pal()
        data = auto_state.get("data") or {}
        rows_all = list(data.get("rows") or [])
        selected = auto_state["selected"]
        header = ft.Row([
            dtext("Автодобавление в акции и минимальная цена", 16, weight=ft.FontWeight.W_800),
            ft.TextButton("Обновить", icon=ft.Icons.REFRESH, on_click=auto_load, disabled=bool(auto_state["busy"])),
        ], spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER)
        controls: list[ft.Control] = [
            header,
            dtext("Автодобавление — Ozon сам ставит товар в акции. Учёт минимальной цены — Ozon не ставит товар в акцию ниже неё, "
                  "но только до сброса таймера; «Продлить таймер» переносит сброс.", 12, pal["muted"]),
        ]
        if auto_actions is None:
            controls.append(dtext("Модуль автоакций не загрузился.", 13, pal["bad"]))
        if auto_state["busy"] and not rows_all:
            controls.append(ft.Row([ft.ProgressRing(width=20, height=20, stroke_width=2.4), dtext("Загружаю состояние товаров…", 13, pal["muted"])], spacing=10))
        if auto_state.get("error"):
            controls.append(dtext(f"Не удалось загрузить: {auto_state['error']}", 13, pal["bad"]))
        if rows_all:
            soon = [r for r in rows_all if r["hours_left"] is not None and r["hours_left"] <= AUTO_SOON_HOURS]
            reverted = [r for r in rows_all if r["reverted_by_ozon"]]
            summary = [
                f"автодобавление включено у {sum(1 for r in rows_all if r['auto_add'])} из {len(rows_all)}",
                f"учёт минимальной цены — у {sum(1 for r in rows_all if r['min_price_on'])}",
            ]
            controls.append(dtext(" · ".join(summary), 13))
            if soon:
                first = min(soon, key=lambda r: r["hours_left"])
                controls.append(dtext(f"⚠ У {len(soon)} товаров учёт минимальной цены сбросится в ближайшие 3 дня "
                                      f"(первый — через {auto_left_text(first['hours_left'])}).", 13, pal["warn"], ft.FontWeight.W_600))
            if reverted:
                controls.append(dtext("⚠ Ozon сам включил обратно автодобавление: " + ", ".join(r["offer_id"] for r in reverted), 13, pal["bad"], ft.FontWeight.W_600))

            def chip(label: str, active: bool, on_click) -> ft.Control:
                return ft.Container(
                    padding=ft.Padding.symmetric(horizontal=12, vertical=6), border_radius=999, data="preserve_excel_fill",
                    bgcolor=pal["accent"] if active else None, border=None if active else ft.Border.all(1, pal["border"]),
                    on_click=on_click, content=dtext(label, 12, "#FFFFFF" if active else pal["ink"], ft.FontWeight.W_600))

            mode = auto_state.get("filter")
            visible = auto_rows_filtered()
            controls.append(ft.Row([
                chip("Все", mode == "all", lambda _e: auto_set_filter("all")),
                chip("Автодобавление вкл.", mode == "auto_add_on", lambda _e: auto_set_filter("auto_add_on")),
                chip("Сброс ≤ 3 дн.", mode == "soon", lambda _e: auto_set_filter("soon")),
                chip("Учёт мин. цены выкл.", mode == "min_off", lambda _e: auto_set_filter("min_off")),
                ft.Container(width=12),
                ft.TextButton("Отметить все видимые", on_click=lambda _e: auto_select({r["product_id"] for r in visible})),
                ft.TextButton("Снять отметки", on_click=lambda _e: auto_select(set()), disabled=not selected),
            ], spacing=8, run_spacing=8, wrap=True, vertical_alignment=ft.CrossAxisAlignment.CENTER))
            controls.append(ft.Row([
                ft.Button(f"{AUTO_ACTIONS[key][0]}" + (f" ({len(selected)})" if selected else ""), bgcolor=color, color="white",
                          disabled=not selected or bool(auto_state["busy"]), on_click=lambda _e, k=key: auto_confirm(k))
                for key, color in (("auto_add_off", DANGER), ("auto_add_on", PRIMARY), ("min_price_on", PRIMARY), ("timer_update", ACCENT))
            ], spacing=8, run_spacing=8, wrap=True))

            def badge(on: bool | None, yes: str = "вкл", no: str = "выкл") -> ft.Control:
                color = pal["good"] if on else pal["muted"]
                return ft.Row([ft.Container(width=8, height=8, border_radius=4, bgcolor=color, data="preserve_excel_fill"),
                               dtext(yes if on else (no if on is not None else "—"), 12, None if on else pal["muted"])], spacing=6, tight=True)

            columns = [("", 1), ("Артикул", 2), ("Цена", 2), ("Мин. цена в Ozon", 2), ("Автодобавление", 3),
                       ("Учёт мин. цены", 5), ("Заметка", 5)]

            def cell(content: ft.Control, weight: int) -> ft.Control:
                return ft.Container(expand=weight, alignment=ft.Alignment(-1, 0), content=content)

            table: list[ft.Control] = [ft.Row([cell(dtext(name, 12, pal["muted"], ft.FontWeight.W_600), w) for name, w in columns], spacing=10)]
            for row in visible:
                def toggle(e: ft.ControlEvent, pid=row["product_id"]) -> None:
                    if e.control.value:
                        auto_state["selected"].add(pid)
                    else:
                        auto_state["selected"].discard(pid)
                    render_auto_panel()

                if row["min_price_on"] and row["expires_at"]:
                    expires = datetime.fromisoformat(row["expires_at"]).astimezone()
                    soon_flag = row["hours_left"] is not None and row["hours_left"] <= AUTO_SOON_HOURS
                    timer = dtext(f"до {expires:%d.%m %H:%M} · осталось {auto_left_text(row['hours_left'])}", 12,
                                  pal["warn"] if soon_flag else None, ft.FontWeight.W_600 if soon_flag else None)
                else:
                    timer = dtext("выкл", 12, pal["muted"])
                notes = []
                if row["reverted_by_ozon"]:
                    notes.append(("Ozon включил автодобавление обратно", pal["bad"]))
                elif row["auto_add_changed_at"] and row["auto_add"] is False:
                    notes.append((f"отключено {row['auto_add_changed_at'][8:10]}.{row['auto_add_changed_at'][5:7]}", pal["muted"]))
                if row["min_price_on"] and not row["min_price"]:
                    notes.append(("мин. цена в Ozon не задана — учитывать нечего", pal["warn"]))
                table.append(ft.Container(
                    padding=ft.Padding.symmetric(vertical=4), data="preserve_excel_fill",
                    border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])),
                    content=ft.Row([
                        cell(ft.Checkbox(value=row["product_id"] in selected, on_change=toggle), 1),
                        cell(dtext(row["offer_id"], 13, weight=ft.FontWeight.W_700), 2),
                        cell(dtext(_format_currency(row["price"]) if row["price"] else "—", 13), 2),
                        cell(dtext(_format_currency(row["min_price"]) if row["min_price"] else "—", 13, None if row["min_price"] else pal["muted"]), 2),
                        cell(badge(row["auto_add"]), 3),
                        cell(timer, 5),
                        cell(ft.Column([dtext(text, 12, color) for text, color in notes], spacing=0, tight=True) if notes else dtext("", 12), 5),
                    ], spacing=10, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ))
            controls.append(ft.Column(table, spacing=0))
            controls.append(dtext(f"Обновлено {str(data.get('loaded') or '')[11:16]}. Журнал изменений — actions history/auto_actions_log.jsonl.", 11, pal["muted"]))
        auto_panel.content = ft.Column(controls, spacing=10)
        auto_panel.bgcolor = pal["surface"]
        auto_panel.border = ft.Border.all(1, pal["border"])
        if push_update and page.controls and host.controls and host.controls[0] is actions:
            page.update(auto_panel)

    def launch_supply_plan() -> None:
        lead, cover, level = (supply_lead.value or "").strip(), (supply_cover.value or "").strip(), (supply_level.value or "").strip().replace(",", ".")
        error = valid_supply(lead, cover, level)
        if error:
            toast(error, True)
            return
        job(
            "План поставки FBO",
            [str(PYTHON), str(SCRIPTS / "fbo_supply_plan.py"), "--lead-days", lead, "--cover-days", cover, "--service-level", level],
        )

    supply = ft.Column(
        [
            ft.Text("Поставки", size=26, weight=ft.FontWeight.W_800, color=TEXT),
            ft.Text("Сколько штук каждого артикула везти на склады FBO и сколько держать на своём складе (FBS) — по спросу, сезонности и текущим остаткам.", size=13, color=MUTED),
            card(
                "План поставки FBO",
                "Прогноз по продажам того же периода год назад с поправкой на рост, разложенный по зонам складов. "
                "Excel с расчётом и объяснением каждого шага сохраняется в stocks reports.",
                [
                    ft.Row([supply_lead, supply_cover, supply_level], wrap=True),
                    ft.Row(
                        [
                            ft.Button("Рассчитать поставку", bgcolor=PRIMARY, color="white", on_click=lambda _e: launch_supply_plan()),
                            ft.TextButton("Открыть stocks reports", on_click=lambda _e: reveal_path(ROOT / "stocks reports", "stocks reports")),
                        ]
                    ),
                ],
                tone="#EAF0F5",
            ),
            supply_results,
        ],
        spacing=18,
    )

    # ---------------------------------------------------------------- Поставки WB
    wb_supply_results = ft.Column(spacing=16, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)

    def latest_wb_supply_summary() -> dict | None:
        folder = ROOT / "stocks reports"
        files = sorted(folder.glob("Поставка WB *.json"), key=lambda path: path.stat().st_mtime) if folder.exists() else []
        if not files:
            return None
        try:
            return json.loads(files[-1].read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return None

    def wb_supply_defaults() -> dict:
        defaults = {"lead_days": 14, "cover_days": 30, "service_level": 0.9, "wb_fbw_min_monthly": 3}
        try:
            defaults.update(json.loads((ROOT / "wb_supply_settings.json").read_text(encoding="utf-8")))
        except (OSError, ValueError):
            pass
        return defaults

    wb_defaults = wb_supply_defaults()
    wb_supply_lead = ft.TextField(label="Срок поставки, дней", value=str(wb_defaults["lead_days"]), width=190,
                                  tooltip="От решения о поставке до продажи на складе WB: сборка, слот, дорога, приёмка")
    wb_supply_cover = ft.TextField(label="Цикл, дней", value=str(wb_defaults["cover_days"]), width=150,
                                   tooltip="На сколько дней продаж должен хватать запас")
    wb_supply_level = ft.TextField(label="Уровень сервиса", value=f"{float(wb_defaults['service_level']):g}", width=170,
                                   tooltip="Вероятность не уйти в ноль за цикл: 0.8 — меньше товара, 0.95 — больше")
    wb_supply_min = ft.TextField(label="На склады WB от, шт./мес.", value=f"{float(wb_defaults['wb_fbw_min_monthly']):g}", width=210,
                                 tooltip="Везти на склады WB только артикулы с прогнозом от стольких штук в месяц")

    def launch_wb_supply_plan() -> None:
        lead, cover = (wb_supply_lead.value or "").strip(), (wb_supply_cover.value or "").strip()
        level = (wb_supply_level.value or "").strip().replace(",", ".")
        minimum = (wb_supply_min.value or "").strip().replace(",", ".")
        error = valid_supply(lead, cover, level)
        try:
            if float(minimum) <= 0:
                raise ValueError
        except ValueError:
            error = error or "Порог для складов WB — положительное число штук в месяц, например 3."
        if error:
            toast(error, True)
            return
        job(WB_SUPPLY_TITLE, [str(PYTHON), str(SCRIPTS / "wb_supply_plan.py"), "--lead-days", lead, "--cover-days", cover,
                              "--service-level", level, "--fbw-min-monthly", minimum])

    def render_wb_supply_results(push_update: bool = True) -> None:
        pal = dash_pal()
        summary = latest_wb_supply_summary()
        if not summary:
            wb_supply_results.controls = [dpanel(ft.Column([
                ft.Icon(ft.Icons.INVENTORY_2_OUTLINED, size=44, color=pal["muted"]),
                dtext("Расчёта пока нет", 18, weight=ft.FontWeight.W_700),
                dtext("Нажмите «Рассчитать для WB» — появятся запас своего склада, поставка на склады WB и сводка вместе с Ozon.", 13, pal["muted"]),
            ], spacing=6, horizontal_alignment=ft.CrossAxisAlignment.CENTER, tight=True), padding=48, alignment=ft.Alignment(0, 0))]
            if push_update and page.controls:
                page.update(wb_supply_results)
            return
        settings = summary["settings"]
        window_start, window_end = summary["window"]
        last_day = (datetime.strptime(window_end, "%Y-%m-%d") - timedelta(days=1)).strftime("%Y-%m-%d")
        season = summary["season"]
        fbw, fbs = summary["fbw"], summary["fbs"]
        season_note = (f"сезон ×{season['factor']:.2f} по истории WB".replace(".", ",") if season["method"] == "season"
                       else "истории мало — без сезонности")
        order_days = summary.get("order_days", 90)  # нет в расчётах до исправления темпа
        bad, good = pal["bad"], pal["good"]
        colors = {"wb_fbw": "#C08CFF" if pal["dark"] else "#8A3FD1"}

        def table(columns: list[tuple[str, int, bool]], rows: list[list[ft.Control]]) -> ft.Control:
            def cell(content: ft.Control, weight: int, right: bool) -> ft.Control:
                return ft.Container(expand=weight, alignment=ft.Alignment(1 if right else -1, 0), content=content)
            lines = [ft.Row([cell(dtext(name, 12, pal["muted"], ft.FontWeight.W_600), w, r) for name, w, r in columns], spacing=12)]
            for values in rows:
                lines.append(ft.Container(
                    padding=ft.Padding.symmetric(vertical=8), data="preserve_excel_fill",
                    border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])),
                    content=ft.Row([cell(v, w, r) for v, (_n, w, r) in zip(values, columns)], spacing=12,
                                   vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ))
            return ft.Column(lines, spacing=0)

        def num(value, color=None, weight=None) -> ft.Control:
            return dtext("—" if not value else str(value), 13, color if value else pal["muted"], weight)

        def money(value) -> ft.Control:
            return dtext(_format_currency(value) if value else "—", 13, None if value else pal["muted"])

        controls: list[ft.Control] = [
            dtext(f"Расчёт от {supply_date(summary['created'], True)} · с {supply_date(summary['created'])} по {supply_date(last_day)} · "
                  f"поставка на WB придёт ~{supply_date(summary['arrival'])} · срок {settings['lead_days']} дн., цикл {settings['cover_days']} дн., "
                  f"уровень сервиса {settings['service_level']:.0%}", 12, pal["muted"]),
            dash_stats_panel([
                dash_stat("Прогноз WB за цикл", f"{summary['window_units']:.0f} шт.",
                          f"{summary['recent_units']:.0f} заказов за {order_days} дн. · {season_note}"),
                dash_stat("На склады WB", f"{fbw['ship']} шт.",
                          f"{fbw['articles']} арт. · {_format_currency(fbw['cost'])}" if fbw["ship"] else "пока везти не нужно"),
                dash_stat("Свой склад для WB", f"{fbs['need']} шт.", f"выставлено на WB {fbs['stock']} шт."),
                dash_stat("Не хватает для WB", f"{fbs['shortage']} шт.",
                          f"{_format_currency(fbs['cost'])} по себестоимости WB" if fbs["shortage"] else "остатка хватит", bad if fbs["shortage"] else good),
            ]),
        ]

        # Свой склад для WB
        idle = [row for row in fbs["rows"] if not row["need"]]
        fbs_rows = [[
            dtext(row["article"], 13, weight=ft.FontWeight.W_700), num(row["orders"]),
            dtext(f"{row['demand']:.1f}".replace(".", ","), 13, pal["muted"]),
            dtext(str(row["need"]), 13, weight=ft.FontWeight.W_700), num(row["stock"]),
            num(row["shortage"], bad, ft.FontWeight.W_700), money(row["cost"]),
        ] for row in fbs["rows"] if row["need"]]
        controls.append(dpanel(ft.Column([
            dtext(f"Свой склад для WB (FBS) — нужно {fbs['need']} шт., есть {fbs['stock']}, не хватает {fbs['shortage']}", 15, weight=ft.FontWeight.W_700),
            dtext(f"Сейчас {summary['own_share']:.0%} заказов WB уходят со своего склада. Сюда входит спрос артикулов, которые на WB не возим, "
                  "и спрос до прихода поставки сверх остатка на складах WB. «Есть» — остаток, выставленный на WB; остатки и поставки Ozon сюда не входят.",
                  12, pal["muted"]),
            table([("Артикул", 2, False), ("Заказов", 2, True), ("Спрос на дни", 2, True), ("Нужно", 2, True),
                   ("Есть", 2, True), ("Не хватает", 2, True), ("Себестоимость", 3, True)], fbs_rows),
            *([dtext(f"Ещё {len(idle)} арт. без спроса на WB на эти дни, на складе по ним {sum(r['stock'] or 0 for r in idle)} шт.: "
                     + ", ".join(r["article"] for r in idle), 12, pal["muted"])] if idle else []),
        ], spacing=10), padding=ft.Padding.only(left=20, right=20, top=16, bottom=12)))

        # Склады WB
        min_monthly = settings.get("wb_fbw_min_monthly", 3)
        if fbw["rows"]:
            fbw_body: ft.Control = table(
                [("Артикул", 2, False), ("В месяц", 2, True), ("На складах WB", 2, True), ("Возвраты в пути", 2, True),
                 ("Цель", 2, True), ("К поставке", 2, True), ("Себестоимость", 3, True)],
                [[dtext(row["article"], 13, weight=ft.FontWeight.W_700), dtext(f"{row['monthly']:.1f}".replace(".", ","), 13),
                  num(row["stock"]), num(row["returns"]), num(row["target"]), num(row["ship"], None, ft.FontWeight.W_700), money(row["cost"])]
                 for row in fbw["rows"]],
            )
        else:
            fbw_body = dtext(f"Ни один артикул не продаётся на WB от {min_monthly:g} шт. в месяц — на склады WB пока везти не нужно, "
                             f"заказы закрывает свой склад. Сейчас на складах WB {fbw['stock']} шт.", 13, pal["muted"])
        controls.append(dpanel(ft.Column([
            dtext(f"Склады WB (FBW) — к поставке {fbw['ship']} шт.", 15, weight=ft.FontWeight.W_700),
            dtext("WB отдаёт остатки только общей суммой «Склад WB», поэтому считается общий объём без разбивки по складам. "
                  f"Возим артикулы с прогнозом от {min_monthly:g} шт. в месяц.", 12, pal["muted"]),
            fbw_body,
        ], spacing=10), padding=ft.Padding.only(left=20, right=20, top=16, bottom=14)))

        # Выкупы по месяцам
        months = summary.get("monthly_sales") or []
        while months and months[0]["units"] == 0:
            months = months[1:]
        top = max((m["units"] for m in months), default=1) or 1
        bars = [ft.Container(
            expand=1, tooltip=f"{m['month'][5:]}.{m['month'][:4]}: {m['units']} шт.",
            content=ft.Column([
                dtext(str(m["units"]), 10, pal["muted"]),
                ft.Container(height=max(2, m["units"] / top * 120), border_radius=ft.BorderRadius.only(top_left=4, top_right=4),
                             bgcolor=colors["wb_fbw"], data="preserve_excel_fill"),
                dtext(f"{m['month'][5:]}.{m['month'][2:4]}", 9, pal["muted"]),
            ], spacing=4, alignment=ft.MainAxisAlignment.END, horizontal_alignment=ft.CrossAxisAlignment.STRETCH),
        ) for m in months]
        season_text = (f"Год назад в окне продаж выкупили {season['window_last_year']:.0f} шт., в базовые {season['base_days']} дн. — "
                       f"{season['base_last_year']:.0f} шт.: коэффициент " + f"{season['factor']:.2f}".replace(".", ",")
                       + (" (истории меньше двух лет — ослаблен вдвое)." if season.get("weight", 1) < 1 else ".")
                       if season["method"] == "season" else "Истории за прошлый год мало — прогноз без сезонности.")
        controls.append(dpanel(ft.Column([
            dtext("Выкупы WB по месяцам", 15, weight=ft.FontWeight.W_700),
            dtext(season_text, 12, pal["muted"]),
            ft.Container(height=160, content=ft.Row(bars, spacing=4, vertical_alignment=ft.CrossAxisAlignment.END)),
        ], spacing=10)))
        wb_supply_results.controls = controls
        if push_update and page.controls:
            page.update(wb_supply_results)

    wb_supply = ft.Column(
        [
            ft.Text("Wildberries", size=13, weight=ft.FontWeight.W_600, color=MUTED),
            ft.Text("Поставки", size=26, weight=ft.FontWeight.W_800, color=TEXT),
            ft.Text("Сколько держать на своём складе под WB и везти ли товар на склады WB. Считается только по данным WB — отдельно от Ozon.", size=13, color=MUTED),
            card(
                "План запасов WB",
                "Свои настройки WB, заказы и остатки WB, себестоимость из wb_costs.xlsx. Excel с расчётом сохраняется в stocks reports.",
                [ft.Row([wb_supply_lead, wb_supply_cover, wb_supply_level, wb_supply_min], wrap=True), ft.Row([
                    ft.Button("Рассчитать для WB", bgcolor=PRIMARY, color="white", on_click=lambda _e: launch_wb_supply_plan()),
                    ft.TextButton("Открыть stocks reports", on_click=lambda _e: reveal_path(ROOT / "stocks reports", "stocks reports")),
                ])],
                tone="#EAF0F5",
            ),
            wb_supply_results,
        ],
        spacing=18,
        horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
    )

    def ai_quick_prompt(prompt_text: str):
        return lambda _e: ai_submit(prompt_text)

    ai_input.on_submit = lambda _e: ai_submit(ai_input.value or "")
    ai_send_button.on_click = lambda _e: ai_submit(ai_input.value or "")
    ai_reset_button.on_click = ai_reset
    ai_refresh_button.on_click = lambda _e: ai_refresh_status()
    ai_new_chat_button.on_click = ai_reset
    ai_model_select.on_change = ai_model_changed
    ai_quick_month_button.on_click = ai_quick_prompt("Как дела в магазине за последний месяц и что улучшить в первую очередь?")
    ai_quick_pressure_button.on_click = ai_quick_prompt("Что сильнее всего давит на прибыль за последние месяцы?")
    ai_quick_brief_button.on_click = ai_quick_prompt("Какие товары срочно пополнить и сколько денег на это нужно?")
    ai_quick_risk_button.on_click = ai_quick_prompt("Какие товары теряют продажи и почему?")

    ai_messages_shell = ft.Container(
        expand=True,
        padding=ft.Padding.only(left=28, right=28, top=18, bottom=12),
        content=ai_messages,
        visible=False,
    )
    ai_composer_card = ft.Container(
        padding=ft.Padding.only(left=18, right=14, top=14, bottom=12),
        border_radius=30,
        bgcolor=ai_surface_tokens()["composer"],
        border=ft.Border.all(1, ai_surface_tokens()["composer_border"]),
        shadow=ft.BoxShadow(
            blur_radius=28,
            spread_radius=0,
            color="#1A000000" if not current_dark["value"] else "#22000000",
            offset=ft.Offset(0, 10),
        ),
        content=ft.Column(
            [
                ft.Row(
                    [
                        ai_input,
                        ai_send_button,
                    ],
                    vertical_alignment=ft.CrossAxisAlignment.END,
                    spacing=10,
                ),
                ft.Row(
                    [
                        ft.Row([ai_busy_ring, ai_busy_label], spacing=8),
                        ft.Row(
                            [
                                ft.Container(content=ai_model_select, width=180),
                                ft.Row([ai_status_dot, ai_status_label], spacing=8),
                            ],
                            spacing=14,
                        ),
                    ],
                    alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                ),
            ],
            spacing=8,
        ),
    )
    ai_composer_shell = ft.Container(
        padding=ft.Padding.only(left=28, right=28, top=10, bottom=24),
        content=ai_centered(ai_composer_card),
    )
    ai_welcome_shell = ft.Container(
        expand=True,
        alignment=ft.Alignment(0, 0),
        content=ft.Column(
            [
                ai_empty_state,
            ],
            spacing=28,
            horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            alignment=ft.MainAxisAlignment.CENTER,
        ),
    )
    ai_body = ft.Column(expand=True, spacing=0)

    def ai_sync_layout() -> None:
        active = ai_has_conversation()
        ai_welcome_shell.visible = not active
        ai_messages_shell.visible = active
        ai_messages_shell.expand = active
        ai_composer_shell.padding = ft.Padding.only(left=24, right=24, top=6, bottom=20 if active else 0)
        ai_body.controls = [ai_messages_shell, ai_composer_shell] if active else [ai_welcome_shell, ai_composer_shell]

    ai_sync_layout()

    ai_header_card = ft.Container(
        padding=ft.Padding.symmetric(horizontal=20, vertical=16),
        border_radius=26,
        bgcolor=ai_surface_tokens()["assistant_bg"],
        border=ft.Border.all(1, ai_surface_tokens()["assistant_border"]),
        content=ft.Row(
            [
                ft.Column(
                    [
                        ai_top_title,
                        ai_top_subtitle,
                    ],
                    spacing=2,
                ),
                ft.Row(
                    [
                        ai_new_chat_button,
                    ],
                    spacing=4,
                ),
            ],
            alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
        ),
    )
    ai_view = ft.Container(
        expand=True,
        bgcolor=ai_surface_tokens()["canvas"],
        padding=0,
        content=ft.Column(
            [
                ft.Container(
                    padding=ft.Padding.only(left=28, right=28, top=20, bottom=6),
                    content=ai_centered(ai_header_card),
                ),
                ai_body,
            ],
            spacing=0,
            expand=True,
        ),
    )

    def open_marketplace_settings(_e=None):
        current = read_settings(ROOT)
        fields = {
            key: ft.TextField(
                label=label, value=current.get(key) or "", password=secret,
                can_reveal_password=secret,
            ) for key, label, secret in FIELDS
        }
        feedback = ft.Text(color="#C62828")

        def close_setup(_e=None):
            setup_dialog.open = False
            page.update()
            page.overlay.remove(setup_dialog)

        def save_setup(_e):
            if job_lock.locked() or discount_dialog_state["busy"] or ai_state["busy"]:
                feedback.value = "Дождитесь завершения текущей операции перед сменой ключей."
                page.update()
                return
            try:
                save_settings(ROOT, {key: field.value or "" for key, field in fields.items()})
            except ValueError as exc:
                feedback.value = str(exc)
                page.update()
                return
            except Exception:
                feedback.value = "Не удалось сохранить настройки. Проверьте доступ к .env и повторите."
                page.update()
                return
            for module_name in ("recommended_prices", "scripts.recommended_prices"):
                module = sys.modules.get(module_name)
                if module is not None:
                    module.refresh_api_credentials()
            close_setup()
            set_status("Настройки сохранены", "Ключи Ozon и WB сохранены. Проверка подключения не выполнялась.", PRIMARY, busy=False)

        setup_dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("Подключение Ozon и Wildberries"),
            content=ft.Column(
                [ft.Text(HELP, size=13),
                 ft.Text("Заполните данные нужных маркетплейсов. Остальные поля можно оставить пустыми. Ключи сохраняются локально в .env.", size=13),
                 *fields.values(), feedback],
                width=600, height=490, scroll=ft.ScrollMode.AUTO, spacing=12,
            ),
            actions=[ft.TextButton("Позже", on_click=close_setup),
                     ft.Button("Сохранить", on_click=save_setup)],
        )
        page.overlay.append(setup_dialog)
        setup_dialog.open = True
        page.update()

    def open_balance_check(_e=None):
        today = date.today()
        month_field = ft.Dropdown(label="Месяц", value=str(today.month), width=200,
                                  options=[ft.dropdown.Option(str(i), title) for i, title in enumerate(MONTHS, 1)])
        year_field = ft.Dropdown(label="Год", value=str(today.year), width=140,
                                 options=[ft.dropdown.Option(str(y)) for y in range(today.year, 2023, -1)])
        result_column = ft.Column(spacing=6)
        status_text = ft.Text("Выберите месяц и нажмите «Проверить».", size=12, color=MUTED)
        checking = {"value": False}
        recent_column = ft.Column(spacing=10)
        recent_status = ft.Text("Загружаем последние 3 месяца…", size=12, color=MUTED)
        recent_checking = {"value": False}

        def render_result(summary):
            if not summary:
                result_column.controls = [ft.Text("Ozon не вернул данных о балансе за этот период.", size=13, color=MUTED)]
                return
            def row(label, value):
                return ft.Text(f"{label}: {_format_currency(value)}", size=13)
            result_column.controls = [
                row("Входящий баланс на начало периода", summary.get("opening_balance")),
                row("Начислено Ozon за период", summary.get("accrued")),
                row("Выплачено Ozon (реальный перевод)", summary.get("payments")),
                row("Исходящий баланс на конец периода", summary.get("closing_balance")),
                row("Комиссия за ранний вывод средств", summary.get("early_payment_fee")),
            ]

        async def check_balance():
            if checking["value"]:
                return
            checking["value"] = True
            check_button.disabled = True
            status_text.value = "Запрашиваем баланс Ozon…"
            result_column.controls = []
            page.update()
            try:
                month, year = int(month_field.value), int(year_field.value)
                summary = await asyncio.to_thread(balance_report.summarize_month, month, year)
                render_result(summary)
                status_text.value = f"{MONTHS[month - 1]} {year} · из раздела Ozon «Финансы → Баланс»."
            except Exception as exc:
                status_text.value = f"Не удалось получить баланс: {exc}"
                result_column.controls = []
            finally:
                checking["value"] = False
                check_button.disabled = False
            page.update()

        def load_one_period(month: int, year: int) -> dict:
            try:
                summary = balance_report.summarize_month(month, year) or {}
            except Exception as exc:
                summary = {"error": str(exc)}
            path = report_path_for_period(month, year)
            metrics, built_at = read_business_metrics(path) if path.exists() else ({}, None)
            return {"month": month, "year": year, "summary": summary, "metrics": metrics, "built_at": built_at}

        async def load_recent_periods():
            if recent_checking["value"]:
                return
            recent_checking["value"] = True
            recent_status.value = "Загружаем последние 3 месяца…"
            recent_column.controls = []
            page.update()
            try:
                periods = iter_periods(today.month, today.year, 3)
                results = await asyncio.to_thread(lambda: [load_one_period(m, y) for m, y in periods])
                cards = []
                for entry in reversed(results):
                    summary = entry["summary"]
                    metrics = entry["metrics"]
                    title = f"{MONTHS[entry['month'] - 1]} {entry['year']}"
                    if summary.get("error"):
                        cards.append(ft.Text(f"{title}: не удалось получить баланс ({summary['error']})", size=13, color=DANGER))
                        continue
                    received = summary.get("payments")
                    cost = metrics.get("Итоговая себестоимость")
                    profit = metrics.get("Чистая прибыль")
                    lines = [
                        ft.Text(title, size=15, weight=ft.FontWeight.W_700, color=TEXT),
                        ft.Text(f"Получено от Ozon (реальный перевод): {_format_currency(received)}", size=13),
                    ]
                    if metrics:
                        lines.append(ft.Text(f"Себестоимость к отправке (по отчёту): {_format_currency(abs(as_float(cost) or 0.0) if cost is not None else None)}", size=13))
                        lines.append(ft.Text(f"Прибыль оставить себе (по отчёту): {_format_currency(profit)}", size=13))
                        pending = metrics.get("Заказы, ожидающие расчёта Ozon (не учтены в прибыли/себестоимости)")
                        if pending:
                            lines.append(ft.Text(f"⚠ Ещё {pending} заказ(ов) не досчитаны Ozon — цифры за месяц могут подрасти позже.", size=12, color=ACCENT))
                    else:
                        lines.append(ft.Text("Отчёт за этот месяц не сформирован — себестоимость и прибыль неизвестны. Сформируйте отчёт на вкладке «Бизнес-сводка».", size=12, color=ACCENT))
                    cards.append(ft.Container(
                        content=ft.Column(lines, spacing=4),
                        padding=12, border_radius=10, bgcolor=semantic_surface("neutral", dark=current_dark["value"]),
                    ))
                recent_column.controls = cards
                recent_status.value = "Готово. «Получено от Ozon» — реальные деньги; себестоимость/прибыль — из уже сформированных месячных отчётов."
            except Exception as exc:
                recent_status.value = f"Не удалось загрузить последние месяцы: {exc}"
            finally:
                recent_checking["value"] = False
            page.update()

        def close_balance_dialog(_e=None):
            balance_dialog.open = False
            page.update()
            page.overlay.remove(balance_dialog)

        check_button = ft.Button("Проверить", on_click=lambda _e: page.run_task(check_balance))
        balance_dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("Баланс Ozon и сверка по месяцам"),
            content=ft.Column(
                [
                    ft.Text("Последние 3 месяца: сколько реально пришло от Ozon и сколько из этого — себестоимость к отправке, а сколько — ваша прибыль.", size=12, color=MUTED),
                    recent_status, recent_column,
                    ft.Divider(),
                    ft.Text("Проверить конкретный месяц отдельно (только баланс Ozon, без отчёта):", size=12, color=MUTED),
                    ft.Row([month_field, year_field, check_button], wrap=True),
                    status_text, result_column,
                ],
                width=560, height=560, scroll=ft.ScrollMode.AUTO, spacing=12,
            ),
            actions=[ft.TextButton("Закрыть", on_click=close_balance_dialog)],
        )
        page.overlay.append(balance_dialog)
        balance_dialog.open = True
        page.update()
        page.run_task(load_recent_periods)

    def data_folder_note() -> str:
        if DATA_ROOT_SOURCE == "env":
            return "Папка задана переменной OZONREPORTX_HOME при запуске — выбор ниже начнёт действовать без неё."
        if DATA_ROOT_SOURCE == "chosen":
            return "Выбрана вами. Программа запомнила её и открывает данные отсюда."
        return "По умолчанию. Можно выбрать другую — например, папку, где уже лежат ваши отчёты и costs.xlsx."

    async def choose_data_folder(_e=None) -> None:
        folder = await asyncio.to_thread(pick_folder, "Папка с данными OzonReportX", ROOT)
        if not folder:
            return
        if not os.access(folder, os.W_OK):
            toast("В эту папку нельзя записывать — выберите другую.", True)
            return
        save_data_folder(folder)
        has_data = any((folder / name).exists() for name in (".env", "costs.xlsx", "reports"))
        toast(f"Папка с данными: {folder}. Перезапустите программу, чтобы открыть её."
              + ("" if has_data else " В ней пока нет данных OzonReportX — программа начнёт с чистого листа."))

    def reset_data_folder(_e=None) -> None:
        save_data_folder(None)
        toast(f"Папка с данными: {DEFAULT_DATA_ROOT}. Перезапустите программу, чтобы открыть её.")

    data_folder_card = card(
        "Папка с данными",
        "Здесь программа хранит отчёты, costs.xlsx, ключи и настройки. Новая папка начинает действовать после перезапуска.",
        [
            ft.Text(str(ROOT), size=12, color=TEXT, selectable=True),
            ft.Text(data_folder_note(), size=12, color=MUTED),
            ft.Button("Выбрать папку…", on_click=choose_data_folder),
            ft.TextButton("Открыть папку", on_click=lambda _e: reveal_path(ROOT, "папка с данными")),
            ft.TextButton("Вернуть «Документы/OzonReportX»", on_click=reset_data_folder,
                          visible=DATA_ROOT_SOURCE == "chosen"),
        ],
        width=320,
    )

    settings = ft.Column(
        [
            ft.Text("Настройки", size=26, weight=ft.FontWeight.W_800, color=TEXT),
            ft.Text("Системные действия и рабочие артефакты проекта.", size=13, color=MUTED),
            ft.ResponsiveRow(
                run_spacing=14,
                spacing=14,
                controls=[
                    card(
                        "Ozon и Wildberries",
                        "Подключение маркетплейсов и настройка ключей API.",
                        [
                            ft.Button("Открыть корень проекта", on_click=lambda _e: reveal_path(ROOT, "корень проекта")),
                            ft.Button("Настроить подключения", on_click=open_marketplace_settings),
                            ft.TextButton("Проверить баланс Ozon", on_click=open_balance_check),
                            ft.Text("Для бизнес-сводки WB нужен доступ токена к категории «Финансы».", size=12, color=MUTED),
                        ],
                        width=320,
                    ),
                    *([data_folder_card] if FROZEN else []),
                    card(
                        "Рабочие файлы",
                        "Быстрый доступ к ключевым папкам и costs.xlsx.",
                        [
                            ft.Button("Открыть costs.xlsx", on_click=lambda _e: reveal_path(ROOT / "costs.xlsx", "costs.xlsx")),
                            ft.TextButton("reports", on_click=lambda _e: reveal_path(ROOT / "reports", "reports")),
                            ft.TextButton("ABC&XYZ reports", on_click=lambda _e: reveal_path(ROOT / "ABC&XYZ reports", "ABC&XYZ reports")),
                        ],
                        width=320,
                    ),
                    card(
                        "Проверка обновлений",
                        "«Проверить» показывает новую версию и что в ней изменилось. «Скачать» открывает страницу новой "
                        "версии: установите её поверх старой, данные в папке с данными сохранятся." if FROZEN else
                        "«Проверить» показывает новую версию и что в ней изменилось. «Установить» делает резервную копию, "
                        "ставит новую версию из GitHub и сохраняет ваши данные; после установки перезапустите приложение.",
                        [
                            ft.Button("Проверить обновления", bgcolor=ACCENT, color="white",
                                      on_click=lambda _e: job("Проверка обновлений", [str(PYTHON), str(SCRIPTS / "_auto_update.py")])),
                            ft.TextButton("Скачать обновление" if FROZEN else "Установить обновление",
                                          on_click=lambda _e: job("Установка обновления",
                                                                  [str(PYTHON), str(SCRIPTS / "_auto_update.py"), "--yes", "--no-restart"])),
                        ],
                        width=320,
                    ),
                ],
            ),
        ],
        spacing=18,
    )

    file_view = ft.Column(
        [
            ft.Text("Файлы", size=26, weight=ft.FontWeight.W_800, color=TEXT),
            ft.Text("Быстрый доступ ко всем результатам работы приложения.", size=13, color=MUTED),
            files_col,
        ],
        spacing=18,
    )

    # ---------------------------------------------------------------- Остатки (общая таблица Ozon + WB)
    stock_state: dict[str, object] = {"data": None, "busy": False, "error": None, "started": False, "filter": "all"}
    stock_subtitle = ft.Text("", size=13, color=MUTED)
    stock_search = ft.TextField(hint_text="Найти артикул", prefix_icon=ft.Icons.SEARCH, dense=True, width=240)
    stock_search.on_change = stock_search.on_submit = lambda _e: render_stock_view()
    stock_refresh_button = ft.Button("Обновить", icon=ft.Icons.REFRESH, bgcolor=PRIMARY, color="white")
    stock_body = ft.Column(spacing=14)

    def stock_load(_e=None) -> None:
        if stock_state["busy"] or stock_overview is None:
            return
        stock_state.update(busy=True, error=None, started=True)
        render_stock_view()

        def worker() -> None:
            try:
                stock_state["data"] = stock_overview.load(ROOT)
            except Exception as exc:
                stock_state["error"] = str(exc)
            finally:
                stock_state["busy"] = False
            ui_queue.put(("call", render_stock_view))

        page.run_thread(worker)

    stock_refresh_button.on_click = stock_load

    def render_stock_view(push_update: bool = True) -> None:
        pal = dash_pal()
        data = stock_state.get("data") or {}
        rows = list(data.get("rows") or [])
        loaded = str(data.get("loaded") or "")
        parts = ["Ozon и Wildberries", "сколько держать на ближайшие 30 дней"]
        if loaded:
            parts.append(f"обновлено {loaded[8:10]}.{loaded[5:7]} в {loaded[11:16]}")
        if data:
            parts.append(f"уровень сервиса {float(data.get('level') or 0):.0%}")
        stock_subtitle.value = " · ".join(parts)
        stock_refresh_button.disabled = bool(stock_state["busy"])

        controls: list[ft.Control] = []
        if stock_overview is None:
            controls.append(dpanel(dtext(f"Модуль остатков не загрузился: {STOCK_IMPORT_ERROR}", 13, pal["bad"])))
        if stock_state["busy"]:
            controls.append(dpanel(ft.Row([
                ft.ProgressRing(width=20, height=20, stroke_width=2.4),
                dtext("Загружаю остатки и историю продаж Ozon и Wildberries… В первый раз это может занять несколько минут: "
                      "история сохраняется и дальше грузится быстро.", 13, pal["muted"], expand=True),
            ], spacing=12)))
        if stock_state.get("error"):
            controls.append(dpanel(dtext(f"Не удалось загрузить остатки: {stock_state['error']}", 13, pal["bad"])))
        for store, message in (data.get("errors") or {}).items():
            controls.append(dpanel(dtext(f"{store} не загрузился — его колонки пустые. {message}", 13, pal["bad"])))
        if rows:
            def total(key: str) -> int:
                return sum(int(row.get(key) or 0) for row in rows)

            def need_stat(label: str, need_key: str, stock_keys: tuple[str, ...], short_key: str) -> ft.Control:
                short = total(short_key)
                have = sum(total(key) for key in stock_keys)
                return dash_stat(label, _format_number(total(need_key)), f"есть {_format_number(have)}, не хватает {_format_number(short)}",
                                 pal["bad"] if short else pal["good"])

            controls.append(dpanel(ft.ResponsiveRow([
                need_stat("Нужно держать FBO", "need_fbo", ("fbo", "fbo_transit"), "short_fbo"),
                need_stat("Нужно держать FBW", "need_fbw", ("fbw", "fbw_transit"), "short_fbw"),
                need_stat("Нужно держать FBS Ozon", "need_fbs_ozon", ("fbs_ozon",), "short_fbs_ozon"),
                need_stat("Нужно держать FBS WB", "need_fbs_wb", ("fbs_wb",), "short_fbs_wb"),
            ], spacing=24, run_spacing=12), padding=ft.Padding.symmetric(horizontal=24, vertical=16)))

            mode = str(stock_state["filter"])
            shown = stock_overview.filter_rows(rows, stock_search.value or "", mode)
            toolbar = ft.Row([
                stock_search,
                *[s_chip(label, mode == key, lambda _e, k=key: (stock_state.update(filter=k), render_stock_view()))
                  for key, label in stock_overview.STOCK_FILTERS.items()],
                dtext(f"Показано {len(shown)} из {len(rows)}", 12, pal["muted"]),
            ], wrap=True, spacing=10, run_spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER)

            columns = [("Артикул", 3), ("FBS Ozon", 2), ("FBS WB", 2), ("FBO", 2), ("FBW", 2), ("Нужно держать FBO", 3), ("Нужно держать FBW", 3),
                       ("Нужно держать FBS Ozon", 3), ("Нужно держать FBS WB", 3)]

            def cell(content: ft.Control, weight: int, right: bool = True, tooltip: str | None = None) -> ft.Control:
                return ft.Container(expand=weight, alignment=ft.Alignment(1 if right else -1, 0), tooltip=tooltip, content=content)

            def stock_text(value, transit: int = 0, need=None) -> ft.Control:
                if value is None:
                    return dtext("—", 13, pal["muted"])
                text = _format_number(value) + (f" +{transit}" if transit else "")
                if value or transit:
                    return dtext(text, 13)
                # Пусто на складе, где товар нужен, — красным; где он и не нужен — просто серым.
                return dtext(text, 13, pal["bad"], ft.FontWeight.W_700) if need else dtext(text, 13, pal["muted"])

            def need_text(need, short: int) -> ft.Control:
                if need is None:
                    return dtext("—", 13, pal["muted"])
                return dtext(str(need) + (f"  (−{short})" if short else ""), 13, pal["bad"] if short else (None if need else pal["muted"]),
                             ft.FontWeight.W_700 if short else None)

            table: list[ft.Control] = [ft.Container(
                padding=ft.Padding.only(bottom=8),
                content=ft.Row([cell(dtext(name, 12, pal["muted"], ft.FontWeight.W_600, text_align=ft.TextAlign.RIGHT if index else None), w, bool(index))
                                for index, (name, w) in enumerate(columns)], spacing=14, vertical_alignment=ft.CrossAxisAlignment.END),
            )]
            if not shown:
                table.append(ft.Container(padding=ft.Padding.symmetric(vertical=12),
                                          content=dtext("Под фильтр ничего не подходит.", 13, pal["muted"])))
            for row in shown:
                table.append(ft.Container(
                    padding=ft.Padding.symmetric(vertical=9), data="preserve_excel_fill",
                    border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])),
                    content=ft.Row([
                        cell(dtext(row["article"], 13, weight=ft.FontWeight.W_700), 3, right=False),
                        cell(stock_text(row["fbs_ozon"], need=row["need_fbs_ozon"]), 2),
                        cell(stock_text(row["fbs_wb"], need=row["need_fbs_wb"]), 2),
                        cell(stock_text(row["fbo"], row["fbo_transit"], row["need_fbo"]), 2,
                             tooltip=f"В пути и в заявках на склады Ozon: {row['fbo_transit']} шт." if row["fbo_transit"] else None),
                        cell(stock_text(row["fbw"], row["fbw_transit"], row["need_fbw"]), 2,
                             tooltip=f"Возвраты в пути на склады WB: {row['fbw_transit']} шт." if row["fbw_transit"] else None),
                        cell(need_text(row["need_fbo"], row["short_fbo"]), 3),
                        cell(need_text(row["need_fbw"], row["short_fbw"]), 3),
                        cell(need_text(row["need_fbs_ozon"], row["short_fbs_ozon"]), 3,
                             tooltip=f"Спрос Ozon на своём складе за 30 дней: {row['fbs_demand_ozon']:g} шт." if row["fbs_demand_ozon"] else None),
                        cell(need_text(row["need_fbs_wb"], row["short_fbs_wb"]), 3,
                             tooltip=f"Спрос WB на своём складе за 30 дней: {row['fbs_demand_wb']:g} шт." if row["fbs_demand_wb"] else None),
                    ], spacing=14, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ))
            notes = [dtext("«Нужно держать» считается так же, как в «Поставках», только на ближайшие 30 дней без срока поставки. "
                           "FBO и FBW — склады маркетплейсов, FBS — ваш склад: туда идёт спрос, который не покроют FBO и FBW. "
                           "На своём складе товар зарезервирован отдельно под Ozon и под WB, поэтому FBS Ozon и FBS WB "
                           "считаются раздельно: товар WB не закрывает нехватку Ozon и наоборот. "
                           "Красным — не хватает с учётом товара в пути.", 12, pal["muted"])]
            notes += [dtext(note, 12, pal["muted"]) for note in data.get("notes") or []]
            controls.append(dpanel(ft.Column([
                dtext("Товары", 15, weight=ft.FontWeight.W_700),
                *notes,
                ft.Container(height=4),
                toolbar,
                ft.Container(height=4),
                ft.Column(table, spacing=0),
            ], spacing=4), padding=ft.Padding.only(left=20, right=20, top=14, bottom=10)))
        elif not stock_state["busy"] and not stock_state.get("error") and stock_state["started"] and not data.get("errors"):
            controls.append(dpanel(dtext("Нет артикулов с остатками или продажами.", 13, pal["muted"])))
        stock_body.controls = controls
        if push_update and page.controls and host.controls and host.controls[0] is stock_view:
            page.update(stock_view)

    stock_view = ft.Column(
        [
            ft.ResponsiveRow(
                [
                    ft.Container(col={"xs": 12, "lg": 8},
                                 content=ft.Column([ft.Text("Остатки", size=26, weight=ft.FontWeight.W_800, color=TEXT), stock_subtitle], spacing=2)),
                    ft.Container(col={"xs": 12, "lg": 4}, alignment=ft.Alignment(1, 0), content=stock_refresh_button),
                ],
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
                run_spacing=12,
            ),
            stock_body,
        ],
        spacing=14,
    )

    # ---------------------------------------------------------------- Поиск и ключевые слова
    SEARCH_COLLECT_TITLE = "Сбор поисковой аналитики"
    SEARCH_TABS = [("overview", "Обзор"), ("queries", "Запросы"), ("products", "Товары"), ("positions", "Позиции")]
    SEARCH_PAGE = 40
    search_state: dict[str, object] = {"data": None, "busy": False, "error": None, "started": False, "tab": "overview",
                                       "filter": "", "targets_only": False, "open": set(), "cluster_rows": {}, "limit": SEARCH_PAGE,
                                       "sku": None, "product": None, "product_busy": False}
    search_subtitle = ft.Text("", size=13, color=MUTED)
    search_refresh_button = ft.Button("Обновить данные", icon=ft.Icons.REFRESH, bgcolor=PRIMARY, color="white",
                                      tooltip="Загрузить из Ozon свежую поисковую аналитику, карточки и продажи (1–3 минуты)")
    search_body = ft.Column(spacing=14)
    search_filter = ft.TextField(hint_text="Найти запрос или кластер", prefix_icon=ft.Icons.SEARCH, dense=True, width=320)

    def search_load(_e=None) -> None:
        if search_state["busy"] or search_view is None:
            return
        search_state.update(busy=True, error=None, started=True, cluster_rows={})
        render_search_view()

        def worker() -> None:
            try:
                search_state["data"] = search_view.load()
                if search_state["sku"] is not None:
                    search_state["product"] = search_view.load_product(int(search_state["sku"]))
            except Exception as exc:
                search_state["error"] = str(exc)
            finally:
                search_state["busy"] = False
            ui_queue.put(("call", render_search_view))

        page.run_thread(worker)

    def search_collect(_e=None) -> None:
        job(SEARCH_COLLECT_TITLE, [str(PYTHON), str(SCRIPTS / "search_collect.py")])

    search_refresh_button.on_click = search_collect

    def search_set_tab(tab: str) -> None:
        search_state["tab"] = tab
        render_search_view()

    def search_open_product(sku: int | None) -> None:
        search_state.update(sku=sku, product=None, product_busy=sku is not None, tab="products")
        render_search_view()
        if sku is None:
            return

        def worker() -> None:
            try:
                search_state["product"] = search_view.load_product(int(sku))
            except Exception as exc:
                search_state["error"] = f"Не удалось открыть товар: {exc}"
            finally:
                search_state["product_busy"] = False
            ui_queue.put(("call", render_search_view))

        page.run_thread(worker)

    def search_toggle_cluster(cluster_id: int) -> None:
        opened = search_state["open"]
        if cluster_id in opened:
            opened.discard(cluster_id)
            render_search_view()
            return
        opened.add(cluster_id)
        render_search_view()

        def worker() -> None:
            try:
                search_state["cluster_rows"][cluster_id] = search_view.load_cluster(cluster_id)
            except Exception as exc:
                search_state["error"] = f"Не удалось открыть кластер: {exc}"
            ui_queue.put(("call", render_search_view))

        page.run_thread(worker)

    def search_toggle_target(keyword_id: int, target: bool) -> None:
        def worker() -> None:
            try:
                search_view.toggle_target(keyword_id, target)
                clusters = {c["id"]: c for c in (search_state.get("data") or {}).get("clusters") or []}
                for cluster_id, rows in search_state["cluster_rows"].items():
                    for row in rows:
                        if row["id"] == keyword_id:
                            row["target"] = target
                    if cluster_id in clusters:
                        clusters[cluster_id]["target"] = any(row["target"] for row in rows)
                product = search_state.get("product")
                if isinstance(product, dict):
                    for row in product["queries"]:
                        if row["id"] == keyword_id:
                            row["target"] = target
            except Exception as exc:
                search_state["error"] = f"Не удалось изменить целевой запрос: {exc}"
            ui_queue.put(("call", render_search_view))

        page.run_thread(worker)

    def search_filter_changed(_e=None) -> None:
        search_state["filter"] = (search_filter.value or "").strip().lower()
        search_state["limit"] = SEARCH_PAGE
        render_search_view()

    search_filter.on_submit = search_filter_changed
    search_filter.on_change = search_filter_changed

    def s_change(value: float | None, inverse: bool = False) -> tuple[str, str | None]:
        """Изменение к прошлой неделе: текст и цвет (рост — зелёный, падение — красный; мелкие колебания — серые)."""
        pal = dash_pal()
        if value is None:
            return "—", pal["muted"]
        text = f"{value * 100:+.0f}%".replace("-", "−")
        if abs(value) < 0.05:
            return text, pal["muted"]
        good = value > 0 if not inverse else value < 0
        return text, pal["good"] if good else pal["bad"]

    def s_cell(content: ft.Control, weight: int, right: bool = True, tooltip: str | None = None) -> ft.Control:
        return ft.Container(expand=weight, alignment=ft.Alignment(1 if right else -1, 0), tooltip=tooltip, content=content)

    def s_header(columns: list[tuple]) -> ft.Control:
        """Заголовок таблицы: (название, вес, подсказка[, влево]). Первая колонка и текстовые (4-й элемент True) — влево."""
        pal = dash_pal()
        cells = []
        for index, column in enumerate(columns):
            name, w, tip = column[:3]
            right = bool(index) and not (len(column) > 3 and column[3])
            cells.append(s_cell(dtext(name, 12, pal["muted"], ft.FontWeight.W_600, text_align=ft.TextAlign.RIGHT if right else None), w, right, tip))
        return ft.Container(padding=ft.Padding.only(bottom=8), content=ft.Row(cells, spacing=12, vertical_alignment=ft.CrossAxisAlignment.END))

    def s_row(cells: list[ft.Control], on_click=None, highlight: bool = False) -> ft.Control:
        pal = dash_pal()
        return ft.Container(
            padding=ft.Padding.symmetric(vertical=9, horizontal=4), data="preserve_excel_fill", on_click=on_click, ink=on_click is not None,
            bgcolor=semantic_surface("warning", dark=pal["dark"]) if highlight else None, border_radius=8 if highlight else None,
            border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])),
            content=ft.Row(cells, spacing=12, vertical_alignment=ft.CrossAxisAlignment.CENTER))

    def s_chip(label: str, selected: bool, on_click) -> ft.Control:
        pal = dash_pal()
        return ft.Container(
            padding=ft.Padding.symmetric(horizontal=14, vertical=8), border_radius=999, ink=True, data="preserve_excel_fill",
            bgcolor=semantic_surface("positive", dark=pal["dark"]) if selected else None,
            border=ft.Border.all(2 if selected else 1, PRIMARY if selected else pal["border"]),
            on_click=on_click, content=dtext(label, 12, weight=ft.FontWeight.W_700))

    def s_star(keyword_id: int, target: bool) -> ft.Control:
        pal = dash_pal()
        return ft.IconButton(
            icon=ft.Icons.STAR if target else ft.Icons.STAR_BORDER, icon_size=18, icon_color=pal["warn"] if target else pal["muted"],
            tooltip="Целевой запрос: парсер выдачи замеряет его первым. Нажмите, чтобы " + ("снять отметку." if target else "отметить."),
            on_click=lambda _e, k=keyword_id, t=target: search_toggle_target(k, not t))

    def s_position(value: float | None, measured: bool) -> ft.Control:
        pal = dash_pal()
        if not measured:
            return dtext("—", 13, pal["muted"])
        if value is None:
            return dtext("нет", 13, pal["bad"], tooltip="Не найден в просмотренной глубине выдачи")
        color = pal["good"] if value <= 10 else (pal["warn"] if value <= 50 else pal["bad"])
        return dtext(f"{value:.0f}", 13, color, ft.FontWeight.W_700)

    def s_bars(title: str, caption: str, points: list[dict], key: str, color: str, unit: str) -> ft.Control:
        """Столбики по дням (как «Заказы по неделям» в «Анализе»): наведите — точное значение."""
        pal = dash_pal()
        top = max((p[key] for p in points), default=0) or 1
        bars = []
        for p in points:
            day = datetime.strptime(p["date"], "%Y-%m-%d")
            bars.append(ft.Container(
                expand=1, tooltip=f"{day:%d.%m}: {_format_number(p[key])} {unit}",
                content=ft.Column([
                    ft.Container(height=max(2, p[key] / top * 110), bgcolor=color, border_radius=ft.BorderRadius.only(top_left=3, top_right=3),
                                 data="preserve_excel_fill"),
                    dtext(f"{day:%d}", 9, pal["muted"], text_align=ft.TextAlign.CENTER),
                ], spacing=2, alignment=ft.MainAxisAlignment.END, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)))
        return ft.Column([dtext(title, 14, weight=ft.FontWeight.W_700), dtext(caption, 12, pal["muted"]),
                          ft.Container(height=135, content=ft.Row(bars, spacing=3, vertical_alignment=ft.CrossAxisAlignment.END))], spacing=4)

    def s_per_100k(value: float | None) -> str:
        """Заказов на 100 тыс. искавших (на входе — на 1000): у товара на 10 тыс. почти всегда «0,0», сравнивать нечего."""
        return "—" if value is None else f"{value * 100:.1f}".replace(".", ",")

    def search_overview(data: dict) -> list[ft.Control]:
        pal = dash_pal()
        k = data["kpis"]
        w7 = data["windows"]["7"]
        period = f"{w7[0][8:10]}.{w7[0][5:7]}–{w7[1][8:10]}.{w7[1][5:7]}"

        def stat(label: str, value: str, key: str, extra: str = "", tooltip: str = "") -> ft.Control:
            text, color = s_change(k[key]["change"])
            control = dash_stat(label, value, f"{text} к прошлой неделе{extra}", None)
            control.content.controls[2].color = color
            control.tooltip = tooltip or None
            return control

        controls: list[ft.Control] = [dpanel(ft.ResponsiveRow([
            stat("Спрос в поиске", _format_number(k["searches"]["value"]), "searches",
                 tooltip=f"Сколько раз за {period} искали запросы, по которым Ozon показывал ваши товары (по каждому запросу — максимум по товарам)."),
            stat("Запросов с вашими товарами", _format_number(k["queries"]["value"]), "queries",
                 tooltip="Разных поисковых запросов, по которым Ozon показывал товары магазина. Больше запросов — шире охват."),
            stat("Заказано", _format_number(k["orders"]["value"]), "orders", f" · {_format_number(round(k['revenue']['value']))} ₽",
                 tooltip="Заказано штук за те же дни (все источники: поиск, реклама, рекомендации)."),
            stat("Заказов на 100 тыс. поисков", s_per_100k(k["per_1000"]["value"]), "per_1000",
                 tooltip="Конверсия спроса в заказы. Падает при неизменном спросе — проблема в карточке, цене или наличии, а не в поиске."),
        ], spacing=24, run_spacing=12), padding=ft.Padding.symmetric(horizontal=24, vertical=16))]

        icons = {"conversion": (ft.Icons.WARNING_AMBER_ROUNDED, pal["bad"]), "demand": (ft.Icons.TRENDING_DOWN, pal["bad"]),
                 "no_sales": (ft.Icons.WARNING_AMBER_ROUNDED, pal["warn"]),
                 "gap": (ft.Icons.TEXT_FIELDS, pal["accent"]), "rising": (ft.Icons.TRENDING_UP, pal["good"]),
                 "drop": (ft.Icons.TRENDING_DOWN, pal["warn"])}
        lines: list[ft.Control] = []
        for item in data["actions"]:
            icon, color = icons.get(item["kind"], (ft.Icons.LIGHTBULB_OUTLINE, pal["muted"]))
            offers = item.get("offers") or []
            chips = [ft.Container(padding=ft.Padding.symmetric(horizontal=8, vertical=2), border_radius=999, ink=True,
                                  border=ft.Border.all(1, pal["border"]), data="preserve_excel_fill", tooltip="Открыть товар",
                                  on_click=lambda _e, o=offer: search_open_offer(o), content=dtext(offer, 11, weight=ft.FontWeight.W_600))
                     for offer in offers[:8]]
            if len(offers) > 8:
                chips.append(dtext(f"и ещё {len(offers) - 8}", 11, pal["muted"]))
            lines.append(ft.Container(
                padding=ft.Padding.symmetric(vertical=8), border=ft.Border.only(top=ft.BorderSide(1, pal["separator"])) if lines else None,
                data="preserve_excel_fill",
                content=ft.Row([ft.Icon(icon, size=18, color=color),
                                ft.Column([dtext(item["text"], 13), *([ft.Row(chips, wrap=True, spacing=4, run_spacing=4)] if chips else [])],
                                          spacing=4, expand=True)],
                               spacing=10, vertical_alignment=ft.CrossAxisAlignment.START)))
        if not lines:
            lines.append(dtext("Явных проблем не видно: спрос и заказы стабильны, главные слова запросов есть в названиях.", 13, pal["muted"]))
        controls.append(dpanel(ft.Column([dtext("Что сделать в первую очередь", 15, weight=ft.FontWeight.W_700),
                                           dtext(f"По данным за {period}; важнее то, что выше. Нажмите на артикул, чтобы открыть товар.", 12, pal["muted"]),
                                           ft.Container(height=4), *lines], spacing=2),
                               padding=ft.Padding.only(left=20, right=20, top=14, bottom=12)))
        controls.append(dpanel(ft.ResponsiveRow([
            ft.Container(col={"xs": 12, "lg": 6}, content=s_bars("Спрос по дням", "Поисков по запросам с вашими товарами", data["daily"], "searches",
                                                                pal["accent"], "поисков")),
            ft.Container(col={"xs": 12, "lg": 6}, content=s_bars("Заказы по дням", "Заказано штук, все источники", data["daily"], "orders",
                                                                pal["good"], "шт.")),
        ], spacing=24, run_spacing=16), padding=ft.Padding.only(left=20, right=20, top=14, bottom=12)))
        return controls

    def search_open_offer(offer: str) -> None:
        data = search_state.get("data") or {}
        for p in data.get("products") or []:
            if str(p["offer_id"]) == str(offer) or str(p["sku"]) == str(offer):
                search_open_product(p["sku"])
                return

    def search_queries(data: dict) -> list[ft.Control]:
        pal = dash_pal()
        text = str(search_state["filter"])
        rows = [c for c in data["clusters"] if (not text or text in c["name"].lower()) and (not search_state["targets_only"] or c["target"])]
        toolbar = ft.Row([
            search_filter,
            s_chip("Только целевые ★", bool(search_state["targets_only"]),
                   lambda _e: (search_state.update(targets_only=not search_state["targets_only"], limit=SEARCH_PAGE), render_search_view())),
            dtext(f"Кластеров: {len(rows)}", 12, pal["muted"]),
        ], wrap=True, spacing=10, vertical_alignment=ft.CrossAxisAlignment.CENTER)
        columns = [("Кластер запросов", 5, None),
                   ("Спрос, 28 дн.", 2, "Сколько раз искали запросы кластера за 28 дней (по каждому запросу — максимум по товарам)"),
                   ("Неделя", 2, "Спрос последних 7 дней к предыдущим 7. «—» — мало данных: Ozon отдаёт по товару только 15 главных "
                                 "запросов в день, и редкий запрос был в данных меньше 4 дней из 7"),
                   ("Запросов", 1, "Сколько формулировок в кластере"),
                   ("Товаров", 1, "Сколько ваших товаров Ozon показывал по этим запросам"),
                   ("В названии", 2, "У скольких из этих товаров все слова кластера есть в названии"),
                   ("Лучшая позиция", 2, "Медиана позиции в выдаче за 3 дня (лучший из ваших товаров); «—» — не замерялась")]
        table: list[ft.Control] = [s_header(columns)]
        for cluster in rows[: int(search_state["limit"])]:
            trend, trend_color = s_change(cluster["trend"])
            opened = cluster["id"] in search_state["open"]
            weak = cluster["products"] and cluster["in_title"] == 0 and cluster["demand_28"] >= 1000
            table.append(s_row([
                s_cell(ft.Row([ft.Icon(ft.Icons.EXPAND_LESS if opened else ft.Icons.EXPAND_MORE, size=16, color=pal["muted"]),
                               ft.Icon(ft.Icons.STAR, size=14, color=pal["warn"], visible=bool(cluster["target"])),
                               dtext(cluster["name"], 13, weight=ft.FontWeight.W_600, expand=True, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)],
                              spacing=4), 5, right=False),
                s_cell(dtext(_format_number(round(cluster["demand_28"])), 13), 2),
                s_cell(dtext(trend, 13, trend_color), 2),
                s_cell(dtext(str(cluster["queries"]), 13), 1),
                s_cell(dtext(str(cluster["products"]), 13), 1),
                s_cell(dtext(f"{cluster['in_title']} из {cluster['products']}" if cluster["products"] else "—", 13,
                             pal["bad"] if weak else None, ft.FontWeight.W_700 if weak else None), 2,
                       tooltip="Частый кластер, но ни в одном названии его слов нет целиком" if weak else None),
                s_cell(s_position(cluster["best_position"], cluster["best_position"] is not None), 2),
            ], on_click=lambda _e, cid=cluster["id"]: search_toggle_cluster(cid)))
            if opened:
                members = search_state["cluster_rows"].get(cluster["id"])
                if members is None:
                    table.append(ft.Container(padding=ft.Padding.only(left=28, top=6, bottom=6), content=ft.ProgressRing(width=16, height=16, stroke_width=2)))
                    continue
                for q in members:
                    q_trend, q_color = s_change(q["trend"])
                    table.append(ft.Container(
                        padding=ft.Padding.only(left=28, top=2, bottom=2), data="preserve_excel_fill",
                        content=ft.Row([
                            s_cell(ft.Row([s_star(q["id"], q["target"]), dtext(q["query"], 12, expand=True, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS),
                                           dtext("вручную" if q["source"] in ("manual", "csv") else "", 11, pal["muted"])], spacing=2), 5, right=False),
                            s_cell(dtext(_format_number(round(q["demand_28"])), 12), 2),
                            s_cell(dtext(q_trend, 12, q_color), 2),
                            s_cell(dtext(", ".join(q["offers"][:4]) + (f" +{len(q['offers']) - 4}" if len(q["offers"]) > 4 else ""), 11, pal["muted"],
                                         max_lines=1, overflow=ft.TextOverflow.ELLIPSIS), 6, tooltip=", ".join(q["offers"]) or None),
                        ], spacing=12, vertical_alignment=ft.CrossAxisAlignment.CENTER)))
        if len(rows) > int(search_state["limit"]):
            table.append(ft.Container(padding=ft.Padding.only(top=8), content=ft.TextButton(
                f"Показать ещё (осталось {len(rows) - int(search_state['limit'])})",
                on_click=lambda _e: (search_state.update(limit=int(search_state["limit"]) + SEARCH_PAGE), render_search_view()))))
        return [dpanel(ft.Column([
            dtext("Семантическое ядро", 15, weight=ft.FontWeight.W_700),
            dtext("Запросы, по которым Ozon показывал ваши товары, склеены по смыслу: разные формы и порядок слов — один кластер. "
                  "Нажмите на кластер, чтобы увидеть формулировки и отметить целевые ★. Свои запросы и CSV добавляются командой "
                  "scripts/search_semantics.py (см. README).", 12, pal["muted"]),
            ft.Container(height=6), toolbar, ft.Container(height=6), ft.Column(table, spacing=0),
        ], spacing=4), padding=ft.Padding.only(left=20, right=20, top=14, bottom=10))]

    def search_products(data: dict) -> list[ft.Control]:
        pal = dash_pal()
        if search_state["sku"] is not None:
            return search_product_detail()
        columns = [("Артикул", 2, None), ("Название", 6, None, True),
                   ("Спрос, 7 дн.", 2, "Сколько покупателей искали запросы, по которым Ozon показывал этот товар"),
                   ("Неделя", 2, "Спрос последних 7 дней к предыдущим 7"),
                   ("Заказы, 28 дн.", 2, None),
                   ("На 100 тыс.", 2, "Заказов на 100 тыс. искавших за 28 дней — сравнивайте товары между собой"),
                   ("В названии", 2, "Доля спроса товара, все слова которого есть в названии"),
                   ("Пробелы", 2, "Частые кластеры, слов которых нет в названии")]
        table: list[ft.Control] = [s_header(columns)]
        for p in data["products"]:
            trend, trend_color = s_change(p["trend"])
            share = p["title_coverage"]
            table.append(s_row([
                s_cell(dtext(str(p["offer_id"] or p["sku"]), 13, weight=ft.FontWeight.W_700), 2, right=False),
                s_cell(dtext(p["name"] or "", 12, pal["muted"], max_lines=1, overflow=ft.TextOverflow.ELLIPSIS), 6, right=False, tooltip=p["name"]),
                s_cell(dtext(_format_number(round(p["demand_7"])), 13), 2),
                s_cell(dtext(trend, 13, trend_color), 2),
                s_cell(dtext(_format_number(round(p["orders_28"])), 13, pal["bad"] if p["no_sales"] else None,
                             ft.FontWeight.W_700 if p["no_sales"] else None), 2,
                       tooltip="Ищут, но не заказывают: большой спрос и ни одного заказа за 28 дней" if p["no_sales"] else None),
                s_cell(dtext(s_per_100k(p["per_1000"]), 13), 2),
                s_cell(dtext("—" if share is None else f"{share * 100:.0f}%", 13, pal["bad"] if share is not None and share < 0.5 else None), 2),
                s_cell(dtext(str(p["gaps"]) if p["gaps"] else "—", 13, pal["warn"] if p["gaps"] else pal["muted"], ft.FontWeight.W_700), 2),
            ], on_click=lambda _e, sku=p["sku"]: search_open_product(sku)))
        return [dpanel(ft.Column([
            dtext("Товары в поиске", 15, weight=ft.FontWeight.W_700),
            dtext("Нажмите на товар: по каким запросам его находят, каких слов не хватает в карточке и что в ней менялось.", 12, pal["muted"]),
            ft.Container(height=6), ft.Column(table, spacing=0),
        ], spacing=4), padding=ft.Padding.only(left=20, right=20, top=14, bottom=10))]

    def search_product_detail() -> list[ft.Control]:
        pal = dash_pal()
        back = ft.TextButton("Все товары", icon=ft.Icons.ARROW_BACK, on_click=lambda _e: search_open_product(None))
        product = search_state.get("product")
        if search_state["product_busy"] or not isinstance(product, dict):
            return [back, dpanel(ft.Row([ft.ProgressRing(width=20, height=20, stroke_width=2.4), dtext("Открываю товар…", 13, pal["muted"])], spacing=12))]
        summary = next((p for p in (search_state.get("data") or {}).get("products") or [] if p["sku"] == product["sku"]), {})
        trend, trend_color = s_change(summary.get("trend"))
        share = product["title_coverage"]
        stats = [
            dash_stat("Спрос, 7 дней", _format_number(round(summary.get("demand_7") or 0)), f"{trend} к прошлой неделе"),
            dash_stat("Заказы, 28 дней", _format_number(round(summary.get("orders_28") or 0)), f"на 100 тыс. искавших: {s_per_100k(summary.get('per_1000'))}"),
            dash_stat("Название покрывает", "—" if share is None else f"{share * 100:.0f}%", "спроса товара",
                      pal["bad"] if share is not None and share < 0.5 else None),
            dash_stat("Карточка", f"{product['card']['images']} фото", f"описание {product['card']['description_length']} знаков · "
                      f"{product['card']['attributes']} характеристик"),
        ]
        stats[0].content.controls[2].color = trend_color
        controls: list[ft.Control] = [
            back,
            ft.Column([dtext(str(product["offer_id"] or product["sku"]), 20, weight=ft.FontWeight.W_800),
                       dtext(product["name"] or "", 13, pal["muted"])], spacing=2),
            dpanel(ft.ResponsiveRow(stats, spacing=24, run_spacing=12), padding=ft.Padding.symmetric(horizontal=24, vertical=16)),
        ]
        if summary.get("no_sales"):
            controls.append(dpanel(ft.Row([ft.Icon(ft.Icons.WARNING_AMBER_ROUNDED, size=18, color=pal["warn"]), dtext(
                "Товар ищут, но не заказывают: за 28 дней ни одного заказа. Слова в названии этого не исправят. Сначала проверьте, "
                "видят ли его (позиция в выдаче), затем — цену и индекс цен, главное фото, рейтинг и отзывы, срок доставки.", 13, expand=True)],
                spacing=10, vertical_alignment=ft.CrossAxisAlignment.START), padding=ft.Padding.symmetric(horizontal=20, vertical=12)))
        if product["gaps"]:
            lines = []
            for gap in product["gaps"][:6]:
                text = (f"Добавьте «{', '.join(gap['missing'])}» — этих слов нет нигде в карточке" if gap["kind"] == "missing"
                        else "Слова есть в описании или характеристиках — перенесите в название")
                lines.append(ft.Row([ft.Icon(ft.Icons.TEXT_FIELDS, size=16, color=pal["accent"]),
                                     dtext(f"«{gap['cluster']}»", 13, weight=ft.FontWeight.W_700),
                                     dtext(f"{_format_number(round(gap['demand']))} искавших за 28 дн. · {text}", 13, pal["muted"], expand=True)],
                                    spacing=8, vertical_alignment=ft.CrossAxisAlignment.START))
            controls.append(dpanel(ft.Column([dtext("Пробелы в карточке", 15, weight=ft.FontWeight.W_700),
                                              dtext("Кластеры, по которым товар ищут, но их слов нет в названии. Добавляйте, только если товар "
                                                    "действительно подходит под запрос: нерелевантные слова снижают конверсию.", 12, pal["muted"]),
                                              ft.Container(height=4), *lines], spacing=6),
                                   padding=ft.Padding.only(left=20, right=20, top=14, bottom=12)))
        mark = lambda flag, share_value: dtext("✓", 14, pal["good"], ft.FontWeight.W_800) if flag else dtext(  # noqa: E731
            f"{share_value * 100:.0f}%" if share_value else "—", 12, pal["muted"])
        cov_table: list[ft.Control] = [s_header([("Кластер", 5, None, True), ("Спрос товара, 28 дн.", 2, "Сколько покупателей искали запросы кластера и видели этот товар"),
                                                 ("Название", 1, "✓ — все слова кластера есть в названии; % — доля слов"),
                                                 ("Описание", 1, None), ("Характеристики", 2, None), ("Нет нигде", 3, "Слова кластера, которых нет ни в названии, ни в описании, ни в характеристиках")])]
        for row in product["coverage"][:25]:
            weak = not row["in_title"] and row["demand"] >= 500
            cov_table.append(s_row([
                s_cell(dtext(row["cluster"], 13, weight=ft.FontWeight.W_600, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS), 5, right=False),
                s_cell(dtext(_format_number(round(row["demand"])), 13), 2),
                s_cell(mark(row["in_title"], row["shares"]["title"]), 1),
                s_cell(mark(row["in_description"], row["shares"]["description"]), 1),
                s_cell(mark(row["in_attributes"], row["shares"]["attributes"]), 2),
                s_cell(dtext(", ".join(row["missing"]) or "—", 12, pal["bad"] if row["missing"] else pal["muted"]), 3),
            ], highlight=weak))
        controls.append(dpanel(ft.Column([dtext("Покрытие запросов карточкой", 15, weight=ft.FontWeight.W_700),
                                          dtext("Сверху — то, по чему товар ищут чаще всего. Подсвечены частые кластеры, которых нет в названии.", 12, pal["muted"]),
                                          ft.Container(height=6), ft.Column(cov_table, spacing=0)], spacing=4),
                               padding=ft.Padding.only(left=20, right=20, top=14, bottom=10)))
        q_table: list[ft.Control] = [s_header([("Запрос", 6, None), ("Спрос товара, 28 дн.", 2, None), ("Неделя", 2, None),
                                               ("Позиция", 2, "Медиана позиции в выдаче за 3 дня по всем регионам")])]
        for q in product["queries"]:
            q_trend, q_color = s_change(q["trend"])
            q_table.append(s_row([
                s_cell(ft.Row([s_star(q["id"], q["target"]), dtext(q["query"], 13, expand=True, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS)], spacing=2),
                       6, right=False),
                s_cell(dtext(_format_number(round(q["demand_28"])), 13), 2),
                s_cell(dtext(q_trend, 13, q_color), 2),
                s_cell(s_position(q["position"], q["measured"]), 2),
            ]))
        controls.append(dpanel(ft.Column([dtext("Главные запросы товара", 15, weight=ft.FontWeight.W_700), ft.Container(height=6),
                                          ft.Column(q_table, spacing=0)], spacing=4),
                               padding=ft.Padding.only(left=20, right=20, top=14, bottom=10)))
        controls.append(dpanel(ft.ResponsiveRow([
            ft.Container(col={"xs": 12, "lg": 6}, content=s_bars("Спрос товара по дням", "Покупателей искали и могли увидеть товар",
                                                                product["daily"], "searches", pal["accent"], "искавших")),
            ft.Container(col={"xs": 12, "lg": 6}, content=s_bars("Заказы товара по дням", "Заказано штук", product["daily"], "orders",
                                                                pal["good"], "шт.")),
        ], spacing=24, run_spacing=16), padding=ft.Padding.only(left=20, right=20, top=14, bottom=12)))
        changes = product["changes"]
        change_lines = [dtext(f"{c['date'][8:10]}.{c['date'][5:7]} — изменено: {', '.join(c['parts'])}", 13) for c in changes[:10]]
        controls.append(dpanel(ft.Column([
            dtext("Изменения карточки", 15, weight=ft.FontWeight.W_700),
            *(change_lines or [dtext(f"Программа сохраняет снимок карточки каждый день (первый — {product['card']['date'] or '—'}). "
                                     "Когда название, описание, характеристики или фото изменятся, здесь появится дата — по ней будет видно, "
                                     "как правка повлияла на спрос и позиции.", 13, pal["muted"])]),
        ], spacing=6), padding=ft.Padding.only(left=20, right=20, top=14, bottom=12)))
        return controls

    def search_positions(data: dict) -> list[ft.Control]:
        pal = dash_pal()
        serp = data["serp"]
        status_lines: list[ft.Control] = []

        def status(ok: bool, text: str) -> ft.Control:
            return ft.Row([ft.Icon(ft.Icons.CHECK_CIRCLE if ok else ft.Icons.REMOVE, size=16, color=pal["good"] if ok else pal["muted"]),
                           dtext(text, 13, expand=True)], spacing=8)

        status_lines.append(status(serp["installed"], "Браузер для парсера установлен" if serp["installed"]
                                   else "Браузер для парсера не установлен: scripts/serp_profile.py setup"))
        ready_regions = [r for r in serp["regions"] if r["profile"] and r["active"]]
        status_lines.append(status(bool(ready_regions), "Профили регионов: " + (", ".join(r["name"] for r in ready_regions) if ready_regions
                                                                                else "нет — scripts/serp_profile.py create msk \"Москва\"")))
        status_lines.append(status(serp["enabled"], "Парсер включён" if serp["enabled"] else "Парсер выключен: добавьте SERP_PARSER_ENABLED=1 в .env"))
        if serp["blocked_until"]:
            until = datetime.fromisoformat(serp["blocked_until"]).astimezone()
            status_lines.append(ft.Row([ft.Icon(ft.Icons.WARNING_AMBER_ROUNDED, size=16, color=pal["bad"]),
                                        dtext(f"Ozon показал капчу — замеры на паузе до {until:%d.%m %H:%M}.", 13, pal["bad"])], spacing=8))
        if serp["last_run"]:
            run = serp["last_run"]
            status_lines.append(dtext(f"Последний замер: {str(run['started_at'])[:16].replace('T', ' ')} UTC — {run['status']}"
                                      + (f" ({run['message']})" if run.get("message") else ""), 12, pal["muted"]))
        controls: list[ft.Control] = [dpanel(ft.Column([
            dtext("Позиции в выдаче по регионам", 15, weight=ft.FontWeight.W_700),
            dtext("Без подписки Premium Ozon не отдаёт позиции через API, поэтому их раз в день замеряет парсер сайта: до 15 запросов "
                  "в каждом регионе, с паузами. Позиция — медиана за 3 дня: один удачный замер не делает запрос «в топе».", 12, pal["muted"]),
            ft.Container(height=4), *status_lines,
        ], spacing=6), padding=ft.Padding.only(left=20, right=20, top=14, bottom=12))]
        rows = data["positions"]
        if not rows:
            controls.append(dpanel(dtext("Замеров пока нет. Когда парсер заработает, здесь будет таблица «запрос × регион» с изменением "
                                         "за неделю и подсветкой запросов, где позиции в регионах сильно расходятся — это сигнал "
                                         "перераспределить остатки по кластерам.", 13, pal["muted"])))
            return controls
        regions = [r["name"] for r in serp["regions"] if r["active"]]
        columns = [("Запрос", 5, None), ("Артикул", 2, None)] + [(name, 2, "Медиана за 3 дня; в скобках — изменение к замеру неделю назад") for name in regions]
        table: list[ft.Control] = [s_header(columns)]
        for row in rows:
            cells = [s_cell(dtext(row["query"] or "", 13, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS), 5, right=False),
                     s_cell(dtext(str(row["offer_id"] or row["sku"]), 13, weight=ft.FontWeight.W_700), 2)]
            for name in regions:
                value = row["regions"].get(name)
                if not value:
                    cells.append(s_cell(dtext("—", 13, pal["muted"]), 2))
                    continue
                moved = value["change"]
                suffix = "" if not moved else f" ({'+' if moved > 0 else '−'}{abs(moved):.0f})"
                cells.append(s_cell(ft.Row([s_position(value["position"], True), dtext(suffix, 11, pal["good"] if moved and moved > 0 else pal["bad"])],
                                           spacing=2, tight=True, alignment=ft.MainAxisAlignment.END), 2))
            spread = row["spread"]
            table.append(s_row(cells, highlight=spread is not None and spread >= 20))
        controls.append(dpanel(ft.Column([dtext("Запрос × регион", 15, weight=ft.FontWeight.W_700),
                                          dtext("Подсвечены запросы, где позиции в регионах расходятся на 20+ мест или товар найден не везде.", 12, pal["muted"]),
                                          ft.Container(height=6), ft.Column(table, spacing=0)], spacing=4),
                               padding=ft.Padding.only(left=20, right=20, top=14, bottom=10)))
        return controls

    def render_search_view(push_update: bool = True) -> None:
        pal = dash_pal()
        data = search_state.get("data") or {}
        parts = ["по каким запросам находят товары и чего не хватает карточкам"]
        if data and not data.get("empty"):
            last = data["last_date"]
            parts.append(f"данные Ozon по {last[8:10]}.{last[5:7]} (Ozon отдаёт их через 2 дня)")
            if data.get("collected_at"):
                collected = datetime.fromisoformat(data["collected_at"]).astimezone()
                parts.append(f"загружено {collected:%d.%m в %H:%M}")
        search_subtitle.value = " · ".join(parts)
        search_refresh_button.disabled = bool(search_state["busy"]) or job_lock.locked()

        controls: list[ft.Control] = []
        if search_view is None:
            controls.append(dpanel(dtext(f"Модуль поиска не загрузился: {SEARCH_IMPORT_ERROR}", 13, pal["bad"])))
        if search_state.get("error"):
            controls.append(dpanel(dtext(str(search_state["error"]), 13, pal["bad"])))
        if search_state["busy"] and not data:
            controls.append(dpanel(ft.Row([ft.ProgressRing(width=20, height=20, stroke_width=2.4),
                                           dtext("Считаю спрос, кластеры и покрытие карточек…", 13, pal["muted"])], spacing=12)))
        elif data.get("empty"):
            controls.append(dpanel(dtext("Данных пока нет. Нажмите «Обновить данные» — программа загрузит из Ozon запросы за последний месяц "
                                         "(1–3 минуты). Дальше данные обновляются сами раз в день при запуске программы.", 13, pal["muted"])))
        elif data:
            controls.append(ft.Row([s_chip(label, search_state["tab"] == key, lambda _e, k=key: search_set_tab(k)) for key, label in SEARCH_TABS],
                                   spacing=8, wrap=True))
            tab = search_state["tab"]
            builder = {"overview": search_overview, "queries": search_queries, "products": search_products, "positions": search_positions}[tab]
            controls.extend(builder(data))
        search_body.controls = controls
        apply_control_theme(search_body, current_dark["value"])
        if push_update and page.controls and host.controls and host.controls[0] is search_view_section:
            page.update(search_view_section)

    search_view_section = ft.Column(
        [
            ft.ResponsiveRow(
                [
                    ft.Container(col={"xs": 12, "lg": 8},
                                 content=ft.Column([ft.Text("Поиск и ключевые слова", size=26, weight=ft.FontWeight.W_800, color=TEXT),
                                                    search_subtitle], spacing=2)),
                    ft.Container(col={"xs": 12, "lg": 4}, alignment=ft.Alignment(1, 0), content=search_refresh_button),
                ],
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
                run_spacing=12,
            ),
            search_body,
        ],
        spacing=14,
    )

    views = [dash, ai_view, reports, pricing, actions, ads_view, analysis_view, search_view_section, supply, stock_view, settings, file_view]
    active_store = {"value": "ozon"}
    store_sections = {"ozon": 0, "wb": 0}
    # Разделы меню каждого магазина (индексы в views): у WB только то, что уже работает.
    store_nav = {
        "ozon": list(range(len(views))),
        "wb": [views.index(view) for view in (dash, pricing, actions, supply, stock_view, settings)],
    }
    section_names = ["Бизнес-сводка", "AI-ассистент", "ABC/XYZ-анализ", "Цены и скидки", "Акции", "Реклама", "Анализ", "Поиск", "Поставки", "Остатки", "Настройки", "Файлы"]

    DASHBOARD_HELP.update(WB_DASHBOARD_HELP)
    wb_dash = WbDashboard(
        root=ROOT, page=page, months=MONTHS,
        ui=SimpleNamespace(dpanel=dpanel, dtext=dtext, dash_stat=dash_stat, dash_stats_panel=dash_stats_panel,
                           dash_list=dash_list, dash_section_title=dash_section_title, dash_pal=dash_pal,
                           money=dash_money, percent=_format_percent, number=_format_number, primary=PRIMARY),
        log=lambda message: ui_queue.put(("log_append", message)),
        reveal=reveal_path, set_sync=set_dashboard_sync, toast=toast,
        on_change=lambda: render_dashboard() if is_wb_store() and host.controls and host.controls[0] is dash else None,
    )

    wb_prices_view = WbPricesView(
        root=ROOT, page=page, log=lambda message: ui_queue.put(("log_append", message)), reveal=reveal_path, toast=toast,
        ui=SimpleNamespace(dpanel=dpanel, dtext=dtext, dash_stat=dash_stat, dash_pal=dash_pal, money=dash_money,
                           percent=_format_percent, primary=PRIMARY, title_color=TEXT, muted_color=MUTED,
                           cell=lambda column, payload, expand: _build_costs_cell(column, payload, current_dark["value"], expand)),
    )

    wb_promotions_view = WbPromotionsView(
        root=ROOT, page=page, log=lambda message: ui_queue.put(("log_append", message)), toast=toast,
        ui=SimpleNamespace(dpanel=dpanel, dtext=dtext, dash_stat=dash_stat, dash_pal=dash_pal, money=dash_money,
                           percent=_format_percent, primary=PRIMARY, title_color=TEXT, muted_color=MUTED,
                           cell=lambda column, payload, expand: _build_costs_cell(column, payload, current_dark["value"], expand)),
    )

    def wb_section(index):
        if index == 0:
            return dash
        if views[index] is actions:
            wb_promotions_view.enter()
            return wb_promotions_view.view
        if views[index] is pricing:
            wb_prices_view.enter()
            return wb_prices_view.view
        if views[index] is settings:
            return settings
        if views[index] is supply:
            render_wb_supply_results(push_update=False)
            return wb_supply
        if views[index] is stock_view:
            render_stock_view(push_update=False)
            return stock_view
        return ft.Column(
            [
                ft.Text("Wildberries", size=13, weight=ft.FontWeight.W_600, color=MUTED),
                ft.Text(section_names[index], size=26, weight=ft.FontWeight.W_800, color=TEXT),
                card(
                    "Раздел Wildberries готовится",
                    "Здесь появятся данные вашего магазина Wildberries. Сейчас можно сохранить API-токен в настройках.",
                    [ft.Button("Настроить Wildberries", on_click=open_marketplace_settings)],
                ),
            ], spacing=18,
        )

    host = ft.Column([views[0]], expand=True, spacing=0, scroll=ft.ScrollMode.AUTO)

    # Синхронизация цен при входе во вкладку «Цены и скидки» — не чаще раза в 5 минут.
    PRICING_AUTO_SYNC_INTERVAL = 300
    pricing_auto_sync = {"last": 0.0}

    def auto_sync_pricing() -> None:
        if time.monotonic() - pricing_auto_sync["last"] < PRICING_AUTO_SYNC_INTERVAL:
            return
        if job_lock.locked() or pricing_costs_dirty["value"]:
            return  # идёт другая задача или в таблице несохранённые правки
        pricing_auto_sync["last"] = time.monotonic()
        job(PRICING_SYNC_TITLE, [str(PYTHON), str(SCRIPTS / "ui_actions.py"), "refresh-pricing"])

    def enter_pricing_tab() -> None:
        """При входе во вкладку: сначала свежие заявки на скидку, затем синхронизация цен."""
        if discount_dialog_state["busy"]:
            auto_sync_pricing()
            return
        open_discount_requests_dialog(quiet=True, on_done=auto_sync_pricing)

    def show_store_section(selected_index: int) -> None:
        store_sections[active_store["value"]] = selected_index
        is_ozon = active_store["value"] == "ozon"
        host.controls = [views[selected_index] if is_ozon else wb_section(selected_index)]
        if views[selected_index] is dash:
            # Один экран сводки на оба магазина: при входе показываем данные выбранного магазина.
            dashboard_more.items[2].visible = not is_ozon
            if dashboard_state.get("store") != active_store["value"]:
                dashboard_state["store"] = active_store["value"]
                dashboard_chart.visible = False
                dashboard_chart_note.value = ""
                set_dashboard_chart_feedback("empty", "Постройте график", "Выберите метрику и период, затем постройте график.")
            refresh_dashboard(push_update=False)
        if is_ozon and views[selected_index] is ads_view:
            ads_nav["screen"] = "types"
            render_ads_view(push_update=False)
        if is_ozon and views[selected_index] is analysis_view:
            render_analysis_view(push_update=False)
        if is_ozon and views[selected_index] is actions:
            render_limit_panel(push_update=False)
            render_actions_calendar(push_update=False)
            render_auto_panel(push_update=False)
        if is_ozon and views[selected_index] is stock_view:
            render_stock_view(push_update=False)
        if is_ozon and views[selected_index] is search_view_section:
            render_search_view(push_update=False)
        is_ai_view = is_ozon and selected_index == 1
        host.scroll = ft.ScrollMode.HIDDEN if is_ai_view else ft.ScrollMode.AUTO
        host_shell.padding = 0 if is_ai_view else ft.Padding.only(left=8, right=10, top=20, bottom=20)
        host_shell.bgcolor = ai_surface_tokens()["canvas"] if is_ai_view else ("#FBFCFF" if not current_dark["value"] else "#1E2328")
        host_shell.gradient = ft.LinearGradient(
            begin=ft.Alignment(0, -1),
            end=ft.Alignment(0, 1),
            colors=[ai_surface_tokens()["canvas"], ai_surface_tokens()["canvas"]]
            if is_ai_view
            else (["#FBFCFF", "#F3F6FB"] if not current_dark["value"] else ["#1F252A", "#191D21"]),
        )
        apply_control_theme(host.controls[0], current_dark["value"])
        if is_ai_view:
            ai_apply_chat_theme()
        store_name = "Ozon" if is_ozon else "Wildberries"
        page.title = f"{store_name} — OzonReportX"
        section_name = "Цены" if not is_ozon and views[selected_index] is pricing else section_names[selected_index]
        set_status("Раздел открыт", f"{store_name} · {section_name}", PRIMARY, busy=False)
        page.update()
        if is_ozon and views[selected_index] is pricing:
            enter_pricing_tab()
        if is_ozon and views[selected_index] is actions and not calendar_state["started"]:
            calendar_load()
        if is_ozon and views[selected_index] is actions and not auto_state["started"]:
            auto_load()
        if is_ozon and views[selected_index] is actions and not limit_state["started"]:
            limit_load()
        if views[selected_index] is stock_view and not stock_state["started"]:
            stock_load()
        if is_ozon and views[selected_index] is search_view_section and not search_state["started"]:
            search_load()

    def open_section(view) -> None:
        index = views.index(view)
        rail.selected_index = store_nav[active_store["value"]].index(index)
        show_store_section(index)

    def nav_change(e: ft.ControlEvent) -> None:
        show_store_section(store_nav[active_store["value"]][e.control.selected_index])

    theme_button = ft.IconButton(
        icon=ft.Icons.DARK_MODE_OUTLINED,
        selected_icon=ft.Icons.LIGHT_MODE_OUTLINED,
        tooltip="Переключить тему",
        selected=start_dark,
    )

    sidebar_button = ft.IconButton(
        icon=ft.Icons.MENU_OPEN_ROUNDED,
        selected_icon=ft.Icons.MENU_ROUNDED,
        tooltip="Свернуть меню",
        selected=False,
    )

    store_buttons = {}

    def style_store_tabs():
        dark = current_dark["value"]
        rail_brand_card.bgcolor = "#252A2F" if dark else "#F7F7FA"
        rail_brand_card.border = ft.Border.all(1, "#394247" if dark else BORDER)
        for key, button in store_buttons.items():
            selected = active_store["value"] == key
            button.bgcolor = ("#0D6B5B" if dark else "#DCEFE9") if selected else "transparent"
            button.content.controls[0].color = ("#FFFFFF" if dark else PRIMARY) if selected else ("#A1A1A6" if dark else MUTED)
            button.content.controls[1].color = ("#FFFFFF" if dark else TEXT) if selected else ("#A1A1A6" if dark else MUTED)
            button.content.controls[2].visible = selected and not sidebar_collapsed["value"]
            button.content.controls[2].color = "#FFFFFF" if dark else PRIMARY

    def select_store(key):
        if active_store["value"] == key:
            return
        active_store["value"] = key
        # У WB нет заявок на скидку — раздел называется просто «Цены».
        rail_destinations[views.index(pricing)].label = "Цены" if key == "wb" else "Цены и скидки"
        rail.destinations = [rail_destinations[index] for index in store_nav[key]]
        rail.selected_index = store_nav[key].index(store_sections[key])
        show_store_section(store_sections[key])
        style_store_tabs()
        page.update()

    for key, title in (("ozon", "Ozon"), ("wb", "Wildberries")):
        store_buttons[key] = ft.Container(
            content=ft.Row([
                ft.Icon(ft.Icons.STOREFRONT_OUTLINED, size=20),
                ft.Text(title, size=16, weight=ft.FontWeight.W_700, expand=True),
                ft.Icon(ft.Icons.CHECK_ROUNDED, size=16, color="white", visible=key == "ozon"),
            ], spacing=10),
            padding=ft.Padding.symmetric(horizontal=12, vertical=11),
            border_radius=16,
            on_click=lambda _e, store=key: select_store(store),
            tooltip=f"Открыть магазин {title}",
        )
    rail_brand_card = ft.Container(
        width=216,
        padding=7,
        border_radius=22,
        bgcolor="#F7F7FA",
        border=ft.Border.all(1, BORDER),
        content=ft.Column(list(store_buttons.values()), spacing=4),
    )

    rail_destinations = [
        ft.NavigationRailDestination(icon=ft.Icons.SPACE_DASHBOARD_OUTLINED, selected_icon=ft.Icons.SPACE_DASHBOARD, label="Бизнес-сводка"),
        ft.NavigationRailDestination(icon=ft.Icons.AUTO_AWESOME_OUTLINED, selected_icon=ft.Icons.AUTO_AWESOME, label="AI-ассистент"),
        ft.NavigationRailDestination(icon=ft.Icons.GRID_VIEW, selected_icon=ft.Icons.GRID_VIEW_ROUNDED, label="ABC/XYZ-анализ"),
        ft.NavigationRailDestination(icon=ft.Icons.LOCAL_OFFER_OUTLINED, selected_icon=ft.Icons.LOCAL_OFFER, label="Цены и скидки"),
        ft.NavigationRailDestination(icon=ft.Icons.LOCAL_ACTIVITY_OUTLINED, selected_icon=ft.Icons.LOCAL_ACTIVITY, label="Акции"),
        ft.NavigationRailDestination(icon=ft.Icons.CAMPAIGN_OUTLINED, selected_icon=ft.Icons.CAMPAIGN, label="Реклама"),
        ft.NavigationRailDestination(icon=ft.Icons.INSIGHTS_OUTLINED, selected_icon=ft.Icons.INSIGHTS, label="Анализ"),
        ft.NavigationRailDestination(icon=ft.Icons.MANAGE_SEARCH_OUTLINED, selected_icon=ft.Icons.MANAGE_SEARCH, label="Поиск"),
        ft.NavigationRailDestination(icon=ft.Icons.INVENTORY_2_OUTLINED, selected_icon=ft.Icons.INVENTORY_2, label="Поставки"),
        ft.NavigationRailDestination(icon=ft.Icons.WAREHOUSE_OUTLINED, selected_icon=ft.Icons.WAREHOUSE, label="Остатки"),
        ft.NavigationRailDestination(icon=ft.Icons.SETTINGS_OUTLINED, selected_icon=ft.Icons.SETTINGS, label="Settings"),
        ft.NavigationRailDestination(icon=ft.Icons.FOLDER_OPEN_OUTLINED, selected_icon=ft.Icons.FOLDER_OPEN, label="Files"),
    ]

    rail = ft.NavigationRail(
        selected_index=0,
        extended=True,
        min_width=92,
        min_extended_width=240,
        group_alignment=-0.88,
        label_type=ft.NavigationRailLabelType.ALL,
        bgcolor=SURFACE,
        indicator_color=PRIMARY_SOFT,
        on_change=nav_change,
        leading=ft.Container(
            content=rail_brand_card,
            padding=ft.Padding.only(bottom=12),
        ),
        destinations=list(rail_destinations),
    )

    # NavigationRail не умеет прокручиваться сам: даём ему естественную высоту
    # и прокручиваем оболочку, чтобы на низких экранах меню не выдавливало статус.
    rail_shell = ft.Container(
        content=ft.Column([ft.Container(content=rail, height=680)], scroll=ft.ScrollMode.AUTO, spacing=0),
        expand=True,
        padding=12,
        bgcolor=SURFACE,
        border_radius=34,
        border=ft.Border.all(1, BORDER),
        shadow=ft.BoxShadow(blur_radius=34, spread_radius=0, color="#12000000", offset=ft.Offset(0, 14)),
    )
    host_shell = ft.Container(
        content=host,
        expand=True,
        padding=ft.Padding.only(left=8, right=10, top=20, bottom=20),
        bgcolor="#FBFCFF",
        gradient=ft.LinearGradient(
            begin=ft.Alignment(0, -1),
            end=ft.Alignment(0, 1),
            colors=["#1F252A", "#191D21"],
        ),
        border_radius=30,
        border=ft.Border.all(1, "#ECECF1"),
        shadow=ft.BoxShadow(blur_radius=34, spread_radius=0, color="#12000000", offset=ft.Offset(0, 14)),
    )
    cancel_row = ft.Row([cancel_button], alignment=ft.MainAxisAlignment.END, visible=False)
    status_shell = ft.Container(
        content=ft.Column(
            [
                ft.Row(
                    [
                        status_indicator,
                        state,
                        spin,
                    ],
                    spacing=10,
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                ),
                status_detail,
                cancel_row,
            ],
            spacing=6,
        ),
        padding=ft.Padding.symmetric(horizontal=14, vertical=10),
        bgcolor=SURFACE,
        border_radius=16,
        border=ft.Border.all(1, BORDER),
        shadow=ft.BoxShadow(blur_radius=12, spread_radius=0, color="#0D000000", offset=ft.Offset(0, 4)),
    )
    cancel_button.on_click = cancel_current_action
    def apply_theme(dark: bool) -> None:
        current_dark["value"] = dark
        page.theme_mode = ft.ThemeMode.DARK if dark else ft.ThemeMode.LIGHT
        page.bgcolor = "#111214" if dark else BG
        rail.bgcolor = "#20262B" if dark else SURFACE
        rail.indicator_color = "#2F4A45" if dark else PRIMARY_SOFT
        rail_shell.bgcolor = "#20262B" if dark else SURFACE
        rail_shell.border = ft.Border.all(1, "#394247" if dark else "#E1E5EE")
        rail_shell.shadow = ft.BoxShadow(
            blur_radius=28,
            spread_radius=0,
            color="#22000000" if dark else "#14000000",
            offset=ft.Offset(0, 12),
        )
        host_shell.bgcolor = ai_surface_tokens()["canvas"] if active_store["value"] == "ozon" and rail.selected_index == 1 else ("#1E2328" if dark else "#FBFCFF")
        host_shell.gradient = ft.LinearGradient(
            begin=ft.Alignment(0, -1),
            end=ft.Alignment(0, 1),
            colors=[ai_surface_tokens()["canvas"], ai_surface_tokens()["canvas"]]
            if active_store["value"] == "ozon" and rail.selected_index == 1
            else (["#1F252A", "#191D21"] if dark else ["#FBFCFF", "#F3F6FB"]),
        )
        host_shell.border = ft.Border.all(1, "#2C2C2E" if dark else "#ECECF1")
        host_shell.shadow = ft.BoxShadow(
            blur_radius=30,
            spread_radius=0,
            color="#18000000" if dark else "#12000000",
            offset=ft.Offset(0, 12),
        )
        status_shell.bgcolor = "#1F252A" if dark else "#FBFCFF"
        status_shell.border = ft.Border.all(1, "#353D42" if dark else "#E1E5EE")
        status_shell.shadow = ft.BoxShadow(
            blur_radius=12,
            spread_radius=0,
            color="#14000000" if dark else "#0D000000",
            offset=ft.Offset(0, 4),
        )
        status_detail.color = "#A1A1A6" if dark else "#6E6E73"
        top_tools_shell.bgcolor = "#D020252A" if dark else "#FAFBFE"
        top_tools_shell.border = ft.Border.all(1, "#3A4348" if dark else "#E1E5EE")
        splash_overlay.bgcolor = "#CC111214" if dark else "#D8F5F5F7"
        splash_card.bgcolor = "#20262B" if dark else "#FFFFFF"
        splash_card.border = ft.Border.all(1, "#354148" if dark else BORDER)
        splash_card.shadow = ft.BoxShadow(
            blur_radius=40,
            spread_radius=0,
            color="#28000000" if dark else "#16000000",
            offset=ft.Offset(0, 16),
        )
        splash_title.color = "#F5F5F7" if dark else "#1D1D1F"
        splash_tagline.color = "#A1A1A6" if dark else "#6E6E73"
        splash_note.color = "#A1A1A6" if dark else "#6E6E73"
        theme_button.selected = dark
        theme_button.icon_color = "#B6C0BC" if dark else "#4F5B66"
        theme_button.selected_icon_color = "#B6C0BC" if dark else "#4F5B66"
        sidebar_button.icon_color = "#B6C0BC" if dark else "#4F5B66"
        sidebar_button.selected_icon_color = "#B6C0BC" if dark else "#4F5B66"
        menu_bar_inner.bgcolor = "#23282D" if dark else "#FFFFFF"
        menu_bar_inner.border = ft.Border.all(1, "#384046" if dark else "#E5E7EF")
        apply_control_theme(rail_brand_card, dark)
        apply_control_theme(rail_shell, dark)
        _render_log_lines(log_state["lines"])
        apply_prompt_dialog_theme()
        apply_discount_dialog_theme()
        refresh_pricing_costs_preview()
        for root_control in page.controls:
            apply_control_theme(root_control, dark)
        ai_apply_chat_theme()
        render_supply_results(push_update=False)
        render_ads_results(push_update=False)
        render_ads_view(push_update=False)
        render_analysis_view(push_update=False)
        render_actions_calendar(push_update=False)
        render_dashboard(push_update=False)
        render_abc_view(push_update=False)
        render_wb_supply_results(push_update=False)
        wb_prices_view.render(push_update=False)
        wb_promotions_view.render(push_update=False)
        render_auto_panel(push_update=False)
        render_limit_panel(push_update=False)
        render_stock_view(push_update=False)
        style_store_tabs()
        set_status(state.value, status_detail.value, status_tone["value"], busy=spin.visible, cancellable=cancel_button.visible)
        page.update()

    def toggle_theme(_e=None) -> None:
        dark = not theme_button.selected
        # Выбор, совпадающий с системой, снова включает автоматическое следование за ней.
        follow_system_theme["value"] = dark == system_prefers_dark(page.platform_brightness)
        apply_theme(dark)

    def on_platform_brightness_change(_e) -> None:
        if follow_system_theme["value"]:
            # Flet 0.84 присылает обычный Event без поля brightness, но к этому моменту обновляет страницу.
            dark = system_prefers_dark(page.platform_brightness)
            if dark != current_dark["value"]:
                apply_theme(dark)

    theme_button.on_click = toggle_theme

    def apply_sidebar_layout() -> None:
        collapsed = sidebar_collapsed["value"]
        sidebar.width = 92 if collapsed else 264
        sidebar_button.selected = collapsed
        sidebar_button.tooltip = "Развернуть меню" if collapsed else "Свернуть меню"
        tools = [menu_bar_inner, theme_button, sidebar_button]
        top_tools_shell.content = (
            ft.Column(tools, spacing=4, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
            if collapsed
            else ft.Row(tools, spacing=10, vertical_alignment=ft.CrossAxisAlignment.CENTER)
        )
        menu_bar_inner.padding = ft.Padding.symmetric(horizontal=4 if collapsed else 10, vertical=6)
        rail.extended = not collapsed
        rail.label_type = ft.NavigationRailLabelType.NONE if collapsed else ft.NavigationRailLabelType.ALL
        rail.min_width = 56 if collapsed else 92
        for destination in rail_destinations:
            destination.tooltip = destination.label if collapsed else None
        rail_shell.padding = 6 if collapsed else 12
        rail_brand_card.width = 56 if collapsed else 216
        rail_brand_card.padding = 4 if collapsed else 7
        for button in store_buttons.values():
            button.content.controls[1].visible = not collapsed
            button.padding = ft.Padding.symmetric(horizontal=14 if collapsed else 12, vertical=11)
        state.visible = not collapsed
        status_detail.visible = not collapsed
        cancel_button.content = ft.Icon(ft.Icons.CLOSE_ROUNDED, size=16, color="white") if collapsed else "Отмена"
        cancel_button.tooltip = "Отмена" if collapsed else None
        status_shell.padding = ft.Padding.symmetric(horizontal=8 if collapsed else 14, vertical=10)
        style_store_tabs()
        ai_apply_chat_theme()
        set_status(state.value, status_detail.value, status_tone["value"], busy=spin.visible, cancellable=cancel_button.visible)

    def toggle_sidebar(_e=None) -> None:
        sidebar_collapsed["value"] = not sidebar_collapsed["value"]
        save_ui_state({"sidebar_collapsed": sidebar_collapsed["value"]})
        apply_sidebar_layout()
        page.update()

    sidebar_button.on_click = toggle_sidebar
    page.on_platform_brightness_change = on_platform_brightness_change

    menu_bar_inner = ft.Container(
        content=ft.Row(
            [
                ai_tools_button,
            ],
            spacing=4,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
            tight=True,
        ),
        padding=ft.Padding.symmetric(horizontal=10, vertical=6),
        border_radius=14,
        bgcolor="#23282D",
        border=ft.Border.all(1, "#384046"),
    )
    top_tools_shell = ft.Container(
        content=ft.Row(
            [
                menu_bar_inner,
                theme_button,
                sidebar_button,
            ],
            spacing=10,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
        ),
        padding=ft.Padding.symmetric(horizontal=8, vertical=8),
        bgcolor="#D020252A" if current_dark["value"] else "#FAFBFE",
        border_radius=18,
        border=ft.Border.all(1, "#3A4348" if current_dark["value"] else "#E1E5EE"),
        shadow=ft.BoxShadow(blur_radius=16, spread_radius=0, color="#12000000", offset=ft.Offset(0, 6)),
    )
    splash_progress = ft.ProgressRing(width=26, height=26, stroke_width=2.6, color=PRIMARY)
    splash_title = ft.Text("OzonReportX", size=22, weight=ft.FontWeight.W_800, color="#F5F5F7")
    splash_tagline = ft.Text("Контроль. Аналитика. Автоматизация.", size=12, color="#A1A1A6")
    splash_note = ft.Text("Подготавливаем рабочее пространство", size=12, color="#A1A1A6")
    splash_card = ft.Container(
        width=360,
        padding=28,
        border_radius=32,
        bgcolor="#20262B",
        border=ft.Border.all(1, "#354148"),
        shadow=ft.BoxShadow(blur_radius=40, spread_radius=0, color="#28000000", offset=ft.Offset(0, 16)),
        content=ft.Column(
            [
                ft.Row(
                    [
                        ft.Container(
                            content=ft.Image(
                                src=str((CODE_ROOT / "img" / "ozonreportx-mark.svg").resolve()),
                                width=42,
                                height=42,
                            ),
                            width=64,
                            height=64,
                            bgcolor="#0D6B5B",
                            border_radius=22,
                            alignment=ft.Alignment(0, 0),
                        ),
                    ],
                    alignment=ft.MainAxisAlignment.CENTER,
                ),
                ft.Column(
                    [
                        splash_title,
                        splash_tagline,
                    ],
                    spacing=4,
                    horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                ),
                ft.Container(height=1, bgcolor="#354148", border_radius=10),
                ft.Row(
                    [
                        splash_progress,
                        splash_note,
                    ],
                    alignment=ft.MainAxisAlignment.CENTER,
                    spacing=12,
                ),
            ],
            spacing=18,
            horizontal_alignment=ft.CrossAxisAlignment.CENTER,
        ),
    )
    splash_overlay = ft.Container(
        expand=True,
        bgcolor="#D8F5F5F7",
        alignment=ft.Alignment(0, 0),
        content=splash_card,
        visible=True,
    )
    # Левая колонка в обычном потоке: меню сжимается/прокручивается,
    # а статус задачи всегда остаётся видимым внизу.
    sidebar = ft.Column(
        [top_tools_shell, rail_shell, status_shell],
        width=264,
        spacing=12,
        horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
    )
    main_layout = ft.Row(
        [sidebar, host_shell],
        expand=True,
        spacing=16,
        vertical_alignment=ft.CrossAxisAlignment.STRETCH,
    )
    page.add(main_layout)
    page.overlay.append(splash_overlay)
    status_ui_ready["value"] = True
    set_status("Запуск", "Подготавливаем интерфейс и данные.", ACCENT, busy=True)
    apply_sidebar_layout()
    apply_theme(start_dark)

    def auto_collect_search() -> None:
        """
        Ежедневный сбор поисковой аналитики (scripts/search_collect.py) при запуске, если сегодня его ещё не было,
        затем — если включён SERP_PARSER_ENABLED — замер позиций в выдаче (scripts/serp_collect.py).
        Без job_lock: оба только читают Ozon и держат свои блокировки, пользователю не мешают.
        """
        settings = read_settings(ROOT)
        if search_db is None or not settings.get("OZON_CLIENT_ID") or not settings.get("OZON_API_KEY"):
            return
        serp_enabled = str(settings.get("SERP_PARSER_ENABLED") or "").strip().lower() in ("1", "true", "yes", "да")
        try:
            collect_due = not search_db.collected_today()
        except Exception as exc:
            log_store.append(f"Поисковая аналитика: не удалось открыть базу — {exc}")
            return
        # Замер позиций сам пропускает уже замеренное сегодня, поэтому запускаем его, даже если аналитика уже собрана.
        if not collect_due and not serp_enabled:
            return

        def report(title: str, code: int, output: str, message: str | None) -> None:
            def done() -> None:
                log_store.append(f"{title} (фон):\n" + output)
                # Плашку статуса не трогаем, пока идёт задача пользователя: сообщение осталось бы поверх её статуса.
                if message and not job_lock.locked():
                    toast(message, error=True)
                if search_state["started"]:
                    search_load()  # раздел «Поиск» уже открывали — показываем свежие данные
            ui_queue.put(("call", done))

        def worker() -> None:
            if collect_due:
                code, output = stream_process([PYTHON, str(SCRIPTS / "search_collect.py")], on_line=lambda _line: None)
                report("Сбор поисковой аналитики", code, output,
                       "Поисковая аналитика собрана не полностью — подробности в журнале." if code == 1 else None)
            if not serp_enabled:
                return
            # Замер позиций идёт 20–30 минут: паузы между запросами, чтобы Ozon не показал капчу.
            code, output = stream_process([PYTHON, str(SCRIPTS / "serp_collect.py")], on_line=lambda _line: None)
            messages = {1: "Замер позиций в выдаче прошёл с ошибками — подробности в журнале.",
                        3: "Ozon показал капчу — замер позиций остановлен на сутки. Подробности в журнале."}
            report("Замер позиций в выдаче", code, output, messages.get(code))

        page.run_thread(worker)

    async def initialize_app() -> None:
        page.update()
        await asyncio.sleep(0.12)
        refresh_dashboard()
        refresh_files()
        refresh_pricing_costs_preview()
        render_supply_results(push_update=False)
        render_ads_results(push_update=False)
        render_ads_view(push_update=False)
        render_analysis_view(push_update=False)
        render_actions_calendar(push_update=False)
        splash_overlay.visible = False
        set_status("Готово", "Интерфейс готов к работе.", PRIMARY, busy=False)
        page.update()
        if needs_setup(ROOT):
            open_marketplace_settings()
        auto_collect_search()

    page.run_task(ui_event_loop)
    page.run_task(initialize_app)


if __name__ == "__main__":
    ft.run(main)
